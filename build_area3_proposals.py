import openpyxl
from pathlib import Path

excel_dir = Path('/Users/arqovalle/Library/CloudStorage/OneDrive-Personal/Software/Antigravity/Cuentas Kinal/Programas de Curso')
output_dir = Path('Revision_Temarios')

# -------------------------------------------------------------
# 1. MANTENIMIENTO MECÁNICO INDUSTRIAL (2026_MMIM.xlsx)
# -------------------------------------------------------------
def generate_mmi():
    wb = openpyxl.load_workbook(excel_dir / '2026_MMIM.xlsx', data_only=True)
    tm = wb['Temas Módulos']

    modules_data = []
    for c in range(1, tm.max_column + 1, 2):
        m_name = tm.cell(row=1, column=c).value
        m_num = tm.cell(row=2, column=c).value
        if m_name:
            topics = []
            for r in range(3, tm.max_row + 1):
                val = tm.cell(row=r, column=c).value
                if val and str(val).strip() and str(val).strip() != '.':
                    topics.append(str(val).strip())
            modules_data.append({
                'num': str(m_num).strip(),
                'name': str(m_name).strip(),
                'topics': topics
            })

    classified_modules = []
    for mod in modules_data:
        m_num = mod['num']
        m_name = mod['name']
        mod_topics = []
        for t in mod['topics']:
            status = 'keep'
            category = 'Mecánica Industrial'
            badge = '🟢 Mantener'
            badge_color = 'bg-emerald-100 text-emerald-800'
            reason = 'Fundamento mecánico y operativo indispensable para el mantenimiento preventivo en plantas productivas de Guatemala.'

            t_lower = t.lower()
            if 'el petróleo' in t_lower or 'proceso de destilación' in t_lower:
                status = 'delete'
                badge = '🔴 Eliminar'
                badge_color = 'bg-red-100 text-red-800'
                category = 'Teoría Petroquímica Excesiva'
                reason = '🔴 Eliminar: Teoría petroquímica abstracta sin impacto práctico en planta; el mecánico requiere dominar tablas de viscosidad ISO VG, compatibilidad de grasas y puntos de goteo.'
            elif 'historia de la seguridad' in t_lower:
                status = 'delete'
                badge = '🔴 Eliminar'
                badge_color = 'bg-red-100 text-red-800'
                category = 'Historia Teórica'
                reason = '🔴 Eliminar: Reseña histórica prescindible; sustituir directamente por normativas vigentes del Reglamento de SSO (Acuerdo Gub. 229-2014).'
            elif 'bomba periférica' in t_lower and 'módulo 7' in m_num.lower():
                status = 'delete'
                badge = '🔴 Eliminar'
                badge_color = 'bg-red-100 text-red-800'
                category = 'Equipo Doméstico'
                reason = '🔴 Eliminar: Las bombas periféricas pequeñas son de uso casero; enfocar las horas en bombas centrífugas de procesos, bombas de cavidad progresiva y bombas de engranajes.'
            elif 'alineación de poleas' in t_lower or 'proceso de alineación' in t_lower:
                status = 'update'
                badge = '🟡 Actualizar'
                badge_color = 'bg-amber-100 text-amber-800'
                category = 'Transmisión de Precisión'
                reason = '🟡 Actualizar: Incorporar alineadores láser magnéticos de poleas y tensiómetros de frecuencia sónicos para evitar sobretensión de fajas y desgaste prematuro de descanseras.'
            elif 'alineación de engranajes' in t_lower or 'cadenas de rodillos' in t_lower:
                status = 'update'
                badge = '🟡 Actualizar'
                badge_color = 'bg-amber-100 text-amber-800'
                category = 'Elementos de Máquina'
                reason = '🟡 Actualizar: Inspección de contacto de dientes con azul de Prusia, medición de huelgo (backlash) y lubricación por goteo/inmersión según AGMA.'
            elif 'instalación correcta de cojinetes' in t_lower or 'mantenimiento de cojinetes' in t_lower:
                status = 'update'
                badge = '🟡 Actualizar'
                badge_color = 'bg-amber-100 text-amber-800'
                category = 'Rodamientos de Alta Exigencia'
                reason = '🟡 Actualizar: Ajustes y tolerancias ISO (h6, js6), montaje por calentamiento inductivo controlado (< 110°C) y desmontaje con extractores mecánicos e hidráulicos sin golpear pistas.'
            elif 'técnicas de alineación de rodamientos y acoples' in t_lower or 'acoplamientos' in t_lower:
                status = 'update'
                badge = '🟡 Actualizar'
                badge_color = 'bg-amber-100 text-amber-800'
                category = 'Alineación de Ejes'
                reason = '🟡 Actualizar: Modernizar del reloj comparador manual a sistemas computarizados de alineación láser de ejes con detección de pata coja (soft foot).'
            elif 'bomba centrífuga' in t_lower or 'sistema de bombeo' in t_lower:
                status = 'update'
                badge = '🟡 Actualizar'
                badge_color = 'bg-amber-100 text-amber-800'
                category = 'Turbomaquinaria y Sellado'
                reason = '🟡 Actualizar: Diagnóstico de cavitación (NPSH), inspección de impulsores, balanceo dinámico y reemplazo/armado de sellos mecánicos con caras de carburo de silicio.'
            elif 'sistemas computarizados de gestión' in t_lower or 'oee' in t_lower:
                status = 'update'
                badge = '🟡 Actualizar'
                badge_color = 'bg-amber-100 text-amber-800'
                category = 'Gestión Digital del Mantenimiento'
                reason = '🟡 Actualizar: Uso de plataformas CMMS/GMAO en tablets móviles para recepción de órdenes de trabajo, control de repuestos y cálculo automático de OEE y MTBF.'

            mod_topics.append({
                'title': t,
                'status': status,
                'badge': badge,
                'badge_color': badge_color,
                'category': category,
                'reason': reason
            })

        if 'Rodamientos y Acoplamientos' in m_name:
            mod_topics.append({
                'title': '[NUEVO] Alineación Láser de Precisión de Ejes y Detección de Pie Cojo (Soft Foot)',
                'status': 'new',
                'badge': '🔵 Propuesta Nueva',
                'badge_color': 'bg-blue-100 text-blue-800',
                'category': 'Metrología Láser de Montaje',
                'reason': '🔵 Propuesta Nueva: Uso de alineadores láser inalámbricos con compensación térmica para eliminar vibraciones en trenes motrices y bombas multietapa.'
            })
            mod_topics.append({
                'title': '[NUEVO] Montaje Térmico de Rodamientos con Calentadores de Inducción',
                'status': 'new',
                'badge': '🔵 Propuesta Nueva',
                'badge_color': 'bg-blue-100 text-blue-800',
                'category': 'Montaje de Precisión SKF/FAG',
                'reason': '🔵 Propuesta Nueva: Desmagnetización automática, control por sonda magnética a 110°C y eliminación total del calentamiento destructivo con soplete o baño de aceite sucio.'
            })
        if 'Mantenimiento Productivo Total' in m_name:
            mod_topics.append({
                'title': '[NUEVO] Mantenimiento Centrado en la Confiabilidad (RCM) y Análisis Causa Raíz (RCA)',
                'status': 'new',
                'badge': '🔵 Propuesta Nueva',
                'badge_color': 'bg-blue-100 text-blue-800',
                'category': 'Ingeniería de Confiabilidad',
                'reason': '🔵 Propuesta Nueva: Metodología de análisis de modos y efectos de fallas (FMEA), árboles de causas de paro y planes de prevención proactivos para activos críticos.'
            })

        classified_modules.append({
            'num': m_num,
            'name': m_name,
            'topics': mod_topics
        })

    total_orig = sum(len(m['topics']) for m in modules_data)
    count_del = sum(sum(1 for t in m['topics'] if t['status'] == 'delete') for m in classified_modules)
    count_upd = sum(sum(1 for t in m['topics'] if t['status'] == 'update') for m in classified_modules)
    count_new = sum(sum(1 for t in m['topics'] if t['status'] == 'new') for m in classified_modules)
    count_keep = sum(sum(1 for t in m['topics'] if t['status'] == 'keep') for m in classified_modules)

    html = f'''<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Auditoría Curricular y Propuesta: Mantenimiento Mecánico Industrial | Kinal 2026</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <script src="https://unpkg.com/lucide@latest"></script>
    <style>
        @import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700;800&family=Space+Grotesk:wght@500;700&display=swap');
        body {{ font-family: 'Inter', sans-serif; background-color: #f8fafc; color: #0f172a; }}
        h1, h2, h3, h4, .font-heading {{ font-family: 'Space Grotesk', sans-serif; }}
        .item-delete {{ background-color: #fef2f2; border-left: 4px solid #ef4444; color: #991b1b; }}
        .item-update {{ background-color: #fefce8; border-left: 4px solid #eab308; color: #854d0e; }}
        .item-new {{ background-color: #eff6ff; border-left: 4px solid #3b82f6; color: #1e40af; }}
        .item-keep {{ background-color: #f0fdf4; border-left: 4px solid #22c55e; color: #166534; }}
        @media print {{
            .no-print {{ display: none !important; }}
            body {{ background-color: white !important; font-size: 11pt; }}
            .shadow-lg, .shadow-md, .shadow-xl {{ box-shadow: none !important; }}
            .break-inside-avoid {{ break-inside: avoid; }}
        }}
    </style>
</head>
<body class="antialiased text-slate-800">

    <header class="bg-slate-900 text-white border-b border-slate-800 sticky top-0 z-50 shadow-md">
        <div class="max-w-7xl mx-auto px-4 sm:px-6 lg:px-8 py-3 flex flex-wrap justify-between items-center gap-4">
            <div class="flex items-center gap-3">
                <a href="index.html" title="Ir al Portal Maestro" class="bg-blue-600 hover:bg-blue-500 text-white font-bold text-xs uppercase px-2.5 py-1 rounded transition cursor-pointer">Kinal ETS 2026</a>
                <span class="text-slate-500 text-sm hidden sm:inline">|</span>
                <h1 class="text-base sm:text-lg font-semibold tracking-tight text-white flex items-center gap-2">
                    <i data-lucide="wrench" class="w-5 h-5 text-blue-400"></i>
                    Mantenimiento Mecánico Industrial
                </h1>
            </div>

            <div class="flex items-center gap-2 text-xs no-print">
                <a href="index.html" class="px-2.5 py-1.5 rounded-lg bg-slate-800 hover:bg-slate-700 text-slate-300 transition">Portal Maestro</a>
                <span class="px-2.5 py-1.5 rounded-lg bg-blue-600 text-white font-bold">MMI</span>
                <a href="Soldadura_Industrial_Propuesta.html" class="px-2.5 py-1.5 rounded-lg bg-slate-800 hover:bg-slate-700 text-slate-300 transition">Soldadura Ind.</a>
                <a href="Calderas_de_Vapor_Propuesta.html" class="px-2.5 py-1.5 rounded-lg bg-slate-800 hover:bg-slate-700 text-slate-300 transition">Calderas de Vapor</a>
                <button onclick="window.print()" class="inline-flex items-center gap-1 px-3 py-1.5 rounded-lg bg-slate-800 hover:bg-slate-700 text-slate-200 border border-slate-700 transition">
                    <i data-lucide="printer" class="w-3.5 h-3.5"></i> Imprimir
                </button>
            </div>
        </div>
    </header>

    <main class="max-w-7xl mx-auto px-4 sm:px-6 lg:px-8 py-8 space-y-10">

        <div>
            <div class="flex items-center gap-2 text-xs font-medium text-slate-500 mb-2">
                <span>Programas de Curso</span>
                <i data-lucide="chevron-right" class="w-3.5 h-3.5"></i>
                <span>2026_MMIM.xlsx</span>
                <i data-lucide="chevron-right" class="w-3.5 h-3.5"></i>
                <span class="text-blue-600 font-semibold">Auditoría Curricular & Propuesta 2026</span>
            </div>
            <h2 class="text-3xl font-extrabold text-slate-900 tracking-tight">
                Auditoría Curricular: Mantenimiento Mecánico Industrial (180 Horas)
            </h2>
            <p class="text-slate-600 mt-2 text-base max-w-4xl">
                Modernización integral orientada a paradas de planta, montaje electromecánico y confiabilidad de activos en Guatemala: Alineación Láser de Ejes, Calentamiento Inductivo de Rodamientos, Análisis Vibracional Básico, Sellos Mecánicos en Bombas Centrífugas y Metodología RCM en sistemas CMMS.
            </p>
        </div>

        <!-- Metric Badges Overview -->
        <div class="grid grid-cols-2 md:grid-cols-4 gap-4">
            <div class="bg-white p-5 rounded-xl border border-slate-200 shadow-sm flex items-center gap-4">
                <div class="w-12 h-12 rounded-lg bg-red-100 flex items-center justify-center text-red-600 font-bold text-xl">
                    <i data-lucide="trash-2" class="w-6 h-6"></i>
                </div>
                <div>
                    <div class="text-2xl font-bold text-slate-900">{count_del} Temas</div>
                    <div class="text-xs font-medium text-red-600 uppercase tracking-wider">A Eliminar (Petróleo/Historia/Periféricas)</div>
                </div>
            </div>

            <div class="bg-white p-5 rounded-xl border border-slate-200 shadow-sm flex items-center gap-4">
                <div class="w-12 h-12 rounded-lg bg-amber-100 flex items-center justify-center text-amber-600 font-bold text-xl">
                    <i data-lucide="refresh-cw" class="w-6 h-6"></i>
                </div>
                <div>
                    <div class="text-2xl font-bold text-slate-900">{count_upd} Temas</div>
                    <div class="text-xs font-medium text-amber-600 uppercase tracking-wider">A Actualizar (Láser/Sellos/CMMS)</div>
                </div>
            </div>

            <div class="bg-white p-5 rounded-xl border border-slate-200 shadow-sm flex items-center gap-4">
                <div class="w-12 h-12 rounded-lg bg-blue-100 flex items-center justify-center text-blue-600 font-bold text-xl">
                    <i data-lucide="plus-circle" class="w-6 h-6"></i>
                </div>
                <div>
                    <div class="text-2xl font-bold text-slate-900">{count_new} Temas</div>
                    <div class="text-xs font-medium text-blue-600 uppercase tracking-wider">Nuevos Propuestos (Láser / Inducción / RCM)</div>
                </div>
            </div>

            <div class="bg-white p-5 rounded-xl border border-slate-200 shadow-sm flex items-center gap-4">
                <div class="w-12 h-12 rounded-lg bg-emerald-100 flex items-center justify-center text-emerald-600 font-bold text-xl">
                    <i data-lucide="check-circle-2" class="w-6 h-6"></i>
                </div>
                <div>
                    <div class="text-2xl font-bold text-slate-900">{count_keep} Temas</div>
                    <div class="text-xs font-medium text-emerald-600 uppercase tracking-wider">Fundamentos Conservados</div>
                </div>
            </div>
        </div>

        <!-- SECCIÓN 1: Contexto Industrial y Demanda Laboral en Guatemala -->
        <section class="bg-white rounded-2xl border border-slate-200 p-6 sm:p-8 shadow-sm space-y-6">
            <div class="border-b border-slate-100 pb-4">
                <span class="text-xs font-bold uppercase tracking-wider text-blue-600 bg-blue-50 px-2.5 py-1 rounded-full border border-blue-200">Sección 1</span>
                <h3 class="text-xl font-bold text-slate-900 mt-2 flex items-center gap-2">
                    <i data-lucide="trending-up" class="w-5 h-5 text-blue-600"></i>
                    Contexto Industrial y Demanda Laboral en Guatemala (2026)
                </h3>
                <p class="text-slate-600 text-sm mt-1">
                    Exigencias técnicas en ingenios azucareros, plantas papeleras, envasadoras de bebidas, cementeras y molinos harineros en Guatemala.
                </p>
            </div>

            <div class="grid md:grid-cols-3 gap-6 text-sm">
                <div class="p-5 rounded-xl bg-slate-50 border border-slate-200 space-y-3">
                    <div class="font-bold text-slate-900 flex items-center gap-2 text-base">
                        <i data-lucide="crosshair" class="w-5 h-5 text-blue-600"></i>
                        Alineación Láser de Ejes y Poleas
                    </div>
                    <p class="text-slate-600 text-xs leading-relaxed">
                        Más del 50% de las fallas prematuras de rodamientos y acoples en plantas industriales guatemaltecas son causadas por desalineación de ejes. El método tradicional de regla y calibrador de espesores ya no es tolerado en equipos críticos de más de 25 HP; las empresas exigen alineación láser con reporte digital.
                    </p>
                </div>

                <div class="p-5 rounded-xl bg-slate-50 border border-slate-200 space-y-3">
                    <div class="font-bold text-slate-900 flex items-center gap-2 text-base">
                        <i data-lucide="flame" class="w-5 h-5 text-amber-600"></i>
                        Cero Daño Térmico en Montaje de Cojinetes
                    </div>
                    <p class="text-slate-600 text-xs leading-relaxed">
                        El calentamiento de rodamientos con soplete de acetileno o baños de aceite sucio degrada la metalurgia del acero cromo y quema los sellos de fábrica. La industria exige el uso estricto de calentadores de inducción electromagnética con control térmico por sonda magnética (máximo 110 °C).
                    </p>
                </div>

                <div class="p-5 rounded-xl bg-slate-50 border border-slate-200 space-y-3">
                    <div class="font-bold text-slate-900 flex items-center gap-2 text-base">
                        <i data-lucide="clipboard-check" class="w-5 h-5 text-emerald-600"></i>
                        Gestión Digital con CMMS y Cultura RCM
                    </div>
                    <p class="text-slate-600 text-xs leading-relaxed">
                        Las empresas han migrado las órdenes de trabajo de papel a sistemas computarizados (SAP PM, Maximo, Fracttal). El mecánico debe registrar tiempos de paro, causas de fallo según taxonomía ISO 14224 y calcular indicadores de confiabilidad como el OEE y el MTTR.
                    </p>
                </div>
            </div>
        </section>

        <!-- SECCIÓN 2: Ficha Oficial de Venta y Diagnóstico Crítico -->
        <section class="bg-white rounded-2xl border border-slate-200 p-6 sm:p-8 shadow-sm space-y-6">
            <div class="border-b border-slate-100 pb-4">
                <span class="text-xs font-bold uppercase tracking-wider text-blue-600 bg-blue-50 px-2.5 py-1 rounded-full border border-blue-200">Sección 2</span>
                <h3 class="text-xl font-bold text-slate-900 mt-2 flex items-center gap-2">
                    <i data-lucide="file-spreadsheet" class="w-5 h-5 text-blue-600"></i>
                    Ficha Oficial de Venta ("Reporte") y Diagnóstico Crítico de Anomalías
                </h3>
            </div>

            <div class="grid md:grid-cols-2 gap-4 text-xs sm:text-sm text-slate-700 bg-slate-50 p-4 rounded-xl border border-slate-200">
                <div><strong>Programa Académico:</strong> Mantenimiento Mecánico Industrial (2026_MMIM.xlsx)</div>
                <div><strong>Duración Total:</strong> 180 Horas distribuidas en 10 Módulos cronológicos</div>
                <div><strong>Jornadas:</strong> Matutina y Vespertina (Sábados de 8:00 a 12:30 y 13:00 a 17:30 hrs)</div>
                <div><strong>Acreditación:</strong> Diploma de Técnico en Mantenimiento Mecánico Industrial</div>
                <div class="md:col-span-2"><strong>Perfil de Ingreso:</strong> Tercer grado de educación básica, destreza manual e interés en maquinaria pesada y mecanismos industriales.</div>
                <div class="md:col-span-2"><strong>Perfil de Egreso:</strong> Competencia para desmontar, inspeccionar, reparar y alinear bombas, reductores, rodamientos y sistemas de transmisión en plantas manufactureras.</div>
            </div>

            <div class="bg-amber-50 border border-amber-300 rounded-xl p-5 space-y-3">
                <h4 class="text-amber-900 font-bold text-sm flex items-center gap-2">
                    <i data-lucide="alert-triangle" class="w-5 h-5 text-amber-600"></i>
                    Diagnóstico Crítico / Anomalías Detectadas en el Programa Original
                </h4>
                <ul class="text-xs text-amber-950 space-y-2 list-disc list-inside leading-relaxed">
                    <li><strong>Contenido Petroquímico Fuera de Enfoque:</strong> El Módulo 5 dedica clases teóricas a <em>"El petróleo"</em> y <em>"El proceso de destilación"</em>, conocimientos de refinería que no aportan a la labor del mecánico de planta, quien necesita diagnosticar compatibilidad de grasas con jabón de litio/poliurea y viscosidades ISO VG.</li>
                    <li><strong>Falta de Metrología Láser en Alineación:</strong> El Módulo 6 se limitaba a métodos tradicionales de alineación de acoples con regla y palpador manual, desfasado de los equipos láser con compensación térmica utilizados en la industria guatemalteca moderna.</li>
                    <li><strong>Incongruencia en Tipos de Bombas:</strong> El Módulo 7 incluye <em>"Bomba periférica"</em> (equipo doméstico de bajo caudal), restando tiempo al mantenimiento de bombas centrífugas de procesos químicos y de alimentos con sellos mecánicos dobles.</li>
                </ul>
            </div>
        </section>

        <!-- SECCIÓN 3: Distribución Modular Detallada y Auditoría Tema por Tema -->
        <section class="bg-white rounded-2xl border border-slate-200 p-6 sm:p-8 shadow-sm space-y-6">
            <div class="flex flex-wrap justify-between items-center gap-4 border-b border-slate-100 pb-4">
                <div>
                    <span class="text-xs font-bold uppercase tracking-wider text-blue-600 bg-blue-50 px-2.5 py-1 rounded-full border border-blue-200">Sección 3</span>
                    <h3 class="text-xl font-bold text-slate-900 mt-2 flex items-center gap-2">
                        <i data-lucide="list-checks" class="w-5 h-5 text-blue-600"></i>
                        Distribución Modular y Auditoría Analítica Tema por Tema
                    </h3>
                    <p class="text-slate-600 text-sm mt-1">
                        Total de temas auditados en el pensum original: <strong>{total_orig} temas</strong> desglosados en sus 10 módulos oficiales.
                    </p>
                </div>

                <!-- Filtros Interactivos -->
                <div class="flex flex-wrap items-center gap-1.5 text-xs no-print">
                    <button onclick="filterItems('all')" id="btn-all" class="px-3 py-1.5 rounded-lg bg-slate-900 text-white font-medium hover:bg-slate-800 transition">Todos ({total_orig + count_new})</button>
                    <button onclick="filterItems('delete')" id="btn-delete" class="px-3 py-1.5 rounded-lg bg-red-100 text-red-800 font-medium hover:bg-red-200 transition">🔴 Eliminar ({count_del})</button>
                    <button onclick="filterItems('update')" id="btn-update" class="px-3 py-1.5 rounded-lg bg-amber-100 text-amber-800 font-medium hover:bg-amber-200 transition">🟡 Actualizar ({count_upd})</button>
                    <button onclick="filterItems('new')" id="btn-new" class="px-3 py-1.5 rounded-lg bg-blue-100 text-blue-800 font-medium hover:bg-blue-200 transition">🔵 Nuevos ({count_new})</button>
                    <button onclick="filterItems('keep')" id="btn-keep" class="px-3 py-1.5 rounded-lg bg-emerald-100 text-emerald-800 font-medium hover:bg-emerald-200 transition">🟢 Mantener ({count_keep})</button>
                </div>
            </div>

            <!-- Listado de Módulos y Temas -->
            <div class="space-y-8">
'''

    for mod in classified_modules:
        m_num = mod['num']
        m_name = mod['name']
        topics = mod['topics']
        html += f'''
                <div class="border border-slate-200 rounded-xl p-5 bg-slate-50/50 space-y-4">
                    <div class="flex items-center justify-between border-b border-slate-200 pb-3">
                        <div class="flex items-center gap-2">
                            <span class="w-7 h-7 rounded-lg bg-blue-600 text-white font-bold text-xs flex items-center justify-center">{m_num.replace('Módulo ', '') if 'Módulo' in m_num else m_num}</span>
                            <h4 class="font-bold text-slate-900 text-base">{m_name}</h4>
                        </div>
                        <span class="text-xs text-slate-500 font-medium">{len(topics)} temas analizados</span>
                    </div>

                    <div class="grid md:grid-cols-2 gap-3">
'''
        for t in topics:
            status_class = f"item-{t['status']}"
            html += f'''
                        <div class="content-item {status_class} p-3.5 rounded-lg text-xs space-y-1.5" data-status="{t['status']}">
                            <div class="flex justify-between items-start gap-2">
                                <span class="font-bold text-slate-900">{t['title']}</span>
                                <span class="px-2 py-0.5 rounded text-[10px] font-extrabold uppercase whitespace-nowrap {t['badge_color']}">{t['badge']}</span>
                            </div>
                            <div class="text-[11px] font-medium text-slate-500 uppercase tracking-wider">{t['category']}</div>
                            <p class="text-slate-700 leading-relaxed">{t['reason']}</p>
                        </div>
'''
        html += '''
                    </div>
                </div>
'''

    html += f'''
            </div>
        </section>

        <!-- SECCIÓN 4: Matriz Comparativa -->
        <section class="bg-white rounded-2xl border border-slate-200 p-6 sm:p-8 shadow-sm space-y-6">
            <div class="border-b border-slate-100 pb-4">
                <span class="text-xs font-bold uppercase tracking-wider text-blue-600 bg-blue-50 px-2.5 py-1 rounded-full border border-blue-200">Sección 4</span>
                <h3 class="text-xl font-bold text-slate-900 mt-2 flex items-center gap-2">
                    <i data-lucide="arrow-left-right" class="w-5 h-5 text-blue-600"></i>
                    Matriz Comparativa: Pensum Tradicional vs. Pensum Modernizado 2026
                </h3>
            </div>

            <div class="overflow-x-auto">
                <table class="w-full text-xs text-left border-collapse border border-slate-200 rounded-lg">
                    <thead class="bg-slate-100 text-slate-700 font-bold uppercase text-[11px]">
                        <tr>
                            <th class="p-3 border border-slate-200">Eje Temático</th>
                            <th class="p-3 border border-slate-200 text-red-800 bg-red-50/50">Enfoque Anterior (Obsolescencia / Brechas)</th>
                            <th class="p-3 border border-slate-200 text-emerald-800 bg-emerald-50/50">Enfoque Modernizado 2026 (Estándar Kinal ETS)</th>
                        </tr>
                    </thead>
                    <tbody class="divide-y divide-slate-200">
                        <tr>
                            <td class="p-3 font-bold text-slate-900 border border-slate-200">Alineación de Ejes</td>
                            <td class="p-3 text-slate-600 border border-slate-200">Uso exclusivo de regla de pelo y reloj comparador manual susceptible a error de paralelismo.</td>
                            <td class="p-3 text-slate-800 border border-slate-200 font-medium">Sistemas láser de alineación con corrección de pata coja (soft foot) y generación de reporte de tolerancia digital.</td>
                        </tr>
                        <tr>
                            <td class="p-3 font-bold text-slate-900 border border-slate-200">Montaje de Rodamientos</td>
                            <td class="p-3 text-slate-600 border border-slate-200">Golpeteo mecánico o calentamiento informal con soplete que destempla las pistas de acero.</td>
                            <td class="p-3 text-slate-800 border border-slate-200 font-medium">Calentadores de inducción electromagnética con control de temperatura a 110 °C y desmagnetización automática.</td>
                        </tr>
                        <tr>
                            <td class="p-3 font-bold text-slate-900 border border-slate-200">Lubricación y Tribología</td>
                            <td class="p-3 text-slate-600 border border-slate-200">Clases teóricas sobre extracción de petróleo crudo y procesos de refinación.</td>
                            <td class="p-3 text-slate-800 border border-slate-200 font-medium">Selección de lubricantes bajo tablas ISO VG / NLGI, compatibilidad de espesantes y cálculo de frecuencias de reengrase.</td>
                        </tr>
                        <tr>
                            <td class="p-3 font-bold text-slate-900 border border-slate-200">Bombas y Sellado</td>
                            <td class="p-3 text-slate-600 border border-slate-200">Enfoque en bombas periféricas domésticas y prensaestopas tradicionales con fugas continuas.</td>
                            <td class="p-3 text-slate-800 border border-slate-200 font-medium">Bombas centrífugas industriales de procesos, cálculo de NPSH para evitar cavitación y montaje de sellos mecánicos de cartucho.</td>
                        </tr>
                        <tr>
                            <td class="p-3 font-bold text-slate-900 border border-slate-200">Gestión del Mantenimiento</td>
                            <td class="p-3 text-slate-600 border border-slate-200">Registros en libretas manuales de papel y mantenimiento puramente correctivo.</td>
                            <td class="p-3 text-slate-800 border border-slate-200 font-medium">Gestión con software CMMS/GMAO en la nube, cálculo de confiabilidad OEE/MTBF y metodología RCM para análisis de fallas.</td>
                        </tr>
                    </tbody>
                </table>
            </div>
        </section>

        <!-- SECCIÓN 5: Innovación Tecnológica: Inteligencia Artificial Aplicada -->
        <section class="bg-gradient-to-r from-blue-900 via-slate-900 to-indigo-950 text-white rounded-2xl p-6 sm:p-8 shadow-xl space-y-6">
            <div class="flex items-center gap-3 border-b border-white/10 pb-4">
                <div class="p-2.5 rounded-lg bg-blue-500/20 text-blue-300">
                    <i data-lucide="bot" class="w-6 h-6"></i>
                </div>
                <div>
                    <span class="bg-blue-500/30 text-blue-200 text-xs px-2.5 py-0.5 rounded-full font-bold uppercase tracking-wider border border-blue-400/30">Módulo Adicional / Opcional</span>
                    <h3 class="text-xl font-bold text-white">Innovación Tecnológica: Inteligencia Artificial en Mantenimiento Mecánico</h3>
                </div>
            </div>

            <div class="grid md:grid-cols-3 gap-6">
                <div class="bg-white/10 backdrop-blur rounded-xl p-5 border border-white/10 space-y-3">
                    <div class="text-blue-300 font-bold text-sm flex items-center gap-2">
                        <i data-lucide="activity" class="w-4 h-4"></i>
                        Análisis Predictivo de Vibraciones Mecánicas Asistido por IA
                    </div>
                    <p class="text-xs text-slate-200 leading-relaxed">
                        Sensores inalámbricos triaxiales acoplados a motores y bombas que transmiten datos a modelos de Machine Learning, identificando desbalance, holguras o falla en pista exterior de rodamientos con semanas de anticipación.
                    </p>
                    <div class="pt-2 border-t border-white/10 text-[11px] text-blue-200 font-mono">
                        Herramienta: SKF Enlight AI / Prüftechnik OMNITREND AI.
                    </div>
                </div>

                <div class="bg-white/10 backdrop-blur rounded-xl p-5 border border-white/10 space-y-3">
                    <div class="text-blue-300 font-bold text-sm flex items-center gap-2">
                        <i data-lucide="camera" class="w-4 h-4"></i>
                        Inspección Visual y Detección de Fricción con Cámaras Acústicas IA
                    </div>
                    <p class="text-xs text-slate-200 leading-relaxed">
                        Cámaras de ultrasonido con IA que visualizan en pantalla la procedencia exacta de ruidos de cavitación en tuberías y chirridos de rodamientos no perceptibles al oído humano en entornos ruidosos de fábrica.
                    </p>
                    <div class="pt-2 border-t border-white/10 text-[11px] text-blue-200 font-mono">
                        Herramienta: Fluke ii910 Acoustic Imager AI.
                    </div>
                </div>

                <div class="bg-white/10 backdrop-blur rounded-xl p-5 border border-white/10 space-y-3">
                    <div class="text-blue-300 font-bold text-sm flex items-center gap-2">
                        <i data-lucide="file-text" class="w-4 h-4"></i>
                        Asistente RAG para Procedimientos de Desarme y Manuales OEM
                    </div>
                    <p class="text-xs text-slate-200 leading-relaxed">
                        Copilot técnico que analiza planos seccionales y manuales de mantenimiento de reductores Sew-Eurodrive y bombas Goulds, entregando de inmediato torques de apriete, tolerancias de montaje y secuencias seguras de desarme.
                    </p>
                    <div class="pt-2 border-t border-white/10 text-[11px] text-blue-200 font-mono">
                        Herramienta: Mechanical OEM Copilot / NotebookLM Industrial.
                    </div>
                </div>
            </div>
        </section>

    </main>

    <footer class="bg-white border-t border-slate-200 mt-12 py-6 text-center text-xs text-slate-500">
        Fundación Kinal • Escuela Técnica Superior • Coordinación Académica • Propuesta Curricular Mantenimiento Mecánico Industrial 2026
    </footer>

    <script>
        lucide.createIcons();

        function filterItems(status) {{
            const items = document.querySelectorAll('.content-item');
            const buttons = ['all', 'delete', 'update', 'new', 'keep'];

            buttons.forEach(btn => {{
                const buttonElement = document.getElementById('btn-' + btn);
                if (buttonElement) {{
                    if (btn === status) {{
                        buttonElement.classList.add('ring-2', 'ring-offset-2', 'ring-slate-900');
                    }} else {{
                        buttonElement.classList.remove('ring-2', 'ring-offset-2', 'ring-slate-900');
                    }}
                }}
            }});

            items.forEach(item => {{
                if (status === 'all') {{
                    item.style.display = 'block';
                }} else {{
                    if (item.getAttribute('data-status') === status) {{
                        item.style.display = 'block';
                    }} else {{
                        item.style.display = 'none';
                    }}
                }}
            }});
        }}
    </script>
</body>
</html>
'''
    (output_dir / 'Mantenimiento_Mecanico_Industrial_Propuesta.html').write_text(html, encoding='utf-8')
    print('Generated Mantenimiento_Mecanico_Industrial_Propuesta.html')

generate_mmi()
