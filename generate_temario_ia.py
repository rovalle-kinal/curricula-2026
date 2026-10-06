import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=70, bottom=70, left=90, right=90):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E0", sz="4"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def style_p(p, space_before=0, space_after=6, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def add_callout(doc, text, bold_prefix="", fill_hex="F7FAFC", border_color="003366"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    
    cell = tbl.cell(0, 0)
    set_cell_background(cell, fill_hex)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=120)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 51, 102)
        
    r_txt = p.add_run(text)
    r_txt.font.name = "Calibri"
    r_txt.font.size = Pt(10)
    r_txt.font.italic = True
    r_txt.font.color.rgb = RGBColor(45, 55, 72)
    
    p_post = doc.add_paragraph()
    style_p(p_post, space_before=0, space_after=6)

def style_heading(p, text, level=1, color=RGBColor(0, 51, 102)):
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.bold = True
    r.font.color.rgb = color
    if level == 1:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        r.font.size = Pt(15)
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        r.font.size = Pt(13)
    elif level == 3:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r.font.size = Pt(11)

def generate_temario_doc():
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    COLOR_PRIMARY = RGBColor(0, 51, 102)      # Azul Kinal
    COLOR_SECONDARY = RGBColor(74, 85, 104)    # Gris pizarra
    COLOR_DARK = RGBColor(26, 32, 44)         # Texto oscuro

    # ==========================================
    # ENCABEZADO INSTITUCIONAL
    # ==========================================
    title_p = doc.add_paragraph()
    style_p(title_p, space_before=0, space_after=6)
    run_title = title_p.add_run("Inteligencia Artificial Aplicada, Ingeniería de Prompts y Productividad Ética")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    style_p(sub_p, space_before=0, space_after=14)
    run_sub = sub_p.add_run("Programa Oficial de Formación y Temario Analítico Detallado  |  Duración: 16 Horas Formativas (1 Mes)\nFundación Kinal — Escuela Técnica Superior | Nivel DQR 3 - 4 | Umbral Aprobatorio: 75/100 Puntos")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_SECONDARY

    add_callout(doc, 
        "«the ability and readiness of the individual to use knowledge, skills and personal, social and methodological competences and to behave in a considered, individual and socially responsible manner. Competence is understood in this sense as the comprehensive ability to act». El programa enfoca la inteligencia artificial como un amplificador de las capacidades humanas, integrando la destreza práctica para formular prompts estructurados con la honestidad profesional y el ideario del «trabajo bien hecho» de Fundación Kinal.", 
        bold_prefix="Definición de Competencia DQR e Ideario Kinal:")

    # ==========================================
    # 1. OBJETIVOS FORMATIVOS Y COMPETENCIAS
    # ==========================================
    h1_1 = doc.add_paragraph()
    style_heading(h1_1, "1. Competencia General y Resultados de Aprendizaje por Módulo", level=1)

    p_cg = doc.add_paragraph()
    style_p(p_cg, space_before=0, space_after=6)
    r = p_cg.add_run("Competencia General del Programa:\n"
                     "El participante interactúa de manera autónoma, fluida y crítica con herramientas de inteligencia artificial generativa, formulando prompts estructurados y profesionales para optimizar tareas de comunicación, análisis de información y productividad cotidiana, aplicando estrictos criterios éticos de privacidad, veracidad y el principio del «trabajo bien hecho» de Fundación Kinal, preservando en todo momento el valor insustituible del criterio humano.")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = COLOR_PRIMARY

    mod_comps = [
        ("Módulo 1: Descubriendo la IA: Fundamentos y Primeros Pasos Prácticos", 
         "Comprende el funcionamiento operativo de los modelos generativos y navega con fluidez en ChatGPT, Gemini y Copilot, configurando su espacio de trabajo digital sin barreras técnicas."),
        ("Módulo 2: Ingeniería de Prompts: El Arte de Instruir con Precisión (Método RC-TRF)", 
         "Diseña prompts de alto impacto utilizando la estructura Rol, Contexto, Tarea, Restricción y Formato, aplicando técnicas de refinamiento iterativo y ejemplos para obtener respuestas profesionales."),
        ("Módulo 3: Casos Reales en la Oficina, Negocio y Vida Profesional", 
         "Aplica herramientas de IA para redactar correspondencia profesional, resumir informes extensos, extraer datos a tablas y estructurar planes de trabajo cotidianos en minutos."),
        ("Módulo 4: Ética, Privacidad, Caza de Alucinaciones y Proyecto Integrador", 
         "Audita críticamente los contenidos generados detectando sesgos y alucinaciones, protege la privacidad de datos sensibles y construye un asistente personal de productividad evaluado con rúbrica institucional (>= 75 pts).")
    ]

    for m_tit, m_desc in mod_comps:
        p_m = doc.add_paragraph()
        style_p(p_m, space_before=2, space_after=4)
        p_m.paragraph_format.left_indent = Inches(0.2)
        r_b = p_m.add_run(f"• {m_tit}: ")
        r_b.font.name = "Calibri"
        r_b.font.size = Pt(9.5)
        r_b.font.bold = True
        r_b.font.color.rgb = COLOR_PRIMARY
        r_d = p_m.add_run(m_desc)
        r_d.font.name = "Calibri"
        r_d.font.size = Pt(9.5)
        r_d.font.color.rgb = COLOR_DARK

    # ==========================================
    # 2. TEMARIO ANALÍTICO DETALLADO POR MÓDULOS
    # ==========================================
    h1_2 = doc.add_paragraph()
    style_heading(h1_2, "2. Temario Analítico Jerárquico y Prácticas Guiadas", level=1)

    # -------------------------------------------------------------
    # MODULO 1
    # -------------------------------------------------------------
    h2_m1 = doc.add_paragraph()
    style_heading(h2_m1, "MÓDULO 1: DESCUBRIENDO LA IA: FUNDAMENTOS Y PRIMEROS PASOS PRÁCTICOS (4 Horas — Sesiones 1 y 2)", level=2)

    temas_m1 = [
        ("1.1 Fundamentos de la IA Generativa sin Tecnicismos", [
            "De la búsqueda tradicional en Google a la conversación inteligente: diferencias conceptuales.",
            "Cómo piensa un modelo de lenguaje (LLM): predicción probabilística de palabras y ventanas de contexto explicadas con metáforas cotidianas.",
            "Qué puede hacer y qué NO puede hacer la IA: alcances reales vs expectativas poco realistas.",
            "El rol insustituible del ser humano: el modelo genera opciones, pero la persona aporta juicio, empatía, contexto y decisión."
        ]),
        ("1.2 El Ecosistema Actual de Herramientas Gratuitas", [
            "Panorama comparativo de las principales plataformas: ChatGPT (OpenAI - GPT-4o mini), Google Gemini, Microsoft Copilot y Claude (Anthropic).",
            "Puntos fuertes de cada herramienta: razonamiento estructurado, búsqueda web en tiempo real, integración con ofimática y redacción creativa.",
            "Creación segura de perfiles de usuario, verificación en dos pasos y configuración de preferencias básicas.",
            "Navegación por la interfaz: cómo iniciar un nuevo chat, organizar conversaciones en carpetas/historiales y uso desde el teléfono móvil (dictado por voz)."
        ]),
        ("1.3 Primeras Interacciones y Dinámicas de Descubrimiento", [
            "El ejercicio del 'Explicador Universal' (Técnica ELI5): pedir que explique conceptos complejos adaptados a diferentes niveles (niño de 10 años, estudiante, directivo).",
            "Lluvia de ideas guiada para situaciones cotidianas (planificación de actividades, recetas con ingredientes disponibles, listas de compras optimizadas).",
            "Superación de la respuesta genérica: por qué las preguntas vagas producen respuestas vacías.",
            "Taller de inicio: personalización de la primera conversación y autodiagnóstico de necesidades laborales."
        ])
    ]

    for sub_tit, sub_items in temas_m1:
        p_st = doc.add_paragraph()
        style_heading(p_st, sub_tit, level=3)
        for it in sub_items:
            p_it = doc.add_paragraph()
            style_p(p_it, space_before=1, space_after=2)
            p_it.paragraph_format.left_indent = Inches(0.25)
            r_b = p_it.add_run("• ")
            r_b.font.bold = True
            r_b.font.color.rgb = COLOR_PRIMARY
            r_txt = p_it.add_run(it)
            r_txt.font.name = "Calibri"
            r_txt.font.size = Pt(9.5)
            r_txt.font.color.rgb = COLOR_DARK

    # -------------------------------------------------------------
    # MODULO 2
    # -------------------------------------------------------------
    h2_m2 = doc.add_paragraph()
    style_heading(h2_m2, "MÓDULO 2: INGENIERÍA DE PROMPTS: EL ARTE DE INSTRUIR CON PRECISIÓN (4 Horas — Sesiones 3 y 4)", level=2)

    temas_m2 = [
        ("2.1 La Estructura Maestra del Prompt (Método RC-TRF)", [
            "Definición y propósito de la ingeniería de prompts: comunicarse con precisión para obtener valor inmediato.",
            "Desglose de los 5 elementos de la fórmula RC-TRF:",
            "  - Rol (R): Asignación de identidad profesional ('Actúa como un experto en servicio al cliente guatemalteco...').",
            "  - Contexto (C): Situación, antecedentes y audiencia ('Trabajo en una distribuidora y un cliente está molesto por una entrega demorada...').",
            "  - Tarea (T): Verbo de acción e instrucción concreta ('Redacta una propuesta de solución y mensaje de disculpa...').",
            "  - Restricciones (R): Límites estrictos ('No excedas de 80 palabras, no ofrezcas descuentos monetarios directos...').",
            "  - Formato (F): Estructura visual de salida ('Entrega en tabla de 2 columnas o con viñetas numeradas...').",
            "Comparación en vivo: resultado antes vs resultado después de aplicar la fórmula."
        ]),
        ("2.2 Técnicas de Precisión y Aprendizaje con Ejemplos (Few-Shot Prompting)", [
            "Zero-Shot vs Few-Shot Prompting: cómo enseñarle a la IA tu propio estilo de comunicación dándole 1 o 2 ejemplos previos.",
            "Cadena de Pensamiento ('Chain of Thought'): la instrucción mágica 'Piensa paso a paso antes de responder' para resolver problemas lógicos.",
            "Técnicas de delimitación de contenido: uso de comillas triples (\"\"\") o corchetes para aislar textos que la IA debe analizar sin confundir con la orden.",
            "Asignación de tono y voz: formal, diplomático, cercano, persuasivo, técnico o pedagógico."
        ]),
        ("2.3 Diálogo Iterativo y Refinamiento Continuo", [
            "La conversación como proceso de esculpido: cómo repreguntar sin reiniciar la ventana de chat.",
            "Instrucciones de ajuste fino: 'Hazlo 30% más conciso', 'Cambia el segundo párrafo para que suene más empático', 'Añade una llamada a la acción clara'.",
            "Creación de una biblioteca personal de prompts reutilizables en notas del celular o procesador de texto.",
            "Taller aplicado: transformación de 3 peticiones deficientes reales en prompts profesionales de alto rendimiento."
        ])
    ]

    for sub_tit, sub_items in temas_m2:
        p_st = doc.add_paragraph()
        style_heading(p_st, sub_tit, level=3)
        for it in sub_items:
            p_it = doc.add_paragraph()
            style_p(p_it, space_before=1, space_after=2)
            p_it.paragraph_format.left_indent = Inches(0.25)
            r_b = p_it.add_run("• ")
            r_b.font.bold = True
            r_b.font.color.rgb = COLOR_PRIMARY
            r_txt = p_it.add_run(it)
            r_txt.font.name = "Calibri"
            r_txt.font.size = Pt(9.5)
            r_txt.font.color.rgb = COLOR_DARK

    # -------------------------------------------------------------
    # MODULO 3
    # -------------------------------------------------------------
    h2_m3 = doc.add_paragraph()
    style_heading(h2_m3, "MÓDULO 3: CASOS REALES EN LA OFICINA, NEGOCIO Y VIDA PROFESIONAL (4 Horas — Sesiones 5 y 6)", level=2)

    temas_m3 = [
        ("3.1 Comunicación Corporativa y Atención al Cliente", [
            "Redacción de correos electrónicos difíciles en minutos: cobranza cordial, felicitaciones de equipo, justificación de retrasos y solicitud de cotizaciones.",
            "Gestión asertiva de mensajes para WhatsApp Business y redes sociales comerciales: respuestas rápidas para preguntas frecuentes (FAQ) manteniendo el calor humano.",
            "Elaboración de minutas y actas de reunión a partir de notas desordenadas tomadas a mano o dictadas por audio.",
            "Traducción contextualizada y corrección ortográfica y gramatical avanzada de informes laborales."
        ]),
        ("3.2 Síntesis, Análisis de Documentos Extensos y Extracción de Datos", [
            "Carga y análisis de archivos (PDFs, manuales, circulares institucionales): cómo extraer resúmenes ejecutivos en 5 puntos clave.",
            "Preguntas directas a documentos: 'Según este reglamento, ¿cuáles son los pasos para solicitar una devolución?'.",
            "Conversión de texto libre desordenado en tablas comparativas estructuradas listas para copiar y pegar en Microsoft Excel o Google Sheets.",
            "Extracción de datos clave: fechas, montos, personas responsables y compromisos a partir de textos extensos."
        ]),
        ("3.3 Planificación, Creatividad y Productividad Operativa", [
            "Estructuración de listas de verificación operativas (checklists) para eventos, inventarios o control de calidad.",
            "Lluvia de ideas estructurada para pequeños negocios: nombres de campañas, promociones de temporada, ideas de contenido para redes sociales.",
            "Preparación de guiones y esquemas para presentaciones en reuniones de trabajo.",
            "Taller de aplicación práctica: resolución de 3 retos de oficina simulados con cronómetro en mano."
        ])
    ]

    for sub_tit, sub_items in temas_m3:
        p_st = doc.add_paragraph()
        style_heading(p_st, sub_tit, level=3)
        for it in sub_items:
            p_it = doc.add_paragraph()
            style_p(p_it, space_before=1, space_after=2)
            p_it.paragraph_format.left_indent = Inches(0.25)
            r_b = p_it.add_run("• ")
            r_b.font.bold = True
            r_b.font.color.rgb = COLOR_PRIMARY
            r_txt = p_it.add_run(it)
            r_txt.font.name = "Calibri"
            r_txt.font.size = Pt(9.5)
            r_txt.font.color.rgb = COLOR_DARK

    # -------------------------------------------------------------
    # MODULO 4
    # -------------------------------------------------------------
    h2_m4 = doc.add_paragraph()
    style_heading(h2_m4, "MÓDULO 4: ÉTICA, PRIVACIDAD, CAZA DE ALUCINACIONES Y PROYECTO INTEGRADOR (4 Horas — Sesiones 7 y 8)", level=2)

    temas_m4 = [
        ("4.1 Caza de Alucinaciones, Verificación y Pensamiento Crítico", [
            "Por qué la IA 'alucina': la naturaleza predictiva del modelo frente a la exactitud fáctica.",
            "La regla de oro de Fundación Kinal: 'Nunca firmes, envíes ni presentes un texto generado por IA que tú mismo no hayas leído, comprendido y verificado'.",
            "Técnicas activas para cazar errores: pedir que cite la fuente, solicitar que contraste con datos oficiales o pedirle que busque contradicciones en su propia respuesta.",
            "Detección de sesgos en respuestas automatizadas y preservación del sentido común humano."
        ]),
        ("4.2 Privacidad, Ciberseguridad e Higiene de Datos Personales y Empresariales", [
            "Qué información NUNCA debe compartirse con una IA pública: contraseñas, DPI, números de tarjetas de crédito, estados financieros confidenciales o expedientes médicos.",
            "Técnica de anonimización de datos: sustituir nombres reales y cifras delicadas por marcadores de posición (ej. '[Cliente A]', '[Monto X]').",
            "Configuración de privacidad en ChatGPT, Gemini y Copilot: cómo desactivar el historial para que no usen tus datos en reentrenamiento.",
            "Derechos de autor, plagio y honestidad académica y laboral: dar crédito y actuar con transparencia."
        ]),
        ("4.3 Proyecto Integrador: Mi Asistente Personal de Productividad y Cierre", [
            "Integración de aprendizajes: definición de un rol especializado adaptado a la ocupación real del participante (asistente de oficina, consultor de ventas, tutor personal).",
            "Diseño de un Prompt de Sistema ('System Prompt') maestro con directrices de tono, restricciones y formato.",
            "Presentación y sustentación del catálogo personal de prompts: demostración en vivo de resolución de una tarea real en menos de 5 minutos.",
            "Evaluación de certificación institucional: aplicación de rúbrica Kinal (nota mínima 75/100) y clausura del curso."
        ])
    ]

    for sub_tit, sub_items in temas_m4:
        p_st = doc.add_paragraph()
        style_heading(p_st, sub_tit, level=3)
        for it in sub_items:
            p_it = doc.add_paragraph()
            style_p(p_it, space_before=1, space_after=2)
            p_it.paragraph_format.left_indent = Inches(0.25)
            r_b = p_it.add_run("• ")
            r_b.font.bold = True
            r_b.font.color.rgb = COLOR_PRIMARY
            r_txt = p_it.add_run(it)
            r_txt.font.name = "Calibri"
            r_txt.font.size = Pt(9.5)
            r_txt.font.color.rgb = COLOR_DARK

    # ==========================================
    # 3. REFERENCIAS Y GUÍAS DE CONSULTA
    # ==========================================
    h1_3 = doc.add_paragraph()
    style_heading(h1_3, "3. Referencias Bibliográficas y Guías de Buenas Prácticas", level=1)

    referencias = [
        ("OpenAI:", "Prompt Engineering Guide — Estrategias y tácticas para interactuar con modelos de lenguaje."),
        ("Google DeepMind / Google Cloud:", "Generative AI Prompting Best Practices for Gemini."),
        ("Anthropic:", "Interactive Prompt Engineering Tutorial & Claude Documentation."),
        ("UNESCO:", "Recomendación sobre la Ética de la Inteligencia Artificial (París, 2021)."),
        ("Fundación Kinal:", "Ideario Institucional: El Principio del «Trabajo Bien Hecho» y Ética en el Trabajo."),
        ("Fundación Kinal:", "Normativa de Convivencia, Asistencia y Evaluación de la Escuela Técnica Superior.")
    ]

    for org, det in referencias:
        p_r = doc.add_paragraph()
        style_p(p_r, space_before=1, space_after=2)
        p_r.paragraph_format.left_indent = Inches(0.2)
        r_b = p_r.add_run(f"• {org} ")
        r_b.font.name = "Calibri"
        r_b.font.size = Pt(9)
        r_b.font.bold = True
        r_b.font.color.rgb = COLOR_PRIMARY
        r_d = p_r.add_run(det)
        r_d.font.name = "Calibri"
        r_d.font.size = Pt(9)
        r_d.font.color.rgb = COLOR_DARK

    section = doc.sections[0]
    footer = section.footer
    footer_p = footer.paragraphs[0]
    style_p(footer_p, space_before=0, space_after=0)
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_f = footer_p.add_run("Fundación Kinal | Inteligencia Artificial Aplicada y Productividad Ética — Temario Oficial")
    run_f.font.name = "Calibri"
    run_f.font.size = Pt(8.5)
    run_f.font.color.rgb = COLOR_SECONDARY

    out_dir = os.path.join("Cursos", "Inteligencia_Artificial")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "Temario_Curso_IA_Kinal.docx")
    doc.save(out_path)
    print(f"Temario guardado con éxito en: {out_path}")

if __name__ == "__main__":
    generate_temario_doc()
