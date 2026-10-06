import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E0", sz="4"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def style_p(p, space_before=0, space_after=4, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def style_heading(p, text, level=1, color=RGBColor(0, 51, 102)):
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.bold = True
    r.font.color.rgb = color
    if level == 1:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        r.font.size = Pt(15)
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        r.font.size = Pt(13)
    elif level == 3:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r.font.size = Pt(11)

def add_callout(doc, text, bold_prefix="", fill_hex="F7FAFC", border_color="003366"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    
    cell = tbl.cell(0, 0)
    set_cell_background(cell, fill_hex)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=120)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 51, 102)
        
    r_txt = p.add_run(text)
    r_txt.font.name = "Calibri"
    r_txt.font.size = Pt(10)
    r_txt.font.italic = True
    r_txt.font.color.rgb = RGBColor(45, 55, 72)
    
    p_post = doc.add_paragraph()
    style_p(p_post, space_before=0, space_after=6)

def generate_dosificacion_doc():
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    COLOR_PRIMARY = RGBColor(0, 51, 102)
    COLOR_SECONDARY = RGBColor(74, 85, 104)
    COLOR_DARK = RGBColor(26, 32, 44)

    # =========================================================================
    # ENCABEZADO TÉCNICO INSTITUCIONAL
    # =========================================================================
    p_inst = doc.add_paragraph()
    style_p(p_inst, space_before=20, space_after=2)
    r_inst = p_inst.add_run("FUNDACIÓN KINAL")
    r_inst.font.name = "Calibri"
    r_inst.font.size = Pt(14)
    r_inst.font.bold = True
    r_inst.font.color.rgb = COLOR_PRIMARY

    p_esc = doc.add_paragraph()
    style_p(p_esc, space_before=0, space_after=20)
    r_esc = p_esc.add_run("Escuela Técnica Superior — Coordinación de Formación Continua y Empleabilidad")
    r_esc.font.name = "Calibri"
    r_esc.font.size = Pt(11)
    r_esc.font.color.rgb = COLOR_SECONDARY

    p_tit = doc.add_paragraph()
    style_p(p_tit, space_before=10, space_after=8)
    r_tit = p_tit.add_run("DOSIFICACIÓN CURRICULAR Y SECUENCIA DIDÁCTICA DETALLADA")
    r_tit.font.name = "Calibri"
    r_tit.font.size = Pt(20)
    r_tit.font.bold = True
    r_tit.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    style_p(p_sub, space_before=0, space_after=18)
    r_sub = p_sub.add_run("Inteligencia Artificial Aplicada, Ingeniería de Prompts y Productividad Ética (8 Sesiones de 2 Horas — 16 Horas Formativas)")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True
    r_sub.font.color.rgb = COLOR_SECONDARY

    add_callout(doc, 
        "«El aprender en el proceso del trabajo es base y a la vez objetivo de la acción formativa». Las 16 horas pedagógicas se distribuyen en 4 semanas (1 mes), a razón de 2 sesiones nocturnas semanales de 120 minutos (Martes y Jueves de 19:00 a 21:00 hrs). Cada sesión contempla 10 minutos de pausa activa; por tanto, el tiempo de actividad formativa neta por encuentro es de 110 minutos (14 horas y 40 minutos en total). El programa está enfocado positivamente en potenciar capacidades y multiplicar la productividad personal y laboral de jóvenes y adultos sin formación especializada, preservando en todo momento el valor del lado humano, la ética y el «trabajo bien hecho» de Fundación Kinal. Nota mínima aprobatoria institucional: 75 puntos sobre 100.",
        bold_prefix="Organización Pedagógica Dual e Ideario Kinal:")

    # =========================================================================
    # 1. MATRIZ GENERAL CRONOLÓGICA DE LAS 8 SESIONES (16 HORAS)
    # =========================================================================
    h1_1 = doc.add_paragraph()
    style_heading(h1_1, "1. Matriz General de Dosificación de las 8 Sesiones (16 Horas — 1 Mes)", level=1)

    sesiones_info = [
        ("Sesión 1", "Semana 1 / Mód. 1", "Qué es la IA generativa, cómo piensa un LLM y primer diálogo efectivo", "Primer diálogo guiado y autodiagnóstico"),
        ("Sesión 2", "Semana 1 / Mód. 1", "Comparativa en vivo: ChatGPT, Gemini, Copilot y Claude; configuración segura", "Entorno multimodelo configurado"),
        ("Sesión 3", "Semana 2 / Mód. 2", "La fórmula maestra del prompt: Rol, Contexto, Tarea, Restricción y Formato", "3 prompts estructurados con método RC-TRF"),
        ("Sesión 4", "Semana 2 / Mód. 2", "Técnicas avanzadas: Ejemplos (Few-Shot), Cadena de Pensamiento y refinamiento", "Prompt con ejemplos y biblioteca personal"),
        ("Sesión 5", "Semana 3 / Mód. 3", "Comunicación profesional: correos difíciles, WhatsApp Business y minutas", "Borrador de correo difícil y minuta lista"),
        ("Sesión 6", "Semana 3 / Mód. 3", "Síntesis de documentos extensos (PDFs) y conversión de datos a tablas de Excel", "Resumen ejecutivo y tabla para Excel"),
        ("Sesión 7", "Semana 4 / Mód. 4", "Caza de alucinaciones, verificación de hechos, privacidad y ética de Kinal", "Ejercicio de detección de alucinaciones"),
        ("Sesión 8", "Semana 4 / Mód. 4", "PROYECTO INTEGRADOR: Mi Asistente Personal de Productividad y Sustentación", "Asistente personal evaluado (>= 75 pts)")
    ]

    tbl_matriz = doc.add_table(rows=len(sesiones_info) + 1, cols=4)
    tbl_matriz.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_matriz.autofit = False
    tbl_matriz.columns[0].width = Inches(1.1)
    tbl_matriz.columns[1].width = Inches(1.8)
    tbl_matriz.columns[2].width = Inches(2.3)
    tbl_matriz.columns[3].width = Inches(1.3)
    set_table_borders(tbl_matriz, color="CBD5E0", sz="4")

    hdr_m = tbl_matriz.rows[0]
    for c in hdr_m.cells:
        set_cell_background(c, "003366")
    headers_mat = ["Sesión (2h)", "Módulo / Semana", "Contenido Central del Encuentro", "Producto Observable"]
    for idx, text in enumerate(headers_mat):
        p = hdr_m.cells[idx].paragraphs[0]
        style_p(p, space_before=0, space_after=0)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for idx, (ses, mod, cont, prod) in enumerate(sesiones_info, start=1):
        row = tbl_matriz.rows[idx]
        if idx % 2 == 1:
            for c in row.cells:
                set_cell_background(c, "F7FAFC")
        for c_idx, val in enumerate([ses, mod, cont, prod]):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=45, bottom=45, left=55, right=55)
            p = cell.paragraphs[0]
            style_p(p, space_before=0, space_after=0)
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(8)
            r.font.color.rgb = COLOR_DARK
            if c_idx in [0, 3]:
                r.font.bold = True

    doc.add_page_break()

    # =========================================================================
    # 2. MICRODISEÑO DIDÁCTICO DETALLADO SESIÓN POR SESIÓN (8 SESIONES)
    # =========================================================================
    h1_2 = doc.add_paragraph()
    style_heading(h1_2, "2. Microdiseño Didáctico Detallado por Sesión (120 Minutos por Encuentro)", level=1)

    p_micro_intro = doc.add_paragraph()
    style_p(p_micro_intro, space_before=0, space_after=8)
    r = p_micro_intro.add_run("Cada una de las 8 sesiones formativas de 2 horas (120 minutos) se desglosa rigurosamente en tres momentos didácticos (Apertura 15 min, Desarrollo 90 min con Pausa de 10 min, y Cierre 15 min), con enfoque 100% interactivo y práctico:")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_DARK

    sesiones_detalladas = [
        # SEMANA 1
        ("Sesión 1", "Módulo 1: Descubriendo la IA", "Fundamentos de la IA Generativa sin Tecnicismos y Primer Diálogo Efectivo",
         "Comprende la naturaleza de los modelos de lenguaje y ejecuta con éxito su primera interacción guiada para resolver una duda cotidiana.",
         "Bienvenida, presentación del curso, ideario de Kinal y visión positiva: la IA como copiloto que potencia nuestras capacidades humanas. Pregunta disparadora: ¿En qué tarea repetitiva de tu semana te gustaría tener un asistente personal?",
         "Exposición dialogada sobre cómo funciona un modelo de lenguaje (predicción inteligente de texto explicada con analogías sencillas, sin fórmulas). Demostración en vivo de interacción. Pausa activa (10 min). Taller práctico individual: técnica del 'Explicador Universal' (pedir a la IA que explique un tema complejo a tres públicos distintos: un niño, un estudiante y un cliente).",
         "Puesta en común voluntaria en el aula virtual de los resultados más sorprendentes. Registro de la primera lección en la Bitácora de Aprendizaje y asignación del reto personal.",
         "Computadora o teléfono con navegador web, acceso a ChatGPT / Gemini, aula virtual Microsoft Teams.",
         "Captura de pantalla o texto del primer diálogo guiado y autodiagnóstico de tareas a potenciar."),

        ("Sesión 2", "Módulo 1: Descubriendo la IA", "Comparativa Práctica: ChatGPT, Gemini, Copilot y Claude en Acción",
         "Compara las fortalezas de los 4 grandes modelos gratuitos y configura su perfil de trabajo personal con medidas básicas de privacidad.",
         "Retroalimentación del reto de la Sesión 1. Pregunta reflexiva: ¿Por qué una herramienta nos da respuestas más creativas y otra busca mejor en internet?",
         "Presentación comparativa en vivo: ChatGPT (claridad y lógica), Gemini (integración con Google y actualidad), Copilot (búsqueda web estructurada) y Claude (calidez de redacción y textos largos). Pausa activa (10 min). Práctica guiada simultánea: los participantes abren dos herramientas en paralelo y envían la misma consulta para contrastar estilos, velocidad y enfoque.",
         "Consolidación de conclusiones: elaboración de un cuadro mental sobre cuándo conviene usar cada herramienta. Registro en la Bitácora de Aprendizaje.",
         "Cuentas gratuitas en Google Gemini, ChatGPT y Microsoft Copilot, computadora o smartphone.",
         "Cuadro comparativo personal de herramientas configuradas y listas para uso cotidiano."),

        # SEMANA 2
        ("Sesión 3", "Módulo 2: Ingeniería de Prompts", "La Estructura Maestra del Prompt Profesional: Método RC-TRF",
         "Aplica la fórmula Rol, Contexto, Tarea, Restricción y Formato para transformar peticiones vagas en instrucciones de alto rendimiento.",
         "Dinámica de apertura: 'El teléfono descompuesto digital'. Se muestra cómo una orden vaga genera un texto genérico e inútil vs cómo un prompt bien estructurado genera oro en polvo.",
         "Desglose paso a paso de la fórmula RC-TRF con ejemplos de la vida real (oficina, comercio, estudio). Demostración del instructor transformando un prompt deficiente. Pausa activa (10 min). Taller en vivo por parejas o individual: redacción de 3 prompts aplicando rigurosamente los 5 componentes para resolver situaciones laborales concretas.",
         "Comparativa en pantalla compartida del 'antes y después' de los textos obtenidos. Reflexión sobre cómo la claridad humana condiciona la calidad de la IA. Registro en bitácora.",
         "Plantilla digital del Método RC-TRF, procesador de texto o notas del celular, ChatGPT / Gemini.",
         "Documento con 3 prompts estructurados bajo el método RC-TRF con sus respuestas validadas."),

        ("Sesión 4", "Módulo 2: Ingeniería de Prompts", "Técnicas Avanzadas Accesibles: Ejemplos (Few-Shot) y Diálogo Iterativo",
         "Enseña su propio estilo a la IA proporcionando ejemplos previos y refina respuestas mediante repreguntas sin reiniciar la conversación.",
         "Reflexión de apertura: la IA es como un nuevo colaborador en la oficina; para que redacte como nosotros queremos, debemos mostrarle ejemplos de nuestro trabajo.",
         "Explicación de Few-Shot Prompting (dar 1 o 2 ejemplos del formato deseado antes de pedir la tarea). La técnica de la 'Cadena de Pensamiento' ('Piensa paso a paso'). Pausa activa (10 min). Práctica guiada: proceso de esculpido conversacional: ajustar tono, reducir longitud en un 30% y solicitar opciones alternativas en la misma conversación.",
         "Creación de la 'Caja de Herramientas de Prompts': guardado de las mejores fórmulas en un documento personal reutilizable. Registro en bitácora.",
         "Modelos generativos (ChatGPT, Gemini, Claude), guía de comandos de refinamiento, plantilla de biblioteca de prompts.",
         "Biblioteca personal de prompts inicial con al menos 4 plantillas reutilizables documentadas."),

        # SEMANA 3
        ("Sesión 5", "Módulo 3: Casos Reales en la Industria", "Comunicación Profesional: Correos Asertivos, WhatsApp y Minutas",
         "Redacta correspondencia laboral asertiva, mensajes comerciales empáticos y minutas de reunión a partir de notas desordenadas.",
         "Apertura: ¿Cuánto tiempo perdemos al día redactando correos complicados o buscando las palabras diplomáticas adecuadas para responder a un reclamo?",
         "Casos prácticos de comunicación: 1. Correo de cobranza formal pero cordial; 2. Respuesta empática a cliente insatisfecho en WhatsApp Business; 3. Solicitud de cotización a proveedores. Pausa activa (10 min). Ejercicio práctico: dictar notas desordenadas de una reunión simulada y pedir a la IA que elabore un acta formal con compromisos, fechas y responsables.",
         "Lectura y validación crítica de los textos: verificar que el tono suene humano, cálido y auténtico, no robótico. Registro en bitácora.",
         "Borradores de notas de reunión ficticia, situaciones de servicio al cliente, herramientas de IA.",
         "Borrador de correo corporativo diplomático y acta de reunión formal generada y revisada."),

        ("Sesión 6", "Módulo 3: Casos Reales en la Industria", "Síntesis de Documentos Extensos (PDFs) y Extracción de Datos a Tablas",
         "Extrae resúmenes ejecutivos de documentos densos y convierte texto no estructurado en tablas limpias listas para copiar en Excel.",
         "Problema cotidiano: recibir un documento de 25 páginas (reglamento, cotización larga, informe) y necesitar los 3 puntos clave en 5 minutos.",
         "Técnicas para cargar y consultar archivos en Gemini y ChatGPT. Formulación de preguntas de extracción precisa ('¿Cuáles son los plazos máximos de entrega según este texto?'). Pausa activa (10 min). Taller de datos: tomar un párrafo con datos dispersos (precios, clientes, fechas) y solicitar su conversión inmediata a tabla de filas y columnas lista para Excel.",
         "Demostración de copia y pegado directo en hojas de cálculo sin descuadre de formato. Reflexión sobre el ahorro de tiempo operativo. Registro en bitácora.",
         "Documentos de muestra en PDF (manuales, reglamentos breves), hojas de cálculo (Excel / Google Sheets), IA con carga de archivos.",
         "Resumen ejecutivo estructurado de un documento PDF y tabla comparativa exportada a Excel."),

        # SEMANA 4
        ("Sesión 7", "Módulo 4: Ética y Verificación", "Caza de Alucinaciones, Privacidad de Datos y el Principio de Verdad de Kinal",
         "Audita textos generados detectando datos inventados y aplica protocolos estrictos de privacidad y anonimización de información confidencial.",
         "Apertura con el ideario de Kinal: el valor del 'trabajo bien hecho' exige verdad y honradez. Un texto con datos falsos destruye la confianza profesional.",
         "Por qué ocurren las alucinaciones: la naturaleza matemática predictiva del modelo. Estrategias activas de verificación: pedir fuentes, contrastar con documentos oficiales y desconfiar de cifras no verificadas. Pausa activa (10 min). Taller de privacidad: qué datos NUNCA subir a la nube (DPI, contraseñas, salarios) y técnica de anonimización (uso de marcadores como '[Empresa X]').",
         "Ejercicio de detección de 'trampas': los estudiantes auditan un texto con 3 alucinaciones inducidas y redactan el informe corregido. Registro en bitácora.",
         "Textos con alucinaciones inducidas, lista de verificación de privacidad de datos, plataforma Kinal.",
         "Auditoría completada de un texto con alucinaciones corregidas y protocolo de privacidad firmado."),

        ("Sesión 8", "Módulo 4: Proyecto Integrador", "PROYECTO INTEGRADOR: Mi Asistente Personal de Productividad y Clausura",
         "Construye y sustenta en vivo su sistema de prompts maestros personalizados para su actividad real, cumpliendo con la rúbrica institucional (>= 75 pts).",
         "Apertura de la sesión de clausura: consolidación del viaje formativo. De no animarse a usar IA a convertirse en directores estratégicos de su propio copiloto digital.",
         "Presentación y sustentación individual o por parejas del Proyecto Integrador: cada participante demuestra en vivo cómo su asistente personal resuelve una tarea de su trabajo o negocio real en menos de 5 minutos, aplicando el método RC-TRF y verificación crítica. Pausa activa (10 min). Rondas de retroalimentación docente y entre compañeros.",
         "Evaluación sumativa mediante rúbrica analítica Kinal (nota mínima 75/100). Mensaje de cierre institucional sobre la primacía de la persona sobre la técnica y entrega de constancias.",
         "Proyectos integradores de los estudiantes, rúbrica institucional de evaluación, Microsoft Teams.",
         "Dossier del Asistente Personal de Productividad con prompts maestros y rúbrica aprobada (>= 75 pts).")
    ]

    for s_info in sesiones_detalladas:
        ses_num, mod_name, tit, ind, ap, des, cie, rec, prod = s_info
        
        h2_s = doc.add_paragraph()
        style_heading(h2_s, f"{ses_num}: {tit}", level=2)

        p_mod = doc.add_paragraph()
        style_p(p_mod, space_before=0, space_after=4)
        r = p_mod.add_run(f"Unidad: {mod_name}")
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = COLOR_SECONDARY

        # Tabla de Microdiseño
        tbl_s = doc.add_table(rows=6, cols=2)
        tbl_s.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_s.autofit = False
        tbl_s.columns[0].width = Inches(1.8)
        tbl_s.columns[1].width = Inches(4.7)
        set_table_borders(tbl_s, color="CBD5E0", sz="4")

        filas_data = [
            ("Indicador de Logro:", ind),
            ("Apertura (15 min):", ap),
            ("Desarrollo (90 min):", des),
            ("Cierre (15 min):", cie),
            ("Recursos y Equipos:", rec),
            ("Producto Observable:", prod)
        ]

        for r_idx, (etiq, cont) in enumerate(filas_data):
            row = tbl_s.rows[r_idx]
            set_cell_background(row.cells[0], "F7FAFC")
            set_cell_margins(row.cells[0], top=50, bottom=50, left=70, right=70)
            set_cell_margins(row.cells[1], top=50, bottom=50, left=70, right=70)
            
            p0 = row.cells[0].paragraphs[0]
            style_p(p0, space_before=0, space_after=0)
            r0 = p0.add_run(etiq)
            r0.font.name = "Calibri"
            r0.font.size = Pt(8.5)
            r0.font.bold = True
            r0.font.color.rgb = COLOR_PRIMARY

            p1 = row.cells[1].paragraphs[0]
            style_p(p1, space_before=0, space_after=0)
            r1 = p1.add_run(cont)
            r1.font.name = "Calibri"
            r1.font.size = Pt(8.5)
            r1.font.color.rgb = COLOR_DARK
            if r_idx in [0, 5]:
                r1.font.bold = True

        p_sp = doc.add_paragraph()
        style_p(p_sp, space_before=2, space_after=4)

    doc.add_page_break()

    # =========================================================================
    # 3. INSTRUMENTOS INSTITUCIONALES DE EVALUACIÓN
    # =========================================================================
    h1_3 = doc.add_paragraph()
    style_heading(h1_3, "3. Instrumentos de Evaluación Institucional y Rúbrica", level=1)

    add_callout(doc,
        "La evaluación del curso se rige por la rúbrica analítica institucional de Fundación Kinal. Todo estudiante debe alcanzar al menos 75 puntos sobre 100 para aprobar y obtener su certificación oficial, demostrando calidad técnica, pensamiento crítico y apego al principio del «trabajo bien hecho».",
        bold_prefix="Criterio de Evaluación y Umbral de Aprobación:")

    p_rub = doc.add_paragraph()
    style_heading(p_rub, "3.1 Rúbrica Analítica para el Proyecto Integrador de Productividad con IA (100 Puntos)", level=2)

    rubrica_data = [
        ("Estructura de Prompts (Método RC-TRF)", "30 Puntos", "Define con precisión Rol, Contexto, Tarea, Restricciones y Formato de salida. Usa lenguaje claro y delimitadores efectivos.", "Prompts vagos, sin restricciones claras o que generan respuestas genéricas sin estructura (< 22.5 pts)."),
        ("Pensamiento Crítico y Verificación", "25 Puntos", "Audita minuciosamente las respuestas, detecta posibles alucinaciones y verifica datos antes de considerarlos definitivos.", "Acepta y copia las respuestas de la IA sin revisar ni comprobar exactitud fáctica (< 18.7 pts)."),
        ("Aplicación Real y Ahorro de Tiempo", "20 Puntos", "El asistente resuelve un problema o tarea laboral/personal concreta, demostrando ahorro medible de tiempo y esfuerzo.", "El ejercicio es puramente teórico o desconectado de la actividad práctica del participante (< 15 pts)."),
        ("Privacidad, Ética y Lado Humano", "15 Puntos", "Demuestra respeto a la confidencialidad de datos (anonimización) y aporta calidez humana y empatía a la comunicación.", "Incluye datos sensibles en prompts o entrega textos fríos con tono robótico evidente (< 11.2 pts)."),
        ("Dossier de Prompts y Bitácora", "10 Puntos", "Presenta su catálogo de prompts ordenado, documentado y listo para ser reutilizado en su trabajo diario.", "Presentación desordenada, incompleta o sin justificación de uso (< 7.5 pts).")
    ]

    tbl_rub = doc.add_table(rows=6, cols=4)
    tbl_rub.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_rub.autofit = False
    tbl_rub.columns[0].width = Inches(1.8)
    tbl_rub.columns[1].width = Inches(1.0)
    tbl_rub.columns[2].width = Inches(2.2)
    tbl_rub.columns[3].width = Inches(1.5)
    set_table_borders(tbl_rub, color="CBD5E0", sz="4")

    hdr_r = tbl_rub.rows[0]
    for c in hdr_r.cells:
        set_cell_background(c, "003366")
    for idx, text in enumerate(["Criterio Evaluado", "Ponderación", "Desempeño Excelente (Trabajo Bien Hecho)", "Desempeño Insuficiente (< 75%)"]):
        p = hdr_r.cells[idx].paragraphs[0]
        style_p(p, space_before=0, space_after=0)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for idx, (crit, pond, exc, ins) in enumerate(rubrica_data, start=1):
        row = tbl_rub.rows[idx]
        if idx % 2 == 1:
            for c in row.cells:
                set_cell_background(c, "F7FAFC")
        for c_idx, val in enumerate([crit, pond, exc, ins]):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
            p = cell.paragraphs[0]
            style_p(p, space_before=0, space_after=0)
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(8)
            r.font.color.rgb = COLOR_DARK
            if c_idx in [0, 1]:
                r.font.bold = True

    p_bf = doc.add_paragraph()
    style_heading(p_bf, "3.2 Formato de la Ficha Maestra de Prompts del Estudiante", level=2)

    tbl_b_fmt = doc.add_table(rows=6, cols=2)
    tbl_b_fmt.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_b_fmt.autofit = False
    tbl_b_fmt.columns[0].width = Inches(2.2)
    tbl_b_fmt.columns[1].width = Inches(4.3)
    set_table_borders(tbl_b_fmt, color="CBD5E0", sz="4")

    formato_data = [
        ("Datos del Estudiante:", "Nombre: _______________________ | Carné: _________ | Sesión: ____ | Fecha: ________"),
        ("Objetivo del Prompt:", "Propósito laboral o personal (ej. 'Redacción de cotizaciones para clientes de ferretería')."),
        ("Texto del Prompt Estructurado (RC-TRF):", "Rol asignado, contexto, instrucción clara, restricciones delimitadas y formato de salida."),
        ("Resultado Obtenido y Refinamiento:", "Respuesta de la IA y ajustes aplicados en segunda iteración para alcanzar excelencia."),
        ("Verificación Crítica Realizada:", "Datos comprobados, corrección de posibles alucinaciones y aporte del toque humano."),
        ("Firma del Participante y Vo.Bo. Docente:", "Firma Estudiante: ___________________  |  Vo.Bo. Instructor Kinal: ___________________")
    ]

    for idx, (campo, lineas) in enumerate(formato_data):
        row = tbl_b_fmt.rows[idx]
        set_cell_background(row.cells[0], "F7FAFC")
        set_cell_margins(row.cells[0], top=55, bottom=55, left=75, right=75)
        set_cell_margins(row.cells[1], top=55, bottom=55, left=75, right=75)

        p0 = row.cells[0].paragraphs[0]
        style_p(p0, space_before=0, space_after=0)
        r0 = p0.add_run(campo)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_PRIMARY

        p1 = row.cells[1].paragraphs[0]
        style_p(p1, space_before=0, space_after=0)
        r1 = p1.add_run(lineas)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9)
        r1.font.color.rgb = COLOR_DARK

    section = doc.sections[0]
    footer = section.footer
    footer_p = footer.paragraphs[0]
    style_p(footer_p, space_before=0, space_after=0)
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_f = footer_p.add_run("Fundación Kinal | Inteligencia Artificial Aplicada y Productividad Ética — Dosificación Oficial")
    run_f.font.name = "Calibri"
    run_f.font.size = Pt(8.5)
    run_f.font.color.rgb = COLOR_SECONDARY

    out_dir = os.path.join("Cursos", "Inteligencia_Artificial")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "Dosificacion_y_Secuencia_Didactica_IA_Kinal.docx")
    doc.save(out_path)
    print(f"Dosificación guardada con éxito en: {out_path}")

if __name__ == "__main__":
    generate_dosificacion_doc()
