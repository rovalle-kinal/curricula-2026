import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=70, bottom=70, left=90, right=90):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E0", sz="4"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_callout(doc, text, bold_prefix="", fill_hex="F7FAFC", border_color="003366"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    
    cell = tbl.cell(0, 0)
    set_cell_background(cell, fill_hex)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=120)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 51, 102)
        
    r_txt = p.add_run(text)
    r_txt.font.name = "Calibri"
    r_txt.font.size = Pt(10)
    r_txt.font.italic = True
    r_txt.font.color.rgb = RGBColor(45, 55, 72)
    
    p_post = doc.add_paragraph()
    style_p(p_post, space_before=0, space_after=6)

def style_heading(p, text, level=1, color=RGBColor(0, 51, 102)):
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.bold = True
    r.font.color.rgb = color
    if level == 1:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        r.font.size = Pt(15)
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        r.font.size = Pt(13)
    elif level == 3:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r.font.size = Pt(11)

def style_p(p, space_before=0, space_after=6, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def generate_propuesta_doc():
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    COLOR_PRIMARY = RGBColor(0, 51, 102)
    COLOR_SECONDARY = RGBColor(74, 85, 104)
    COLOR_DARK = RGBColor(26, 32, 44)

    # =========================================================================
    # PORTADA INSTITUCIONAL
    # =========================================================================
    p_inst = doc.add_paragraph()
    style_p(p_inst, space_before=20, space_after=2)
    r_inst = p_inst.add_run("FUNDACIÓN KINAL")
    r_inst.font.name = "Calibri"
    r_inst.font.size = Pt(14)
    r_inst.font.bold = True
    r_inst.font.color.rgb = COLOR_PRIMARY

    p_esc = doc.add_paragraph()
    style_p(p_esc, space_before=0, space_after=20)
    r_esc = p_esc.add_run("Escuela Técnica Superior — Coordinación de Formación Continua y Empleabilidad")
    r_esc.font.name = "Calibri"
    r_esc.font.size = Pt(11)
    r_esc.font.color.rgb = COLOR_SECONDARY

    p_tit = doc.add_paragraph()
    style_p(p_tit, space_before=14, space_after=10)
    r_tit = p_tit.add_run("PLANIFICACIÓN Y PROPUESTA FORMATIVA INSTITUCIONAL:\nCIBERSEGURIDAD Y FUNDAMENTOS DE SEGURIDAD DE LA INFORMACIÓN")
    r_tit.font.name = "Calibri"
    r_tit.font.size = Pt(20)
    r_tit.font.bold = True
    r_tit.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    style_p(p_sub, space_before=0, space_after=20)
    r_sub = p_sub.add_run("Programa de Capacitación Híbrida para la Reconversión Laboral y Empleabilidad Inmediata en el Sector Tecnológico de Guatemala (80 Horas — 20 Sesiones)")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(11.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = COLOR_SECONDARY

    # Ficha Técnica
    tbl_ficha = doc.add_table(rows=8, cols=2)
    tbl_ficha.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ficha.autofit = False
    tbl_ficha.columns[0].width = Inches(2.2)
    tbl_ficha.columns[1].width = Inches(4.3)
    set_table_borders(tbl_ficha, color="CBD5E0", sz="6")

    ficha_data = [
        ("Nivel Formativo de Referencia:", "Equivalencia DQR Nivel 4 - 5 (Competencia técnica especializada y autonomía operativa)"),
        ("Población Destinataria:", "Jóvenes y adultos desempleados de 18 a 45 años con conocimientos previos de informática"),
        ("Duración Total:", "80 horas pedagógicas (20 sesiones de 4 horas / 240 min; 73h 20min de actividad neta)"),
        ("Modalidad Híbrida:", "80% Virtual (16 sesiones / 64 horas) y 20% Presencial (4 sesiones / 16 horas en Kinal)"),
        ("Estructura de Horarios:", "Por convenir institucionalmente (secuencia por sesiones, sin imponer días u horas fijas)"),
        ("Orientación de Empleabilidad:", "Soporte técnico, operación tecnológica y funciones iniciales de apoyo a seguridad"),
        ("Nota Mínima Aprobatoria:", "75 puntos sobre 100 en todas las comprobaciones, productos y proyecto integrador"),
        ("Sede Presencial:", "Fundación Kinal — Laboratorios de Cómputo y Redes, Sede Central, Zona 7, Ciudad de Guatemala")
    ]

    for idx, (label, val) in enumerate(ficha_data):
        row = tbl_ficha.rows[idx]
        set_cell_background(row.cells[0], "F7FAFC")
        set_cell_margins(row.cells[0], top=65, bottom=65, left=90, right=90)
        set_cell_margins(row.cells[1], top=65, bottom=65, left=90, right=90)
        
        p0 = row.cells[0].paragraphs[0]
        style_p(p0, space_before=0, space_after=0)
        r0 = p0.add_run(label)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_PRIMARY

        p1 = row.cells[1].paragraphs[0]
        style_p(p1, space_before=0, space_after=0)
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_DARK

    doc.add_page_break()

    # =========================================================================
    # 1. MARCO FILOSÓFICO E IDEARIO DE KINAL: ÉTICA Y EL TRABAJO BIEN HECHO
    # =========================================================================
    h1_1 = doc.add_paragraph()
    style_heading(h1_1, "1. Marco Filosófico e Ideario Institucional de Fundación Kinal", level=1)

    p_id_desc = doc.add_paragraph()
    style_p(p_id_desc, space_before=0, space_after=6)
    r = p_id_desc.add_run("El diseño formativo se cimenta en la visión humanística y ética de Fundación Kinal, garantizando que el aprendizaje técnico esté indisolublemente ligado a la responsabilidad personal, la honradez y el servicio a la comunidad:")
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = COLOR_DARK

    add_callout(doc, "«Formar a jóvenes y adultos a través de una educación integral, con énfasis en las áreas técnicas y tecnológicas, influyendo positivamente en su trabajo, su familia y la sociedad».", bold_prefix="Misión Institucional de Kinal:")

    add_callout(doc, "«Visión cristiana de la vida. Respeto a la dignidad de la persona. Espíritu de servicio. Trabajo bien hecho. Libertad personal responsable».", bold_prefix="Valores Nucleares de Kinal:")

    p_et = doc.add_paragraph()
    style_heading(p_et, "1.1 El Principio del «Trabajo Bien Hecho» Aplicado a la Seguridad de la Información", level=2)

    p_et_txt = doc.add_paragraph()
    style_p(p_et_txt, space_before=0, space_after=6)
    r = p_et_txt.add_run("En este programa, el principio rector del «trabajo bien hecho» se manifiesta en estándares profesionales innegociables:\n"
                          "• Rigor en la asignación de controles: Un control no es una declaración abstracta; debe contener alcance explícito, responsable asignado, procedimiento de ejecución y evidencia verificable de cumplimiento.\n"
                          "• Trazabilidad y verdad en los registros: Separar hechos comprobables de meras hipótesis técnicas al analizar eventos en SIEM o servidores, preservando la exactitud de fechas, actores y resultados sin alterar evidencias.\n"
                          "• Confidencialidad y protección cotidiana de datos: Cero tolerancia a compartir credenciales o anotarlas en lugares visibles; uso estricto de cuentas individuales y canales aprobados desde la primera sesión.")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_DARK

    # =========================================================================
    # 2. FUNDAMENTACIÓN, ENFOQUE PEDAGÓGICO Y CASO TRANSVERSAL
    # =========================================================================
    h1_2 = doc.add_paragraph()
    style_heading(h1_2, "2. Fundamentación Técnica, Enfoque Pedagógico y Caso Transversal", level=1)

    p_met = doc.add_paragraph()
    style_p(p_met, space_before=0, space_after=6)
    r = p_met.add_run("A diferencia de propuestas puramente teóricas o cursos desarticulados de herramientas ofensivas, este programa adopta un enfoque constructivo orientado a la realidad laboral de las organizaciones en Guatemala:\n"
                      "• Caso Práctico Transversal de Empresa Ficticia: Cada sesión conecta un concepto técnico con una situación laboral concreta a través de una empresa ficticia que cuenta con estaciones de trabajo, red local, un servidor corporativo, una aplicación web y servicios SaaS. El mismo caso se amplía progresivamente durante los 8 módulos, evitando ejercicios aislados o controles descontextualizados.\n"
                      "• Enfoque de Análisis y Verificación (No Explotativo): Las actividades prácticas se centran en leer, comparar, representar, justificar y documentar. Las demostraciones se ejecutan sobre escenarios autorizados y datos simulados, sin exigir explotación de vulnerabilidades ni ataques intrusivos.\n"
                      "• Delimitación Profesional del Alcance: El curso forma el criterio técnico para interpretar sistemas, dialogar con proveedores, diseñar salvaguardas y escalar incidentes; no acredita especialización en pentesting, auditoría avanzada ni ingeniería de redes.")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_DARK

    # =========================================================================
    # 3. MARCO DQR Y SISTEMA DUAL DE APRENDIZAJE
    # =========================================================================
    h1_3 = doc.add_paragraph()
    style_heading(h1_3, "3. Marco Alemán de Cualificaciones (DQR Nivel 4-5) y Enfoque Dual", level=1)

    add_callout(doc, "«concepto de aprendizaje para jóvenes que tiene como objetivo preparar a los aprendices para su vida profesional, por lo que la formación se realiza en régimen de alternancia entre la escuela o el centro de formación y la empresa. El aprender en el proceso del trabajo es base y a la vez objetivo de la acción formativa».", bold_prefix="Definición de Formación Profesional Dual:")

    add_callout(doc, "«the ability and readiness of the individual to use knowledge, skills and personal, social and methodological competences and to behave in a considered, individual and socially responsible manner. Competence is understood in this sense as the comprehensive ability to act».", bold_prefix="Concepto de Competencia en el DQR:")

    tbl_dqr = doc.add_table(rows=5, cols=3)
    tbl_dqr.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_dqr.autofit = False
    tbl_dqr.columns[0].width = Inches(1.8)
    tbl_dqr.columns[1].width = Inches(1.5)
    tbl_dqr.columns[2].width = Inches(3.2)
    set_table_borders(tbl_dqr, color="CBD5E0", sz="4")

    hdr_dqr = tbl_dqr.rows[0]
    for c in hdr_dqr.cells:
        set_cell_background(c, "003366")
    for idx, text in enumerate(["Dimensión DQR", "Subdimensión", "Evidencia de Desempeño en el Programa"]):
        p = hdr_dqr.cells[idx].paragraphs[0]
        style_p(p, space_before=0, space_after=0)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    dqr_items = [
        ("Competencia Profesional\n(Fachkompetenz)", "Conocimientos\n(Wissen)", "Arquitectura de sistemas/nube, protocolos web (HTTP/TLS), marcos de autenticación (SAML/OAuth), ciclo de vulnerabilidades (CVE/CVSS) y estándares normativos (ISO 27001, PCI DSS, SOC, HIPAA)."),
        ("Competencia Profesional\n(Fachkompetenz)", "Destrezas\n(Fertigkeiten)", "Diagramación de flujos de sistemas, construcción de matrices RBAC, diseño de controles verificables, correlación de logs en SIEM, lectura de avisos de seguridad y redacción de SOPs."),
        ("Competencia Personal\n(Personale Kompetenz)", "Competencia Social\n(Sozialkompetenz)", "Capacidad de comunicación asertiva entre el área técnica y las áreas usuarias, fundamentación de decisiones ante comités de cambio y reporte oportuno y sereno de incidentes."),
        ("Competencia Personal\n(Personale Kompetenz)", "Autonomía\n(Selbständigkeit)", "Criterio técnico independiente para evaluar riesgos, responsabilidad en la custodia de accesos, autoinspección de calidad técnica y adhesión inquebrantable a la ética digital.")
    ]

    for idx, (dim, sub, evid) in enumerate(dqr_items, start=1):
        row = tbl_dqr.rows[idx]
        if idx % 2 == 1:
            for c in row.cells:
                set_cell_background(c, "F7FAFC")
        for c_idx, val in enumerate([dim, sub, evid]):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=65, bottom=65, left=75, right=75)
            p = cell.paragraphs[0]
            style_p(p, space_before=0, space_after=0)
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            r.font.color.rgb = COLOR_DARK

    doc.add_page_break()

    # =========================================================================
    # 4. PERFILES INSTITUCIONALES DE INGRESO Y EGRESO
    # =========================================================================
    h1_4 = doc.add_paragraph()
    style_heading(h1_4, "4. Perfiles de Ingreso y Egreso Institucional", level=1)

    p_ing = doc.add_paragraph()
    style_heading(p_ing, "4.1 Perfil de Ingreso", level=2)
    p_ing_txt = doc.add_paragraph()
    style_p(p_ing_txt, space_before=0, space_after=6)
    r = p_ing_txt.add_run("Dirigido a personas de 18 a 45 años desempleadas con conocimientos previos de informática que buscan fortalecer su preparación para funciones iniciales de soporte, operación tecnológica y apoyo a seguridad de la información. Se requiere manejo básico de computadora, navegador, archivos y documentos; no se requiere programación ni administración avanzada de redes.")
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = COLOR_DARK

    p_egr = doc.add_paragraph()
    style_heading(p_egr, "4.2 Perfil de Egreso", level=2)
    p_egr_txt = doc.add_paragraph()
    style_p(p_egr_txt, space_before=0, space_after=12)
    r = p_egr_txt.add_run("El egresado interpreta sistemas, redes y servicios digitales delimitando fronteras y riesgos; diseña controles de identidad y acceso bajo mínimo privilegio; analiza la protección del transporte web; correlaciona eventos en SIEM; gestiona el ciclo de vida de vulnerabilidades; redacta políticas y SOPs ejecutables; y reporta señales de incidentes bajo las mejores prácticas de la industria y los valores de Kinal.")
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = COLOR_DARK

    # =========================================================================
    # 5. ARQUITECTURA CURRICULAR Y ESTRUCTURA MODULAR (8 MÓDULOS / 80 HORAS)
    # =========================================================================
    h1_5 = doc.add_paragraph()
    style_heading(h1_5, "5. Arquitectura Curricular y Distribución Modular (80 Horas — 20 Sesiones)", level=1)

    p_mod_desc = doc.add_paragraph()
    style_p(p_mod_desc, space_before=0, space_after=6)
    r = p_mod_desc.add_run("Las 80 horas corresponden a 20 encuentros de 240 minutos (con 20 minutos de receso por encuentro, sumando 73h 20min de actividad formativa efectiva). El programa se distribuye en 8 módulos especializados:")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_DARK

    tbl_mods = doc.add_table(rows=9, cols=5)
    tbl_mods.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_mods.autofit = False
    tbl_mods.columns[0].width = Inches(0.9)
    tbl_mods.columns[1].width = Inches(2.2)
    tbl_mods.columns[2].width = Inches(0.6)
    tbl_mods.columns[3].width = Inches(0.9)
    tbl_mods.columns[4].width = Inches(1.9)
    set_table_borders(tbl_mods, color="CBD5E0", sz="4")

    hdr_m = tbl_mods.rows[0]
    for c in hdr_m.cells:
        set_cell_background(c, "003366")
    headers_m = ["Módulo", "Área de Formación", "Horas", "Sesiones", "Evidencia Principal del Módulo"]
    for idx, text in enumerate(headers_m):
        p = hdr_m.cells[idx].paragraphs[0]
        style_p(p, space_before=0, space_after=0)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    mods_data = [
        ("Módulo 1", "Sistemas y gestión informática para la seguridad", "12 h", "Sesiones 1-3", "Diagrama del sistema con equipos, red, servidor, SaaS y registro de incidente."),
        ("Módulo 2", "Identidad y control de acceso en aplicaciones", "12 h", "Sesiones 4-6", "Flujo de sesión local/federado, matriz RBAC y procedimiento de baja de usuario."),
        ("Módulo 3", "Diseño de controles y responsabilidades", "8 h", "Sesiones 7-8", "Diseño de acceso temporal, roles separados y reconstrucción con ticket y logs SIEM."),
        ("Módulo 4", "Seguridad de comunicaciones web y APIs", "12 h", "Sesiones 9-11", "Trazado del recorrido web, revisión de certificados y petición de API segura."),
        ("Módulo 5", "Ciclo de vida y gestión de vulnerabilidades", "12 h", "Sesiones 12-14", "Plan de ciclo de vida de inventario y tratamiento justificado de aviso CVE real."),
        ("Módulo 6", "Documentación y gestión del riesgo de seguridad", "8 h", "Sesiones 15-16", "Política de accesos, SOP de baja, evaluación de riesgo y solicitud de aceptación."),
        ("Módulo 7", "Estándares, regulaciones e informes de aseguramiento", "8 h", "Sesiones 17-18", "Comparativa de aplicabilidad entre PCI DSS, ISO 27001, HIPAA y SOC 1/2."),
        ("Módulo 8", "Higiene digital y respuesta inicial a incidentes", "8 h", "Sesiones 19-20", "Reporte de phishing, preservación y sustentación del Expediente de Seguridad.")
    ]

    for idx, (m, a, h, s, e) in enumerate(mods_data, start=1):
        row = tbl_mods.rows[idx]
        if idx % 2 == 1:
            for c in row.cells:
                set_cell_background(c, "F7FAFC")
        for c_idx, val in enumerate([m, a, h, s, e]):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=55, bottom=55, left=65, right=65)
            p = cell.paragraphs[0]
            style_p(p, space_before=0, space_after=0)
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(8.5)
            r.font.color.rgb = COLOR_DARK
            if c_idx in [0, 2]:
                r.font.bold = True

    # =========================================================================
    # 6. ESQUEMA DE EVALUACIÓN Y EXPEDIENTE DEL CASO (UMBRAL ≥ 75 PUNTOS)
    # =========================================================================
    h1_6 = doc.add_paragraph()
    style_heading(h1_6, "6. Esquema de Evaluación Institucional y Expediente del Caso", level=1)

    p_ev_desc = doc.add_paragraph()
    style_p(p_ev_desc, space_before=0, space_after=6)
    r = p_ev_desc.add_run("El sistema de evaluación exige explicar por qué se propone cada control y cómo se comprobaría su resultado operativo. Conforme a las normas de Fundación Kinal, la nota mínima de suficiencia técnica es de 75 puntos sobre 100:")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_DARK

    tbl_ev = doc.add_table(rows=4, cols=3)
    tbl_ev.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ev.autofit = False
    tbl_ev.columns[0].width = Inches(2.2)
    tbl_ev.columns[1].width = Inches(1.1)
    tbl_ev.columns[2].width = Inches(3.2)
    set_table_borders(tbl_ev, color="CBD5E0", sz="4")

    hdr_ev = tbl_ev.rows[0]
    for c in hdr_ev.cells:
        set_cell_background(c, "003366")
    for idx, text in enumerate(["Componente Evaluativo", "Ponderación", "Evidencia y Criterio de Verificación"]):
        p = hdr_ev.cells[idx].paragraphs[0]
        style_p(p, space_before=0, space_after=0)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    ev_items = [
        ("Comprobaciones de aprendizaje", "20%", "Respuestas breves, interpretación de flujos y defensa de decisiones técnicas."),
        ("Productos de los módulos", "40%", "Diagramas de arquitectura, matrices RBAC, fichas y documentos revisados."),
        ("Proyecto integrador (Expediente)", "40%", "Expediente de seguridad del caso transversal y sustentación oral final.")
    ]

    for idx, (comp, pond, crit) in enumerate(ev_items, start=1):
        row = tbl_ev.rows[idx]
        if idx % 2 == 1:
            for c in row.cells:
                set_cell_background(c, "F7FAFC")
        for c_idx, val in enumerate([comp, pond, crit]):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=60, bottom=60, left=70, right=70)
            p = cell.paragraphs[0]
            style_p(p, space_before=0, space_after=0)
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(8.5)
            r.font.color.rgb = COLOR_DARK
            if c_idx in [0, 1]:
                r.font.bold = True

    p_exp = doc.add_paragraph()
    style_p(p_exp, space_before=8, space_after=4)
    r = p_exp.add_run("Contenido del Expediente de Seguridad del Caso Transversal (Entregable Terminal):")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = COLOR_PRIMARY

    exp_items = [
        "Alcance del sistema y clasificación de información; diagrama con componentes, datos y dependencias.",
        "Flujo de autenticación y sesión; matriz de roles y permisos (RBAC); alta, cambio y baja de usuarios.",
        "Asignación de responsabilidades y registro trazable de solicitud, aprobación, ejecución y cierre.",
        "Recorrido web y ficha de certificados; controles básicos sobre una API.",
        "Inventario con ciclo de vida; hallazgo de vulnerabilidad y plan de tratamiento con verificación.",
        "Política de accesos y SOP de baja; evaluación de riesgo y solicitud de aceptación temporal.",
        "Referencia de cumplimiento justificada y evidencia requerida; reporte formal de comunicación sospechosa."
    ]
    for e in exp_items:
        p_e = doc.add_paragraph()
        style_p(p_e, space_before=1, space_after=2)
        p_e.paragraph_format.left_indent = Inches(0.2)
        r_b = p_e.add_run("• ")
        r_b.font.bold = True
        r_b.font.color.rgb = COLOR_PRIMARY
        r_t = p_e.add_run(e)
        r_t.font.name = "Calibri"
        r_t.font.size = Pt(9.5)
        r_t.font.color.rgb = COLOR_DARK

    section = doc.sections[0]
    footer = section.footer
    footer_p = footer.paragraphs[0]
    style_p(footer_p, space_before=0, space_after=0)
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_f = footer_p.add_run("Fundación Kinal | Ciberseguridad y Fundamentos de Seguridad de la Información — Propuesta Oficial")
    run_f.font.name = "Calibri"
    run_f.font.size = Pt(8.5)
    run_f.font.color.rgb = COLOR_SECONDARY

    out_dir = os.path.join("Cursos", "Ciberseguridad")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "Propuesta_Curso_Ciberseguridad_Kinal.docx")
    doc.save(out_path)
    print(f"Propuesta guardada con éxito en: {out_path}")

if __name__ == "__main__":
    generate_propuesta_doc()
