import os
import re
import html
import json
import docx

WORKSPACE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CURSOS_DIR = os.path.join(WORKSPACE_DIR, "Cursos")
WEB_DIR = os.path.join(CURSOS_DIR, "Web")
os.makedirs(WEB_DIR, exist_ok=True)

def parse_docx(path):
    doc = docx.Document(path)
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    tables = []
    for t in doc.tables:
        rows = []
        for r in t.rows:
            rows.append([c.text.strip() for c in r.cells])
        tables.append(rows)
    return paras, tables

def clean_html(text):
    if not text:
        return ""
    return html.escape(text).replace('\n', '<br>')

# ==============================================================================
# 1. PARSE INTELIGENCIA ARTIFICIAL
# ==============================================================================
ia_prop_p, ia_prop_t = parse_docx(os.path.join(CURSOS_DIR, "Inteligencia_Artificial", "Propuesta_Curso_IA_Kinal.docx"))
ia_dos_p, ia_dos_t = parse_docx(os.path.join(CURSOS_DIR, "Inteligencia_Artificial", "Dosificacion_y_Secuencia_Didactica_IA_Kinal.docx"))

ia_ficha = ia_prop_t[0]
ia_rubric = ia_dos_t[11]
ia_bitacora = ia_dos_t[12]

ia_sessions = []
# 4 Módulos de 2 sesiones cada uno (total 8 sesiones de 2h en 4 semanas)
for idx in range(2, 10):
    t = ia_dos_t[idx]
    s_num = idx - 1
    t1_row = ia_dos_t[1][s_num] if s_num < len(ia_dos_t[1]) else ["", "", "", ""]
    
    obj = t[0][1]
    
    # Map remaining rows safely
    ap, des, cie, eq, evi = "", "", "", "", ""
    for r in t[1:]:
        lbl = r[0].lower()
        val = r[1]
        if "apertura" in lbl:
            ap = val
        elif "desarrollo" in lbl:
            des = val
        elif "cierre" in lbl:
            cie = val
        elif "recurso" in lbl or "equipo" in lbl:
            eq = val
        elif "producto" in lbl or "evidencia" in lbl:
            evi = val
    
    mod_sem = t1_row[1] if len(t1_row) > 1 else f"Módulo {((s_num-1)//2)+1}"
    title_short = t1_row[2] if len(t1_row) > 2 else f"Sesión {s_num}"
    prod_obs = t1_row[3] if len(t1_row) > 3 else evi
    
    semana_num = ((s_num - 1) // 2) + 1
    dia_semana = "Martes" if (s_num % 2 == 1) else "Jueves"
    mod_num = ((s_num - 1) // 2) + 1

    ia_sessions.append({
        "num": s_num,
        "code": f"S-{s_num:02d}",
        "mod_num": mod_num,
        "semana": f"Semana {semana_num}",
        "dia": dia_semana,
        "horario": f"{dia_semana} de 19:00 a 21:00 hrs",
        "module": f"Módulo {mod_num}",
        "title": title_short,
        "duration": "2 Horas (120 min)",
        "objective": obj,
        "apertura": ap,
        "desarrollo": des,
        "cierre": cie,
        "evidencias": prod_obs,
        "equipamiento": eq
    })

ia_modules_data = [
    {
        "num": 1,
        "title": "Descubriendo la IA: Fundamentos y Primeros Pasos Prácticos",
        "period": "MÓDULO 1 (Semana 1 • Sesiones 1 y 2)",
        "hours": "4 Horas Pedagógicas",
        "logro": "Comprender el funcionamiento operativo de los modelos generativos y navegar con fluidez en ChatGPT, Gemini y Copilot, configurando su espacio de trabajo digital sin barreras técnicas.",
        "topics": [
            {
                "num": 1,
                "title": "Fundamentos de la IA Generativa sin Tecnicismos",
                "cat": "fundamentos",
                "cat_name": "🧠 Fundamentos & Modelos LLM",
                "cat_badge": "bg-blue-100 text-blue-900 border border-blue-200",
                "desc": "De la búsqueda tradicional a la conversación inteligente. Cómo piensa un modelo de lenguaje (LLM): predicción probabilística de palabras y ventanas de contexto explicadas con metáforas cotidianas. Alcances reales vs expectativas infundadas."
            },
            {
                "num": 2,
                "title": "El Ecosistema Actual de Herramientas Gratuitas",
                "cat": "fundamentos",
                "cat_name": "🧠 Fundamentos & Modelos LLM",
                "cat_badge": "bg-blue-100 text-blue-900 border border-blue-200",
                "desc": "Panorama comparativo: ChatGPT (GPT-4o mini), Google Gemini, Microsoft Copilot y Claude. Puntos fuertes de cada uno: razonamiento, búsqueda web en vivo y ofimática. Creación de perfiles seguros y navegación por la interfaz."
            },
            {
                "num": 3,
                "title": "Primeras Interacciones y Dinámicas de Descubrimiento",
                "cat": "fundamentos",
                "cat_name": "🧠 Fundamentos & Modelos LLM",
                "cat_badge": "bg-blue-100 text-blue-900 border border-blue-200",
                "desc": "El 'Explicador Universal' (técnica ELI5). Lluvia de ideas guiada para situaciones cotidianas. Superación de la respuesta genérica: por qué preguntas vagas producen respuestas vacías. Autodiagnóstico de necesidades laborales."
            }
        ]
    },
    {
        "num": 2,
        "title": "Ingeniería de Prompts: El Arte de Instruir con Precisión (Método RC-TRF)",
        "period": "MÓDULO 2 (Semana 2 • Sesiones 3 y 4)",
        "hours": "4 Horas Pedagógicas",
        "logro": "Diseñar prompts de alto impacto utilizando la estructura Rol, Contexto, Tarea, Restricción y Formato, aplicando técnicas de refinamiento iterativo y ejemplos para obtener respuestas profesionales.",
        "topics": [
            {
                "num": 4,
                "title": "La Estructura Maestra del Prompt (Método RC-TRF)",
                "cat": "prompts",
                "cat_name": "✍️ Ingeniería de Prompts (RC-TRF)",
                "cat_badge": "bg-purple-100 text-purple-900 border border-purple-200",
                "desc": "Desglose de los 5 elementos de la fórmula RC-TRF: Rol profesional asignado, Contexto y antecedentes, Tarea con verbo de acción, Restricciones estrictas y Formato visual de salida. Comparativa de resultados antes vs después."
            },
            {
                "num": 5,
                "title": "Técnicas de Precisión y Aprendizaje con Ejemplos (Few-Shot Prompting)",
                "cat": "prompts",
                "cat_name": "✍️ Ingeniería de Prompts (RC-TRF)",
                "cat_badge": "bg-purple-100 text-purple-900 border border-purple-200",
                "desc": "Zero-Shot vs Few-Shot Prompting: enseñar tu propio estilo mediante 1 o 2 ejemplos. Cadena de pensamiento ('Piensa paso a paso'). Delimitadores de contenido con comillas triples (\"\"\") y asignación precisa de tono y voz."
            },
            {
                "num": 6,
                "title": "Diálogo Iterativo y Biblioteca de Prompts Reutilizables",
                "cat": "prompts",
                "cat_name": "✍️ Ingeniería de Prompts (RC-TRF)",
                "cat_badge": "bg-purple-100 text-purple-900 border border-purple-200",
                "desc": "La conversación como proceso de esculpido: repreguntar y ajustar sin reiniciar el chat ('Hazlo 30% más conciso', 'Añade una llamada a la acción'). Creación de una biblioteca personal de plantillas maestras reutilizables."
            }
        ]
    },
    {
        "num": 3,
        "title": "Casos Reales en la Oficina, Negocio y Vida Profesional",
        "period": "MÓDULO 3 (Semana 3 • Sesiones 5 y 6)",
        "hours": "4 Horas Pedagógicas",
        "logro": "Aplicar herramientas de IA para redactar correspondencia profesional, resumir informes extensos, extraer datos a tablas y estructurar planes de trabajo cotidianos en minutos.",
        "topics": [
            {
                "num": 7,
                "title": "Comunicación Corporativa y Atención al Cliente",
                "cat": "aplicacion",
                "cat_name": "💼 Oficina & Productividad",
                "cat_badge": "bg-amber-100 text-amber-900 border border-amber-200",
                "desc": "Redacción de correos difíciles en minutos (cobranza cordial, felicitaciones, justificación de retrasos). Mensajes asertivos para WhatsApp Business y redes comerciales. Minutas de reunión a partir de notas desordenadas o audios."
            },
            {
                "num": 8,
                "title": "Síntesis, Análisis de Documentos Extensos y Extracción de Datos",
                "cat": "aplicacion",
                "cat_name": "💼 Oficina & Productividad",
                "cat_badge": "bg-amber-100 text-amber-900 border border-amber-200",
                "desc": "Carga y consulta de archivos PDF y reglamentos densos. Preguntas directas a documentos. Conversión de texto libre desordenado en tablas comparativas estructuradas listas para copiar y pegar en Microsoft Excel o Google Sheets."
            },
            {
                "num": 9,
                "title": "Planificación, Creatividad y Productividad Operativa",
                "cat": "aplicacion",
                "cat_name": "💼 Oficina & Productividad",
                "cat_badge": "bg-amber-100 text-amber-900 border border-amber-200",
                "desc": "Estructuración de listas de verificación (checklists) operativas para eventos o inventarios. Lluvia de ideas para promociones y campañas de temporada. Guiones para presentaciones laborales y resolución de retos contra reloj."
            }
        ]
    },
    {
        "num": 4,
        "title": "Ética, Privacidad, Caza de Alucinaciones y Proyecto Integrador",
        "period": "MÓDULO 4 (Semana 4 • Sesiones 7 y 8)",
        "hours": "4 Horas Pedagógicas",
        "logro": "Auditar críticamente los contenidos generados detectando sesgos y alucinaciones, proteger la privacidad de datos sensibles y construir un asistente personal de productividad evaluado con rúbrica institucional (>= 75 pts).",
        "topics": [
            {
                "num": 10,
                "title": "Caza de Alucinaciones, Verificación y Pensamiento Crítico",
                "cat": "etica",
                "cat_name": "🛡️ Ética, Privacidad & Verificación",
                "cat_badge": "bg-emerald-100 text-emerald-900 border border-emerald-200",
                "desc": "Por qué la IA alucina: naturaleza probabilística frente a verdad fáctica. La regla de oro de Kinal: 'Nunca firmes ni envíes un texto de IA que tú mismo no hayas verificado'. Técnicas de contraste de fuentes y detección de sesgos."
            },
            {
                "num": 11,
                "title": "Privacidad, Ciberseguridad e Higiene de Datos Personales",
                "cat": "etica",
                "cat_name": "🛡️ Ética, Privacidad & Verificación",
                "cat_badge": "bg-emerald-100 text-emerald-900 border border-emerald-200",
                "desc": "Datos que NUNCA deben compartirse con una IA pública (DPI, contraseñas, tarjetas, estados financieros). Técnicas de anonimización con marcadores de posición ([Cliente A], [Monto X]). Desactivación de historiales de entrenamiento."
            },
            {
                "num": 12,
                "title": "Proyecto Integrador: Mi Asistente Personal de Productividad y Cierre",
                "cat": "etica",
                "cat_name": "🛡️ Ética, Privacidad & Verificación",
                "cat_badge": "bg-emerald-100 text-emerald-900 border border-emerald-200",
                "desc": "Diseño de un Prompt de Sistema ('System Prompt') maestro personalizado a la ocupación real del estudiante. Demostración en vivo de resolución de una tarea laboral compleja en < 5 minutos. Evaluación con rúbrica Kinal (>= 75 pts)."
            }
        ]
    }
]

# ==============================================================================
# GENERADORES DE TARJETAS MODULARES PARA TEMARIO
# ==============================================================================
def build_modular_temario_html(modules):
    cards_html = ""
    for mod in modules:
        topics_items = ""
        for t in mod["topics"]:
            topics_items += f"""
            <div class="content-item item-{t['cat']} p-3.5 rounded-xl text-xs space-y-1.5 transition border border-slate-200/90 bg-white hover:bg-slate-50 shadow-sm" data-category="{t['cat']}">
                <div class="flex flex-wrap justify-between items-center gap-1.5 font-bold">
                    <span class="text-slate-900 text-sm flex items-center gap-2">
                        <span class="w-5 h-5 rounded-full bg-slate-100 text-slate-700 text-xs font-black flex items-center justify-center shrink-0">{t['num']}</span>
                        {clean_html(t['title'])}
                    </span>
                    <span class="px-2.5 py-0.5 rounded-full text-[10px] font-extrabold uppercase {t['cat_badge']}">
                        {clean_html(t['cat_name'])}
                    </span>
                </div>
                <p class="text-slate-600 leading-relaxed pl-7">{clean_html(t['desc'])}</p>
            </div>
            """
        
        cards_html += f"""
        <div class="bg-white rounded-2xl border border-slate-300 shadow-sm overflow-hidden flex flex-col hover:shadow-md transition">
            <div class="bg-slate-900 text-white p-4 sm:p-5 border-b border-slate-800 space-y-1">
                <div class="flex flex-wrap justify-between items-center text-xs text-amber-400 font-semibold mb-1 gap-2">
                    <span>{clean_html(mod['period'])}</span>
                    <span class="px-2 py-0.5 rounded bg-purple-900/60 text-purple-200 border border-purple-400/30 text-[11px] font-bold">{clean_html(mod['hours'])}</span>
                </div>
                <h4 class="text-base sm:text-lg font-bold text-white leading-snug">{clean_html(mod['title'])}</h4>
                <p class="text-xs text-slate-300 pt-1 leading-relaxed">
                    <strong class="text-amber-300">Logro esperado:</strong> {clean_html(mod['logro'])}
                </p>
            </div>
            <div class="p-4 sm:p-5 divide-y divide-slate-100 flex-1 space-y-2.5 bg-slate-50/50">
                {topics_items}
            </div>
        </div>
        """
    return cards_html

ia_modular_cards = build_modular_temario_html(ia_modules_data)

# ==============================================================================
# GENERADOR DE FICHA TÉCNICA UNIFICADA
# ==============================================================================
def build_single_card_ficha(ficha_rows, code):
    rows_html = ""
    for r in ficha_rows:
        rows_html += f"""
        <div class="py-3 sm:grid sm:grid-cols-3 sm:gap-4 flex flex-col">
            <dt class="font-bold text-slate-900 text-xs sm:col-span-1">{clean_html(r[0])}</dt>
            <dd class="text-slate-700 text-xs sm:col-span-2 mt-0.5 sm:mt-0 leading-relaxed">{clean_html(r[1])}</dd>
        </div>
        """
    
    return f"""
    <div class="bg-white rounded-3xl border border-slate-200 p-6 sm:p-8 shadow-sm space-y-5">
        <div class="flex flex-wrap items-center justify-between border-b border-slate-200 pb-4 gap-2">
            <div class="flex items-center gap-2">
                <span class="w-8 h-8 rounded-lg bg-purple-50 text-purple-900 flex items-center justify-center font-black">
                    <i data-lucide="file-badge" class="w-4 h-4"></i>
                </span>
                <h3 class="text-xl font-black text-slate-900">
                    Ficha Técnica Oficial del Programa
                </h3>
            </div>
            <span class="px-3 py-1 rounded-full bg-purple-50 text-purple-900 text-xs font-bold border border-purple-200">
                Código: {clean_html(code)}
            </span>
        </div>
        <div class="divide-y divide-slate-100">
            {rows_html}
        </div>
    </div>
    """

ia_ficha_single_card = build_single_card_ficha(ia_ficha, "IA-2026-KINAL")

# ==============================================================================
# GENERADOR DEL CALENDARIO CRONOLÓGICO VISUAL (4 SEMANAS)
# ==============================================================================
def build_ia_calendar_html(sessions):
    weeks_html = ""
    for w in range(1, 5):
        w_sessions = [s for s in sessions if s["semana"] == f"Semana {w}"]
        cards = ""
        for s in w_sessions:
            cards += f"""
            <div class="bg-white rounded-xl border border-slate-200 p-4 shadow-sm hover:border-purple-600 hover:shadow-md transition flex flex-col justify-between space-y-3">
                <div class="space-y-1.5">
                    <div class="flex justify-between items-center text-[11px]">
                        <span class="font-black px-2 py-0.5 rounded bg-slate-900 text-purple-300">{s['code']}</span>
                        <span class="font-bold text-purple-800 bg-purple-50 px-2 py-0.5 rounded">{s['dia']} (19:00 - 21:00)</span>
                    </div>
                    <h5 class="font-bold text-slate-900 text-xs leading-snug">{clean_html(s['title'])}</h5>
                    <p class="text-[11px] text-slate-500 line-clamp-2">{clean_html(s['objective'])}</p>
                </div>
                <div class="pt-2 border-t border-slate-100 flex items-center justify-between text-[11px]">
                    <span class="text-slate-600 font-semibold">{s['duration']}</span>
                    <button onclick="goToSession({s['num']})" class="text-purple-700 font-bold hover:underline flex items-center gap-0.5">
                        Ver Secuencia →
                    </button>
                </div>
            </div>
            """
        weeks_html += f"""
        <div class="bg-slate-50 rounded-2xl border border-slate-200 p-4 space-y-3">
            <div class="flex items-center justify-between border-b border-slate-200 pb-2">
                <h4 class="font-black text-slate-900 text-sm flex items-center gap-2">
                    <i data-lucide="calendar" class="w-4 h-4 text-purple-700"></i> Semana {w} (Módulo {w})
                </h4>
                <span class="text-xs font-bold text-slate-500">2 Sesiones (4 Horas)</span>
            </div>
            <div class="grid sm:grid-cols-2 gap-3">
                {cards}
            </div>
        </div>
        """
    return f'<div class="grid lg:grid-cols-2 gap-4">{weeks_html}</div>'

ia_calendar_content = build_ia_calendar_html(ia_sessions)

# Rúbrica de Evaluación
def build_rubric_rows_html(rubric):
    return "".join([f"""
    <tr class="hover:bg-slate-50/80 transition">
        <td class="py-3.5 px-4 font-bold text-slate-900 bg-slate-50 border-r border-slate-200 align-top">{clean_html(row[0])}</td>
        <td class="py-3 px-4 text-purple-950 font-bold bg-purple-50/20 border-r border-slate-200 align-top text-center">{clean_html(row[1])}</td>
        <td class="py-3 px-4 text-emerald-950 bg-emerald-50/30 border-r border-slate-200 align-top">{clean_html(row[2])}</td>
        <td class="py-3 px-4 text-rose-950 bg-rose-50/30 align-top text-xs">{clean_html(row[3])}</td>
    </tr>
    """ for r_idx, row in enumerate(rubric) if r_idx > 0])

ia_rubric_rows = build_rubric_rows_html(ia_rubric)

# Bitácora de Prompts
def build_bitacora_table_html(bitacora):
    rows = ""
    for r in bitacora:
        rows += f"""
        <tr class="hover:bg-slate-50 transition">
            <td class="py-2.5 px-4 font-bold text-slate-900 bg-slate-50 border-r border-slate-200 align-top text-xs w-1/3">{clean_html(r[0])}</td>
            <td class="py-2.5 px-4 text-slate-700 align-top text-xs leading-relaxed">{clean_html(r[1])}</td>
        </tr>
        """
    return f"""
    <div class="bg-white rounded-3xl border border-slate-200 p-6 sm:p-8 shadow-sm space-y-4">
        <div class="flex items-center justify-between border-b border-slate-200 pb-3">
            <h4 class="font-black text-slate-900 text-base flex items-center gap-2">
                <i data-lucide="book-open" class="w-5 h-5 text-purple-700"></i>
                Estructura de la Bitácora de Prompts & Reflexión Ética
            </h4>
            <span class="px-2.5 py-1 rounded-full bg-purple-100 text-purple-900 font-bold text-xs">Portafolio Reutilizable</span>
        </div>
        <p class="text-xs text-slate-600">
            Formato oficial de registro de plantillas de prompts de alto impacto, refinamiento iterativo y verificación crítica que cada estudiante compila a lo largo del curso.
        </p>
        <div class="overflow-x-auto border border-slate-200 rounded-2xl">
            <table class="w-full text-left border-collapse divide-y divide-slate-100">
                <tbody>
                    {rows}
                </tbody>
            </table>
        </div>
    </div>
    """

ia_bitacora_card = build_bitacora_table_html(ia_bitacora)

print("Datos de IA parseados correctamente. Generando HTML...")

# ==============================================================================
# HTML BUILDER: IA.HTML
# ==============================================================================
ia_sessions_json = json.dumps(ia_sessions, ensure_ascii=False)

ia_html = f"""<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Inteligencia Artificial Aplicada y Productividad Ética | Kinal 2026</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <script src="https://unpkg.com/lucide@latest"></script>
    <style>
        @import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700;800&family=Space+Grotesk:wght@500;700;800;900&display=swap');
        body {{ font-family: 'Inter', sans-serif; background-color: #f8fafc; color: #0f172a; }}
        h1, h2, h3, h4, .font-heading {{ font-family: 'Space Grotesk', sans-serif; }}
        .hero-pattern {{
            background: linear-gradient(135deg, #090d16 0%, #1e1b4b 50%, #4338ca 100%);
        }}
        .tab-btn.active {{
            background-color: #003366;
            color: #ffffff;
            font-weight: 700;
        }}
        .session-item.active {{
            background-color: #003366;
            color: #ffffff;
            border-color: #003366;
        }}
        .session-item.active span, .session-item.active p {{
            color: #e2e8f0;
        }}
        .session-item.active .badge-code {{
            background-color: #a855f7;
            color: #ffffff;
        }}
    </style>
</head>
<body class="antialiased min-h-screen flex flex-col">

    <!-- Top Navigation Header -->
    <header class="bg-slate-950 text-white border-b border-slate-800 sticky top-0 z-50 shadow-md">
        <div class="max-w-7xl mx-auto px-4 sm:px-6 lg:px-8 py-3 flex flex-wrap justify-between items-center gap-4">
            <div class="flex items-center gap-3">
                <a href="index.html" class="bg-slate-800 hover:bg-slate-700 text-slate-200 text-xs font-bold px-3 py-1.5 rounded-lg transition flex items-center gap-1.5">
                    <i data-lucide="arrow-left" class="w-3.5 h-3.5"></i> Cursos
                </a>
                <span class="text-slate-600">|</span>
                <div class="text-slate-200 font-bold text-xs sm:text-sm flex items-center gap-2">
                    <i data-lucide="sparkles" class="w-4 h-4 text-purple-400"></i>
                    Inteligencia Artificial Aplicada, Prompts y Productividad Ética
                </div>
            </div>
            
            <!-- Quick Download Buttons -->
            <div class="flex items-center gap-2 text-xs">
                <span class="text-slate-400 hidden sm:inline">Word:</span>
                <a href="../Inteligencia_Artificial/Propuesta_Curso_IA_Kinal.docx" download class="px-2.5 py-1 rounded bg-blue-900/60 hover:bg-blue-800 text-blue-200 transition font-medium" title="Descargar Propuesta Oficial">
                    Propuesta
                </a>
                <a href="../Inteligencia_Artificial/Temario_Curso_IA_Kinal.docx" download class="px-2.5 py-1 rounded bg-emerald-900/60 hover:bg-emerald-800 text-emerald-200 transition font-medium" title="Descargar Temario Oficial">
                    Temario
                </a>
                <a href="../Inteligencia_Artificial/Dosificacion_y_Secuencia_Didactica_IA_Kinal.docx" download class="px-2.5 py-1 rounded bg-purple-900/60 hover:bg-purple-800 text-purple-200 transition font-medium" title="Descargar Dosificación Oficial">
                    Dosificación
                </a>
            </div>
        </div>
    </header>

    <!-- Hero Course Banner -->
    <section class="hero-pattern text-white py-10 px-4 sm:px-6 lg:px-8">
        <div class="max-w-7xl mx-auto space-y-3">
            <div class="flex flex-wrap items-center gap-2">
                <span class="px-3 py-1 rounded-full bg-purple-500/20 border border-purple-400/30 text-purple-300 text-xs font-bold tracking-wide uppercase">
                    Modalidad Virtual Sincrónica (Teams + Kinal.academy)
                </span>
                <span class="px-3 py-1 rounded-full bg-blue-500/20 border border-blue-400/30 text-blue-300 text-xs font-bold">
                    DQR Nivel 3 - 4
                </span>
                <span class="px-3 py-1 rounded-full bg-emerald-500/20 border border-emerald-400/30 text-emerald-300 text-xs font-bold">
                    Aprobación: ≥ 75 Pts
                </span>
            </div>
            <h1 class="text-2xl sm:text-4xl font-black tracking-tight leading-tight">
                Inteligencia Artificial Aplicada, Ingeniería de Prompts y Productividad Ética
            </h1>
            <p class="text-slate-300 text-xs sm:text-sm max-w-4xl leading-relaxed">
                Dominio práctico de herramientas de IA generativa (ChatGPT, Google Gemini, Copilot, Claude), ingeniería de prompts mediante el método estructurado RC-TRF (Rol, Contexto, Tarea, Restricción, Formato), automatización de tareas de oficina y negocio, detección de alucinaciones y ética con centralidad en la persona.
            </p>
        </div>
    </section>

    <!-- Tab Bar Navigation (UNA SOLA PALABRA, SIN NUMERACIÓN) -->
    <div class="bg-white border-b border-slate-200 sticky top-14 z-40 shadow-sm">
        <div class="max-w-7xl mx-auto px-4 sm:px-6 lg:px-8">
            <nav class="flex space-x-2 sm:space-x-3 py-2.5 overflow-x-auto text-xs font-bold">
                <button onclick="switchTab('propuesta')" id="btn-propuesta" class="tab-btn active px-4 py-2 rounded-xl transition flex items-center gap-1.5 shrink-0">
                    <i data-lucide="file-text" class="w-3.5 h-3.5"></i> Propuesta
                </button>
                <button onclick="switchTab('temario')" id="btn-temario" class="tab-btn px-4 py-2 rounded-xl bg-slate-100 hover:bg-slate-200 text-slate-700 transition flex items-center gap-1.5 shrink-0">
                    <i data-lucide="list-tree" class="w-3.5 h-3.5"></i> Temario
                </button>
                <button onclick="switchTab('calendario')" id="btn-calendario" class="tab-btn px-4 py-2 rounded-xl bg-slate-100 hover:bg-slate-200 text-slate-700 transition flex items-center gap-1.5 shrink-0">
                    <i data-lucide="calendar" class="w-3.5 h-3.5"></i> Calendario
                </button>
                <button onclick="switchTab('secuencias')" id="btn-secuencias" class="tab-btn px-4 py-2 rounded-xl bg-slate-100 hover:bg-slate-200 text-slate-700 transition flex items-center gap-1.5 shrink-0">
                    <i data-lucide="layers" class="w-3.5 h-3.5"></i> Secuencias
                </button>
                <button onclick="switchTab('evaluacion')" id="btn-evaluacion" class="tab-btn px-4 py-2 rounded-xl bg-slate-100 hover:bg-slate-200 text-slate-700 transition flex items-center gap-1.5 shrink-0">
                    <i data-lucide="award" class="w-3.5 h-3.5"></i> Evaluación
                </button>
            </nav>
        </div>
    </div>

    <!-- Main Container -->
    <main class="max-w-7xl mx-auto px-4 sm:px-6 lg:px-8 py-8 flex-1 space-y-8">

        <!-- ================================================================== -->
        <!-- PESTAÑA 1: PROPUESTA -->
        <!-- ================================================================== -->
        <div id="tab-propuesta" class="tab-content space-y-8">
            
            <!-- Ficha Técnica (UNA SOLA TARJETA) -->
            {ia_ficha_single_card}

            <!-- Marco Filosófico Kinal -->
            <div class="grid md:grid-cols-2 gap-6">
                <div class="bg-blue-50 border-l-4 border-blue-900 p-6 rounded-r-3xl space-y-3 shadow-sm">
                    <div class="flex items-center gap-2 text-blue-900 font-extrabold text-sm uppercase">
                        <i data-lucide="flag" class="w-4 h-4"></i> Misión Institucional de Fundación Kinal
                    </div>
                    <p class="text-slate-700 text-xs italic leading-relaxed">
                        «Formar a jóvenes y adultos a través de una educación integral, con énfasis en las áreas técnicas y tecnológicas, influyendo positivamente en su trabajo, su familia y la sociedad».
                    </p>
                    <div class="pt-2 text-xs font-bold text-blue-950">
                        Valores Nucleares: <span class="font-normal italic">Visión cristiana de la vida. Respeto a la dignidad de la persona. Espíritu de servicio. Trabajo bien hecho. Libertad personal responsable.</span>
                    </div>
                </div>

                <div class="bg-purple-50 border-l-4 border-purple-600 p-6 rounded-r-3xl space-y-3 shadow-sm">
                    <div class="flex items-center gap-2 text-purple-900 font-extrabold text-sm uppercase">
                        <i data-lucide="sparkle" class="w-4 h-4"></i> El Principio del «Trabajo Bien Hecho» en la Era de la IA
                    </div>
                    <p class="text-slate-700 text-xs leading-relaxed">
                        En Fundación Kinal, la inteligencia artificial no sustituye el ingenio ni el corazón humano: <strong>los potencia</strong>. El trabajo bien hecho implica que la persona mantiene el criterio final, la empatía y la responsabilidad ética, aplicando una regla de oro irrenunciable: <em>«Nunca firmes, envíes ni presentes un texto generado por IA que tú mismo no hayas leído, comprendido y verificado críticamente»</em>.
                    </p>
                </div>
            </div>

            <!-- Matriz DQR -->
            <div class="bg-white rounded-3xl border border-slate-200 p-8 shadow-sm space-y-6">
                <div class="space-y-1">
                    <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                        <i data-lucide="award" class="w-5 h-5 text-blue-700"></i>
                        Alineación con el Marco Alemán de Cualificaciones (DQR Nivel 3-4)
                    </h3>
                    <p class="text-slate-600 text-xs">
                        Desglose de competencias integrales de acción (Handlungskompetenz) para la habilitación tecnológica y autonomía operativa.
                    </p>
                </div>

                <div class="overflow-x-auto">
                    <table class="w-full text-xs text-left border-collapse">
                        <thead>
                            <tr class="bg-slate-100 text-slate-700 uppercase font-black border-b border-slate-200">
                                <th class="py-3 px-4">Dimensión DQR</th>
                                <th class="py-3 px-4">Subdimensión</th>
                                <th class="py-3 px-4">Evidencia de Desempeño en el Curso</th>
                            </tr>
                        </thead>
                        <tbody class="divide-y divide-slate-100 text-slate-700">
                            <tr>
                                <td class="py-3 px-4 font-bold text-blue-900 bg-blue-50/50">Competencia Profesional (Fachkompetenz)</td>
                                <td class="py-3 px-4 font-bold">Conocimiento (Wissen)</td>
                                <td class="py-3 px-4">Comprende la naturaleza probabilística de los LLMs, ventanas de contexto, arquitectura de los principales modelos (ChatGPT, Gemini, Copilot, Claude) y principios de privacidad de datos.</td>
                            </tr>
                            <tr>
                                <td class="py-3 px-4 font-bold text-blue-900 bg-blue-50/50">Competencia Profesional (Fachkompetenz)</td>
                                <td class="py-3 px-4 font-bold">Destrezas (Fertigkeiten)</td>
                                <td class="py-3 px-4">Formula prompts estructurados con la fórmula RC-TRF, aplica técnicas Few-Shot y Cadena de Pensamiento, extrae resúmenes ejecutivos de PDFs y convierte datos desordenados en tablas limpias para Excel.</td>
                            </tr>
                            <tr>
                                <td class="py-3 px-4 font-bold text-purple-900 bg-purple-50/50">Competencia Personal (Personale Kompetenz)</td>
                                <td class="py-3 px-4 font-bold">Autonomía (Selbständigkeit)</td>
                                <td class="py-3 px-4">Construye de forma independiente su Asistente Personal de Productividad y catálogo de prompts maestros adaptados a su trabajo real, auditando críticamente la veracidad de los resultados.</td>
                            </tr>
                            <tr>
                                <td class="py-3 px-4 font-bold text-purple-900 bg-purple-50/50">Competencia Personal (Personale Kompetenz)</td>
                                <td class="py-3 px-4 font-bold">Competencia Social y Ética</td>
                                <td class="py-3 px-4">Aplica principios de anonimización de datos sensibles de clientes y empresa, respeta la propiedad intelectual y cultiva una comunicación profesional cálida y empática en sus mensajes asistidos por IA.</td>
                            </tr>
                        </tbody>
                    </table>
                </div>
            </div>

        </div>

        <!-- ================================================================== -->
        <!-- PESTAÑA 2: TEMARIO (ESTÉTICA DE CONSTRUCCIÓN APARTADO 3) -->
        <!-- ================================================================== -->
        <div id="tab-temario" class="tab-content hidden space-y-6">
            
            <!-- Barra Superior de Filtros por Categoría -->
            <div class="bg-white rounded-3xl p-5 border border-slate-200 shadow-sm space-y-3">
                <div class="flex flex-col sm:flex-row justify-between items-start sm:items-center gap-2">
                    <div>
                        <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                            <i data-lucide="layers" class="w-5 h-5 text-purple-700"></i>
                            Distribución Modular con Código de Color
                        </h3>
                        <p class="text-slate-600 text-xs">
                            Estructura de 4 módulos formativos divididos por áreas clave. Filtra por categoría para enfocar el estudio.
                        </p>
                    </div>
                    <span class="text-xs font-bold text-slate-500 bg-slate-100 px-3 py-1 rounded-full">
                        Total: 12 Contenidos Clave
                    </span>
                </div>

                <!-- Botones de Filtro -->
                <div class="flex flex-wrap gap-2 pt-2 border-t border-slate-100">
                    <button onclick="filterIaCategory('all')" id="btn-ia-all" class="px-3.5 py-1.5 rounded-xl text-xs font-extrabold bg-slate-900 text-white shadow-sm transition">
                        Todos los Temas (12)
                    </button>
                    <button onclick="filterIaCategory('fundamentos')" id="btn-ia-fundamentos" class="px-3.5 py-1.5 rounded-xl text-xs font-extrabold bg-blue-50 hover:bg-blue-100 text-blue-900 border border-blue-200 transition">
                        🧠 Fundamentos & Modelos LLM
                    </button>
                    <button onclick="filterIaCategory('prompts')" id="btn-ia-prompts" class="px-3.5 py-1.5 rounded-xl text-xs font-extrabold bg-purple-50 hover:bg-purple-100 text-purple-900 border border-purple-200 transition">
                        ✍️ Ingeniería de Prompts (RC-TRF)
                    </button>
                    <button onclick="filterIaCategory('aplicacion')" id="btn-ia-aplicacion" class="px-3.5 py-1.5 rounded-xl text-xs font-extrabold bg-amber-50 hover:bg-amber-100 text-amber-900 border border-amber-200 transition">
                        💼 Oficina & Productividad
                    </button>
                    <button onclick="filterIaCategory('etica')" id="btn-ia-etica" class="px-3.5 py-1.5 rounded-xl text-xs font-extrabold bg-emerald-50 hover:bg-emerald-100 text-emerald-900 border border-emerald-200 transition">
                        🛡️ Ética, Privacidad & Verificación
                    </button>
                </div>
            </div>

            <!-- Cuadrícula Modular de Temas -->
            <div class="grid lg:grid-cols-2 gap-6" id="ia-modules-grid">
                {ia_modular_cards}
            </div>

        </div>

        <!-- ================================================================== -->
        <!-- PESTAÑA 3: CALENDARIO CRONOLÓGICO -->
        <!-- ================================================================== -->
        <div id="tab-calendario" class="tab-content hidden space-y-6">
            
            <div class="flex flex-col sm:flex-row justify-between items-start sm:items-center gap-3 bg-white rounded-3xl p-6 border border-slate-200 shadow-sm">
                <div>
                    <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                        <i data-lucide="calendar" class="w-5 h-5 text-purple-700"></i>
                        Calendario Cronológico de Encuentros Virtuales
                    </h3>
                    <p class="text-slate-600 text-xs mt-1">
                        Programa de 4 semanas (1 mes calendario) • 2 sesiones nocturnas por semana (Martes y Jueves de 19:00 a 21:00 hrs • 2h por encuentro).
                    </p>
                </div>
                <div class="flex items-center gap-2 text-xs font-bold text-purple-900 bg-purple-50 px-3 py-1.5 rounded-xl border border-purple-200">
                    <i data-lucide="video" class="w-4 h-4 text-purple-700"></i> Microsoft Teams & Kinal.academy
                </div>
            </div>

            <!-- Grid de Calendario -->
            <div class="space-y-6">
                {ia_calendar_content}
            </div>

        </div>

        <!-- ================================================================== -->
        <!-- PESTAÑA 4: SECUENCIAS (MASTER-DETAIL SIN ALARGAR LA PÁGINA) -->
        <!-- ================================================================== -->
        <div id="tab-secuencias" class="tab-content hidden space-y-6">
            
            <div class="flex flex-col sm:flex-row justify-between items-start sm:items-center gap-3 bg-white rounded-3xl p-5 border border-slate-200 shadow-sm">
                <div>
                    <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                        <i data-lucide="layers" class="w-5 h-5 text-purple-700"></i>
                        Secuencia Didáctica Interactiva Sesión a Sesión
                    </h3>
                    <p class="text-slate-600 text-xs mt-0.5">
                        Selecciona una sesión en el panel izquierdo para consultar su microdiseño en 3 momentos (Apertura 15 min, Desarrollo & Taller 90 min, Cierre & Reto 15 min).
                    </p>
                </div>
                <div class="flex items-center gap-2 w-full sm:w-auto">
                    <div class="relative w-full sm:w-56">
                        <i data-lucide="search" class="w-3.5 h-3.5 text-slate-400 absolute left-3 top-2.5"></i>
                        <input type="text" id="ia-sec-search" oninput="filterIaList()" placeholder="Buscar sesión..." class="w-full pl-8 pr-3 py-1.5 text-xs rounded-xl border border-slate-200 bg-slate-50 focus:bg-white focus:outline-none focus:ring-2 focus:ring-purple-600 transition">
                    </div>
                </div>
            </div>

            <!-- Master-Detail Container -->
            <div class="grid lg:grid-cols-12 gap-6 items-start">
                
                <!-- LISTADO LATERAL (MASTER) -->
                <div class="lg:col-span-4 bg-white rounded-3xl border border-slate-200 p-4 shadow-sm space-y-2">
                    <div class="text-[11px] font-black uppercase tracking-wider text-slate-400 px-2 pb-1 border-b border-slate-100 flex justify-between">
                        <span>Listado de Sesiones</span>
                        <span id="ia-count-badge">8 Sesiones</span>
                    </div>
                    <div id="ia-sessions-list" class="space-y-1.5 max-h-[560px] overflow-y-auto pr-1">
                        <!-- Rendered by JS -->
                    </div>
                </div>

                <!-- DETALLE DE LA SESIÓN SELECCIONADA (DETAIL) -->
                <div class="lg:col-span-8 bg-white rounded-3xl border border-slate-200 p-6 sm:p-8 shadow-sm space-y-6" id="ia-detail-card">
                    <!-- Rendered by JS -->
                </div>

            </div>

        </div>

        <!-- ================================================================== -->
        <!-- PESTAÑA 5: EVALUACIÓN -->
        <!-- ================================================================== -->
        <div id="tab-evaluacion" class="tab-content hidden space-y-8">
            
            <!-- Esquema Resumen -->
            <div class="bg-white rounded-3xl border border-slate-200 p-8 shadow-sm space-y-6">
                <div class="flex items-center justify-between border-b border-slate-200 pb-3">
                    <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                        <i data-lucide="check-square" class="w-5 h-5 text-purple-700"></i>
                        Esquema de Evaluación Institucional & Acreditación
                    </h3>
                    <span class="px-3 py-1 rounded-full bg-amber-100 text-amber-900 font-black text-xs">
                        Umbral Mínimo: 75 / 100 Puntos
                    </span>
                </div>
                <div class="grid md:grid-cols-4 gap-4 text-xs">
                    <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 text-center space-y-1">
                        <div class="text-slate-500 font-bold uppercase text-[10px]">Talleres & Retos Prácticos</div>
                        <div class="text-3xl font-black text-purple-900">40%</div>
                        <div class="text-slate-600">Resolución de Casos en Vivo</div>
                    </div>
                    <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 text-center space-y-1">
                        <div class="text-slate-500 font-bold uppercase text-[10px]">Proyecto Integrador</div>
                        <div class="text-3xl font-black text-emerald-700">25%</div>
                        <div class="text-slate-600">Asistente Personal Evaluado</div>
                    </div>
                    <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 text-center space-y-1">
                        <div class="text-slate-500 font-bold uppercase text-[10px]">Pensamiento Crítico</div>
                        <div class="text-3xl font-black text-blue-900">20%</div>
                        <div class="text-slate-600">Caza de Alucinaciones</div>
                    </div>
                    <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 text-center space-y-1">
                        <div class="text-slate-500 font-bold uppercase text-[10px]">Bitácora de Prompts & Ética</div>
                        <div class="text-3xl font-black text-amber-700">15%</div>
                        <div class="text-slate-600">Catálogo Personal y Privacidad</div>
                    </div>
                </div>
            </div>

            <!-- Rúbrica Terminal -->
            <div class="bg-white rounded-3xl border border-slate-200 p-8 shadow-sm space-y-6">
                <div class="flex flex-col sm:flex-row justify-between items-start sm:items-center gap-2 border-b border-slate-200 pb-4">
                    <div>
                        <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                            <i data-lucide="award" class="w-5 h-5 text-purple-700"></i>
                            Rúbrica Analítica Terminal de Evaluación Práctica
                        </h3>
                        <p class="text-slate-600 text-xs mt-1">
                            Evaluación objetiva del desempeño técnico con ponderaciones explícitas y umbral institucional de 75 puntos.
                        </p>
                    </div>
                    <span class="px-3 py-1.5 rounded-xl bg-amber-100 text-amber-900 font-extrabold text-xs">
                        Aprobación Mínima: ≥ 75 Pts
                    </span>
                </div>

                <div class="overflow-x-auto">
                    <table class="w-full text-xs text-left border-collapse">
                        <thead>
                            <tr class="bg-slate-100 text-slate-800 uppercase font-black border-b border-slate-200">
                                <th class="py-3 px-4 w-1/4">Criterio Evaluado</th>
                                <th class="py-3 px-4 w-20 text-center">Ponderación</th>
                                <th class="py-3 px-4 text-emerald-900">Desempeño Excelente (Trabajo Bien Hecho)</th>
                                <th class="py-3 px-4 text-rose-900">Desempeño Insuficiente (&lt; 75%)</th>
                            </tr>
                        </thead>
                        <tbody class="divide-y divide-slate-100">
                            {ia_rubric_rows}
                        </tbody>
                    </table>
                </div>
            </div>

            <!-- Bitácora de Prompts Card -->
            {ia_bitacora_card}

        </div>

    </main>

    <!-- Footer -->
    <footer class="bg-slate-950 text-slate-400 py-8 border-t border-slate-800 text-xs">
        <div class="max-w-7xl mx-auto px-4 sm:px-6 lg:px-8 flex flex-col sm:flex-row justify-between items-center gap-4 text-center sm:text-left">
            <div>
                <p class="font-bold text-slate-200">Fundación Kinal • Escuela Técnica Superior & Coordinación de Formación Continua</p>
                <p class="text-slate-500">Diseño Curricular e Instruccional 2026 | Estándares DQR 3-4 & Inteligencia Artificial Ética</p>
            </div>
            <div class="flex items-center gap-4 flex-wrap">
                <a href="index.html" class="hover:text-amber-400 transition">Catálogo de Cursos</a>
                <span class="text-slate-700">•</span>
                <a href="mecatronica.html" class="hover:text-amber-400 transition">Mecatrónica</a>
                <span class="text-slate-700">•</span>
                <a href="ciberseguridad.html" class="hover:text-amber-400 transition">Ciberseguridad</a>
                <span class="text-slate-700">•</span>
                <a href="automatizacion.html" class="hover:text-amber-400 transition">Automatización</a>
                <span class="text-slate-700">•</span>
                <a href="cableado_estructurado.html" class="hover:text-amber-400 transition">Cableado Estructurado</a>
            </div>
        </div>
    </footer>

    <!-- JavaScript Interactivity -->
    <script>
        const iaSessionsData = {ia_sessions_json};
        let currentIaSessionId = 1;

        function switchTab(tabId) {{
            document.querySelectorAll('.tab-content').forEach(el => el.classList.add('hidden'));
            document.querySelectorAll('.tab-btn').forEach(btn => {{
                btn.classList.remove('active');
                btn.classList.add('bg-slate-100', 'text-slate-700');
            }});

            const targetTab = document.getElementById('tab-' + tabId);
            const targetBtn = document.getElementById('btn-' + tabId);
            if (targetTab && targetBtn) {{
                targetTab.classList.remove('hidden');
                targetBtn.classList.add('active');
                targetBtn.classList.remove('bg-slate-100', 'text-slate-700');
            }}
            window.scrollTo({{ top: 0, behavior: 'smooth' }});
        }}

        function filterIaCategory(cat) {{
            document.querySelectorAll('#tab-temario button[id^="btn-ia-"]').forEach(btn => {{
                btn.className = "px-3.5 py-1.5 rounded-xl text-xs font-extrabold bg-slate-100 hover:bg-slate-200 text-slate-700 transition";
            }});

            const activeBtn = document.getElementById('btn-ia-' + cat);
            if (activeBtn) {{
                activeBtn.className = "px-3.5 py-1.5 rounded-xl text-xs font-extrabold bg-slate-900 text-white shadow-sm transition";
            }}

            const items = document.querySelectorAll('.content-item');
            items.forEach(it => {{
                if (cat === 'all' || it.getAttribute('data-category') === cat) {{
                    it.style.display = 'block';
                }} else {{
                    it.style.display = 'none';
                }}
            }});
        }}

        function renderIaSessionsList(filteredSessions) {{
            const container = document.getElementById('ia-sessions-list');
            container.innerHTML = '';

            filteredSessions.forEach(s => {{
                const item = document.createElement('div');
                item.className = `session-item p-3 rounded-2xl border border-slate-200 cursor-pointer transition flex items-center justify-between text-xs hover:border-purple-800 hover:bg-slate-50 ${{s.num === currentIaSessionId ? 'active' : 'bg-white'}}`;
                item.onclick = () => selectIaSession(s.num);
                
                item.innerHTML = `
                    <div class="space-y-0.5 pr-2">
                        <div class="flex items-center gap-1.5 font-bold">
                            <span class="badge-code px-2 py-0.5 rounded text-[10px] font-black bg-slate-100 text-slate-800">${{s.code}}</span>
                            <span class="line-clamp-1 text-slate-900 text-xs">${{s.title}}</span>
                        </div>
                        <div class="text-[10px] text-slate-500">${{s.semana}} • ${{s.dia}}</div>
                    </div>
                    <i data-lucide="chevron-right" class="w-4 h-4 text-slate-400 shrink-0"></i>
                `;
                container.appendChild(item);
            }});
            lucide.createIcons();
        }}

        function selectIaSession(num) {{
            currentIaSessionId = num;
            const s = iaSessionsData.find(item => item.num === num);
            if (!s) return;

            document.querySelectorAll('.session-item').forEach(el => {{
                el.classList.remove('active');
                el.classList.add('bg-white');
            }});

            renderIaSessionsList(iaSessionsData);

            const detailContainer = document.getElementById('ia-detail-card');
            detailContainer.innerHTML = `
                <div class="border-b border-slate-200 pb-5 space-y-2">
                    <div class="flex flex-wrap justify-between items-center gap-2">
                        <div class="flex items-center gap-2">
                            <span class="px-3 py-1 rounded-xl bg-slate-950 text-purple-400 font-black text-xs tracking-wider">${{s.code}}</span>
                            <span class="px-2.5 py-0.5 rounded-lg bg-purple-50 text-purple-900 font-bold text-xs border border-purple-200">${{s.module}}</span>
                        </div>
                        <span class="text-xs font-bold text-slate-500 bg-slate-100 px-3 py-1 rounded-lg">${{s.duration}}</span>
                    </div>
                    <h3 class="text-xl sm:text-2xl font-black text-slate-900 leading-snug">${{s.title}}</h3>
                    <div class="text-xs text-purple-800 font-bold pt-1">
                        ${{s.semana}} • ${{s.horario}} (Aula Virtual Microsoft Teams)
                    </div>
                </div>

                <div class="bg-purple-50/60 rounded-2xl p-4 border border-purple-200 space-y-1">
                    <div class="text-[11px] uppercase font-black tracking-wider text-purple-900 flex items-center gap-1.5">
                        <i data-lucide="target" class="w-3.5 h-3.5 text-purple-700"></i> Indicador de Logro / Objetivo Formativo
                    </div>
                    <p class="text-slate-800 text-xs leading-relaxed font-medium">${{s.objective}}</p>
                </div>

                <div class="space-y-4">
                    <h4 class="text-sm font-black text-slate-900 uppercase tracking-wide flex items-center gap-2 border-b border-slate-100 pb-2">
                        <i data-lucide="clock" class="w-4 h-4 text-purple-700"></i> Secuencia Didáctica en Tres Momentos
                    </h4>

                    <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 space-y-1.5">
                        <div class="flex items-center justify-between">
                            <span class="font-bold text-xs text-blue-950 flex items-center gap-2">
                                <span class="w-2.5 h-2.5 rounded-full bg-blue-600"></span> Fase 1: Apertura & Motivación
                            </span>
                            <span class="text-[11px] font-bold text-slate-500">15 Minutos</span>
                        </div>
                        <p class="text-slate-700 text-xs leading-relaxed pl-4.5">${{s.apertura}}</p>
                    </div>

                    <div class="p-4 rounded-2xl bg-purple-50/40 border border-purple-200 space-y-1.5">
                        <div class="flex items-center justify-between">
                            <span class="font-bold text-xs text-purple-950 flex items-center gap-2">
                                <span class="w-2.5 h-2.5 rounded-full bg-emerald-600"></span> Fase 2: Desarrollo & Taller Guiado en Vivo
                            </span>
                            <span class="text-[11px] font-bold text-emerald-800 bg-emerald-100 px-2 py-0.5 rounded">90 Minutos (Interacción Sincrónica)</span>
                        </div>
                        <p class="text-slate-700 text-xs leading-relaxed pl-4.5">${{s.desarrollo}}</p>
                    </div>

                    <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 space-y-1.5">
                        <div class="flex items-center justify-between">
                            <span class="font-bold text-xs text-blue-950 flex items-center gap-2">
                                <span class="w-2.5 h-2.5 rounded-full bg-purple-600"></span> Fase 3: Cierre, Verificación & Reto Semanal
                            </span>
                            <span class="text-[11px] font-bold text-slate-500">15 Minutos</span>
                        </div>
                        <p class="text-slate-700 text-xs leading-relaxed pl-4.5">${{s.cierre}}</p>
                    </div>
                </div>

                <div class="grid sm:grid-cols-2 gap-4 pt-2">
                    <div class="bg-emerald-50/60 rounded-2xl p-4 border border-emerald-200 space-y-1.5">
                        <div class="text-[11px] uppercase font-black tracking-wider text-emerald-900 flex items-center gap-1.5">
                            <i data-lucide="check-circle-2" class="w-3.5 h-3.5 text-emerald-700"></i> Evidencias & Producto Observable
                        </div>
                        <p class="text-slate-700 text-xs leading-relaxed">${{s.evidencias}}</p>
                    </div>

                    <div class="bg-slate-100 rounded-2xl p-4 border border-slate-200 space-y-1.5">
                        <div class="text-[11px] uppercase font-black tracking-wider text-slate-900 flex items-center gap-1.5">
                            <i data-lucide="cpu" class="w-3.5 h-3.5 text-purple-700"></i> Recursos Digitales & Plataformas
                        </div>
                        <p class="text-slate-700 text-xs leading-relaxed">${{s.equipamiento}}</p>
                    </div>
                </div>
            `;
            lucide.createIcons();
        }}

        function filterIaList() {{
            const q = document.getElementById('ia-sec-search').value.toLowerCase();
            const filtered = iaSessionsData.filter(s => 
                s.title.toLowerCase().includes(q) || 
                s.code.toLowerCase().includes(q) ||
                s.module.toLowerCase().includes(q) ||
                s.objective.toLowerCase().includes(q)
            );
            document.getElementById('ia-count-badge').textContent = `${{filtered.length}} Sesiones`;
            renderIaSessionsList(filtered);
        }}

        function goToSession(num) {{
            switchTab('secuencias');
            selectIaSession(num);
        }}

        document.addEventListener('DOMContentLoaded', () => {{
            lucide.createIcons();
            selectIaSession(1);
        }});
    </script>
</body>
</html>
"""

with open(os.path.join(WEB_DIR, "ia.html"), "w", encoding="utf-8") as f:
    f.write(ia_html)
print("Generado: Cursos/Web/ia.html")

# ==============================================================================
# 2. ACTUALIZACIÓN DEL CATÁLOGO CENTRAL (INDEX.HTML CON LOS 5 CURSOS)
# ==============================================================================
index_5_cursos_html = """<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Catálogo de Nuevos Cursos en Desarrollo | Ecosistema Kinal 2026</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <script src="https://unpkg.com/lucide@latest"></script>
    <style>
        @import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700;800&family=Space+Grotesk:wght@500;700;800;900&display=swap');
        body { font-family: 'Inter', sans-serif; background-color: #f8fafc; color: #0f172a; }
        h1, h2, h3, h4, .font-heading { font-family: 'Space Grotesk', sans-serif; }
        .hero-gradient {
            background: linear-gradient(135deg, #09172e 0%, #0f2d59 50%, #1e3a8a 100%);
        }
        .card-hover {
            transition: all 0.3s cubic-bezier(0.4, 0, 0.2, 1);
        }
        .card-hover:hover {
            transform: translateY(-4px);
            box-shadow: 0 20px 25px -5px rgba(0, 0, 0, 0.1), 0 10px 10px -5px rgba(0, 0, 0, 0.04);
        }
    </style>
</head>
<body class="antialiased min-h-screen flex flex-col">

    <!-- Top Institutional Header -->
    <header class="bg-slate-950 text-white border-b border-slate-800 sticky top-0 z-50 shadow-md">
        <div class="max-w-7xl mx-auto px-4 sm:px-6 lg:px-8 py-3 flex flex-wrap justify-between items-center gap-4">
            <div class="flex items-center gap-3">
                <a href="../../index.html" class="bg-gradient-to-r from-amber-500 to-amber-600 text-slate-950 font-black text-xs tracking-wider uppercase px-2.5 py-1 rounded shadow-sm hover:opacity-90 transition flex items-center gap-1">
                    <i data-lucide="arrow-left" class="w-3.5 h-3.5"></i> Portal Maestro
                </a>
                <span class="text-slate-600">|</span>
                <div class="text-slate-200 font-bold text-xs sm:text-sm flex items-center gap-2">
                    <i data-lucide="folder-kanban" class="w-4 h-4 text-blue-400"></i>
                    Formación Continua & Empleabilidad • Portafolio de 5 Cursos
                </div>
            </div>
            <nav class="flex items-center gap-2 text-xs">
                <a href="../../TSU/index.html" class="px-3 py-1.5 rounded-lg bg-slate-800/80 hover:bg-slate-700 text-slate-300 transition flex items-center gap-1">
                    <i data-lucide="graduation-cap" class="w-3.5 h-3.5 text-blue-400"></i> TSU (3er Año)
                </a>
                <a href="../../Revision_Temarios/index.html" class="px-3 py-1.5 rounded-lg bg-slate-800/80 hover:bg-slate-700 text-slate-300 transition flex items-center gap-1">
                    <i data-lucide="wrench" class="w-3.5 h-3.5 text-amber-400"></i> Especialidades Técnicas
                </a>
            </nav>
        </div>
    </header>

    <!-- Hero Section -->
    <section class="hero-gradient text-white py-14 px-4 sm:px-6 lg:px-8">
        <div class="max-w-6xl mx-auto text-center space-y-4">
            <div class="inline-flex items-center gap-2 px-3 py-1 rounded-full bg-blue-500/20 border border-blue-400/30 text-blue-300 text-xs font-semibold">
                <i data-lucide="sparkles" class="w-3.5 h-3.5 text-amber-400"></i> Catálogo Curricular & Didáctico Kinal 2026 • Marco Alemán DQR 3-5
            </div>
            <h1 class="text-3xl sm:text-5xl font-black tracking-tight leading-tight text-white">
                Portafolio de Nuevos Cursos Técnicos en Desarrollo
            </h1>
            <p class="text-slate-300 text-sm sm:text-base max-w-3xl mx-auto leading-relaxed">
                Plataforma interactiva para la exploración y gestión de los programas formativos diseñados para la reconversión laboral, especialización industrial y actualización tecnológica bajo los estándares institucionales de <strong>Fundación Kinal</strong>.
            </p>

            <!-- Metrics Summary Grid (5 CURSOS) -->
            <div class="grid grid-cols-2 md:grid-cols-4 gap-4 pt-4 max-w-4xl mx-auto text-left">
                <div class="bg-white/10 backdrop-blur rounded-2xl p-4 border border-white/10">
                    <div class="text-[11px] uppercase tracking-wider text-slate-300 font-bold">Cursos en Diseño</div>
                    <div class="text-3xl font-black text-white">5 Activos</div>
                    <div class="text-[11px] text-amber-300 font-medium">Meca • Ciber • Auto • Cable • IA</div>
                </div>
                <div class="bg-white/10 backdrop-blur rounded-2xl p-4 border border-white/10">
                    <div class="text-[11px] uppercase tracking-wider text-slate-300 font-bold">Carga Formativa Total</div>
                    <div class="text-3xl font-black text-white">386 Horas</div>
                    <div class="text-[11px] text-emerald-300 font-medium">90h + 80h + 120h + 80h + 16h</div>
                </div>
                <div class="bg-white/10 backdrop-blur rounded-2xl p-4 border border-white/10">
                    <div class="text-[11px] uppercase tracking-wider text-slate-300 font-bold">Microdiseño Didáctico</div>
                    <div class="text-3xl font-black text-purple-300">88 Sesiones</div>
                    <div class="text-[11px] text-purple-200 font-medium">Apertura • Desarrollo • Cierre</div>
                </div>
                <div class="bg-white/10 backdrop-blur rounded-2xl p-4 border border-white/10">
                    <div class="text-[11px] uppercase tracking-wider text-slate-300 font-bold">Exigencia Terminal</div>
                    <div class="text-3xl font-black text-amber-400">≥ 75 Pts</div>
                    <div class="text-[11px] text-slate-300 font-medium">Ideario del «Trabajo Bien Hecho»</div>
                </div>
            </div>
        </div>
    </section>

    <!-- Main Content: Course Cards & Hub -->
    <main class="max-w-7xl mx-auto px-4 sm:px-6 lg:px-8 py-12 space-y-12 flex-1">
        
        <div class="flex flex-col sm:flex-row justify-between items-start sm:items-end gap-4 border-b border-slate-200 pb-5">
            <div>
                <h2 class="text-2xl sm:text-3xl font-black text-slate-900 tracking-tight">
                    Cursos Disponibles para Consulta
                </h2>
                <p class="text-slate-600 text-sm mt-1">
                    Selecciona un curso para consultar su propuesta institucional, temario modular completo, dosificación detallada sesión por sesión y rúbricas.
                </p>
            </div>
            <div class="inline-flex items-center gap-2 px-3 py-1.5 rounded-lg bg-emerald-50 text-emerald-800 border border-emerald-200 text-xs font-bold">
                <i data-lucide="check-circle-2" class="w-4 h-4 text-emerald-600"></i> Documentación Word Oficial Sincronizada (15 Documentos)
            </div>
        </div>

        <!-- Course Cards Grid -->
        <div class="grid lg:grid-cols-2 gap-8">
            
            <!-- TARJETA 1: MECATRÓNICA -->
            <div class="bg-white rounded-3xl border border-slate-200 p-7 shadow-sm card-hover flex flex-col justify-between space-y-6">
                <div class="space-y-4">
                    <div class="flex flex-wrap justify-between items-start gap-2">
                        <span class="px-3 py-1 rounded-full bg-blue-50 text-blue-900 font-extrabold text-xs tracking-wide uppercase border border-blue-200 flex items-center gap-1.5">
                            <i data-lucide="cpu" class="w-3.5 h-3.5 text-blue-700"></i> Modalidad 100% Presencial
                        </span>
                        <span class="px-2.5 py-0.5 rounded-md bg-amber-50 text-amber-900 border border-amber-200 text-xs font-bold">
                            DQR Nivel 4 - 5
                        </span>
                    </div>

                    <div>
                        <h3 class="text-2xl font-black text-slate-900 hover:text-blue-900 transition leading-snug">
                            Mecatrónica Industrial y Fabricación Digital Aplicada
                        </h3>
                        <p class="text-slate-600 text-sm mt-2 leading-relaxed">
                            Mecánica aplicada, tolerancias dimensionales, prototipado rápido en cortadora láser CO2, manufactura aditiva FDM, fresado CNC, electroneumática, automatización con PLC Siemens S7-1200 y puesta en marcha de celdas industriales integradas.
                        </p>
                    </div>

                    <!-- Tech Badges -->
                    <div class="flex flex-wrap gap-1.5 pt-1">
                        <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">CAD 3D Paramétrico</span>
                        <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">Láser CO2</span>
                        <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">Impresión 3D FDM</span>
                        <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">Mecanizado CNC</span>
                        <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">PLC Siemens S7-1200</span>
                    </div>

                    <!-- Specific Course Metrics -->
                    <div class="grid grid-cols-3 gap-3 bg-slate-50 rounded-2xl p-3.5 border border-slate-200 text-center">
                        <div>
                            <div class="text-[10px] uppercase font-bold text-slate-500">Duración</div>
                            <div class="text-lg font-black text-blue-900">90 Horas</div>
                            <div class="text-[10px] text-slate-500">20 Sábados (4.5h)</div>
                        </div>
                        <div>
                            <div class="text-[10px] uppercase font-bold text-slate-500">Enfoque</div>
                            <div class="text-lg font-black text-emerald-700">75% Taller</div>
                            <div class="text-[10px] text-slate-500">25% Fundamentos</div>
                        </div>
                        <div>
                            <div class="text-[10px] uppercase font-bold text-slate-500">Proyecto</div>
                            <div class="text-lg font-black text-purple-900">Capstone</div>
                            <div class="text-[10px] text-slate-500">Celda Mecatrónica</div>
                        </div>
                    </div>
                </div>

                <!-- Action Buttons -->
                <div class="space-y-2.5 pt-2 border-t border-slate-100">
                    <a href="mecatronica.html" class="w-full inline-flex justify-center items-center gap-2 px-5 py-3 rounded-xl bg-blue-900 hover:bg-blue-800 text-white font-bold text-sm shadow-md transition">
                        <i data-lucide="eye" class="w-4 h-4"></i> Explorar Curso Web Completo
                    </a>
                    <div class="grid grid-cols-3 gap-2 text-center text-xs">
                        <a href="../Mecatrónica/Propuesta_Curso_Mecatronica_Kinal.docx" download class="p-2 rounded-lg bg-slate-100 hover:bg-slate-200 text-slate-700 font-medium transition flex items-center justify-center gap-1" title="Descargar Propuesta en Word">
                            <i data-lucide="file-text" class="w-3.5 h-3.5 text-blue-700"></i> Propuesta
                        </a>
                        <a href="../Mecatrónica/Temario_Curso_Mecatronica_Kinal.docx" download class="p-2 rounded-lg bg-slate-100 hover:bg-slate-200 text-slate-700 font-medium transition flex items-center justify-center gap-1" title="Descargar Temario en Word">
                            <i data-lucide="list-checks" class="w-3.5 h-3.5 text-emerald-700"></i> Temario
                        </a>
                        <a href="../Mecatrónica/Dosificacion_y_Secuencia_Didactica_Mecatronica_Kinal.docx" download class="p-2 rounded-lg bg-slate-100 hover:bg-slate-200 text-slate-700 font-medium transition flex items-center justify-center gap-1" title="Descargar Dosificación en Word">
                            <i data-lucide="calendar" class="w-3.5 h-3.5 text-amber-700"></i> Dosificación
                        </a>
                    </div>
                </div>
            </div>

            <!-- TARJETA 2: CIBERSEGURIDAD -->
            <div class="bg-white rounded-3xl border border-slate-200 p-7 shadow-sm card-hover flex flex-col justify-between space-y-6">
                <div class="space-y-4">
                    <div class="flex flex-wrap justify-between items-start gap-2">
                        <span class="px-3 py-1 rounded-full bg-indigo-50 text-indigo-900 font-extrabold text-xs tracking-wide uppercase border border-indigo-200 flex items-center gap-1.5">
                            <i data-lucide="shield-check" class="w-3.5 h-3.5 text-indigo-700"></i> Modalidad Híbrida (80/20)
                        </span>
                        <span class="px-2.5 py-0.5 rounded-md bg-amber-50 text-amber-900 border border-amber-200 text-xs font-bold">
                            DQR Nivel 4 - 5
                        </span>
                    </div>

                    <div>
                        <h3 class="text-2xl font-black text-slate-900 hover:text-indigo-900 transition leading-snug">
                            Ciberseguridad y Fundamentos de Seguridad de la Información
                        </h3>
                        <p class="text-slate-600 text-sm mt-2 leading-relaxed">
                            Protección de datos, identidad y accesos (IAM), comunicaciones seguras, gobierno, gestión de vulnerabilidades CVE/CVSS y cumplimiento de normativas internacionales mediante un Caso Práctico Transversal de Empresa Ficticia.
                        </p>
                    </div>

                    <!-- Tech Badges -->
                    <div class="flex flex-wrap gap-1.5 pt-1">
                        <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">Tríada CIA / ITIL</span>
                        <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">IAM / MFA / RBAC</span>
                        <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">Criptografía & PKI</span>
                        <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">CVE / CVSS</span>
                        <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">ISO 27001 & SOC 2</span>
                    </div>

                    <!-- Specific Course Metrics -->
                    <div class="grid grid-cols-3 gap-3 bg-slate-50 rounded-2xl p-3.5 border border-slate-200 text-center">
                        <div>
                            <div class="text-[10px] uppercase font-bold text-slate-500">Duración</div>
                            <div class="text-lg font-black text-indigo-900">80 Horas</div>
                            <div class="text-[10px] text-slate-500">20 Sesiones (4h c/u)</div>
                        </div>
                        <div>
                            <div class="text-[10px] uppercase font-bold text-slate-500">Estructura</div>
                            <div class="text-lg font-black text-emerald-700">80% Virt. / 20% Pres.</div>
                            <div class="text-[10px] text-slate-500">64h Online / 16h Sede</div>
                        </div>
                        <div>
                            <div class="text-[10px] uppercase font-bold text-slate-500">Entregable</div>
                            <div class="text-lg font-black text-purple-900">Expediente</div>
                            <div class="text-[10px] text-slate-500">Caso Transversal</div>
                        </div>
                    </div>
                </div>

                <!-- Action Buttons -->
                <div class="space-y-2.5 pt-2 border-t border-slate-100">
                    <a href="ciberseguridad.html" class="w-full inline-flex justify-center items-center gap-2 px-5 py-3 rounded-xl bg-indigo-900 hover:bg-indigo-800 text-white font-bold text-sm shadow-md transition">
                        <i data-lucide="eye" class="w-4 h-4"></i> Explorar Curso Web Completo
                    </a>
                    <div class="grid grid-cols-3 gap-2 text-center text-xs">
                        <a href="../Ciberseguridad/Propuesta_Curso_Ciberseguridad_Kinal.docx" download class="p-2 rounded-lg bg-slate-100 hover:bg-slate-200 text-slate-700 font-medium transition flex items-center justify-center gap-1" title="Descargar Propuesta en Word">
                            <i data-lucide="file-text" class="w-3.5 h-3.5 text-indigo-700"></i> Propuesta
                        </a>
                        <a href="../Ciberseguridad/Temario_Curso_Ciberseguridad_Kinal.docx" download class="p-2 rounded-lg bg-slate-100 hover:bg-slate-200 text-slate-700 font-medium transition flex items-center justify-center gap-1" title="Descargar Temario en Word">
                            <i data-lucide="list-checks" class="w-3.5 h-3.5 text-emerald-700"></i> Temario
                        </a>
                        <a href="../Ciberseguridad/Dosificacion_y_Secuencia_Didactica_Ciberseguridad.docx" download class="p-2 rounded-lg bg-slate-100 hover:bg-slate-200 text-slate-700 font-medium transition flex items-center justify-center gap-1" title="Descargar Dosificación en Word">
                            <i data-lucide="calendar" class="w-3.5 h-3.5 text-amber-700"></i> Dosificación
                        </a>
                    </div>
                </div>
            </div>

            <!-- TARJETA 3: AUTOMATIZACIÓN Y CONTROL ELÉCTRICO -->
            <div class="bg-white rounded-3xl border border-slate-200 p-7 shadow-sm card-hover flex flex-col justify-between space-y-6">
                <div class="space-y-4">
                    <div class="flex flex-wrap justify-between items-start gap-2">
                        <span class="px-3 py-1 rounded-full bg-amber-50 text-amber-900 font-extrabold text-xs tracking-wide uppercase border border-amber-200 flex items-center gap-1.5">
                            <i data-lucide="zap" class="w-3.5 h-3.5 text-amber-600"></i> Híbrida Asimétrica (80/20)
                        </span>
                        <span class="px-2.5 py-0.5 rounded-md bg-blue-50 text-blue-900 border border-blue-200 text-xs font-bold">
                            DQR Nivel 4 - 5
                        </span>
                    </div>

                    <div>
                        <h3 class="text-2xl font-black text-slate-900 hover:text-amber-700 transition leading-snug">
                            Automatización y Control Eléctrico Industrial con PLC y VFD
                        </h3>
                        <p class="text-slate-600 text-sm mt-2 leading-relaxed">
                            Control electromagnético de potencia, seguridad eléctrica bajo norma NFPA 70E / LOTO, parametrización de variadores de frecuencia comerciales, programación en TIA Portal para PLC Siemens S7-1200, comunicación PROFINET y diagnóstico metódico de averías en líneas continuas.
                        </p>
                    </div>

                    <!-- Tech Badges -->
                    <div class="flex flex-wrap gap-1.5 pt-1">
                        <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">NFPA 70E / LOTO</span>
                        <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">Contactores AC-3</span>
                        <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">Variadores VFD</span>
                        <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">PLC Siemens S7-1200</span>
                        <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">PROFINET / WinCC HMI</span>
                    </div>

                    <!-- Specific Course Metrics -->
                    <div class="grid grid-cols-3 gap-3 bg-slate-50 rounded-2xl p-3.5 border border-slate-200 text-center">
                        <div>
                            <div class="text-[10px] uppercase font-bold text-slate-500">Duración</div>
                            <div class="text-lg font-black text-amber-700">120 Horas</div>
                            <div class="text-[10px] text-slate-500">20 Sesiones (6h c/u)</div>
                        </div>
                        <div>
                            <div class="text-[10px] uppercase font-bold text-slate-500">Práctica</div>
                            <div class="text-lg font-black text-emerald-700">80% Banco</div>
                            <div class="text-[10px] text-slate-500">96h Taller / 24h LMS</div>
                        </div>
                        <div>
                            <div class="text-[10px] uppercase font-bold text-slate-500">Proyecto</div>
                            <div class="text-lg font-black text-blue-900">Certificación</div>
                            <div class="text-[10px] text-slate-500">Troubleshooting en Celda</div>
                        </div>
                    </div>
                </div>

                <!-- Action Buttons -->
                <div class="space-y-2.5 pt-2 border-t border-slate-100">
                    <a href="automatizacion.html" class="w-full inline-flex justify-center items-center gap-2 px-5 py-3 rounded-xl bg-slate-900 hover:bg-slate-800 text-amber-400 font-bold text-sm shadow-md transition">
                        <i data-lucide="eye" class="w-4 h-4"></i> Explorar Curso Web Completo
                    </a>
                    <div class="grid grid-cols-3 gap-2 text-center text-xs">
                        <a href="../Automatización/Propuesta_Curso_Automatizacion_Kinal.docx" download class="p-2 rounded-lg bg-slate-100 hover:bg-slate-200 text-slate-700 font-medium transition flex items-center justify-center gap-1" title="Descargar Propuesta en Word">
                            <i data-lucide="file-text" class="w-3.5 h-3.5 text-blue-700"></i> Propuesta
                        </a>
                        <a href="../Automatización/Temario_Curso_Automatizacion_Kinal.docx" download class="p-2 rounded-lg bg-slate-100 hover:bg-slate-200 text-slate-700 font-medium transition flex items-center justify-center gap-1" title="Descargar Temario en Word">
                            <i data-lucide="list-checks" class="w-3.5 h-3.5 text-emerald-700"></i> Temario
                        </a>
                        <a href="../Automatización/Dosificacion_y_Secuencia_Didactica_Automatizacion_Kinal.docx" download class="p-2 rounded-lg bg-slate-100 hover:bg-slate-200 text-slate-700 font-medium transition flex items-center justify-center gap-1" title="Descargar Dosificación en Word">
                            <i data-lucide="calendar" class="w-3.5 h-3.5 text-amber-700"></i> Dosificación
                        </a>
                    </div>
                </div>
            </div>

            <!-- TARJETA 4: CABLEADO ESTRUCTURADO Y FIBRA ÓPTICA -->
            <div class="bg-white rounded-3xl border border-slate-200 p-7 shadow-sm card-hover flex flex-col justify-between space-y-6">
                <div class="space-y-4">
                    <div class="flex flex-wrap justify-between items-start gap-2">
                        <span class="px-3 py-1 rounded-full bg-cyan-50 text-cyan-900 font-extrabold text-xs tracking-wide uppercase border border-cyan-200 flex items-center gap-1.5">
                            <i data-lucide="network" class="w-3.5 h-3.5 text-cyan-700"></i> Presencial con Apoyo Digital
                        </span>
                        <span class="px-2.5 py-0.5 rounded-md bg-amber-50 text-amber-900 border border-amber-200 text-xs font-bold">
                            DQR Nivel 4 - 5
                        </span>
                    </div>

                    <div>
                        <h3 class="text-2xl font-black text-slate-900 hover:text-cyan-700 transition leading-snug">
                            Cableado Estructurado y Redes de Cobre y Fibra Óptica
                        </h3>
                        <p class="text-slate-600 text-sm mt-2 leading-relaxed">
                            Canalizaciones físicas y cuartos de telecomunicaciones bajo ANSI/TIA-569, montaje de racks de 19'', conectorización Cat 6A, empalme de fibra óptica por fusión (&lt; 0.05 dB), certificación instrumental con escáneres Fluke Networks y rotulado TIA-606.
                        </p>
                    </div>

                    <!-- Tech Badges -->
                    <div class="flex flex-wrap gap-1.5 pt-1">
                        <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">ANSI/TIA-568-E</span>
                        <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">Cat 6A F/UTP</span>
                        <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">Fusión Fibra OS2/OM4</span>
                        <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">Fluke DSX Certificación</span>
                        <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">TIA-606 Rotulado</span>
                    </div>

                    <!-- Specific Course Metrics -->
                    <div class="grid grid-cols-3 gap-3 bg-slate-50 rounded-2xl p-3.5 border border-slate-200 text-center">
                        <div>
                            <div class="text-[10px] uppercase font-bold text-slate-500">Duración</div>
                            <div class="text-lg font-black text-cyan-800">80 Horas</div>
                            <div class="text-[10px] text-slate-500">20 Sesiones (4h c/u)</div>
                        </div>
                        <div>
                            <div class="text-[10px] uppercase font-bold text-slate-500">Horario</div>
                            <div class="text-lg font-black text-emerald-700">10 Semanas</div>
                            <div class="text-[10px] text-slate-500">Martes y Jueves</div>
                        </div>
                        <div>
                            <div class="text-[10px] uppercase font-bold text-slate-500">Entregable</div>
                            <div class="text-lg font-black text-purple-900">Dossier</div>
                            <div class="text-[10px] text-slate-500">Planos & Certificados</div>
                        </div>
                    </div>
                </div>

                <!-- Action Buttons -->
                <div class="space-y-2.5 pt-2 border-t border-slate-100">
                    <a href="cableado_estructurado.html" class="w-full inline-flex justify-center items-center gap-2 px-5 py-3 rounded-xl bg-cyan-900 hover:bg-cyan-800 text-white font-bold text-sm shadow-md transition">
                        <i data-lucide="eye" class="w-4 h-4"></i> Explorar Curso Web Completo
                    </a>
                    <div class="grid grid-cols-3 gap-2 text-center text-xs">
                        <a href="../Cableado_Estructurado/Propuesta_Curso_Cableado_Estructurado_Kinal.docx" download class="p-2 rounded-lg bg-slate-100 hover:bg-slate-200 text-slate-700 font-medium transition flex items-center justify-center gap-1" title="Descargar Propuesta en Word">
                            <i data-lucide="file-text" class="w-3.5 h-3.5 text-blue-700"></i> Propuesta
                        </a>
                        <a href="../Cableado_Estructurado/Temario_Curso_Cableado_Estructurado_Kinal.docx" download class="p-2 rounded-lg bg-slate-100 hover:bg-slate-200 text-slate-700 font-medium transition flex items-center justify-center gap-1" title="Descargar Temario en Word">
                            <i data-lucide="list-checks" class="w-3.5 h-3.5 text-emerald-700"></i> Temario
                        </a>
                        <a href="../Cableado_Estructurado/Dosificacion_y_Secuencia_Didactica_Cableado_Estructurado_Kinal.docx" download class="p-2 rounded-lg bg-slate-100 hover:bg-slate-200 text-slate-700 font-medium transition flex items-center justify-center gap-1" title="Descargar Dosificación en Word">
                            <i data-lucide="calendar" class="w-3.5 h-3.5 text-amber-700"></i> Dosificación
                        </a>
                    </div>
                </div>
            </div>

            <!-- TARJETA 5: INTELIGENCIA ARTIFICIAL APLICADA (DESTACADO / NEW) -->
            <div class="bg-white rounded-3xl border-2 border-purple-300 p-7 shadow-md card-hover flex flex-col justify-between space-y-6 lg:col-span-2 bg-gradient-to-br from-white via-purple-50/20 to-indigo-50/30">
                <div class="space-y-4">
                    <div class="flex flex-wrap justify-between items-start gap-2">
                        <div class="flex items-center gap-2">
                            <span class="px-3 py-1 rounded-full bg-purple-100 text-purple-900 font-extrabold text-xs tracking-wide uppercase border border-purple-200 flex items-center gap-1.5">
                                <i data-lucide="sparkles" class="w-3.5 h-3.5 text-purple-700"></i> 100% Virtual Sincrónica (Teams)
                            </span>
                            <span class="px-2.5 py-0.5 rounded-full bg-emerald-100 text-emerald-800 font-black text-[10px] uppercase">
                                ¡Nuevo Programa 2026!
                            </span>
                        </div>
                        <span class="px-2.5 py-0.5 rounded-md bg-purple-50 text-purple-900 border border-purple-200 text-xs font-bold">
                            DQR Nivel 3 - 4
                        </span>
                    </div>

                    <div class="grid lg:grid-cols-3 gap-6 items-center">
                        <div class="lg:col-span-2">
                            <h3 class="text-2xl sm:text-3xl font-black text-slate-900 hover:text-purple-900 transition leading-snug">
                                Inteligencia Artificial Aplicada, Ingeniería de Prompts y Productividad Ética
                            </h3>
                            <p class="text-slate-600 text-sm mt-2 leading-relaxed">
                                Formulación de prompts estructurados (Método RC-TRF), automatización de correspondencia laboral y minutas, síntesis de documentos extensos (PDFs), extracción de datos a tablas de Excel, caza de alucinaciones y diseño de un Asistente Personal de Productividad centrado en la persona.
                            </p>

                            <!-- Tech Badges -->
                            <div class="flex flex-wrap gap-1.5 pt-3">
                                <span class="px-2.5 py-1 rounded-lg bg-purple-100/70 text-purple-900 text-xs font-semibold">Método RC-TRF</span>
                                <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">ChatGPT & GPT-4o mini</span>
                                <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">Google Gemini</span>
                                <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">Microsoft Copilot</span>
                                <span class="px-2.5 py-1 rounded-lg bg-slate-100 text-slate-700 text-xs font-semibold">Claude (Anthropic)</span>
                                <span class="px-2.5 py-1 rounded-lg bg-emerald-100 text-emerald-900 text-xs font-semibold">Caza de Alucinaciones</span>
                                <span class="px-2.5 py-1 rounded-lg bg-blue-100 text-blue-900 text-xs font-semibold">Privacidad & Datos</span>
                            </div>
                        </div>

                        <!-- Specific Course Metrics -->
                        <div class="grid grid-cols-3 lg:grid-cols-1 gap-2.5 bg-white/80 backdrop-blur rounded-2xl p-4 border border-purple-200 text-center shadow-sm">
                            <div class="py-1">
                                <div class="text-[10px] uppercase font-bold text-slate-500">Duración & Formato</div>
                                <div class="text-lg font-black text-purple-900">16 Horas (1 Mes)</div>
                                <div class="text-[10px] text-slate-500">8 Sesiones Nocturnas (2h c/u)</div>
                            </div>
                            <div class="py-1 border-t lg:border-t-0 border-slate-100">
                                <div class="text-[10px] uppercase font-bold text-slate-500">Horario de Clases</div>
                                <div class="text-base font-black text-slate-800">Martes y Jueves</div>
                                <div class="text-[10px] text-slate-500">19:00 a 21:00 hrs (Virtual)</div>
                            </div>
                            <div class="py-1 border-t border-slate-100">
                                <div class="text-[10px] uppercase font-bold text-slate-500">Proyecto Terminal</div>
                                <div class="text-base font-black text-emerald-700">Asistente Personal</div>
                                <div class="text-[10px] text-slate-500">Dossier de Prompts Evaluado</div>
                            </div>
                        </div>
                    </div>
                </div>

                <!-- Action Buttons -->
                <div class="space-y-2.5 pt-2 border-t border-purple-200">
                    <a href="ia.html" class="w-full inline-flex justify-center items-center gap-2 px-5 py-3 rounded-xl bg-purple-900 hover:bg-purple-800 text-white font-bold text-sm shadow-md transition">
                        <i data-lucide="eye" class="w-4 h-4"></i> Explorar Curso Web Completo de Inteligencia Artificial
                    </a>
                    <div class="grid grid-cols-3 gap-2 text-center text-xs">
                        <a href="../Inteligencia_Artificial/Propuesta_Curso_IA_Kinal.docx" download class="p-2 rounded-lg bg-slate-100 hover:bg-slate-200 text-slate-700 font-medium transition flex items-center justify-center gap-1" title="Descargar Propuesta en Word">
                            <i data-lucide="file-text" class="w-3.5 h-3.5 text-purple-700"></i> Propuesta
                        </a>
                        <a href="../Inteligencia_Artificial/Temario_Curso_IA_Kinal.docx" download class="p-2 rounded-lg bg-slate-100 hover:bg-slate-200 text-slate-700 font-medium transition flex items-center justify-center gap-1" title="Descargar Temario en Word">
                            <i data-lucide="list-checks" class="w-3.5 h-3.5 text-emerald-700"></i> Temario
                        </a>
                        <a href="../Inteligencia_Artificial/Dosificacion_y_Secuencia_Didactica_IA_Kinal.docx" download class="p-2 rounded-lg bg-slate-100 hover:bg-slate-200 text-slate-700 font-medium transition flex items-center justify-center gap-1" title="Descargar Dosificación en Word">
                            <i data-lucide="calendar" class="w-3.5 h-3.5 text-amber-700"></i> Dosificación
                        </a>
                    </div>
                </div>
            </div>

        </div>

        <!-- Comparative Matrix (5 Cursos) -->
        <div class="bg-white rounded-3xl border border-slate-200 p-8 shadow-sm space-y-6">
            <div class="space-y-1">
                <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                    <i data-lucide="sliders-horizontal" class="w-5 h-5 text-blue-700"></i>
                    Matriz Comparativa Ejecutiva del Portafolio de 5 Cursos
                </h3>
                <p class="text-slate-600 text-xs">
                    Comparación transversal de parámetros metodológicos, temporales y evaluativos entre todas las especialidades formativas de Fundación Kinal.
                </p>
            </div>

            <div class="overflow-x-auto">
                <table class="w-full text-xs text-left border-collapse">
                    <thead>
                        <tr class="bg-slate-100 text-slate-700 uppercase font-black tracking-wider border-b border-slate-200">
                            <th class="py-3 px-3">Dimensión Curricular</th>
                            <th class="py-3 px-3 text-blue-900">Mecatrónica</th>
                            <th class="py-3 px-3 text-indigo-900">Ciberseguridad</th>
                            <th class="py-3 px-3 text-amber-800">Automatización</th>
                            <th class="py-3 px-3 text-cyan-800">Cableado Redes</th>
                            <th class="py-3 px-3 text-purple-900 bg-purple-50/50">IA Aplicada & Prompts</th>
                        </tr>
                    </thead>
                    <tbody class="divide-y divide-slate-100 text-slate-700">
                        <tr>
                            <td class="py-3 px-3 font-bold bg-slate-50">Carga Horaria</td>
                            <td class="py-3 px-3">90h (20 sáb. • 4.5h)</td>
                            <td class="py-3 px-3">80h (20 ses. • 4h)</td>
                            <td class="py-3 px-3">120h (20 ses. • 6h)</td>
                            <td class="py-3 px-3">80h (20 ses. • 4h)</td>
                            <td class="py-3 px-3 font-bold text-purple-900 bg-purple-50/30">16h (8 ses. nocturnas • 2h)</td>
                        </tr>
                        <tr>
                            <td class="py-3 px-3 font-bold bg-slate-50">Modalidad</td>
                            <td class="py-3 px-3"><span class="px-2 py-0.5 rounded bg-blue-100 text-blue-800 font-bold">100% Presencial</span></td>
                            <td class="py-3 px-3"><span class="px-2 py-0.5 rounded bg-indigo-100 text-indigo-800 font-bold">Híbrida (80/20)</span></td>
                            <td class="py-3 px-3"><span class="px-2 py-0.5 rounded bg-amber-100 text-amber-800 font-bold">Híbrida (80/20)</span></td>
                            <td class="py-3 px-3"><span class="px-2 py-0.5 rounded bg-cyan-100 text-cyan-800 font-bold">Presencial Taller</span></td>
                            <td class="py-3 px-3 bg-purple-50/30"><span class="px-2 py-0.5 rounded bg-purple-100 text-purple-900 font-bold">100% Virtual Síncrona</span></td>
                        </tr>
                        <tr>
                            <td class="py-3 px-3 font-bold bg-slate-50">Módulos</td>
                            <td class="py-3 px-3">5 Módulos temáticos</td>
                            <td class="py-3 px-3">8 Módulos progresivos</td>
                            <td class="py-3 px-3">4 Módulos de 30h</td>
                            <td class="py-3 px-3">4 Módulos de 20h</td>
                            <td class="py-3 px-3 bg-purple-50/30">4 Módulos semanales (4h c/u)</td>
                        </tr>
                        <tr>
                            <td class="py-3 px-3 font-bold bg-slate-50">Ecosistema & Equipos</td>
                            <td class="py-3 px-3">Láser CO2, FDM, CNC, Festo, S7-1200</td>
                            <td class="py-3 px-3">Linux, Wireshark, SIEM, NIST NVD</td>
                            <td class="py-3 px-3">Bancos potencia, VFDs, S7-1200, KTP HMI</td>
                            <td class="py-3 px-3">Racks 42U, Fusionadora, Fluke DSX</td>
                            <td class="py-3 px-3 bg-purple-50/30">ChatGPT, Gemini, Copilot, Claude, Teams</td>
                        </tr>
                        <tr>
                            <td class="py-3 px-3 font-bold bg-slate-50">Entregable Terminal</td>
                            <td class="py-3 px-3 font-bold">Celda Mecatrónica Funcional</td>
                            <td class="py-3 px-3 font-bold">Expediente Caso Transversal</td>
                            <td class="py-3 px-3 font-bold">Celda Operativa + Troubleshooting</td>
                            <td class="py-3 px-3 font-bold">Certificación Rack + As-Built</td>
                            <td class="py-3 px-3 font-bold text-purple-900 bg-purple-50/30">Asistente Personal + Dossier Prompts</td>
                        </tr>
                        <tr>
                            <td class="py-3 px-3 font-bold bg-slate-50">Nota Mínima</td>
                            <td class="py-3 px-3 font-bold text-amber-700">75 / 100 Pts</td>
                            <td class="py-3 px-3 font-bold text-amber-700">75 / 100 Pts</td>
                            <td class="py-3 px-3 font-bold text-amber-700">75 / 100 Pts</td>
                            <td class="py-3 px-3 font-bold text-amber-700">75 / 100 Pts</td>
                            <td class="py-3 px-3 font-bold text-amber-700 bg-purple-50/30">75 / 100 Pts</td>
                        </tr>
                    </tbody>
                </table>
            </div>
        </div>

        <!-- Institutional Quotes & Philosophy Section -->
        <div class="grid md:grid-cols-2 gap-6 text-xs">
            <div class="bg-blue-50 border-l-4 border-blue-900 p-5 rounded-r-2xl space-y-2">
                <div class="font-extrabold text-blue-900 uppercase tracking-wide">Misión Institucional de Fundación Kinal</div>
                <p class="text-slate-700 italic leading-relaxed">
                    «Formar a jóvenes y adultos a través de una educación integral, con énfasis en las áreas técnicas y tecnológicas, influyendo positivamente en su trabajo, su familia y la sociedad».
                </p>
            </div>
            <div class="bg-amber-50 border-l-4 border-amber-600 p-5 rounded-r-2xl space-y-2">
                <div class="font-extrabold text-amber-900 uppercase tracking-wide">Valores Nucleares de Kinal</div>
                <p class="text-slate-700 italic leading-relaxed">
                    «Visión cristiana de la vida. Respeto a la dignidad de la persona. Espíritu de servicio. Trabajo bien hecho. Libertad personal responsable».
                </p>
            </div>
        </div>

    </main>

    <!-- Footer -->
    <footer class="bg-slate-950 text-slate-400 py-8 border-t border-slate-800 text-xs">
        <div class="max-w-7xl mx-auto px-4 sm:px-6 lg:px-8 flex flex-col sm:flex-row justify-between items-center gap-4 text-center sm:text-left">
            <div>
                <p class="font-bold text-slate-200">Fundación Kinal • Dirección Académica & Escuela Técnica Superior</p>
                <p class="text-slate-500">Diseño Curricular, Didáctico e Instruccional 2026 | Sistema Dual & DQR 3-5</p>
            </div>
            <div class="flex items-center gap-3 flex-wrap">
                <a href="../../index.html" class="hover:text-amber-400 transition">Portal Maestro</a>
                <span class="text-slate-700">•</span>
                <a href="mecatronica.html" class="hover:text-amber-400 transition">Mecatrónica</a>
                <span class="text-slate-700">•</span>
                <a href="ciberseguridad.html" class="hover:text-amber-400 transition">Ciberseguridad</a>
                <span class="text-slate-700">•</span>
                <a href="automatizacion.html" class="hover:text-amber-400 transition">Automatización</a>
                <span class="text-slate-700">•</span>
                <a href="cableado_estructurado.html" class="hover:text-amber-400 transition">Cableado Redes</a>
                <span class="text-slate-700">•</span>
                <a href="ia.html" class="hover:text-purple-400 transition">Inteligencia Artificial</a>
            </div>
        </div>
    </footer>

    <script>
        lucide.createIcons();
    </script>
</body>
</html>
"""

with open(os.path.join(WEB_DIR, "index.html"), "w", encoding="utf-8") as f:
    f.write(index_5_cursos_html)
print("Actualizado: Cursos/Web/index.html con los 5 cursos.")

# ==============================================================================
# 3. ACTUALIZACIÓN DE FOOTERS EN PÁGINAS EXISTENTES PARA INCLUIR IA.HTML
# ==============================================================================
existing_pages = ["mecatronica.html", "ciberseguridad.html", "automatizacion.html", "cableado_estructurado.html"]
for p in existing_pages:
    fpath = os.path.join(WEB_DIR, p)
    if os.path.exists(fpath):
        with open(fpath, "r", encoding="utf-8") as f:
            content = f.read()
        
        # Check if ia.html is already in footer
        if 'ia.html' not in content:
            # Append before closing of the footer flex items
            content = content.replace(
                '</div>\n        </div>\n    </footer>',
                '<span class="text-slate-700">•</span>\n                <a href="ia.html" class="hover:text-purple-400 transition">Inteligencia Artificial</a>\n            </div>\n        </div>\n    </footer>'
            )
            with open(fpath, "w", encoding="utf-8") as f:
                f.write(content)
            print(f"Footer actualizado en: {p}")
