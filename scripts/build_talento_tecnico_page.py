import os
import re
import html
import docx
import json

WORKSPACE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CURSOS_DIR = os.path.join(WORKSPACE_DIR, "Cursos")
TALENTO_DIR = os.path.join(CURSOS_DIR, "Talento_Tecnico")

def clean_html(text):
    if not text:
        return ""
    return html.escape(text).replace('\n', '<br>')

# 1. Parsear los 3 archivos Word oficiales
prop_doc = docx.Document(os.path.join(TALENTO_DIR, "Propuesta_Diplomado_Talento_Tecnico_Kinal.docx"))
tem_doc = docx.Document(os.path.join(TALENTO_DIR, "Temario_Diplomado_Talento_Tecnico_Kinal.docx"))
dos_doc = docx.Document(os.path.join(TALENTO_DIR, "Dosificacion_y_Secuencia_Didactica_Talento_Tecnico_Kinal.docx"))

# Ficha técnica oficial
ficha_rows = []
for r in prop_doc.tables[0].rows:
    ficha_rows.append([c.text.strip() for c in r.cells])

# Dosificación cronológica (Tabla 1)
cron_rows = []
for r in dos_doc.tables[1].rows[1:]:
    cron_rows.append([c.text.strip() for c in r.cells])

# Sesiones didácticas (Tablas 2 a 1+num_sessions)
sessions = []
num_sessions = len(cron_rows)
for idx in range(2, 2 + num_sessions):
    t = dos_doc.tables[idx]
    s_num = idx - 1
    t1_row = cron_rows[s_num - 1] if (s_num - 1) < len(cron_rows) else ["", "", "", ""]
    
    mod_dur = t.rows[0].cells[1].text.strip()
    obj = t.rows[1].cells[1].text.strip()
    ap = t.rows[2].cells[1].text.strip()
    des = t.rows[3].cells[1].text.strip()
    cie = t.rows[4].cells[1].text.strip()
    eq = t.rows[5].cells[1].text.strip()
    evi = t.rows[6].cells[1].text.strip()
    
    enc_name = t1_row[0] if len(t1_row) > 0 else f"Sesión {s_num}"
    mod_name = t1_row[1] if len(t1_row) > 1 else f"Módulo {s_num}"
    title_short = t1_row[2] if len(t1_row) > 2 else f"Sesión {s_num}"
    prod_obs = t1_row[3] if len(t1_row) > 3 else evi
    
    if "Presencial" in enc_name:
        dur_str = "2 Horas (120 min)"
        tipo_str = "Presencial en Campus Kinal"
        semana_str = "Sesión Presencial Inicial" if s_num == 1 else "Sesión Presencial de Cierre"
        horario_str = "Jornada Presencial Especial (Campus Kinal)"
    else:
        dur_str = "2 Horas (120 min)"
        tipo_str = "Virtual en Microsoft Teams"
        semana_str = f"Sesión Virtual {s_num - 1}"
        horario_str = "Sesión Virtual Sincrónica (Microsoft Teams)"

    sessions.append({
        "num": s_num,
        "code": f"S-{s_num:02d}",
        "mod_num": s_num,
        "module": mod_name,
        "title": title_short,
        "duration": dur_str,
        "tipo": tipo_str,
        "semana": semana_str,
        "horario": horario_str,
        "objective": obj,
        "apertura": ap,
        "desarrollo": des,
        "cierre": cie,
        "evidencias": evi,
        "equipamiento": eq,
        "producto_obs": prod_obs
    })

# Rúbrica analítica de acreditación (búsqueda dinámica)
rubrica_rows = []
for t in dos_doc.tables:
    if len(t.rows) > 0 and "Actividad Formativa" in t.rows[0].cells[0].text:
        for r in t.rows[1:]:
            rubrica_rows.append([c.text.strip() for c in r.cells])
        break

# Componentes de evaluación (4 Actividades Formativas Evaluables — Total: 100 Puntos)
eval_prop_rows = [
    ["Módulo 1: Factor Humano & Convivencia", "25 Pts", "Protocolo de respeto, seguridad psicológica y canales de reporte sin temor en planta"],
    ["Módulo 2: Manejo de Personal & Micro-Charlas", "25 Pts", "Descriptor por competencias demostrables, onboarding 30-60-90 y 3 micro-charlas"],
    ["Módulo 3: Validación Técnica Superior TSU", "25 Pts", "Prueba práctica en banco y validación técnica bajo el estándar del Técnico Superior"],
    ["Módulo 4: Sistematización & Acuerdos", "25 Pts", "Sustentación de la ficha diagnóstica empresarial y firma de acuerdos de vinculación"]
]

# Temario analítico estructurado en los 4 Módulos Oficiales (Un módulo por sesión de 2 horas)
temario_modules = [
    {
        "num": 1,
        "semana": "Sesión Presencial 1 (Campus Kinal)",
        "title": "Módulo 1: Apertura Institucional, Liderazgo y Factor Humano en Cuadrillas de Trabajo",
        "hours": "2 Horas (Campus Kinal)",
        "logro": "Conocer la visión formativa integral de Kinal e implementar la cultura del Factor Humano en áreas operativas: respeto a la dignidad de la persona, erradicación tajante del maltrato verbal y fomento de seguridad psicológica.",
        "eval_act": "Código de Convivencia Operativa y Protocolo de Seguridad Psicológica y Cero Maltrato (25 Pts)",
        "topics": [
            {
                "num": 1,
                "title": "Apertura Institucional y Presentación de la Iniciativa de Vinculación",
                "cat": "vinculacion",
                "cat_name": "🤝 Vinculación Kinal",
                "cat_badge": "bg-blue-100 text-blue-900 border border-blue-200",
                "desc": "Propósito institucional: generar un espacio de diálogo con las empresas participantes para identificar oportunidades de mejora y conocer las competencias técnicas y humanas que demanda el sector productivo."
            },
            {
                "num": 2,
                "title": "La Filosofía del Factor Humano de Fundación Kinal y sus 4 Pilares",
                "cat": "factor_humano",
                "cat_name": "⭐ Factor Humano Kinal",
                "cat_badge": "bg-amber-100 text-amber-900 border border-amber-200",
                "desc": "La persona en el centro de los procesos productivos: Respeto a la Dignidad, El Trabajo Bien Hecho (orden 5S, custodia de maquinaria y rigor profesional), Espíritu de Servicio y Libertad Responsable."
            },
            {
                "num": 3,
                "title": "Liderazgo Humanista y Seguridad Psicológica en Planta",
                "cat": "liderazgo",
                "cat_name": "👥 Liderazgo Operativo",
                "cat_badge": "bg-indigo-100 text-indigo-900 border border-indigo-200",
                "desc": "Superación definitiva de la cultura del capataz y maltrato verbal. Creación de seguridad psicológica para que los operarios reporten averías e incidentes inmediatamente sin temor a represalias destructivas."
            }
        ]
    },
    {
        "num": 2,
        "semana": "Sesión Virtual 1 (Microsoft Teams)",
        "title": "Módulo 2: Manejo de Personal Técnico y Desarrollo de Mandos Medios en Planta",
        "hours": "2 Horas (Microsoft Teams)",
        "logro": "Aplicar metodologías contemporáneas para la atracción y selección por competencias demostrables, estructurar planes de onboarding técnico a 30-60-90 días y equipar a mandos medios con feedback formativo y micro-charlas de valores vinculadas a KPIs.",
        "eval_act": "Descriptor de Puesto Técnico por Competencias, Plan Onboarding y Catálogo de 3 Micro-Charlas (25 Pts)",
        "topics": [
            {
                "num": 4,
                "title": "Atracción y Reclutamiento Efectivo de Talento Operativo",
                "cat": "gestion",
                "cat_name": "📋 Selección & Onboarding",
                "cat_badge": "bg-emerald-100 text-emerald-900 border border-emerald-200",
                "desc": "Por qué fracasan los descriptores tradicionales. Matrices de selección por competencias demostrables (saber conceptual, destreza manual e instrumental y actitud de seguridad operativa)."
            },
            {
                "num": 5,
                "title": "Inducción Técnica Acelerada (30-60-90 Días) y Reducción de Rotación",
                "cat": "gestion",
                "cat_name": "📈 Manejo de Personal",
                "cat_badge": "bg-purple-100 text-purple-900 border border-purple-200",
                "desc": "Costos ocultos de la rotación temprana en plantas. Procesos de onboarding técnico paso a paso: conocimiento de maquinaria, procedimientos estándar y rutas de crecimiento operativo."
            },
            {
                "num": 6,
                "title": "Supervisión de Primera Línea y 'Micro-Charlas de Factor Humano'",
                "cat": "liderazgo",
                "cat_name": "👥 Mandos Medios",
                "cat_badge": "bg-indigo-100 text-indigo-900 border border-indigo-200",
                "desc": "El supervisor como puente clave entre la gerencia y el piso de planta. Retroalimentación asertiva sin humillar y conducción de micro-charlas de 10 min de valores vinculadas a KPIs (OEE y SSO)."
            }
        ]
    },
    {
        "num": 3,
        "semana": "Sesión Virtual 2 (Microsoft Teams)",
        "title": "Módulo 3: Fortalecimiento y Validación de Competencias: Nivel Técnico Superior (TSU)",
        "hours": "2 Horas (Microsoft Teams)",
        "logro": "Estructurar matrices de competencias operativas observables (saber, hacer y ser), diseñar pruebas de evaluación práctica en puesto y articular el nivel Técnico Superior (TSU) como el mecanismo institucional de validación técnica y respaldo objetivo para la industria.",
        "eval_act": "Guía de Evaluación Práctica en Puesto y Matriz de Validación Técnica bajo Estándar TSU (25 Pts)",
        "topics": [
            {
                "num": 7,
                "title": "Mapeo y Estandarización de Competencias Operativas de Planta",
                "cat": "competencias",
                "cat_name": "🎯 Competencias Técnicas",
                "cat_badge": "bg-purple-100 text-purple-900 border border-purple-200",
                "desc": "Definición y estructuración de competencias de planta bajo marcos internacionales (DQR). Desglose tridimensional: saber analítico/normativo, saber hacer instrumental y saber ser ético/operativo."
            },
            {
                "num": 8,
                "title": "El Técnico Superior Universitario (TSU) como Mecanismo de Validación y Respaldo Industrial",
                "cat": "competencias",
                "cat_name": "🎓 Validación Técnica TSU",
                "cat_badge": "bg-blue-100 text-blue-900 border border-blue-200",
                "desc": "El TSU como mecanismo riguroso de validación de competencias técnicas y respaldo objetivo para la industria, destacando la relevancia de las competencias a nivel superior en el área técnica para garantizar la fiabilidad operativa y la supervisión rigurosa."
            },
            {
                "num": 9,
                "title": "Evaluación Práctica en Puesto y Transferencia de Habilidades",
                "cat": "evaluacion",
                "cat_name": "🛠️ Evaluación en Puesto",
                "cat_badge": "bg-emerald-100 text-emerald-900 border border-emerald-200",
                "desc": "Diseño de pruebas prácticas en banco de trabajo (lectura de planos, calibración y diagnóstico). Esquemas de semilleros internos y el rol del mentor técnico (Ausbilder) en planta."
            }
        ]
    },
    {
        "num": 4,
        "semana": "Sesión Presencial 2 (Campus Kinal)",
        "title": "Módulo 4: Sistematización de Demandas Industriales, Acuerdos de Vinculación y Clausura",
        "hours": "2 Horas (Campus Kinal)",
        "logro": "Participar en la mesa redonda plenaria de demandas formativas de la industria, formalizar acuerdos de cooperación institucional y recibir el diploma de participación individual y reconocimiento corporativo de Kinal.",
        "eval_act": "Sustentación Plenaria de la Ficha Diagnóstica y Firma de Acuerdos de Vinculación (25 Pts)",
        "topics": [
            {
                "num": 10,
                "title": "Mesa Redonda Plenaria: Sistematización de Demandas Actuales de la Industria",
                "cat": "vinculacion",
                "cat_name": "🤝 Diálogo Industria-Kinal",
                "cat_badge": "bg-blue-100 text-blue-900 border border-blue-200",
                "desc": "Consolidación de las matrices diagnósticas: qué competencias técnicas y humanas demanda prioritariamente el sector productivo para retroalimentar la formación en carreras técnicas y del nivel Técnico Superior."
            },
            {
                "num": 11,
                "title": "Acuerdos de Vinculación y Rutas de Cooperación Kinal-Empresas",
                "cat": "vinculacion",
                "cat_name": "📑 Acuerdos Escuela-Empresa",
                "cat_badge": "bg-emerald-100 text-emerald-900 border border-emerald-200",
                "desc": "Acceso preferente a la bolsa de egresados, programas de formación dual y certificación de competencias laborales de colaboradores de planta en Kinal."
            },
            {
                "num": 12,
                "title": "Acto Solemne de Clausura: Diplomas de Participación y Reconocimientos Institucionales",
                "cat": "inmersion",
                "cat_name": "🎓 Entrega de Diplomas & Clausura",
                "cat_badge": "bg-amber-100 text-amber-900 border border-amber-200",
                "desc": "Entrega protocolaria de Diplomas de Participación Individual a los representantes empresariales y Reconocimientos Institucionales a empresas aliadas, seguido de brindis de honor."
            }
        ]
    }
]

# Construir el JSON de sesiones para interactividad JS
sessions_json = json.dumps(sessions, ensure_ascii=False)

html_content = f"""<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <link rel="icon" type="image/png" href="assets/escudo_kinal_cuadrado.png"><rect width='100' height='100' rx='22' fill='%2322376D'/><path d='M50 15 C55 35 78 45 78 65 C78 80 65 90 50 90 C35 90 22 80 22 65 C22 45 45 35 50 15 Z' fill='%23EC9123'/><path d='M50 38 C53 48 64 54 64 68 C64 76 58 82 50 82 C42 82 36 76 36 68 C36 54 47 48 50 38 Z' fill='%23FFD166'/></svg>">
    <title>Gestión del Talento Técnico y Mandos Medios | Kinal 2026</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <script src="https://unpkg.com/lucide@latest"></script>
    <style>
        body {{ font-family: Arial, Helvetica, sans-serif; background-color: #F8FAFC; color: #0F172A; }}
        h1, h2, h3, h4, .font-heading, .brand-serif {{ font-family: 'Times New Roman', Times, serif; }}
        .hero-kinal-gradient {{
            background: linear-gradient(135deg, #0D162C 0%, #152244 45%, #22376D 100%);
        }}
        .tab-btn.active {{
            background-color: #22376D !important;
            color: #ffffff !important;
            font-weight: 700 !important;
        }}
        .session-item.active {{
            background-color: #22376D !important;
            color: #ffffff !important;
            border-color: #22376D !important;
        }}
        .session-item.active span, .session-item.active p {{
            color: #e2e8f0 !important;
        }}
        .session-item.active .badge-code {{
            background-color: #EC9123 !important;
            color: #0D162C !important;
            font-weight: 800 !important;
        }}
    </style>

    <!-- ========================================================= -->
    <!-- KINAL MOBILE ENGINE - ZERO CLIPPING RESPONSIVE SYSTEM      -->
    <!-- ========================================================= -->
    <style id="kinal-mobile-responsive-system">
        @media (max-width: 767px) {{
            *, *::before, *::after {{
                box-sizing: border-box !important;
            }}
            html, body {{
                width: 100% !important;
                max-width: 100vw !important;
                overflow-x: hidden !important;
                margin: 0 !important;
                padding: 0 !important;
                -webkit-text-size-adjust: 100% !important;
                touch-action: pan-y pinch-zoom;
            }}
            main, section, header, footer, 
            .max-w-7xl, .max-w-6xl, .max-w-5xl, .max-w-4xl, .max-w-3xl, .max-w-2xl {{
                width: 100% !important;
                max-width: 100% !important;
                min-width: 0 !important;
                box-sizing: border-box !important;
            }}
            .grid {{
                width: 100% !important;
                min-width: 0 !important;
            }}
            .card-hover, [class*="rounded-3xl"], [class*="rounded-2xl"], .area-course-card {{
                width: 100% !important;
                max-width: 100% !important;
                min-width: 0 !important;
                padding-left: 1rem !important;
                padding-right: 1rem !important;
                box-sizing: border-box !important;
            }}
            h1, h2, h3, h4, h5, p, span, a, div {{
                overflow-wrap: break-word !important;
                word-break: break-word !important;
            }}
            .no-scrollbar::-webkit-scrollbar, 
            nav::-webkit-scrollbar,
            .overflow-x-auto::-webkit-scrollbar {{
                display: none !important;
                width: 0 !important;
                height: 0 !important;
            }}
            .no-scrollbar, 
            nav,
            .overflow-x-auto {{
                -ms-overflow-style: none !important;
                scrollbar-width: none !important;
                -webkit-overflow-scrolling: touch !important;
            }}
            .tab-btn, button, select, input[type="text"] {{
                min-height: 42px;
                touch-action: manipulation;
            }}
            nav.flex.overflow-x-auto, 
            .tab-nav-wrapper {{
                display: flex !important;
                flex-wrap: nowrap !important;
                overflow-x: auto !important;
                padding-top: 4px !important;
                padding-bottom: 4px !important;
            }}
            nav.flex.overflow-x-auto > * {{
                flex-shrink: 0 !important;
            }}
            main {{
                padding-left: 12px !important;
                padding-right: 12px !important;
                padding-top: 14px !important;
                padding-bottom: 24px !important;
            }}
            header .max-w-7xl, header .max-w-6xl {{
                padding-left: 12px !important;
                padding-right: 12px !important;
                padding-top: 8px !important;
                padding-bottom: 8px !important;
            }}
            .table-responsive-mobile, 
            .curriculum-grid-container {{
                width: 100% !important;
                max-width: 100% !important;
                overflow-x: auto !important;
                -webkit-overflow-scrolling: touch !important;
                border-radius: 12px !important;
            }}
            .tab-btn:active, button:active, .card-hover:active, .area-course-card:active {{
                transform: scale(0.97) !important;
                transition: transform 0.1s ease !important;
            }}
            .btn-doc-group {{
                display: grid !important;
                grid-template-columns: repeat(3, minmax(0, 1fr)) !important;
                gap: 6px !important;
            }}
            .btn-doc-item {{
                display: flex !important;
                flex-direction: column !important;
                align-items: center !important;
                justify-content: center !important;
                padding: 6px 4px !important;
                font-size: 10px !important;
                gap: 2px !important;
                text-align: center !important;
            }}
        }}

        @media (max-width: 1023px) {{
            .session-item {{
                cursor: pointer;
                -webkit-tap-highlight-color: transparent;
                transition: all 0.2s ease;
            }}
            [id$="-detail-card"] {{
                scroll-margin-top: 80px;
            }}
        }}
    </style>
</head>
<body class="min-h-screen flex flex-col antialiased pt-[115px] sm:pt-[125px]">

    <!-- ========================================================= -->
    <!-- ENCABEZADO FIJO PERMANENTE UNIFICADO                      -->
    <!-- ========================================================= -->
    <div class="fixed top-0 left-0 right-0 z-50 bg-white shadow-md border-b border-slate-200">
        
        <!-- BARRA INSTITUCIONAL SUPERIOR -->
        <div class="bg-[#152244] text-white text-xs border-b-2 border-[#EC9123]">
            <div class="max-w-7xl mx-auto px-4 sm:px-6 lg:px-8 py-1.5 flex flex-wrap justify-between items-center gap-2">
                <div class="flex items-center gap-3">
                    <span class="inline-flex items-center gap-1.5 font-bold tracking-wider uppercase text-[#EC9123]">
                        <i data-lucide="award" class="w-3.5 h-3.5"></i> Fundación Kinal
                    </span>
                    <span class="text-slate-500">|</span>
                    <span class="text-slate-300 font-medium hidden sm:inline">
                        Ideario Institucional: <em>«El trabajo bien hecho»</em>
                    </span>
                </div>
                <div class="flex items-center gap-4 text-slate-300 font-medium">
                    <span class="flex items-center gap-1.5">
                        <i data-lucide="calendar" class="w-3.5 h-3.5 text-[#EC9123]"></i> Ciclo Académico 2026
                    </span>
                    <span class="text-slate-600">|</span>
                    <span class="text-emerald-400 font-bold flex items-center gap-1">
                        <i data-lucide="shield-check" class="w-3.5 h-3.5"></i> Aval UNIS-Kinal
                    </span>
                </div>
            </div>
        </div>

        <!-- HEADER CORPORATIVO (LOGO ETS & BOTONES DE NAVEGACIÓN PRINCIPAL) -->
        <header class="max-w-7xl mx-auto px-4 sm:px-6 lg:px-8 py-2 sm:py-2.5 flex flex-wrap justify-between items-center gap-3 border-b border-slate-100">
            <div class="flex items-center">
                <a href="index.html" class="flex items-center group">
                    <img src="assets/escuela_tecnica_superior.png" alt="Escuela Técnica Superior" class="h-8 sm:h-9 w-auto object-contain transition group-hover:opacity-90">
                </a>
            </div>

            <!-- Botones Principales: Inicio y Vuelta a Cursos Nuevos -->
            <nav class="flex items-center gap-2 sm:gap-3 text-xs">
                <a href="cursos_nuevos.html" class="px-3.5 py-2 rounded-lg bg-white hover:bg-slate-50 text-[#22376D] transition flex items-center gap-1.5 shadow-sm whitespace-nowrap font-bold border border-slate-200" title="Volver al catálogo de cursos nuevos">
                    <i data-lucide="arrow-left" class="w-4 h-4 text-[#EC9123]"></i> Nuevos
                </a>
                <a href="index.html" class="px-3.5 py-2 rounded-lg bg-white hover:bg-slate-50 text-[#22376D] transition flex items-center gap-1.5 shadow-sm whitespace-nowrap font-bold border border-slate-200">
                    <i data-lucide="home" class="w-4 h-4 text-[#EC9123]"></i>
                </a>
            </nav>
        </header>

        <!-- SUB-BARRA INTEGRADA DE PESTAÑAS DEL CURSO (UNIFICADA CON EL ENCABEZADO) -->
        <div class="bg-slate-50 border-t border-slate-200">
            <div class="max-w-7xl mx-auto px-4 sm:px-6 lg:px-8">
                <nav class="flex space-x-2 py-2 overflow-x-auto no-scrollbar text-xs font-bold">
                    <button onclick="switchTab('propuesta')" id="btn-propuesta" class="tab-btn active px-3.5 py-1.5 rounded-lg transition flex items-center gap-1.5 shrink-0">
                        <i data-lucide="file-text" class="w-3.5 h-3.5"></i> Propuesta
                    </button>
                    <button onclick="switchTab('temario')" id="btn-temario" class="tab-btn px-3.5 py-1.5 rounded-lg bg-white hover:bg-slate-100 text-slate-700 border border-slate-200 transition flex items-center gap-1.5 shrink-0">
                        <i data-lucide="list-tree" class="w-3.5 h-3.5"></i> Temario
                    </button>
                    <button onclick="switchTab('calendario')" id="btn-calendario" class="tab-btn px-3.5 py-1.5 rounded-lg bg-white hover:bg-slate-100 text-slate-700 border border-slate-200 transition flex items-center gap-1.5 shrink-0">
                        <i data-lucide="calendar" class="w-3.5 h-3.5"></i> Calendario
                    </button>
                    <button onclick="switchTab('secuencias')" id="btn-secuencias" class="tab-btn px-3.5 py-1.5 rounded-lg bg-white hover:bg-slate-100 text-slate-700 border border-slate-200 transition flex items-center gap-1.5 shrink-0">
                        <i data-lucide="layers" class="w-3.5 h-3.5"></i> Secuencias
                    </button>
                    <button onclick="switchTab('evaluacion')" id="btn-evaluacion" class="tab-btn px-3.5 py-1.5 rounded-lg bg-white hover:bg-slate-100 text-slate-700 border border-slate-200 transition flex items-center gap-1.5 shrink-0">
                        <i data-lucide="award" class="w-3.5 h-3.5"></i> Evaluación
                    </button>
                </nav>
            </div>
        </div>

    </div>

    <!-- Hero Course Banner -->
    <section class="hero-kinal-gradient relative overflow-hidden text-white py-10 px-4 sm:px-6 lg:px-8">
        <div class="max-w-7xl mx-auto space-y-3">
            <div class="flex flex-wrap items-center gap-2">
                <span class="px-3 py-1 rounded-full bg-blue-500/20 border border-blue-400/30 text-blue-300 text-xs font-bold tracking-wide uppercase flex items-center gap-1.5">
                    <i data-lucide="handshake" class="w-3.5 h-3.5 text-blue-300"></i> Vinculación Empresarial Kinal
                </span>
                <span class="px-3 py-1 rounded-full bg-amber-500/20 border border-amber-400/30 text-amber-300 text-xs font-bold">
                    Capacitación Gerencial
                </span>
                <span class="px-3 py-1 rounded-full bg-emerald-500/20 border border-emerald-400/30 text-emerald-300 text-xs font-bold">
                    Acreditación: ≥ 75 Pts
                </span>
            </div>
            <h1 class="text-2xl sm:text-4xl font-black tracking-tight leading-tight brand-serif">
                Programa Ejecutivo de Vinculación y Liderazgo: Gestión del Talento Técnico y Mandos Medios
            </h1>
            <p class="text-slate-300 text-xs sm:text-sm max-w-4xl leading-relaxed">
                Iniciativa institucional de Fundación Kinal para fortalecer a líderes de equipos y mandos medios en manejo de personal, liderazgo humanista, competencias técnicas operativas y alineación con las necesidades actuales de la industria guatemalteca.
            </p>
            <div class="pt-2 flex flex-wrap gap-4 text-xs text-slate-200">
                <span class="flex items-center gap-1.5">
                    <i data-lucide="clock" class="w-4 h-4 text-[#EC9123]"></i> Total: 8 Horas Formativas Ejecutivas
                </span>
                <span class="flex items-center gap-1.5">
                    <i data-lucide="map-pin" class="w-4 h-4 text-[#EC9123]"></i> 2 Sesiones Presenciales (Campus Kinal • 4h)
                </span>
                <span class="flex items-center gap-1.5">
                    <i data-lucide="video" class="w-4 h-4 text-[#EC9123]"></i> 2 Sesiones Virtuales (Microsoft Teams • 4h)
                </span>
            </div>
        </div>
    </section>

    <!-- Main Container -->
    <main class="max-w-7xl mx-auto px-4 sm:px-6 lg:px-8 py-8 flex-1 space-y-8">

        <!-- ================================================================== -->
        <!-- PESTAÑA 1: PROPUESTA -->
        <!-- ================================================================== -->
        <div id="tab-propuesta" class="tab-content space-y-8">
            
            <!-- Ficha Técnica (UNA SOLA TARJETA) -->
            <div class="bg-white rounded-3xl border border-slate-200 p-6 sm:p-8 shadow-sm space-y-5">
                <div class="flex flex-wrap items-center justify-between border-b border-slate-200 pb-4 gap-2">
                    <div class="flex items-center gap-2">
                        <span class="w-8 h-8 rounded-lg bg-blue-50 text-blue-900 flex items-center justify-center font-black">
                            <i data-lucide="file-badge" class="w-4 h-4 text-[#22376D]"></i>
                        </span>
                        <h3 class="text-xl font-black text-slate-900">
                            Ficha Técnica Oficial del Programa
                        </h3>
                    </div>
                    <span class="px-3 py-1 rounded-full bg-blue-50 text-blue-900 text-xs font-bold border border-blue-200">
                        Código: VINCULACION-LIDERAZGO-2026
                    </span>
                </div>
                <div class="divide-y divide-slate-100">
"""

for row in ficha_rows:
    html_content += f"""
                    <div class="py-3 sm:grid sm:grid-cols-3 sm:gap-4 flex flex-col">
                        <dt class="font-bold text-slate-900 text-xs sm:col-span-1">{clean_html(row[0])}</dt>
                        <dd class="text-slate-700 text-xs sm:col-span-2 mt-0.5 sm:mt-0 leading-relaxed">{clean_html(row[1])}</dd>
                    </div>
    """

html_content += f"""
                </div>
            </div>

            <!-- Marco Filosófico Kinal -->
            <div class="grid md:grid-cols-2 gap-6">
                <div class="bg-blue-50 border-l-4 border-blue-900 p-6 rounded-r-3xl space-y-3 shadow-sm">
                    <div class="flex items-center gap-2 text-blue-900 font-extrabold text-sm uppercase">
                        <i data-lucide="flag" class="w-4 h-4"></i> Misión Institucional de Fundación Kinal
                    </div>
                    <p class="text-slate-700 text-xs italic leading-relaxed">
                        «Formar a jóvenes y adultos a través de una educación integral, con énfasis en las áreas técnicas y tecnológicas, influyendo positivamente en su trabajo, su familia y la sociedad».
                    </p>
                    <div class="pt-2 text-xs font-bold text-blue-950">
                        Valores Nucleares: <span class="font-normal italic">Visión cristiana de la vida. Respeto a la dignidad de la persona. Espíritu de servicio. Trabajo bien hecho. Libertad personal responsable.</span>
                    </div>
                </div>

                <div class="bg-amber-50 border-l-4 border-amber-600 p-6 rounded-r-3xl space-y-3 shadow-sm">
                    <div class="flex items-center gap-2 text-amber-900 font-extrabold text-sm uppercase">
                        <i data-lucide="heart-handshake" class="w-4 h-4"></i> El Principio del «Trabajo Bien Hecho» y Factor Humano
                    </div>
                    <p class="text-slate-700 text-xs leading-relaxed">
                        En Kinal, el <strong>Factor Humano</strong> es una rama de formación universal que erradica el capataz autoritario, el maltrato verbal y el sarcasmo en plantas y talleres. Fomenta la dignidad, el orgullo por el trabajo pulcro (5S), la seguridad psicológica y la custodia ética de maquinaria y herramientas.
                    </p>
                </div>
            </div>

            <!-- Fundamentación Estratégica: Validación TSU -->
            <div class="bg-white rounded-3xl border border-slate-200 p-8 shadow-sm space-y-4">
                <h3 class="text-xl font-black text-slate-900 brand-serif flex items-center gap-2">
                    <i data-lucide="shield-check" class="w-5 h-5 text-blue-700"></i>
                    Fundamentación: El Nivel Técnico Superior (TSU) como Respaldo Industrial
                </h3>
                <p class="text-xs text-slate-600 leading-relaxed">
                    La manufactura contemporánea (automatización, PLCs, robótica, redes industriales) exige certidumbre operativa. El <strong>Técnico Superior Universitario (TSU)</strong> se constituye en el mecanismo riguroso de validación de competencias técnicas que respalda a la industria, acreditando que el personal operativo y de supervisión domina el diagnóstico analítico, la lectura de planos y el cálculo aplicado.
                </p>
                <p class="text-xs text-slate-600 leading-relaxed">
                    Formar mandos medios con competencias a nivel superior en el área técnica permite cerrar la brecha entre la planeación de la ingeniería de diseño y la ejecución real en piso de planta, garantizando fiabilidad operativa y reducción de costos de mantenimiento.
                </p>
            </div>

            <!-- Estructura Curricular Modular (4 Módulos) -->
            <div class="bg-white rounded-3xl border border-slate-200 p-6 sm:p-8 shadow-sm space-y-6">
                <div class="flex flex-col sm:flex-row justify-between items-start sm:items-center gap-2 border-b border-slate-200 pb-3">
                    <h3 class="text-lg font-black text-slate-900 brand-serif flex items-center gap-2">
                        <i data-lucide="layers" class="w-5 h-5 text-[#22376D]"></i> Estructura Curricular: 4 Módulos y Actividades Evaluables (Total 8 Horas)
                    </h3>
                    <span class="text-xs font-bold text-slate-500">4 Sesiones de 2 Horas • 100 Pts Totales</span>
                </div>
                <div class="grid sm:grid-cols-2 lg:grid-cols-4 gap-4">
                    <div class="p-4 rounded-2xl bg-amber-50/60 border border-amber-200 space-y-2 flex flex-col justify-between">
                        <div class="space-y-1.5">
                            <span class="px-2.5 py-0.5 rounded-full bg-amber-900 text-white text-[10px] font-black uppercase">Módulo 1 • Presencial (2h)</span>
                            <h4 class="font-bold text-sm text-slate-900 leading-snug">Liderazgo y Factor Humano en Cuadrillas</h4>
                            <p class="text-xs text-slate-600">Apertura institucional, cultura de Kinal, respeto a la dignidad y erradicación del maltrato en planta.</p>
                        </div>
                        <div class="pt-2 border-t border-amber-200/60 text-[11px] font-bold text-amber-950 flex items-center gap-1.5">
                            <i data-lucide="file-check-2" class="w-3.5 h-3.5 text-amber-700"></i> Código Convivencia y Cero Maltrato (25 Pts)
                        </div>
                    </div>
                    <div class="p-4 rounded-2xl bg-purple-50/60 border border-purple-200 space-y-2 flex flex-col justify-between">
                        <div class="space-y-1.5">
                            <span class="px-2.5 py-0.5 rounded-full bg-purple-900 text-white text-[10px] font-black uppercase">Módulo 2 • Virtual Teams (2h)</span>
                            <h4 class="font-bold text-sm text-slate-900 leading-snug">Manejo de Personal y Mandos Medios</h4>
                            <p class="text-xs text-slate-600">Atracción por competencias, onboarding a 30-60-90 días, retroalimentación asertiva y micro-charlas.</p>
                        </div>
                        <div class="pt-2 border-t border-purple-200/60 text-[11px] font-bold text-purple-950 flex items-center gap-1.5">
                            <i data-lucide="file-check-2" class="w-3.5 h-3.5 text-purple-700"></i> Descriptor, Onboarding y Micro-Charlas (25 Pts)
                        </div>
                    </div>
                    <div class="p-4 rounded-2xl bg-blue-50/60 border border-blue-200 space-y-2 flex flex-col justify-between">
                        <div class="space-y-1.5">
                            <span class="px-2.5 py-0.5 rounded-full bg-blue-900 text-white text-[10px] font-black uppercase">Módulo 3 • Virtual Teams (2h)</span>
                            <h4 class="font-bold text-sm text-slate-900 leading-snug">Fortalecimiento y Validación TSU</h4>
                            <p class="text-xs text-slate-600">El Técnico Superior Universitario como respaldo técnico formal, matrices operativas y evaluación en banco.</p>
                        </div>
                        <div class="pt-2 border-t border-blue-200/60 text-[11px] font-bold text-blue-950 flex items-center gap-1.5">
                            <i data-lucide="file-check-2" class="w-3.5 h-3.5 text-blue-700"></i> Evaluación Práctica y Validación TSU (25 Pts)
                        </div>
                    </div>
                    <div class="p-4 rounded-2xl bg-emerald-50/60 border border-emerald-200 space-y-2 flex flex-col justify-between">
                        <div class="space-y-1.5">
                            <span class="px-2.5 py-0.5 rounded-full bg-emerald-900 text-white text-[10px] font-black uppercase">Módulo 4 • Presencial (2h)</span>
                            <h4 class="font-bold text-sm text-slate-900 leading-snug">Sistematización, Acuerdos y Clausura</h4>
                            <p class="text-xs text-slate-600">Mesa Redonda Plenaria industria-academia, acuerdos de vinculación y entrega solemne de diplomas.</p>
                        </div>
                        <div class="pt-2 border-t border-emerald-200/60 text-[11px] font-bold text-emerald-950 flex items-center gap-1.5">
                            <i data-lucide="file-check-2" class="w-3.5 h-3.5 text-emerald-700"></i> Sustentación Plenaria y Acuerdos (25 Pts)
                        </div>
                    </div>
                </div>
            </div>

            <!-- Descarga de Documentos Oficiales en Word -->
            <div class="bg-white rounded-3xl border border-slate-200 p-6 sm:p-8 shadow-sm space-y-4">
                <div class="flex items-center justify-between border-b border-slate-200 pb-3">
                    <h4 class="text-base sm:text-lg font-black text-slate-900 brand-serif flex items-center gap-2">
                        <i data-lucide="download" class="w-5 h-5 text-blue-700"></i>
                        Documentación Oficial del Programa en Microsoft Word
                    </h4>
                    <span class="text-xs font-bold text-emerald-700 bg-emerald-50 px-3 py-1 rounded-full border border-emerald-200">
                        Formatos Institucionales
                    </span>
                </div>
                <div class="grid sm:grid-cols-3 gap-3 btn-doc-group">
                    <a href="Cursos/Talento_Tecnico/Propuesta_Diplomado_Talento_Tecnico_Kinal.docx" download class="btn-doc-item p-4 rounded-2xl bg-slate-50 hover:bg-blue-50 border border-slate-200 hover:border-blue-300 transition text-center space-y-1">
                        <i data-lucide="file-text" class="w-6 h-6 text-blue-700 mx-auto"></i>
                        <div class="font-bold text-xs text-slate-900">Propuesta Formativa</div>
                        <div class="text-[10px] text-slate-500">Documento Oficial (.docx)</div>
                    </a>
                    <a href="Cursos/Talento_Tecnico/Temario_Diplomado_Talento_Tecnico_Kinal.docx" download class="btn-doc-item p-4 rounded-2xl bg-slate-50 hover:bg-emerald-50 border border-slate-200 hover:border-emerald-300 transition text-center space-y-1">
                        <i data-lucide="list-checks" class="w-6 h-6 text-emerald-700 mx-auto"></i>
                        <div class="font-bold text-xs text-slate-900">Temario Analítico</div>
                        <div class="text-[10px] text-slate-500">Documento Oficial (.docx)</div>
                    </a>
                    <a href="Cursos/Talento_Tecnico/Dosificacion_y_Secuencia_Didactica_Talento_Tecnico_Kinal.docx" download class="btn-doc-item p-4 rounded-2xl bg-slate-50 hover:bg-amber-50 border border-slate-200 hover:border-amber-300 transition text-center space-y-1">
                        <i data-lucide="calendar" class="w-6 h-6 text-amber-700 mx-auto"></i>
                        <div class="font-bold text-xs text-slate-900">Dosificación Didáctica</div>
                        <div class="text-[10px] text-slate-500">Documento Oficial (.docx)</div>
                    </a>
                </div>
            </div>

        </div>

        <!-- ================================================================== -->
        <!-- PESTAÑA 2: TEMARIO -->
        <!-- ================================================================== -->
        <div id="tab-temario" class="tab-content hidden space-y-6">
            
            <!-- Barra Superior de Filtros por Categoría -->
            <div class="bg-white rounded-3xl p-5 border border-slate-200 shadow-sm space-y-3">
                <div class="flex flex-col sm:flex-row justify-between items-start sm:items-center gap-2">
                    <div>
                        <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                            <i data-lucide="layers" class="w-5 h-5 text-blue-700"></i>
                            Distribución Modular con Código de Color
                        </h3>
                        <p class="text-slate-600 text-xs">
                            Estructura de 4 módulos formativos divididos por áreas clave. Filtra por categoría para enfocar el análisis.
                        </p>
                    </div>
                    <span class="text-xs font-bold text-slate-500 bg-slate-100 px-3 py-1 rounded-full">
                        Total: 12 Contenidos Clave
                    </span>
                </div>

                <!-- Botones de Filtro -->
                <div class="flex flex-wrap gap-2 pt-2 border-t border-slate-100">
                    <button onclick="filterTalentoTemario('all')" id="btn-talento-all" class="px-3.5 py-1.5 rounded-xl text-xs font-extrabold bg-slate-900 text-white shadow-sm transition">
                        Todos los Temas (12)
                    </button>
                    <button onclick="filterTalentoTemario('factor_humano')" id="btn-talento-factor_humano" class="px-3.5 py-1.5 rounded-xl text-xs font-extrabold bg-amber-50 hover:bg-amber-100 text-amber-900 border border-amber-200 transition">
                        ⭐ Factor Humano Kinal
                    </button>
                    <button onclick="filterTalentoTemario('gestion')" id="btn-talento-gestion" class="px-3.5 py-1.5 rounded-xl text-xs font-extrabold bg-purple-50 hover:bg-purple-100 text-purple-900 border border-purple-200 transition">
                        📈 Manejo de Personal & Onboarding
                    </button>
                    <button onclick="filterTalentoTemario('competencias')" id="btn-talento-competencias" class="px-3.5 py-1.5 rounded-xl text-xs font-extrabold bg-blue-50 hover:bg-blue-100 text-blue-900 border border-blue-200 transition">
                        🎯 Competencias & Validación TSU
                    </button>
                    <button onclick="filterTalentoTemario('liderazgo')" id="btn-talento-liderazgo" class="px-3.5 py-1.5 rounded-xl text-xs font-extrabold bg-indigo-50 hover:bg-indigo-100 text-indigo-900 border border-indigo-200 transition">
                        👥 Mandos Medios
                    </button>
                    <button onclick="filterTalentoTemario('vinculacion')" id="btn-talento-vinculacion" class="px-3.5 py-1.5 rounded-xl text-xs font-extrabold bg-emerald-50 hover:bg-emerald-100 text-emerald-900 border border-emerald-200 transition">
                        🤝 Diálogo Industria-Kinal
                    </button>
                </div>
            </div>

            <!-- Cuadrícula Modular de Temarios -->
            <div class="grid lg:grid-cols-2 gap-6">
"""

for mod in temario_modules:
    html_content += f"""
                <div class="bg-white rounded-2xl border border-slate-300 shadow-sm overflow-hidden flex flex-col hover:shadow-md transition">
                    <div class="bg-slate-900 text-white p-4 sm:p-5 border-b border-slate-800 space-y-2">
                        <div class="flex flex-wrap justify-between items-center text-xs text-amber-400 font-semibold gap-2">
                            <span>MÓDULO {mod['num']} • {mod['semana']}</span>
                            <span class="px-2 py-0.5 rounded bg-blue-900/60 text-blue-200 border border-blue-400/30 text-[11px] font-bold">{mod['hours']}</span>
                        </div>
                        <h4 class="text-base sm:text-lg font-bold text-white leading-snug">{mod['title']}</h4>
                        <p class="text-xs text-slate-300 leading-relaxed">
                            <strong class="text-amber-300">Logro esperado:</strong> {mod['logro']}
                        </p>
                        <div class="inline-flex items-center gap-1.5 px-2.5 py-1 rounded-lg bg-emerald-950/80 border border-emerald-500/40 text-emerald-300 text-[11px] font-bold">
                            <i data-lucide="check-circle" class="w-3.5 h-3.5 text-emerald-400"></i> Actividad Evaluable: {mod['eval_act']}
                        </div>
                    </div>
                    <div class="p-4 sm:p-5 divide-y divide-slate-100 flex-1 space-y-2.5 bg-slate-50/50">
    """
    for t in mod['topics']:
        html_content += f"""
                        <div class="content-item item-{t['cat']} p-3.5 rounded-xl text-xs space-y-1.5 transition border border-slate-200/90 bg-white hover:bg-slate-50 shadow-sm" data-category="{t['cat']}">
                            <div class="flex flex-wrap justify-between items-center gap-1.5 font-bold">
                                <span class="text-slate-900 text-sm flex items-center gap-2">
                                    <span class="w-5 h-5 rounded-full bg-slate-100 text-slate-700 text-xs font-black flex items-center justify-center shrink-0">{t['num']}</span>
                                    {t['title']}
                                </span>
                                <span class="px-2.5 py-0.5 rounded-full text-[10px] font-extrabold uppercase {t['cat_badge']}">{t['cat_name']}</span>
                            </div>
                            <p class="text-slate-600 leading-relaxed pl-7">{t['desc']}</p>
                        </div>
        """
    html_content += """
                    </div>
                </div>
    """

html_content += f"""
            </div>
        </div>

        <!-- ================================================================== -->
        <!-- PESTAÑA 3: CALENDARIO CRONOLÓGICO -->
        <!-- ================================================================== -->
        <div id="tab-calendario" class="tab-content hidden space-y-6">
            
            <div class="flex flex-col sm:flex-row justify-between items-start sm:items-center gap-3 bg-white rounded-3xl p-6 border border-slate-200 shadow-sm">
                <div>
                    <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                        <i data-lucide="calendar" class="w-5 h-5 text-blue-700"></i>
                        Calendario Cronológico de Encuentros (Total: 8 Horas)
                    </h3>
                    <p class="text-slate-600 text-xs mt-1">
                        Programa de 4 sesiones de 2 horas (4h presenciales en Campus Kinal + 4h virtuales en Microsoft Teams).
                    </p>
                </div>
                <div class="flex items-center gap-2 text-xs font-bold text-blue-900 bg-blue-50 px-3 py-1.5 rounded-xl border border-blue-200">
                    <i data-lucide="check-circle-2" class="w-4 h-4 text-blue-700"></i> 4 Sesiones Programadas
                </div>
            </div>

            <!-- Grid de Calendario -->
            <div class="grid sm:grid-cols-2 lg:grid-cols-4 gap-4">
"""

for s in sessions:
    badge_bg = "bg-amber-100 text-amber-900" if "Presencial" in s['tipo'] else "bg-purple-100 text-purple-900"
    html_content += f"""
                <div class="bg-white rounded-xl border border-slate-200 p-4 shadow-sm hover:border-[#22376D] hover:shadow-md transition flex flex-col justify-between space-y-3">
                    <div class="space-y-1.5">
                        <div class="flex justify-between items-center text-[11px]">
                            <span class="font-black px-2 py-0.5 rounded bg-slate-900 text-[#EC9123]">{s['code']}</span>
                            <span class="font-bold px-2 py-0.5 rounded {badge_bg} text-[10px]">{s['tipo']}</span>
                        </div>
                        <h5 class="font-bold text-slate-900 text-xs leading-snug line-clamp-2">{s['title']}</h5>
                        <p class="text-[11px] text-slate-500">{s['semana']}</p>
                    </div>
                    <div class="pt-2 border-t border-slate-100 flex items-center justify-between text-[11px]">
                        <span class="text-slate-600 font-semibold">{s['duration']}</span>
                        <button onclick="goToSession({s['num']})" class="text-[#22376D] font-bold hover:underline flex items-center gap-0.5">
                            Ver Secuencia →
                        </button>
                    </div>
                </div>
    """

html_content += f"""
            </div>
        </div>

        <!-- ================================================================== -->
        <!-- PESTAÑA 4: SECUENCIAS (MASTER-DETAIL) -->
        <!-- ================================================================== -->
        <div id="tab-secuencias" class="tab-content hidden space-y-6">
            
            <div class="flex flex-col sm:flex-row justify-between items-start sm:items-center gap-3 bg-white rounded-3xl p-5 border border-slate-200 shadow-sm">
                <div>
                    <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                        <i data-lucide="layers" class="w-5 h-5 text-blue-700"></i>
                        Secuencia Didáctica Interactiva Sesión a Sesión
                    </h3>
                    <p class="text-slate-600 text-xs mt-0.5">
                        Selecciona una sesión en el panel lateral para revisar su microdiseño en tres momentos (Apertura, Desarrollo y Cierre) y su actividad evaluable.
                    </p>
                </div>
                <div class="flex items-center gap-2 w-full sm:w-auto">
                    <div class="relative w-full sm:w-56">
                        <i data-lucide="search" class="w-3.5 h-3.5 text-slate-400 absolute left-3 top-2.5"></i>
                        <input type="text" id="talento-sec-search" oninput="filterTalentoList()" placeholder="Buscar sesión..." class="w-full pl-8 pr-3 py-1.5 text-xs rounded-xl border border-slate-200 bg-slate-50 focus:bg-white focus:outline-none focus:ring-2 focus:ring-[#22376D] transition">
                    </div>
                </div>
            </div>

            <!-- Master-Detail Container -->
            <div class="grid lg:grid-cols-12 gap-6 items-start">
                
                <!-- LISTADO LATERAL (MASTER) -->
                <div class="lg:col-span-4 bg-white rounded-3xl border border-slate-200 p-4 shadow-sm space-y-2">
                    <div class="text-[11px] font-black uppercase tracking-wider text-slate-400 px-2 pb-1 border-b border-slate-100 flex justify-between">
                        <span>Listado de Sesiones</span>
                        <span id="talento-count-badge">4 Sesiones</span>
                    </div>
                    <div id="talento-sessions-list" class="space-y-1.5 max-h-[560px] overflow-y-auto pr-1">
                        <!-- Rendered by JS -->
                    </div>
                </div>

                <!-- DETALLE DE LA SESIÓN SELECCIONADA (DETAIL) -->
                <div class="lg:col-span-8 bg-white rounded-3xl border border-slate-200 p-6 sm:p-8 shadow-sm space-y-6" id="talento-detail-card">
                    <!-- Rendered by JS -->
                </div>

            </div>
        </div>

        <!-- ================================================================== -->
        <!-- PESTAÑA 5: EVALUACIÓN -->
        <!-- ================================================================== -->
        <div id="tab-evaluacion" class="tab-content hidden space-y-8">
            
            <!-- Esquema Resumen -->
            <div class="bg-white rounded-3xl border border-slate-200 p-8 shadow-sm space-y-6">
                <div class="flex items-center justify-between border-b border-slate-200 pb-3">
                    <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                        <i data-lucide="check-square" class="w-5 h-5 text-blue-700"></i>
                        Esquema de Evaluación Institucional & Acreditación
                    </h3>
                    <span class="px-3 py-1 rounded-full bg-amber-100 text-amber-900 font-black text-xs">
                        Umbral Mínimo: 75 / 100 Puntos
                    </span>
                </div>
                <div class="grid sm:grid-cols-2 lg:grid-cols-4 gap-4 text-xs">
"""

for ev in eval_prop_rows:
    html_content += f"""
                    <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 text-center space-y-1">
                        <div class="text-slate-500 font-bold uppercase text-[10px]">{clean_html(ev[0])}</div>
                        <div class="text-3xl font-black text-[#22376D]">{clean_html(ev[1])}</div>
                        <div class="text-slate-600">{clean_html(ev[2])}</div>
                    </div>
    """

html_content += f"""
                </div>
            </div>

            <!-- Rúbrica Terminal -->
            <div class="bg-white rounded-3xl border border-slate-200 p-8 shadow-sm space-y-6">
                <div class="flex flex-col sm:flex-row justify-between items-start sm:items-center gap-2 border-b border-slate-200 pb-4">
                    <div>
                        <h3 class="text-xl font-black text-slate-900 brand-serif flex items-center gap-2">
                            <i data-lucide="award" class="w-5 h-5 text-blue-700"></i>
                            Rúbrica Analítica Terminal de Evaluación Práctica
                        </h3>
                        <p class="text-slate-600 text-xs mt-1">
                            Evaluación objetiva del desempeño técnico con 4 criterios ponderados a 25 puntos y umbral institucional de 75 puntos.
                        </p>
                    </div>
                    <span class="px-3 py-1.5 rounded-xl bg-amber-100 text-amber-900 font-extrabold text-xs">
                        Aprobación Mínima: ≥ 75 Pts
                    </span>
                </div>

                <div class="overflow-x-auto">
                    <table class="w-full text-xs text-left border-collapse">
                        <thead>
                            <tr class="bg-slate-100 text-slate-800 uppercase font-black border-b border-slate-200">
                                <th class="py-3 px-4 w-1/4">Actividad Formativa / Módulo</th>
                                <th class="py-3 px-4 w-24 text-center">Ponderación</th>
                                <th class="py-3 px-4 text-emerald-900">Desempeño Excelente (Trabajo Bien Hecho)</th>
                                <th class="py-3 px-4 text-rose-900">Desempeño Insuficiente (&lt; 75%)</th>
                            </tr>
                        </thead>
                        <tbody class="divide-y divide-slate-100">
"""

for rub in rubrica_rows:
    html_content += f"""
                            <tr class="hover:bg-slate-50/80 transition">
                                <td class="py-3.5 px-4 font-bold text-slate-900 bg-slate-50 border-r border-slate-200 align-top">{clean_html(rub[0])}</td>
                                <td class="py-3 px-4 text-blue-950 font-bold bg-blue-50/20 border-r border-slate-200 align-top text-center">{clean_html(rub[1])}</td>
                                <td class="py-3 px-4 text-emerald-950 bg-emerald-50/30 border-r border-slate-200 align-top">{clean_html(rub[2])}</td>
                                <td class="py-3 px-4 text-rose-950 bg-rose-50/30 align-top text-xs">{clean_html(rub[3])}</td>
                            </tr>
    """

html_content += f"""
                        </tbody>
                    </table>
                </div>
            </div>

        </div>

    </main>

    <!-- Footer Corporativo -->
    <footer class="bg-[#0D162C] text-white py-10 px-4 sm:px-6 lg:px-8 border-t border-slate-800 text-xs mt-12">
        <div class="max-w-7xl mx-auto flex flex-col sm:flex-row justify-between items-center gap-4 text-center sm:text-left">
            <div>
                <div class="font-bold text-sm tracking-wide text-amber-400">FUNDACIÓN KINAL</div>
                <div class="text-slate-400 mt-1">Dirección de Vinculación Empresarial • Programa Ejecutivo de Liderazgo y Talento Técnico (8 Horas)</div>
            </div>
            <div class="text-slate-400 text-xs space-y-1">
                <div>Ideario Institucional: <em>«El trabajo bien hecho»</em> • Ciclo 2026</div>
                <div class="flex justify-center sm:justify-end gap-3 text-slate-400">
                    <a href="cursos_nuevos.html" class="hover:text-white transition">Catálogo General</a>
                    <span>•</span>
                    <a href="https://kinal.org.gt" target="_blank" class="hover:text-white transition">Kinal Oficial</a>
                </div>
            </div>
        </div>
    </footer>

    <!-- Interactive Logic (Tabs, Filtering & Master-Detail) -->
    <script>
        const TALENTO_SESSIONS = {sessions_json};
        let currentTalentoIndex = 1;

        function switchTab(tabId) {{
            document.querySelectorAll('.tab-content').forEach(el => el.classList.add('hidden'));
            document.querySelectorAll('.tab-btn').forEach(btn => {{
                btn.classList.remove('active');
                btn.classList.add('bg-white', 'text-slate-700');
            }});

            const targetTab = document.getElementById('tab-' + tabId);
            const targetBtn = document.getElementById('btn-' + tabId);

            if (targetTab) targetTab.classList.remove('hidden');
            if (targetBtn) {{
                targetBtn.classList.add('active');
                targetBtn.classList.remove('bg-white', 'text-slate-700');
            }}

            if (tabId === 'secuencias') {{
                selectTalentoSession(currentTalentoIndex);
            }}

            window.scrollTo({{ top: 0, behavior: 'smooth' }});
        }}

        function renderTalentoSessionsList(sessionsArray) {{
            const container = document.getElementById('talento-sessions-list');
            if (!container) return;
            container.innerHTML = '';

            sessionsArray.forEach(s => {{
                const item = document.createElement('div');
                const isActive = s.num === currentTalentoIndex;
                item.className = `session-item p-3 rounded-2xl border text-xs cursor-pointer transition flex items-center justify-between gap-2 ${{isActive ? 'active' : 'bg-white hover:bg-slate-50 border-slate-200'}}`;
                item.onclick = () => selectTalentoSession(s.num);
                item.innerHTML = `
                    <div class="space-y-0.5 overflow-hidden">
                        <div class="flex items-center gap-1.5">
                            <span class="badge-code px-2 py-0.5 rounded text-[10px] font-black ${{isActive ? '' : 'bg-slate-900 text-amber-300'}}">${{s.code}}</span>
                            <span class="font-bold text-slate-900 truncate">${{s.title}}</span>
                        </div>
                        <div class="text-[10px] text-slate-500">${{s.semana}} • ${{s.tipo}}</div>
                    </div>
                    <i data-lucide="chevron-right" class="w-4 h-4 text-slate-400 shrink-0"></i>
                `;
                container.appendChild(item);
            }});
            lucide.createIcons();
        }}

        function selectTalentoSession(num) {{
            if (window.innerWidth < 1024) {{
                setTimeout(function() {{
                    const card = document.getElementById('talento-detail-card');
                    if (card) card.scrollIntoView({{ behavior: 'smooth', block: 'start' }});
                }}, 60);
            }}
            currentTalentoIndex = num;
            const s = TALENTO_SESSIONS.find(item => item.num === num);
            if (!s) return;

            document.querySelectorAll('.session-item').forEach(el => {{
                el.classList.remove('active');
                el.classList.add('bg-white');
            }});

            renderTalentoSessionsList(TALENTO_SESSIONS);

            const detailContainer = document.getElementById('talento-detail-card');
            detailContainer.innerHTML = `
                <div class="border-b border-slate-200 pb-5 space-y-2">
                    <div class="flex flex-wrap justify-between items-center gap-2">
                        <div class="flex items-center gap-2">
                            <span class="px-3 py-1 rounded-xl bg-slate-950 text-amber-400 font-black text-xs tracking-wider">${{s.code}}</span>
                            <span class="px-2.5 py-0.5 rounded-lg bg-blue-50 text-blue-900 font-bold text-xs border border-blue-200">${{s.module}}</span>
                        </div>
                        <span class="text-xs font-bold text-slate-500 bg-slate-100 px-3 py-1 rounded-lg">${{s.duration}}</span>
                    </div>
                    <h3 class="text-xl sm:text-2xl font-black text-slate-900 leading-snug">${{s.title}}</h3>
                    <div class="text-xs text-[#22376D] font-bold pt-1">
                        ${{s.semana}} • ${{s.horario}}
                    </div>
                </div>

                <div class="bg-blue-50/60 rounded-2xl p-4 border border-blue-200 space-y-1">
                    <div class="text-[11px] uppercase font-black tracking-wider text-blue-900 flex items-center gap-1.5">
                        <i data-lucide="target" class="w-3.5 h-3.5 text-[#22376D]"></i> Indicador de Logro / Objetivo Formativo
                    </div>
                    <p class="text-slate-800 text-xs leading-relaxed font-medium">${{s.objective}}</p>
                </div>

                <div class="space-y-4">
                    <h4 class="text-sm font-black text-slate-900 uppercase tracking-wide flex items-center gap-2 border-b border-slate-100 pb-2">
                        <i data-lucide="clock" class="w-4 h-4 text-[#22376D]"></i> Secuencia Didáctica en Tres Momentos (120 Minutos)
                    </h4>

                    <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 space-y-1.5">
                        <div class="flex items-center justify-between">
                            <span class="font-bold text-xs text-blue-950 flex items-center gap-2">
                                <span class="w-2.5 h-2.5 rounded-full bg-blue-600"></span> Fase 1: Apertura, Encuadre & Dinámica Inicial
                            </span>
                            <span class="text-[11px] font-bold text-slate-500">15-20 Minutos</span>
                        </div>
                        <p class="text-slate-700 text-xs leading-relaxed pl-4.5">${{s.apertura}}</p>
                    </div>

                    <div class="p-4 rounded-2xl bg-amber-50/40 border border-amber-200 space-y-1.5">
                        <div class="flex items-center justify-between">
                            <span class="font-bold text-xs text-amber-950 flex items-center gap-2">
                                <span class="w-2.5 h-2.5 rounded-full bg-amber-600"></span> Fase 2: Desarrollo, Casos Reales & Taller Aplicativo
                            </span>
                            <span class="text-[11px] font-bold text-amber-800 bg-amber-100 px-2 py-0.5 rounded">85-90 Minutos</span>
                        </div>
                        <p class="text-slate-700 text-xs leading-relaxed pl-4.5">${{s.desarrollo}}</p>
                    </div>

                    <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 space-y-1.5">
                        <div class="flex items-center justify-between">
                            <span class="font-bold text-xs text-blue-950 flex items-center gap-2">
                                <span class="w-2.5 h-2.5 rounded-full bg-emerald-600"></span> Fase 3: Cierre, Síntesis & Asignación de Entregable
                            </span>
                            <span class="text-[11px] font-bold text-slate-500">15-20 Minutos</span>
                        </div>
                        <p class="text-slate-700 text-xs leading-relaxed pl-4.5">${{s.cierre}}</p>
                    </div>
                </div>

                <div class="grid sm:grid-cols-2 gap-4 pt-2">
                    <div class="bg-emerald-50/60 rounded-2xl p-4 border border-emerald-200 space-y-1.5">
                        <div class="text-[11px] uppercase font-black tracking-wider text-emerald-900 flex items-center gap-1.5">
                            <i data-lucide="check-circle-2" class="w-3.5 h-3.5 text-emerald-700"></i> Actividad Formativa Evaluable (25 Pts)
                        </div>
                        <p class="text-slate-700 text-xs leading-relaxed">${{s.evidencias}}</p>
                    </div>

                    <div class="bg-slate-100 rounded-2xl p-4 border border-slate-200 space-y-1.5">
                        <div class="text-[11px] uppercase font-black tracking-wider text-slate-900 flex items-center gap-1.5">
                            <i data-lucide="monitor" class="w-3.5 h-3.5 text-[#22376D]"></i> Recursos Didácticos & Sede
                        </div>
                        <p class="text-slate-700 text-xs leading-relaxed">${{s.equipamiento}}</p>
                    </div>
                </div>
            `;
            lucide.createIcons();
        }}

        function filterTalentoList() {{
            const q = document.getElementById('talento-sec-search').value.toLowerCase();
            const filtered = TALENTO_SESSIONS.filter(s => 
                s.title.toLowerCase().includes(q) || 
                s.code.toLowerCase().includes(q) ||
                s.module.toLowerCase().includes(q) ||
                s.objective.toLowerCase().includes(q)
            );
            document.getElementById('talento-count-badge').textContent = `${{filtered.length}} Sesiones`;
            renderTalentoSessionsList(filtered);
        }}

        function filterTalentoTemario(category) {{
            const items = document.querySelectorAll('.content-item');
            items.forEach(it => {{
                if (category === 'all' || it.getAttribute('data-category') === category) {{
                    it.style.display = 'block';
                }} else {{
                    it.style.display = 'none';
                }}
            }});

            document.querySelectorAll('[id^="btn-talento-"]').forEach(btn => {{
                btn.className = 'px-3.5 py-1.5 rounded-xl text-xs font-extrabold bg-slate-100 hover:bg-slate-200 text-slate-700 border border-slate-200 transition';
            }});

            const activeBtn = document.getElementById('btn-talento-' + category);
            if (activeBtn) {{
                activeBtn.className = 'px-3.5 py-1.5 rounded-xl text-xs font-extrabold bg-slate-900 text-white shadow-sm transition';
            }}
        }}

        function goToSession(num) {{
            switchTab('secuencias');
            selectTalentoSession(num);
        }}

        document.addEventListener('DOMContentLoaded', () => {{
            lucide.createIcons();
            selectTalentoSession(1);
        }});
    </script>
</body>
</html>
"""

# Guardar en raíz y en web
out_path_1 = os.path.join(WORKSPACE_DIR, "curso_diplomado_gerencial.html")
out_path_2 = os.path.join(WORKSPACE_DIR, "curso_talento_tecnico.html")

with open(out_path_1, "w", encoding="utf-8") as f:
    f.write(html_content)
with open(out_path_2, "w", encoding="utf-8") as f:
    f.write(html_content)

print(f"Generated successfully: {out_path_1}")
print(f"Generated successfully: {out_path_2}")
