import docx
import json
import os

def parse_docx_content(docx_path):
    doc = docx.Document(docx_path)
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    tables_data = []
    for t in doc.tables:
        t_rows = []
        for r in t.rows:
            row_cells = [c.text.strip() for c in r.cells]
            t_rows.append(row_cells)
        tables_data.append(t_rows)
    return {"paragraphs": paras, "tables": tables_data}

info_meca_prop = parse_docx_content("Cursos/Mecatrónica/Propuesta_Curso_Mecatronica_Kinal.docx")
info_meca_tem = parse_docx_content("Cursos/Mecatrónica/Temario_Curso_Mecatronica_Kinal.docx")
info_meca_dos = parse_docx_content("Cursos/Mecatrónica/Dosificacion_y_Secuencia_Didactica_Mecatronica_Kinal.docx")

print("Mecatronica parsed successfully:")
print(" Propuesta paragraphs:", len(info_meca_prop["paragraphs"]), "tables:", len(info_meca_prop["tables"]))
print(" Temario paragraphs:", len(info_meca_tem["paragraphs"]), "tables:", len(info_meca_tem["tables"]))
print(" Dosificacion paragraphs:", len(info_meca_dos["paragraphs"]), "tables:", len(info_meca_dos["tables"]))

info_ciber_prop = parse_docx_content("Cursos/Ciberseguridad/Propuesta_Curso_Ciberseguridad_Kinal.docx")
info_ciber_tem = parse_docx_content("Cursos/Ciberseguridad/Temario_Curso_Ciberseguridad_Kinal.docx")
info_ciber_dos = parse_docx_content("Cursos/Ciberseguridad/Dosificacion_y_Secuencia_Didactica_Ciberseguridad.docx")

print("Ciberseguridad parsed successfully:")
print(" Propuesta paragraphs:", len(info_ciber_prop["paragraphs"]), "tables:", len(info_ciber_prop["tables"]))
print(" Temario paragraphs:", len(info_ciber_tem["paragraphs"]), "tables:", len(info_ciber_tem["tables"]))
print(" Dosificacion paragraphs:", len(info_ciber_dos["paragraphs"]), "tables:", len(info_ciber_dos["tables"]))
