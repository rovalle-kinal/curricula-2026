import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_temario():
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    COLOR_PRIMARY = RGBColor(0, 51, 102)      # Azul Kinal
    COLOR_SECONDARY = RGBColor(74, 85, 104)    # Gris pizarra
    COLOR_DARK = RGBColor(26, 32, 44)         # Texto oscuro

    def style_p(p, space_before=0, space_after=6, line_spacing=1.15):
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing

    # ==========================================
    # # Nombre del curso
    # ==========================================
    title_p = doc.add_paragraph()
    style_p(title_p, space_before=0, space_after=12)
    run_title = title_p.add_run("Fundamentos de Ciberseguridad y Operaciones SOC: Hardening, Monitoreo y Respuesta a Incidentes")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    style_p(sub_p, space_before=0, space_after=18)
    run_sub = sub_p.add_run("Programa Técnico para la Empleabilidad Inmediata — Fundación Kinal | Nivel DQR 4-5")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_SECONDARY

    border_p = doc.add_paragraph()
    style_p(border_p, space_before=0, space_after=12)
    pPr = border_p._element.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="4" w:color="003366"/></w:pBdr>')
    pPr.append(pBdr)

    # ==========================================
    # ## Misión de Kinal
    # ==========================================
    h2_mision = doc.add_paragraph()
    style_p(h2_mision, space_before=10, space_after=4)
    run_h2_mision = h2_mision.add_run("Misión de Kinal")
    run_h2_mision.font.name = "Calibri"
    run_h2_mision.font.size = Pt(14)
    run_h2_mision.font.bold = True
    run_h2_mision.font.color.rgb = COLOR_PRIMARY

    p_mision = doc.add_paragraph()
    style_p(p_mision, space_before=0, space_after=12)
    run_mision = p_mision.add_run("«Formar a jóvenes y adultos a través de una educación integral, con énfasis en las áreas técnicas y tecnológicas, influyendo positivamente en su trabajo, su familia y la sociedad».")
    run_mision.font.name = "Calibri"
    run_mision.font.size = Pt(11)
    run_mision.font.italic = True
    run_mision.font.color.rgb = COLOR_DARK

    # ==========================================
    # ## Visión de Kinal
    # ==========================================
    h2_vision = doc.add_paragraph()
    style_p(h2_vision, space_before=10, space_after=4)
    run_h2_vision = h2_vision.add_run("Visión de Kinal")
    run_h2_vision.font.name = "Calibri"
    run_h2_vision.font.size = Pt(14)
    run_h2_vision.font.bold = True
    run_h2_vision.font.color.rgb = COLOR_PRIMARY

    p_vision = doc.add_paragraph()
    style_p(p_vision, space_before=0, space_after=12)
    run_vision = p_vision.add_run("Ser líderes en la formación técnica, tecnológica y humana de la región, propiciando la superación personal, laboral y social de nuestros estudiantes con un alto sentido ético y excelencia profesional.")
    run_vision.font.name = "Calibri"
    run_vision.font.size = Pt(11)
    run_vision.font.color.rgb = COLOR_DARK

    # ==========================================
    # ## Competencia del curso
    # ==========================================
    h2_comp = doc.add_paragraph()
    style_p(h2_comp, space_before=12, space_after=6)
    run_h2_comp = h2_comp.add_run("Competencia del curso")
    run_h2_comp.font.name = "Calibri"
    run_h2_comp.font.size = Pt(14)
    run_h2_comp.font.bold = True
    run_h2_comp.font.color.rgb = COLOR_PRIMARY

    items_generales = [
        ("Competencia general:", "Implementa medidas de aseguramiento y bastionado en sistemas informáticos, correlaciona eventos de seguridad en plataformas SIEM y ejecuta protocolos de contención ante incidentes cibernéticos, actuando con estricto apego al ideario ético de Kinal y las mejores prácticas de la industria."),
        ("Duración total en horas:", "80 horas pedagógicas totales (64 horas en modalidad virtual sincrónica/asincrónica y 16 horas en modalidad presencial práctica)."),
        ("Horario:", "Por convenir (Sesiones virtuales de 3 horas, 2 días entre semana; sesiones presenciales intensivas de 8 horas en día sábado)."),
        ("Modalidad:", "Híbrida: 80% virtual (64 horas) a través de plataforma educativa y laboratorios en la nube, y 20% presencial (16 horas) en los laboratorios de cómputo y redes de Fundación Kinal."),
        ("Perfil de ingreso:", "Dirigido a jóvenes y adultos desempleados de 18 a 45 años con conocimientos previos en informática, administración básica de sistemas o redes, motivados a reconvertirse laboralmente hacia la ciberdefensa con compromiso ético y disciplina."),
        ("Perfil de egreso:", "El egresado opera como Analista SOC Junior o Técnico de Ciberseguridad, capaz de monitorear telemetría, realizar triaje de alertas de seguridad, aplicar hardening en Windows y Linux, analizar amenazas de malware o phishing y ejecutar contención básica de incidentes.")
    ]

    for label, text in items_generales:
        p_item = doc.add_paragraph()
        style_p(p_item, space_before=0, space_after=6)
        run_lbl = p_item.add_run(f"**{label}** ")
        run_lbl.font.name = "Calibri"
        run_lbl.font.size = Pt(11)
        run_lbl.font.bold = True
        run_lbl.font.color.rgb = COLOR_PRIMARY

        run_txt = p_item.add_run(text)
        run_txt.font.name = "Calibri"
        run_txt.font.size = Pt(11)
        run_txt.font.color.rgb = COLOR_DARK

    # ==========================================
    # ## Temario:
    # ==========================================
    doc.add_page_break()

    h2_temario = doc.add_paragraph()
    style_p(h2_temario, space_before=10, space_after=8)
    run_h2_temario = h2_temario.add_run("Temario:")
    run_h2_temario.font.name = "Calibri"
    run_h2_temario.font.size = Pt(16)
    run_h2_temario.font.bold = True
    run_h2_temario.font.color.rgb = COLOR_PRIMARY

    modulos = [
        {
            "num": "1",
            "nombre": "Fundamentos de Redes Seguras, Protocolos y Arquitectura de Ciberdefensa (16 Horas: 12h Virtuales / 4h Presenciales)",
            "indicador": "Analiza el tráfico de red con Wireshark e implementa segmentación lógica mediante firewalls y listas de control de acceso para mitigar vectores de intrusión perimetral en la infraestructura corporativa.",
            "temas": [
                {
                    "titulo": "Conceptos Fundamentales de Ciberseguridad y Amenazas en Guatemala",
                    "subtemas": [
                        "Tríada CIA (Confidencialidad, Integridad y Disponibilidad) y modelo de no repudio",
                        "Diferenciación técnica: amenaza, vulnerabilidad, riesgo e impacto en el negocio",
                        "Panorama de ciberataques en el sector bancario, retail y telecomunicaciones de Guatemala"
                    ]
                },
                {
                    "titulo": "Arquitectura TCP/IP y Análisis de Protocolos de Red",
                    "subtemas": [
                        "Protocolos vulnerables vs. alternos seguros: HTTP vs. HTTPS, Telnet vs. SSH, FTP vs. SFTP",
                        "Ataques basados en protocolos: ARP Spoofing, DNS Poisoning y rogue DHCP",
                        "Captura e inspección profunda de paquetes de red mediante Wireshark y tcpdump"
                    ]
                },
                {
                    "titulo": "Defensa Perimetral y Segmentación Lógica de Red",
                    "subtemas": [
                        "Diseño de zonas de red: Red Interna (LAN), Zona Desmilitarizada (DMZ) y enlaces WAN",
                        "Principios del modelo Zero Trust y microsegmentación mediante VLANs de seguridad",
                        "Configuración de reglas de filtrado de paquetes y NAT en firewalls pfSense / Cisco"
                    ]
                },
                {
                    "titulo": "Taller Presencial de Cableado Seguro y Configuración de Firewalls (4 Horas en Kinal)",
                    "subtemas": [
                        "Montaje físico de infraestructura de red en racks de laboratorio de Kinal",
                        "Conexión física y configuración de interfaces LAN/WAN/DMZ en appliance de seguridad",
                        "Evaluación práctica del Módulo 1 (Prueba de filtrado y bloqueo de escaneos de puertos)"
                    ]
                }
            ]
        },
        {
            "num": "2",
            "nombre": "Hardening de Sistemas Operativos, Control de Identidades y Accesos (16 Horas: 12h Virtuales / 4h Presenciales)",
            "indicador": "Aplica directivas de bastionado (hardening) en servidores Windows y Linux, configurando autenticación multifactor y el principio de mínimo privilegio para reducir la superficie de ataque corporativa.",
            "temas": [
                {
                    "titulo": "Bastionado (Hardening) de Servidores Linux Corporativos",
                    "subtemas": [
                        "Gestión avanzada de permisos de archivos (chmod, chown, SUID/SGID) y auditoría con LinPEAS",
                        "Aseguramiento del servicio OpenSSH: autenticación por llave pública y deshabilitación de root",
                        "Configuración de cortafuegos local (UFW / iptables) y desactivación de servicios obsoletos"
                    ]
                },
                {
                    "titulo": "Seguridad y Directivas en Entornos Windows y Active Directory",
                    "subtemas": [
                        "Implementación de Directivas de Grupo Locales (GPO) y mitigación de elevación de privilegios",
                        "Control de Cuentas de Usuario (UAC) y auditoría de eventos de inicio de sesión",
                        "Aplicación de recomendaciones de seguridad basadas en las guías oficiales CIS Benchmarks"
                    ]
                },
                {
                    "titulo": "Gestión de Identidades, Accesos (IAM) y Criptografía Práctica",
                    "subtemas": [
                        "Principio de menor privilegio (PoLP) y Control de Acceso Basado en Roles (RBAC)",
                        "Autenticación Multifactor (MFA) y protección contra ataques de fuerza bruta y pulverización de contraseñas",
                        "Criptografía aplicada: algoritmos simétricos (AES), asimétricos (RSA) y funciones hash (SHA-256)"
                    ]
                },
                {
                    "titulo": "Taller Presencial de Hardening en Vivo y Pruebas de Cumplimiento (4 Horas en Kinal)",
                    "subtemas": [
                        "Despliegue de máquinas virtuales vulnerables en laboratorio local de Kinal",
                        "Aplicación contra reloj de scripts de hardening y remediación de hallazgos",
                        "Evaluación práctica del Módulo 2 (Auditoría de cumplimiento de bastionado ≥ 75 pts)"
                    ]
                }
            ]
        },
        {
            "num": "3",
            "nombre": "Operaciones de Seguridad (SOC Nivel 1), Análisis de Logs y Monitoreo SIEM (16 Horas Virtuales)",
            "indicador": "Correlaciona registros de auditoría (logs) y gestiona alertas de seguridad en una plataforma SIEM clasificando eventos según severidad y descartando falsos positivos en entornos simulados.",
            "temas": [
                {
                    "titulo": "Arquitectura y Flujo de Trabajo en un SOC Moderno",
                    "subtemas": [
                        "Estructura y roles de un SOC: Nivel 1 (Triaje), Nivel 2 (Investigación) y Nivel 3 (Threat Hunting)",
                        "Ciclo de vida de una alerta de seguridad: recepción, enriquecimiento, clasificación y escalamiento",
                        "Uso de plataformas de gestión de incidentes y sistemas de tickets para trazabilidad de eventos"
                    ]
                },
                {
                    "titulo": "Generación, Recolección y Análisis de Logs",
                    "subtemas": [
                        "Event Viewer de Windows: Análisis de Event IDs críticos (4624, 4625, 4672, 4720, 7045)",
                        "Registros de auditoría en Linux: Syslog, auth.log y registros de servidores web Apache/Nginx",
                        "Normalización y centralización de registros mediante agentes de recolección de eventos"
                    ]
                },
                {
                    "titulo": "Despliegue y Operación Práctica en Plataforma SIEM (Wazuh / Splunk)",
                    "subtemas": [
                        "Instalación de agentes SIEM en clientes Windows y Linux en laboratorio virtual",
                        "Creación de tableros de visualización, filtros de búsqueda y reglas de detección personalizadas",
                        "Correlación en tiempo real de eventos de autenticación sospechosa y modificaciones no autorizadas"
                    ]
                },
                {
                    "titulo": "Triaje de Alertas y Gestión de Falsos Positivos",
                    "subtemas": [
                        "Criterios técnicos para distinguir falsos positivos de intrusiones verídicas",
                        "Enriquecimiento de alertas con Indicadores de Compromiso (IOCs) y bases de reputación",
                        "Evaluación práctica del Módulo 3 (Simulación de turno de guardia SOC y reporte formal de tickets)"
                    ]
                }
            ]
        },
        {
            "num": "4",
            "nombre": "Detección de Amenazas, Análisis de Malware, Phishing y Vulnerabilidades (16 Horas: 12h Virtuales / 4h Presenciales)",
            "indicador": "Identifica campañas de ingeniería social, analiza cabeceras de correos fraudulentos y ejecuta escaneos de vulnerabilidades emitiendo recomendaciones de mitigación técnica viables.",
            "temas": [
                {
                    "titulo": "Anatomía del Malware y Vectores de Infección",
                    "subtemas": [
                        "Clasificación técnica de malware: Ransomware, Troyanos, InfoStealers, Gusanos y Rootkits",
                        "Técnicas de evasión de antivirus: ofuscación, empaquetado y persistencia en el registro",
                        "Análisis estático seguro: cálculo de hashes, extracción de cadenas (strings) e inspección en VirusTotal"
                    ]
                },
                {
                    "titulo": "Análisis de Correos Maliciosos y Campañas de Phishing",
                    "subtemas": [
                        "Inspección de cabeceras SMTP (Received, Return-Path, Message-ID) para rastrear origen real",
                        "Mecanismos de autenticación de correo corporativo: SPF, DKIM y políticas DMARC",
                        "Extracción segura de enlaces fraudulentos y análisis de adjuntos mediante PhishTool y URLhaus"
                    ]
                },
                {
                    "titulo": "Gestión y Escaneo de Vulnerabilidades Corporativas",
                    "subtemas": [
                        "Ciclo de vida de la vulnerabilidad: descubrimiento, evaluación, priorización y remediación",
                        "Puntuación de severidad según el estándar CVSS (Common Vulnerability Scoring System)",
                        "Ejecución de escaneos de vulnerabilidades autenticados y no autenticados con OpenVAS / Greenbone"
                    ]
                },
                {
                    "titulo": "Taller Presencial de Detonación en Sandbox y Análisis de Amenazas (4 Horas en Kinal)",
                    "subtemas": [
                        "Configuración de entorno de análisis en máquinas virtuales aisladas (Sandbox)",
                        "Detonación controlada de muestras de malware didáctico y monitoreo de procesos anómalos",
                        "Evaluación práctica del Módulo 4 (Informe técnico de análisis de phishing y plan de remediación)"
                    ]
                }
            ]
        },
        {
            "num": "5",
            "nombre": "Respuesta a Incidentes, Simulación de Crisis (Blue Team) y Proyecto Capstone (16 Horas: 12h Virtuales / 4h Presenciales)",
            "indicador": "Ejecuta las fases de contención, erradicación y documentación de un incidente de seguridad digital conforme al marco NIST SP 800-61, sustentando el informe técnico de cierre y lecciones aprendidas.",
            "temas": [
                {
                    "titulo": "Marco de Respuesta a Incidentes (NIST SP 800-61 Rev. 2)",
                    "subtemas": [
                        "Fases estándar: Preparación, Detección y Análisis, Contención, Erradicación, Recuperación y Lecciones",
                        "Técnicas de aislamiento urgente: desconexión de red, terminación de procesos y bloqueo en firewall",
                        "Preservación básica de evidencia digital y principios de cadena de custodia"
                    ]
                },
                {
                    "titulo": "Estrategias de Ciberdefensa (Blue Team) y Mitigación de Ransomware",
                    "subtemas": [
                        "Planes de contingencia ante ataques de secuestro de datos (Ransomware)",
                        "Estrategias de respaldo inmutable (Regla 3-2-1) y pruebas periódicas de restauración",
                        "Comunicación efectiva y gestión de crisis institucional bajo el principio de servicio de Kinal"
                    ]
                },
                {
                    "titulo": "Preparación para la Inserción Laboral y Certificaciones",
                    "subtemas": [
                        "Estructuración de hoja de vida técnica para puestos de Analista SOC Junior y Soporte de Seguridad",
                        "Simulación de entrevistas técnicas y preguntas habituales de empleadores en Guatemala",
                        "Ruta de certificación: preparación para exámenes CompTIA Security+ y Cisco CCST Cybersecurity"
                    ]
                },
                {
                    "titulo": "Taller Presencial de Sala de Crisis, Reto Capstone / CTF y Clausura (4 Horas en Kinal)",
                    "subtemas": [
                        "Simulación presencial de ataque coordinado en tiempo real sobre red empresarial simulada",
                        "Contención, erradicación y restauración de servicios por parte de células Blue Team de estudiantes",
                        "Evaluación Terminal Integradora, defensa oral ante panel técnico de Kinal y entrega de diplomas"
                    ]
                }
            ]
        }
    ]

    for mod in modulos:
        h3_mod = doc.add_paragraph()
        style_p(h3_mod, space_before=14, space_after=4)
        run_h3 = h3_mod.add_run(f"Módulo {mod['num']}: {mod['nombre']}")
        run_h3.font.name = "Calibri"
        run_h3.font.size = Pt(13)
        run_h3.font.bold = True
        run_h3.font.color.rgb = COLOR_PRIMARY

        p_ind = doc.add_paragraph()
        style_p(p_ind, space_before=0, space_after=6)
        run_ind_lbl = p_ind.add_run("**Indicador de logro:** ")
        run_ind_lbl.font.name = "Calibri"
        run_ind_lbl.font.size = Pt(11)
        run_ind_lbl.font.bold = True
        run_ind_lbl.font.color.rgb = COLOR_SECONDARY

        run_ind_txt = p_ind.add_run(mod["indicador"])
        run_ind_txt.font.name = "Calibri"
        run_ind_txt.font.size = Pt(11)
        run_ind_txt.font.color.rgb = COLOR_DARK

        for t in mod["temas"]:
            p_tema = doc.add_paragraph()
            style_p(p_tema, space_before=3, space_after=2)
            p_tema.paragraph_format.left_indent = Inches(0.25)
            run_t_bullet = p_tema.add_run("• Tema: ")
            run_t_bullet.font.name = "Calibri"
            run_t_bullet.font.size = Pt(11)
            run_t_bullet.font.bold = True
            run_t_bullet.font.color.rgb = COLOR_PRIMARY

            run_t_title = p_tema.add_run(t["titulo"])
            run_t_title.font.name = "Calibri"
            run_t_title.font.size = Pt(11)
            run_t_title.font.bold = True
            run_t_title.font.color.rgb = COLOR_DARK

            for sub in t["subtemas"]:
                p_sub = doc.add_paragraph()
                style_p(p_sub, space_before=1, space_after=2)
                p_sub.paragraph_format.left_indent = Inches(0.55)
                run_sub_dash = p_sub.add_run("– ")
                run_sub_dash.font.name = "Calibri"
                run_sub_dash.font.size = Pt(10.5)
                run_sub_dash.font.color.rgb = COLOR_SECONDARY

                run_sub_txt = p_sub.add_run(sub)
                run_sub_txt.font.name = "Calibri"
                run_sub_txt.font.size = Pt(10.5)
                run_sub_txt.font.color.rgb = COLOR_DARK

    # ==========================================
    # ## Materiales y Herramientas Necesarias
    # ==========================================
    doc.add_page_break()

    h2_mat = doc.add_paragraph()
    style_p(h2_mat, space_before=12, space_after=6)
    run_h2_mat = h2_mat.add_run("Materiales y Herramientas Necesarias (Precios y Licenciamiento)")
    run_h2_mat.font.name = "Calibri"
    run_h2_mat.font.size = Pt(16)
    run_h2_mat.font.bold = True
    run_h2_mat.font.color.rgb = COLOR_PRIMARY

    p_mat_intro = doc.add_paragraph()
    style_p(p_mat_intro, space_before=0, space_after=8)
    run_intro = p_mat_intro.add_run("El curso ha sido diseñado con un enfoque costo-eficiente basado en software libre y de código abierto (Open Source), complementado con herramientas gratuitas del ecosistema Cisco Networking Academy disponible para Fundación Kinal. A continuación se detalla la lista de herramientas requeridas y su esquema de costos:")
    run_intro.font.name = "Calibri"
    run_intro.font.size = Pt(11)
    run_intro.font.color.rgb = COLOR_DARK

    categorias_herramientas = [
        ("Herramientas Open Source Principales (Costo $0.00 / Gratuitas):", [
            "Wazuh (SIEM & XDR Open Source): Plataforma completa para monitoreo de seguridad, recolección y análisis de logs, monitoreo de integridad de archivos (FIM) y detección de vulnerabilidades. Costo: $0.00 (100% Gratuito y de código abierto).",
            "Wireshark: Analizador de protocolos de red estándar en la industria para inspección profunda de paquetes y diagnóstico de tráfico anómalo. Costo: $0.00 (Gratuito / Open Source).",
            "Oracle VirtualBox: Hipervisor de virtualización local tipo 2 para montaje de laboratorios aislados en la computadora del estudiante. Costo: $0.00 (Gratuito bajo licencia GPLv3).",
            "Distribuciones Linux de Laboratorio (Kali Linux, Ubuntu Server 22.04 LTS, REMnux): Sistemas operativos para ejercicios de defensa y análisis de incidentes. Costo: $0.00 (Gratuitos).",
            "pfSense Community Edition / OPNsense: Firewall y enrutador virtual de código abierto para prácticas de filtrado de paquetes y segmentación DMZ. Costo: $0.00 (Gratuito).",
            "Greenbone Community Edition (OpenVAS): Escáner de vulnerabilidades de red de nivel profesional para detección de fallos y parches pendientes. Costo: $0.00 (Gratuito / Open Source).",
            "LinPEAS y WinPEAS: Scripts comunitarios de auditoría de seguridad para identificación de vectores de escalada de privilegios y malas configuraciones. Costo: $0.00 (Gratuito en GitHub).",
            "CyberChef: Plataforma web de análisis y decodificación de datos (Base64, Hex, Hashing, compresión) creada por el GCHQ británico. Costo: $0.00 (Gratuito / Open Source).",
            "Plataformas Gratuitas de Inteligencia de Amenazas: VirusTotal, AbuseIPDB, URLhaus, Shodan (nivel académico) y Talos Intelligence. Costo: $0.00 (Acceso comunitario gratuito)."
        ]),
        ("Herramientas del Ecosistema Cisco y Esquema de Precios:", [
            "Cisco Packet Tracer: Simulador oficial de redes, routers, switches y firewalls Cisco ASA. Costo: $0.00 / mes (100% Gratuito mediante registro institucional en Cisco Networking Academy / Skills for All).",
            "Cisco Networking Academy (NetAcad) / Skills for All: Currícula oficial y laboratorios interactivos (Cybersecurity Essentials, Network Defense, Endpoint Security). Costo: $0.00 / mes (Gratuito para Academias Cisco registradas como Fundación Kinal).",
            "Snort IDS/IPS (mantenido por Cisco Talos): Sistema de detección y prevención de intrusiones de código abierto basado en firmas. Costo: $0.00 / mes (Open Source Gratuito con reglas comunitarias). Suscripción opcional a reglas de pago en tiempo real (Personal Subscriber): ~$2.42 USD / mes ($29 USD/año), no indispensable para fines didácticos.",
            "Cisco Modeling Labs (CML) - Personal Edition (Herramienta Opcional Avanzada): Emulador virtual avanzado de routers y switches Cisco para topologías empresariales complejas. Costo: $199.00 USD al año (~$16.58 USD / mes) por licencia de instructor. (No requerida para los estudiantes si se utiliza Packet Tracer o GNS3/EVE-NG gratuitos).",
            "Plataforma de Laboratorios en Nube TryHackMe (Opcional): Salas públicas de aprendizaje en ciberseguridad. Nivel básico: $0.00 / mes (Gratuito). Nivel premium educativo opcional para acceso ilimitado a máquinas virtuales en la nube: ~$10.00 a $14.00 USD / mes por usuario."
        ]),
        ("Infraestructura de Laboratorio Presencial en Kinal (16 Horas):", [
            "Laboratorio de Cómputo Kinal: Computadoras con procesador Core i5/i7 o equivalente, 16 GB de RAM, 256 GB SSD y soporte de virtualización habilitado en BIOS.",
            "Equipamiento de Red Físico de Taller: Switches administrables capa 2/3, routers de laboratorio, cables de red UTP Cat 6 ponchados y regletas eléctricas con protección de sobretensión.",
            "Equipo de Protección y Seguridad: Protocolos de conexión en red aislada de pruebas (Air-Gapped o VLAN privada sin salida a internet corporativo para pruebas de malware)."
        ])
    ]

    for cat_title, items in categorias_herramientas:
        p_cat = doc.add_paragraph()
        style_p(p_cat, space_before=8, space_after=3)
        run_cat = p_cat.add_run(f"• {cat_title}")
        run_cat.font.name = "Calibri"
        run_cat.font.size = Pt(11.5)
        run_cat.font.bold = True
        run_cat.font.color.rgb = COLOR_PRIMARY

        for it in items:
            p_it = doc.add_paragraph()
            style_p(p_it, space_before=1, space_after=2)
            p_it.paragraph_format.left_indent = Inches(0.35)
            run_dash = p_it.add_run("– ")
            run_dash.font.name = "Calibri"
            run_dash.font.size = Pt(10)
            run_dash.font.color.rgb = COLOR_SECONDARY

            run_text = p_it.add_run(it)
            run_text.font.name = "Calibri"
            run_text.font.size = Pt(10)
            run_text.font.color.rgb = COLOR_DARK

    section = doc.sections[0]
    footer = section.footer
    footer_p = footer.paragraphs[0]
    style_p(footer_p, space_before=0, space_after=0)
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_f = footer_p.add_run("Fundación Kinal | Fundamentos de Ciberseguridad — Temario Oficial")
    run_f.font.name = "Calibri"
    run_f.font.size = Pt(9)
    run_f.font.color.rgb = COLOR_SECONDARY

    out_dir = os.path.join("Cursos", "Ciberseguridad")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "Temario_Curso_Ciberseguridad_Kinal.docx")
    doc.save(out_path)
    print(f"Temario guardado con éxito en: {out_path}")

if __name__ == "__main__":
    create_temario()
