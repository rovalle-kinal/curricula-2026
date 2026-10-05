import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=70, bottom=70, left=90, right=90):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E0", sz="4"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def style_p(p, space_before=0, space_after=4, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def style_heading(p, text, level=1, color=RGBColor(0, 51, 102)):
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.bold = True
    r.font.color.rgb = color
    if level == 1:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        r.font.size = Pt(15)
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        r.font.size = Pt(13)
    elif level == 3:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r.font.size = Pt(11)

def add_callout(doc, text, bold_prefix="", fill_hex="F7FAFC", border_color="003366"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    
    cell = tbl.cell(0, 0)
    set_cell_background(cell, fill_hex)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=120)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 51, 102)
        
    r_txt = p.add_run(text)
    r_txt.font.name = "Calibri"
    r_txt.font.size = Pt(10)
    r_txt.font.italic = True
    r_txt.font.color.rgb = RGBColor(45, 55, 72)
    
    p_post = doc.add_paragraph()
    style_p(p_post, space_before=0, space_after=6)

def generate_dosificacion():
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    COLOR_PRIMARY = RGBColor(0, 51, 102)
    COLOR_SECONDARY = RGBColor(74, 85, 104)
    COLOR_DARK = RGBColor(26, 32, 44)

    # =========================================================================
    # PORTADA / ENCABEZADO TÉCNICO
    # =========================================================================
    p_inst = doc.add_paragraph()
    style_p(p_inst, space_before=20, space_after=2)
    r_inst = p_inst.add_run("FUNDACIÓN KINAL")
    r_inst.font.name = "Calibri"
    r_inst.font.size = Pt(14)
    r_inst.font.bold = True
    r_inst.font.color.rgb = COLOR_PRIMARY

    p_esc = doc.add_paragraph()
    style_p(p_esc, space_before=0, space_after=20)
    r_esc = p_esc.add_run("Escuela Técnica Superior — Coordinación de Formación Continua y Empleabilidad")
    r_esc.font.name = "Calibri"
    r_esc.font.size = Pt(11)
    r_esc.font.color.rgb = COLOR_SECONDARY

    p_tit = doc.add_paragraph()
    style_p(p_tit, space_before=10, space_after=8)
    r_tit = p_tit.add_run("DOSIFICACIÓN CURRICULAR Y SECUENCIA DIDÁCTICA DETALLADA")
    r_tit.font.name = "Calibri"
    r_tit.font.size = Pt(20)
    r_tit.font.bold = True
    r_tit.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    style_p(p_sub, space_before=0, space_after=18)
    r_sub = p_sub.add_run("Fundamentos de Ciberseguridad y Operaciones SOC (Modalidad Híbrida: 64h Virtuales / 16h Presenciales — 80 Horas)")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True
    r_sub.font.color.rgb = COLOR_SECONDARY

    add_callout(doc, 
        "«El aprender en el proceso del trabajo es base y a la vez objetivo de la acción formativa». Las 80 horas del curso se distribuyen en 21 sesiones virtuales de 3 horas (más 1 hora de asesoría técnica final = 64 horas virtuales) y 2 talleres presenciales de 8 horas los sábados en laboratorios de Kinal (16 horas presenciales). Los horarios están por convenir y no se fijan fechas estimadas de calendario. Cada sesión se desarrolla bajo los tres momentos didácticos: Apertura, Desarrollo y Cierre. Nota mínima aprobatoria institucional: 75 puntos sobre 100.",
        bold_prefix="Estructura Didáctica Dual y Formato de Horarios:")

    # =========================================================================
    # 1. MATRIZ GENERAL DE DOSIFICACIÓN HÍBRIDA (TABLA SINÓPTICA DE SESIONES)
    # =========================================================================
    h1_1 = doc.add_paragraph()
    style_heading(h1_1, "1. Matriz General de Dosificación de Sesiones Híbridas (80 Horas)", level=1)

    sesiones_data = [
        ("Sesión V-01", "Virtual (3h)", "Módulo 1: Redes", "Tríada CIA, vectores de ataque en Guatemala y modelo TCP/IP", "Plataforma / NetAcad"),
        ("Sesión V-02", "Virtual (3h)", "Módulo 1: Redes", "Protocolos vulnerables vs seguros (SSH, HTTPS) y ARP spoofing", "VirtualBox / Wireshark"),
        ("Sesión V-03", "Virtual (3h)", "Módulo 1: Redes", "Captura profunda de paquetes en Wireshark y análisis de puertos", "Wireshark / Kali"),
        ("Sesión V-04", "Virtual (3h)", "Módulo 1: Redes", "Segmentación lógica, DMZ y modelo Zero Trust en Packet Tracer", "Cisco Packet Tracer"),
        ("Sesión V-05", "Virtual (3h)", "Módulo 2: Hardening", "Hardening de Linux: permisos, usuarios y auditoría con LinPEAS", "Ubuntu Server VM"),
        ("Sesión V-06", "Virtual (3h)", "Módulo 2: Hardening", "Aseguramiento de SSH, firewall UFW y bloqueo de accesos root", "Linux / OpenSSH"),
        ("Sesión V-07", "Virtual (3h)", "Módulo 2: Hardening", "Seguridad en Windows: Directivas GPO locales y control UAC", "Windows Server VM"),
        ("Sesión V-08", "Virtual (3h)", "Módulo 2: Hardening", "Gestión de Identidades (IAM), MFA y Criptografía práctica (AES/SHA)", "CyberChef / MFA Labs"),
        ("Sábado P-01", "Presencial (8h)", "Módulos 1 y 2", "Taller de cableado seguro, routers/firewalls físicos y bastionado en vivo", "Lab Físico Kinal"),
        ("Sesión V-09", "Virtual (3h)", "Módulo 3: SOC / SIEM", "Estructura de un SOC, roles N1/N2/N3 y ciclo de vida de alertas", "Plataforma / Diapositivas"),
        ("Sesión V-10", "Virtual (3h)", "Módulo 3: SOC / SIEM", "Análisis de Event IDs críticos en Windows y Syslog de Linux", "Event Viewer / Logs"),
        ("Sesión V-11", "Virtual (3h)", "Módulo 3: SOC / SIEM", "Despliegue de Wazuh SIEM: instalación de agentes y recolección", "Wazuh Server VM"),
        ("Sesión V-12", "Virtual (3h)", "Módulo 3: SOC / SIEM", "Tableros de visualización, filtros de búsqueda y reglas en Wazuh", "Wazuh Dashboard"),
        ("Sesión V-13", "Virtual (3h)", "Módulo 3: SOC / SIEM", "Correlación de eventos de autenticación sospechosa y fuerza bruta", "Wazuh Rules"),
        ("Sesión V-14", "Virtual (3h)", "Módulo 3: SOC / SIEM", "Triaje de alertas, enriquecimiento con IOCs y tickets de incidentes", "TheHive / Tickets"),
        ("Sesión V-15", "Virtual (3h)", "Módulo 4: Amenazas", "Clasificación de malware: Ransomware, Troyanos y análisis de hashes", "VirusTotal / Sandbox"),
        ("Sesión V-16", "Virtual (3h)", "Módulo 4: Amenazas", "Análisis de phishing: cabeceras SMTP, SPF, DKIM y DMARC", "PhishTool / Web"),
        ("Sesión V-17", "Virtual (3h)", "Módulo 4: Amenazas", "Gestión y escaneo de vulnerabilidades con OpenVAS / Greenbone", "OpenVAS / Greenbone"),
        ("Sesión V-18", "Virtual (3h)", "Módulo 4: Amenazas", "Priorización de parches según criticidad CVSS y plan de remediación", "Bases CVE / Reportes"),
        ("Sesión V-19", "Virtual (3h)", "Módulo 5: Incidentes", "Marco NIST SP 800-61: contención, erradicación y preservación", "NIST Framework"),
        ("Sesión V-20", "Virtual (3h)", "Módulo 5: Incidentes", "Estrategias de Blue Team, mitigación de ransomware y respaldos 3-2-1", "Casos Reales / Labs"),
        ("Sesión V-21", "Virtual (3h)", "Módulo 5: Incidentes", "Armado de CV técnico para SOC Jr y simulación de entrevistas", "Taller Empleabilidad"),
        ("Sesión V-22", "Virtual (1h)", "Módulo 5: Cierre", "Asesoría técnica final, preparación de certificaciones y consultas", "Sesión de Consultas"),
        ("Sábado P-02", "Presencial (8h)", "Módulos 4 y 5", "Taller de sala de crisis SOC, Reto Capstone / CTF y Evaluación Terminal", "Lab Físico Kinal")
    ]

    tbl_matriz = doc.add_table(rows=len(sesiones_data) + 1, cols=5)
    tbl_matriz.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_matriz.autofit = False
    tbl_matriz.columns[0].width = Inches(1.1)
    tbl_matriz.columns[1].width = Inches(1.0)
    tbl_matriz.columns[2].width = Inches(1.2)
    tbl_matriz.columns[3].width = Inches(2.2)
    tbl_matriz.columns[4].width = Inches(1.0)
    set_table_borders(tbl_matriz, color="CBD5E0", sz="4")

    hdr_m = tbl_matriz.rows[0]
    for c in hdr_m.cells:
        set_cell_background(c, "003366")
    headers_mat = ["Sesión", "Modalidad", "Módulo", "Contenido Temático Central", "Plataforma / Lab"]
    for idx, text in enumerate(headers_mat):
        p = hdr_m.cells[idx].paragraphs[0]
        style_p(p, space_before=0, space_after=0)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for idx, (ses, mod, eje, cont, plat) in enumerate(sesiones_data, start=1):
        row = tbl_matriz.rows[idx]
        if "Presencial" in mod:
            for c in row.cells:
                set_cell_background(c, "EBF8FF")
        elif idx % 2 == 1:
            for c in row.cells:
                set_cell_background(c, "F7FAFC")
        for c_idx, val in enumerate([ses, mod, eje, cont, plat]):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=40, bottom=40, left=50, right=50)
            p = cell.paragraphs[0]
            style_p(p, space_before=0, space_after=0)
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(8)
            r.font.color.rgb = COLOR_DARK
            if c_idx in [0, 1]:
                r.font.bold = True

    doc.add_page_break()

    # =========================================================================
    # 2. MICRODISEÑO DIDÁCTICO DETALLADO: SESIONES VIRTUALES Y TALLERES PRESENCIALES
    # =========================================================================
    h1_2 = doc.add_paragraph()
    style_heading(h1_2, "2. Microdiseño Didáctico de Sesiones Virtuales y Talleres Presenciales", level=1)

    p_micro_desc = doc.add_paragraph()
    style_p(p_micro_desc, space_before=0, space_after=8)
    r = p_micro_desc.add_run("A continuación se desglosan los momentos didácticos (Apertura, Desarrollo y Cierre) para las sesiones virtuales de 3 horas y los 2 talleres presenciales intensivos de 8 horas en Kinal:")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_DARK

    bloques_detalle = [
        {
            "titulo": "ESTRUCTURA DE LAS SESIONES VIRTUALES (3 HORAS / 180 MINUTOS)",
            "descripcion": "Aplicable a las sesiones virtuales sincrónicas entre semana (V-01 a V-21):",
            "momentos": [
                ("Apertura / Inicio (20 min)", "Puntualidad y pase de lista digital. Presentación de un caso real o incidente de seguridad en empresas guatemaltecas relacionado con el tema del día. Activación de conocimientos previos y planteamiento de la pregunta desafiante de la sesión."),
                ("Desarrollo / Laboratorio Práctico (140 min)", "Bloque 1 (40 min): Demostración conceptual y técnica guiada por el instructor en la máquina virtual o plataforma.\nBloque 2 (70 min): Ejecución individual del laboratorio práctico por parte de los alumnos (captura Wireshark, hardening Linux, correlación Wazuh, etc.) con asistencia técnica en vivo.\nBloque 3 (30 min): Puesta en común, resolución de errores comunes y discusión de impacto en entornos de producción."),
                ("Cierre / Consolidación (20 min)", "Síntesis de aprendizajes clave y validación de logro. Asignación de la actividad de registro en la Bitácora de Incidentes (Berichtsheft) y recomendaciones de autoestudio.")
            ]
        },
        {
            "titulo": "SÁBADO PRESENCIAL 1 (8 HORAS / 480 MINUTOS EN KINAL): REDES SEGURAS Y HARDENING EN VIVO",
            "descripcion": "Integra las 4 horas prácticas del Módulo 1 y las 4 horas prácticas del Módulo 2:",
            "momentos": [
                ("Apertura (45 min)", "Bienvenida en el laboratorio de cómputo y redes de Kinal. Charla de seguridad lógica y física. Firma del Acuerdo de Uso Ético (NDA didáctico). Organización de parejas de trabajo y asignación de puestos y racks."),
                ("Desarrollo - Bloque Mañana: Redes Seguras (195 min)", "Cableado físico de switches y routers de laboratorio. Configuración de interfaces LAN, WAN y DMZ en appliances de seguridad pfSense / Cisco. Creación de reglas de firewall, bloqueo de tráfico ICMP/Telnet y prueba de escaneo con Nmap desde segmento externo para verificar efectividad."),
                ("Receso Formativo / Almuerzo (45 min)", "Almuerzo y convivencia fraterna en instalaciones de Fundación Kinal."),
                ("Desarrollo - Bloque Tarde: Hardening en Vivo (150 min)", "Despliegue de dos servidores vulnerables en red local. Aplicación contra reloj de listas CIS Benchmarks y scripts de bastionado en Linux y Windows. Ejecución de escaneos LinPEAS para comprobar remediación del 100% de vectores críticos."),
                ("Cierre y Evaluación de Medio Término (45 min)", "Inspección técnica por parte del instructor. Evaluación con rúbrica práctica (≥ 75 pts). Registro de hallazgos en la bitácora física Berichtsheft y orden 5S del laboratorio.")
            ]
        },
        {
            "titulo": "SÁBADO PRESENCIAL 2 (8 HORAS / 480 MINUTOS EN KINAL): SALA DE CRISIS SOC Y CAPSTONE",
            "descripcion": "Integra las 4 horas prácticas del Módulo 4 y las 4 horas prácticas del Módulo 5:",
            "momentos": [
                ("Apertura (45 min)", "Instalación de la Sala de Crisis en el laboratorio central de Kinal. Asignación de roles en células Blue Team (Defensores de Red, Analistas de Logs y Especialistas de Contención). Presentación de la infraestructura víctima simulada."),
                ("Desarrollo - Bloque Mañana: Análisis de Malware y Phishing en Sandbox (195 min)", "Configuración de máquinas virtuales en red aislada (Air-Gapped). Extracción de artefactos sospechosos de correos simulados, detonación controlada, análisis de procesos anómalos y extracción de IOCs."),
                ("Receso Formativo / Almuerzo (45 min)", "Almuerzo y preparación de la estrategia de defensa de la célula Blue Team."),
                ("Desarrollo - Bloque Tarde: Reto Capstone / CTF de Ciberdefensa (150 min)", "Simulación de ataque en vivo a la red: inyección de tráfico malicioso y elevación de privilegios. Las células de estudiantes deben detectar la intrusión en Wazuh, aislar los endpoints comprometidos, bloquear al atacante y restaurar servicios críticos sin interrumpir la operación."),
                ("Cierre, Defensa Oral y Clausura (45 min)", "Sustentación oral del informe técnico ante jurado evaluador de Kinal. Revisión final de la Bitácora Berichtsheft. Deliberación de notas finales (aprobatorio ≥ 75 pts) y acto de clausura.")
            ]
        }
    ]

    for b in bloques_detalle:
        h2_b = doc.add_paragraph()
        style_heading(h2_b, b["titulo"], level=2)

        p_desc = doc.add_paragraph()
        style_p(p_desc, space_before=0, space_after=4)
        r = p_desc.add_run(b["descripcion"])
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.italic = True
        r.font.color.rgb = COLOR_SECONDARY

        tbl_b = doc.add_table(rows=len(b["momentos"]), cols=2)
        tbl_b.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_b.autofit = False
        tbl_b.columns[0].width = Inches(2.2)
        tbl_b.columns[1].width = Inches(4.3)
        set_table_borders(tbl_b, color="CBD5E0", sz="4")

        for r_idx, (momento, detalle) in enumerate(b["momentos"]):
            row = tbl_b.rows[r_idx]
            set_cell_background(row.cells[0], "F7FAFC")
            set_cell_margins(row.cells[0], top=60, bottom=60, left=80, right=80)
            set_cell_margins(row.cells[1], top=60, bottom=60, left=80, right=80)

            p0 = row.cells[0].paragraphs[0]
            style_p(p0, space_before=0, space_after=0)
            r0 = p0.add_run(momento)
            r0.font.name = "Calibri"
            r0.font.size = Pt(9)
            r0.font.bold = True
            r0.font.color.rgb = COLOR_PRIMARY

            p1 = row.cells[1].paragraphs[0]
            style_p(p1, space_before=0, space_after=0)
            r1 = p1.add_run(detalle)
            r1.font.name = "Calibri"
            r1.font.size = Pt(9)
            r1.font.color.rgb = COLOR_DARK

        p_sp = doc.add_paragraph()
        style_p(p_sp, space_before=0, space_after=8)

    doc.add_page_break()

    # =========================================================================
    # 3. RÚBRICA ANALÍTICA DE EVALUACIÓN PRÁCTICA INSTITUCIONAL (≥ 75 PUNTOS)
    # =========================================================================
    h1_3 = doc.add_paragraph()
    style_heading(h1_3, "3. Instrumentos y Rúbricas Analíticas de Evaluación (Normativa Kinal)", level=1)

    p_rub_desc = doc.add_paragraph()
    style_p(p_rub_desc, space_before=0, space_after=8)
    r = p_rub_desc.add_run("Toda prueba de laboratorio y proyecto de ciberdefensa se califica sobre 100 puntos, con un umbral de aprobación técnica de 75 puntos. Los criterios analíticos institucionales son:")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_DARK

    tbl_rub = doc.add_table(rows=6, cols=5)
    tbl_rub.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_rub.autofit = False
    tbl_rub.columns[0].width = Inches(1.5)
    tbl_rub.columns[1].width = Inches(1.2)
    tbl_rub.columns[2].width = Inches(1.3)
    tbl_rub.columns[3].width = Inches(1.3)
    tbl_rub.columns[4].width = Inches(1.2)
    set_table_borders(tbl_rub, color="CBD5E0", sz="4")

    hdr_r = tbl_rub.rows[0]
    for c in hdr_r.cells:
        set_cell_background(c, "003366")
    headers_rub = ["Criterio Evaluado", "Excelente\n(90 - 100 pts)", "Muy Bueno\n(80 - 89 pts)", "Aprobado Mínimo\n(75 - 79 pts)", "No Aprobado\n(< 75 pts)"]
    for idx, text in enumerate(headers_rub):
        p = hdr_r.cells[idx].paragraphs[0]
        style_p(p, space_before=0, space_after=0)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    criterios_rub = [
        ("Configuración Técnica y Seguridad\n(Peso: 25%)", 
         "Reglas de firewall y directivas de hardening 100% efectivas. Cero puertos vulnerables expuestos.", 
         "Reglas efectivas con mínimas observaciones no críticas remediadas de inmediato.", 
         "Configuración operativa básica pero con alguna regla permisiva menor que se corrige en sesión.", 
         "Configuraciones inseguras por defecto, servicios vulnerables expuestos o fallo de firewall."),
        
        ("Triaje, Detección y Análisis de Logs\n(Peso: 25%)", 
         "Identifica con precisión el vector de ataque en SIEM/logs, diferencia 100% falsos positivos y extrae IOCs.", 
         "Detecta la alerta correctamente y extrae IOCs principales con mínima demora.", 
         "Identifica el evento malicioso pero presenta dificultad menor al categorizar la severidad.", 
         "Omite alertas críticas, confunde falsos positivos o no logra interpretar los registros de auditoría."),
        
        ("Capacidad de Respuesta y Contención\n(Peso: 25%)", 
         "Aisla el incidente en menos de 10 min, preserva evidencia con hashes y restaura la operatividad.", 
         "Contiene el incidente en tiempo oportuno sin pérdida de evidencia crítica.", 
         "Logra contener la amenaza pero tarda más del tiempo establecido o afecta servicios no involucrados.", 
         "Incapaz de contener la propagación del ataque o causa caída total no programada de la red."),
        
        ("Ética Profesional, NDA y Trabajo Bien Hecho\n(Peso: 15%)", 
         "Cumplimiento absoluto del NDA, uso estrictamente ético del laboratorio, orden y respeto a normas Kinal.", 
         "Conducta ética intachable y respeto permanente al reglamento de laboratorio.", 
         "Cumple las normas éticas pero requiere recordatorios menores sobre orden en bitácoras o reportes.", 
         "Violación de políticas de uso ético, escaneos no autorizados o negligencia en el trato institucional."),
        
        ("Bitácora de Incidentes (Berichtsheft)\n(Peso: 10%)", 
         "Memoria técnica completa, capturas de pantalla con hashes SHA-256, redacción ejecutiva impecable.", 
         "Reporte claro con diagramas de topología y descripciones precisas de las acciones tomadas.", 
         "Documentación con los datos indispensables pero con redacción esquemática básica.", 
         "Bitácora incompleta, sin capturas probatorias, con omisión de fechas o datos erróneos.")
    ]

    for idx, (crit, exc, mb, ap, nap) in enumerate(criterios_rub, start=1):
        row = tbl_rub.rows[idx]
        if idx % 2 == 1:
            for c in row.cells:
                set_cell_background(c, "F7FAFC")
        for c_idx, val in enumerate([crit, exc, mb, ap, nap]):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=60, bottom=60, left=70, right=70)
            p = cell.paragraphs[0]
            style_p(p, space_before=0, space_after=0)
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(8)
            r.font.color.rgb = COLOR_DARK
            if c_idx == 0:
                r.font.bold = True

    # =========================================================================
    # 4. FORMATO GUÍA DE LA BITÁCORA DE INCIDENTES (BERICHTSHEFT)
    # =========================================================================
    h1_4 = doc.add_paragraph()
    style_heading(h1_4, "4. Formato Guía de la Bitácora de Incidentes del Aprendiz (Berichtsheft)", level=1)

    p_ber_desc = doc.add_paragraph()
    style_p(p_ber_desc, space_before=0, space_after=8)
    r = p_ber_desc.add_run("El estudiante debe registrar semanalmente cada sesión práctica en su Bitácora de Incidentes y Operaciones SOC (Berichtsheft), reflejando el aprendizaje en el puesto de trabajo:")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_DARK

    tbl_b_fmt = doc.add_table(rows=7, cols=2)
    tbl_b_fmt.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_b_fmt.autofit = False
    tbl_b_fmt.columns[0].width = Inches(2.2)
    tbl_b_fmt.columns[1].width = Inches(4.3)
    set_table_borders(tbl_b_fmt, color="CBD5E0", sz="4")

    formato_data = [
        ("Nombre del Estudiante y Carné:", "___________________________________________________________"),
        ("Sesión y Modalidad:", "Sesión No. _____  |  Modalidad: [ ] Virtual 3h   [ ] Presencial 8h"),
        ("Módulo y Herramientas Utilizadas:", "[ ] Wireshark   [ ] Wazuh SIEM   [ ] VirtualBox   [ ] Packet Tracer   [ ] OpenVAS"),
        ("Descripción Técnica del Laboratorio:", "Detalle secuencial de la configuración, escaneo, análisis o contención realizada."),
        ("Evidencias Técnicas Recolectadas:", "Hashes SHA-256, direcciones IP origen/destino, Event IDs y capturas de pantalla adjuntas."),
        ("Dificultad Superada («Trabajo Bien Hecho»):", "Registro de errores de configuración encontrados y procedimiento técnico aplicado para resolverlos."),
        ("Firma del Estudiante y Vo.Bo. del Instructor:", "Firma Estudiante: ___________________  |  Firma Instructor: ___________________")
    ]

    for idx, (campo, lineas) in enumerate(formato_data):
        row = tbl_b_fmt.rows[idx]
        set_cell_background(row.cells[0], "F7FAFC")
        set_cell_margins(row.cells[0], top=60, bottom=60, left=80, right=80)
        set_cell_margins(row.cells[1], top=60, bottom=60, left=80, right=80)

        p0 = row.cells[0].paragraphs[0]
        style_p(p0, space_before=0, space_after=0)
        r0 = p0.add_run(campo)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_PRIMARY

        p1 = row.cells[1].paragraphs[0]
        style_p(p1, space_before=0, space_after=0)
        r1 = p1.add_run(lineas)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9)
        r1.font.color.rgb = COLOR_DARK

    section = doc.sections[0]
    footer = section.footer
    footer_p = footer.paragraphs[0]
    style_p(footer_p, space_before=0, space_after=0)
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_f = footer_p.add_run("Fundación Kinal | Dosificación y Secuencia Didáctica — Fundamentos de Ciberseguridad")
    run_f.font.name = "Calibri"
    run_f.font.size = Pt(9)
    run_f.font.color.rgb = COLOR_SECONDARY

    out_dir = os.path.join("Cursos", "Ciberseguridad")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "Dosificacion_y_Secuencia_Didactica_Ciberseguridad.docx")
    doc.save(out_path)
    print(f"Dosificación guardada con éxito en: {out_path}")

if __name__ == "__main__":
    generate_dosificacion()
