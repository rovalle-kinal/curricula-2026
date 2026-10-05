import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E0", sz="4"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_callout(doc, text, bold_prefix="", fill_hex="F7FAFC", border_color="003366"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    
    cell = tbl.cell(0, 0)
    set_cell_background(cell, fill_hex)
    set_cell_margins(cell, top=130, bottom=130, left=180, right=130)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 51, 102)
        
    r_txt = p.add_run(text)
    r_txt.font.name = "Calibri"
    r_txt.font.size = Pt(10)
    r_txt.font.italic = True
    r_txt.font.color.rgb = RGBColor(45, 55, 72)
    
    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_before = Pt(0)
    p_post.paragraph_format.space_after = Pt(6)

def style_heading(p, text, level=1, color=RGBColor(0, 51, 102)):
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.bold = True
    r.font.color.rgb = color
    if level == 1:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        r.font.size = Pt(15)
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        r.font.size = Pt(13)
    elif level == 3:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r.font.size = Pt(11)

def style_p(p, space_before=0, space_after=6, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def generate_propuesta():
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    COLOR_PRIMARY = RGBColor(0, 51, 102)
    COLOR_SECONDARY = RGBColor(74, 85, 104)
    COLOR_DARK = RGBColor(26, 32, 44)

    # =========================================================================
    # PORTADA INSTITUCIONAL
    # =========================================================================
    p_inst = doc.add_paragraph()
    style_p(p_inst, space_before=20, space_after=2)
    r_inst = p_inst.add_run("FUNDACIÓN KINAL")
    r_inst.font.name = "Calibri"
    r_inst.font.size = Pt(14)
    r_inst.font.bold = True
    r_inst.font.color.rgb = COLOR_PRIMARY

    p_esc = doc.add_paragraph()
    style_p(p_esc, space_before=0, space_after=20)
    r_esc = p_esc.add_run("Escuela Técnica Superior — Coordinación de Formación Continua y Empleabilidad")
    r_esc.font.name = "Calibri"
    r_esc.font.size = Pt(11)
    r_esc.font.color.rgb = COLOR_SECONDARY

    p_tit = doc.add_paragraph()
    style_p(p_tit, space_before=14, space_after=10)
    r_tit = p_tit.add_run("PROPUESTA FORMATIVA INSTITUCIONAL:\nCURSO DE FUNDAMENTOS DE CIBERSEGURIDAD Y OPERACIONES SOC")
    r_tit.font.name = "Calibri"
    r_tit.font.size = Pt(20)
    r_tit.font.bold = True
    r_tit.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    style_p(p_sub, space_before=0, space_after=20)
    r_sub = p_sub.add_run("Capacitación Híbrida para la Reconversión Laboral y Empleabilidad Inmediata en el Ecosistema Tecnológico de Guatemala (80 Horas)")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(11.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = COLOR_SECONDARY

    # Ficha Técnica
    tbl_ficha = doc.add_table(rows=8, cols=2)
    tbl_ficha.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ficha.autofit = False
    tbl_ficha.columns[0].width = Inches(2.2)
    tbl_ficha.columns[1].width = Inches(4.3)
    set_table_borders(tbl_ficha, color="CBD5E0", sz="6")

    ficha_data = [
        ("Nivel Formativo de Referencia:", "Equivalencia DQR Nivel 4 - 5 (Competencia técnica especializada y autonomía operativa)"),
        ("Población Destinataria:", "Jóvenes y adultos desempleados de 18 a 45 años con conocimientos previos de informática"),
        ("Duración Total:", "80 horas pedagógicas totales"),
        ("Modalidad Híbrida:", "80% Virtual (64 horas sincrónicas/laboratorios) / 20% Presencial (16 horas prácticas en Kinal)"),
        ("Estructura de Horarios:", "Por convenir (Sesiones virtuales de 3 horas, 2 días entre semana; sesiones presenciales de 8 horas los sábados). Sin fechas fijas."),
        ("Enfoque Laboral Central:", "Inserción laboral inmediata como Analista SOC Junior / Técnico de Soporte en Ciberseguridad"),
        ("Nota Mínima Aprobatoria:", "75 puntos sobre 100 en todas las evaluaciones de taller, proyectos y examen terminal"),
        ("Sede Presencial:", "Fundación Kinal — Laboratorios de Redes y Cómputo, Sede Central, Zona 7, Ciudad de Guatemala")
    ]

    for idx, (label, val) in enumerate(ficha_data):
        row = tbl_ficha.rows[idx]
        set_cell_background(row.cells[0], "F7FAFC")
        set_cell_margins(row.cells[0], top=70, bottom=70, left=90, right=90)
        set_cell_margins(row.cells[1], top=70, bottom=70, left=90, right=90)
        
        p0 = row.cells[0].paragraphs[0]
        style_p(p0, space_before=0, space_after=0)
        r0 = p0.add_run(label)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_PRIMARY

        p1 = row.cells[1].paragraphs[0]
        style_p(p1, space_before=0, space_after=0)
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_DARK

    doc.add_page_break()

    # =========================================================================
    # 1. MARCO FILOSÓFICO E IDEARIO DE KINAL: ÉTICA DIGITAL Y TRABAJO BIEN HECHO
    # =========================================================================
    h1_1 = doc.add_paragraph()
    style_heading(h1_1, "1. Marco Filosófico e Ideario Institucional de Fundación Kinal", level=1)

    p_id_desc = doc.add_paragraph()
    style_p(p_id_desc, space_before=0, space_after=6)
    r = p_id_desc.add_run("En el campo de la ciberseguridad, la solvencia técnica es inseparable de la rectitud moral y la confiabilidad personal. La capacitación en defensa de infraestructuras informáticas exige de los estudiantes una adhesión irrestricta a los principios que sustentan a Fundación Kinal:")
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = COLOR_DARK

    add_callout(doc, "«Formar a jóvenes y adultos a través de una educación integral, con énfasis en las áreas técnicas y tecnológicas, influyendo positivamente en su trabajo, su familia y la sociedad».", bold_prefix="Misión Institucional de Kinal:")

    add_callout(doc, "«Visión cristiana de la vida. Respeto a la dignidad de la persona. Espíritu de servicio. Trabajo bien hecho. Libertad personal responsable».", bold_prefix="Valores Nucleares de Kinal:")

    p_et = doc.add_paragraph()
    style_heading(p_et, "1.1 El Principio del «Trabajo Bien Hecho» y la Ética Profesional en Ciberdefensa", level=2)

    p_et_txt = doc.add_paragraph()
    style_p(p_et_txt, space_before=0, space_after=6)
    r = p_et_txt.add_run("El «trabajo bien hecho» en la seguridad de la información se traduce en disciplina técnica, minuciosidad y sigilo profesional:\n"
                          "• Cero negligencia en configuraciones: Cada directiva de seguridad, regla de firewall o parámetro de autenticación debe aplicarse con precisión exhaustiva, evitando configuraciones por defecto o 'atajos' inseguros.\n"
                          "• Documentación rigurosa de incidentes: El analista debe levantar reportes claros, fidedignos y con evidencia verificable (hashes SHA-256, marcas temporales UTC, capturas de tráfico intactas) preservando la cadena de custodia.\n"
                          "• Acuerdo de Confidencialidad y Uso Ético (NDA Didáctico): Cada participante suscribe al inicio del curso un compromiso formal de no utilizar las herramientas de análisis de vulnerabilidades fuera de los entornos controlados de laboratorio, entendiendo que el conocimiento técnico está puesto al servicio del bien común y la protección de la sociedad.")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_DARK

    # =========================================================================
    # 2. DEMANDA LABORAL EN GUATEMALA Y ENFOQUE EN EMPLEABILIDAD INMEDIATA
    # =========================================================================
    h1_2 = doc.add_paragraph()
    style_heading(h1_2, "2. Demanda del Mercado Laboral Guatemalteco y Empleabilidad Inmediata", level=1)

    p_mkt = doc.add_paragraph()
    style_p(p_mkt, space_before=0, space_after=6)
    r = p_mkt.add_run("Una investigación reciente sobre las ofertas de empleo en Guatemala revela una escasez aguda de perfiles operativos en ciberseguridad. Las empresas demandan cada vez más personal junior capaz de incorporarse de inmediato a turnos de monitoreo y soporte técnico de seguridad:\n"
                      "• Sectores Contratantes en Guatemala: El sistema bancario y financiero nacional (Banco Industrial, Banrural, BAC Credomatic, BAM, G&T Continental), las empresas de telecomunicaciones (Tigo, Claro), firmas de servicios administrados de seguridad SOC (Devel Security, Cyberseg, S-Cube), cadenas de retail y empresas de BPO con servicios tecnológicos hacia Estados Unidos y Latinoamérica.\n"
                      "• Puesto de Entrada Más Demandado: Analista de Centro de Operaciones de Seguridad (Analista SOC Nivel 1) y Técnico de Soporte en Ciberseguridad y Hardening.\n"
                      "• Nivel Salarial Promedio: Las vacantes de entrada para técnicos con destrezas prácticas comprobadas oscilan entre Q6,000.00 y Q9,000.00 mensuales, ofreciendo una vía concreta y rápida de superación socioeconómica para personas actualmente desempleadas.\n"
                      "• Competencias de Empleabilidad Priorizadas: Triaje rápido de alertas en plataformas SIEM, lectura e interpretación de logs en Windows/Linux, filtrado de tráfico en firewalls, análisis de correos phishing y aplicación de medidas de contención inmediata.")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_DARK

    # =========================================================================
    # 3. MARCO DQR Y FORMACIÓN EN ALTERNANCIA (SISTEMA DUAL)
    # =========================================================================
    h1_3 = doc.add_paragraph()
    style_heading(h1_3, "3. Marco Alemán de Cualificaciones (DQR Nivel 4-5) y Metodología Dual", level=1)

    add_callout(doc, "«concepto de aprendizaje para jóvenes que tiene como objetivo preparar a los aprendices para su vida profesional, por lo que la formación se realiza en régimen de alternancia entre la escuela o el centro de formación y la empresa. El aprender en el proceso del trabajo es base y a la vez objetivo de la acción formativa».", bold_prefix="Definición de Formación Profesional Dual:")

    add_callout(doc, "«the ability and readiness of the individual to use knowledge, skills and personal, social and methodological competences and to behave in a considered, individual and socially responsible manner. Competence is understood in this sense as the comprehensive ability to act».", bold_prefix="Concepto de Competencia en el DQR:")

    tbl_dqr = doc.add_table(rows=5, cols=3)
    tbl_dqr.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_dqr.autofit = False
    tbl_dqr.columns[0].width = Inches(1.8)
    tbl_dqr.columns[1].width = Inches(1.5)
    tbl_dqr.columns[2].width = Inches(3.2)
    set_table_borders(tbl_dqr, color="CBD5E0", sz="4")

    hdr_dqr = tbl_dqr.rows[0]
    for c in hdr_dqr.cells:
        set_cell_background(c, "003366")
    for idx, text in enumerate(["Dimensión DQR", "Subdimensión", "Evidencia Operativa en el Curso"]):
        p = hdr_dqr.cells[idx].paragraphs[0]
        style_p(p, space_before=0, space_after=0)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    dqr_items = [
        ("Competencia Profesional\n(Fachkompetenz)", "Conocimientos\n(Wissen)", "Arquitectura de protocolos TCP/IP, modelos de defensa perimetral, funcionamiento de malware, directivas GPO/Linux y marcos de respuesta NIST SP 800-61."),
        ("Competencia Profesional\n(Fachkompetenz)", "Destrezas\n(Fertigkeiten)", "Configuración de firewalls, inspección de paquetes en Wireshark, correlación de logs en Wazuh SIEM, bastionado de servidores y análisis de correos phishing."),
        ("Competencia Personal\n(Personale Kompetenz)", "Competencia Social\n(Sozialkompetenz)", "Comunicación técnica serena y clara durante incidentes críticos, trabajo coordinado en células Blue Team y elaboración de informes para gerencia."),
        ("Competencia Personal\n(Personale Kompetenz)", "Autonomía\n(Selbständigkeit)", "Criterio analítico para distinguir falsos positivos de intrusiones reales, apego ético incondicional y autoexigencia bajo el principio del 'trabajo bien hecho'.")
    ]

    for idx, (dim, sub, evid) in enumerate(dqr_items, start=1):
        row = tbl_dqr.rows[idx]
        if idx % 2 == 1:
            for c in row.cells:
                set_cell_background(c, "F7FAFC")
        for c_idx, val in enumerate([dim, sub, evid]):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
            p = cell.paragraphs[0]
            style_p(p, space_before=0, space_after=0)
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            r.font.color.rgb = COLOR_DARK

    doc.add_page_break()

    # =========================================================================
    # 4. PERFILES DE INGRESO Y EGRESO
    # =========================================================================
    h1_4 = doc.add_paragraph()
    style_heading(h1_4, "4. Perfiles de Ingreso y Egreso Institucional", level=1)

    p_ing = doc.add_paragraph()
    style_heading(p_ing, "4.1 Perfil de Ingreso", level=2)
    p_ing_txt = doc.add_paragraph()
    style_p(p_ing_txt, space_before=0, space_after=6)
    r = p_ing_txt.add_run("Dirigido a jóvenes y adultos desempleados de 18 a 45 años con conocimientos previos en informática, administración básica de sistemas o redes, motivados a reconvertirse laboralmente hacia la ciberdefensa con compromiso ético y disciplina.")
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = COLOR_DARK

    p_egr = doc.add_paragraph()
    style_heading(p_egr, "4.2 Perfil de Egreso", level=2)
    p_egr_txt = doc.add_paragraph()
    style_p(p_egr_txt, space_before=0, space_after=12)
    r = p_egr_txt.add_run("El egresado opera como Analista SOC Junior o Técnico de Ciberseguridad, capaz de monitorear telemetría, realizar triaje de alertas de seguridad, aplicar hardening en Windows y Linux, analizar amenazas de malware o phishing y ejecutar contención básica de incidentes.")
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = COLOR_DARK

    # =========================================================================
    # 5. ESTRUCTURA MODULAR, MODALIDAD HÍBRIDA Y PROYECTOS POR MÓDULO
    # =========================================================================
    h1_5 = doc.add_paragraph()
    style_heading(h1_5, "5. Arquitectura Curricular y Modalidad Híbrida (80 Horas)", level=1)

    p_mod_desc = doc.add_paragraph()
    style_p(p_mod_desc, space_before=0, space_after=6)
    r = p_mod_desc.add_run("El programa distribuye sus 80 horas en un 80% virtual (64 horas de clases sincrónicas y laboratorios guiados) y un 20% presencial (16 horas de talleres intensivos en 2 sábados de 8 horas en Kinal):")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_DARK

    tbl_mods = doc.add_table(rows=6, cols=5)
    tbl_mods.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_mods.autofit = False
    tbl_mods.columns[0].width = Inches(1.1)
    tbl_mods.columns[1].width = Inches(1.8)
    tbl_mods.columns[2].width = Inches(1.2)
    tbl_mods.columns[3].width = Inches(1.8)
    tbl_mods.columns[4].width = Inches(0.6)
    set_table_borders(tbl_mods, color="CBD5E0", sz="4")

    hdr_m = tbl_mods.rows[0]
    for c in hdr_m.cells:
        set_cell_background(c, "003366")
    headers_m = ["Módulo", "Eje Formativo Central", "Distribución Híbrida", "Proyecto Práctico Evaluado (≥ 75 pts)", "Total"]
    for idx, text in enumerate(headers_m):
        p = hdr_m.cells[idx].paragraphs[0]
        style_p(p, space_before=0, space_after=0)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    mods_data = [
        ("Módulo 1", "Redes Seguras y Arquitectura de Ciberdefensa", "12h Virtuales\n4h Presenciales*", "Auditoría de tráfico con Wireshark y configuración de reglas de firewall perimetral", "16 h"),
        ("Módulo 2", "Hardening de S.O., Identidad y Accesos", "12h Virtuales\n4h Presenciales*", "Plantilla de bastionado seguro (GPO / scripts) verificada con escáner de cumplimiento", "16 h"),
        ("Módulo 3", "Operaciones SOC y Monitoreo SIEM", "16h Virtuales\n(Intensivo SIEM)", "Despliegue de agentes Wazuh, triaje de alertas y elaboración de tickets formales", "16 h"),
        ("Módulo 4", "Detección de Amenazas y Vulnerabilidades", "12h Virtuales\n4h Presenciales**", "Informe técnico de análisis de phishing y plan de remediación de vulnerabilidades", "16 h"),
        ("Módulo 5", "Respuesta a Incidentes y Reto Capstone", "12h Virtuales\n4h Presenciales**", "Proyecto Capstone: Simulación de crisis SOC (Blue Team), contención y defensa oral", "16 h")
    ]

    for idx, (m, eje, dist, proy, tot) in enumerate(mods_data, start=1):
        row = tbl_mods.rows[idx]
        if idx % 2 == 1:
            for c in row.cells:
                set_cell_background(c, "F7FAFC")
        for c_idx, val in enumerate([m, eje, dist, proy, tot]):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=60, bottom=60, left=70, right=70)
            p = cell.paragraphs[0]
            style_p(p, space_before=0, space_after=0)
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(8.5)
            r.font.color.rgb = COLOR_DARK
            if c_idx in [0, 4]:
                r.font.bold = True

    p_note_pres = doc.add_paragraph()
    style_p(p_note_pres, space_before=4, space_after=12)
    r = p_note_pres.add_run("* Las 4h presenciales del Módulo 1 y las 4h presenciales del Módulo 2 se integran en el Sábado Presencial 1 (8 horas de taller en Kinal).\n"
                           "** Las 4h presenciales del Módulo 4 y las 4h presenciales del Módulo 5 se integran en el Sábado Presencial 2 (8 horas de taller y Capstone en Kinal).")
    r.font.name = "Calibri"
    r.font.size = Pt(8.5)
    r.font.italic = True
    r.font.color.rgb = COLOR_SECONDARY

    # =========================================================================
    # 6. ECOSISTEMA DE HERRAMIENTAS, SOFTWARE LIBRE Y PRECIOS
    # =========================================================================
    h1_6 = doc.add_paragraph()
    style_heading(h1_6, "6. Ecosistema de Herramientas Open Source y Cisco (Precios y Licenciamiento)", level=1)

    p_her_desc = doc.add_paragraph()
    style_p(p_her_desc, space_before=0, space_after=6)
    r = p_her_desc.add_run("El curso optimiza la inversión institucional utilizando un 90% de herramientas Open Source de nivel corporativo, articuladas con la membresía gratuita de Fundación Kinal como Cisco Networking Academy:")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_DARK

    tbl_her = doc.add_table(rows=8, cols=4)
    tbl_her.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_her.autofit = False
    tbl_her.columns[0].width = Inches(1.8)
    tbl_her.columns[1].width = Inches(1.5)
    tbl_her.columns[2].width = Inches(1.4)
    tbl_her.columns[3].width = Inches(1.8)
    set_table_borders(tbl_her, color="CBD5E0", sz="4")

    hdr_h = tbl_her.rows[0]
    for c in hdr_h.cells:
        set_cell_background(c, "003366")
    headers_h = ["Herramienta Tecnológica", "Tipo / Licencia", "Costo Mensual / Anual", "Función en el Programa"]
    for idx, text in enumerate(headers_h):
        p = hdr_h.cells[idx].paragraphs[0]
        style_p(p, space_before=0, space_after=0)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    her_data = [
        ("Wazuh SIEM & XDR", "Open Source (GPL)", "$0.00 / mes (Gratuito)", "Recolección de logs, correlación y triaje de alertas de seguridad"),
        ("Wireshark & tcpdump", "Open Source (GPL)", "$0.00 / mes (Gratuito)", "Inspección profunda de paquetes y auditoría de protocolos de red"),
        ("Greenbone (OpenVAS)", "Open Source", "$0.00 / mes (Gratuito)", "Escaneo de vulnerabilidades y evaluación de parches de seguridad"),
        ("Oracle VirtualBox / VMs", "Open Source (GPL)", "$0.00 / mes (Gratuito)", "Montaje de laboratorios virtuales aislados (Kali, Ubuntu, Windows)"),
        ("Cisco Packet Tracer", "Cisco NetAcad", "$0.00 / mes (Gratuito)", "Simulación de topologías de red, switches, routers y firewalls ASA"),
        ("Cisco Snort IDS/IPS", "Cisco Talos (GPL)", "$0.00 / mes (Gratuito)", "Detección de intrusiones basada en firmas de tráfico de red"),
        ("Cisco Modeling Labs (CML)\n[Herramienta Opcional]", "Cisco Comercial", "$199.00 USD / año\n(~$16.58 USD / mes)", "Emulación avanzada para instructor (no requerida para alumnos)")
    ]

    for idx, (nom, tip, cos, fun) in enumerate(her_data, start=1):
        row = tbl_her.rows[idx]
        if idx % 2 == 1:
            for c in row.cells:
                set_cell_background(c, "F7FAFC")
        for c_idx, val in enumerate([nom, tip, cos, fun]):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=60, bottom=60, left=70, right=70)
            p = cell.paragraphs[0]
            style_p(p, space_before=0, space_after=0)
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(8.5)
            r.font.color.rgb = COLOR_DARK
            if c_idx == 0:
                r.font.bold = True

    # =========================================================================
    # 7. ESQUEMA DE EVALUACIÓN Y ACREDITACIÓN INSTITUCIONAL (KINAL)
    # =========================================================================
    h1_7 = doc.add_paragraph()
    style_heading(h1_7, "7. Esquema de Evaluación y Acreditación (Umbral ≥ 75 Puntos)", level=1)

    p_ev_desc = doc.add_paragraph()
    style_p(p_ev_desc, space_before=0, space_after=6)
    r = p_ev_desc.add_run("Para acreditar el curso y obtener el diploma de certificación técnica emitido por Fundación Kinal, el estudiante debe obtener una nota mínima de 75 puntos sobre 100, demostrando solvencia teórica y competencia operativa en incidentes reales:")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_DARK

    tbl_ev = doc.add_table(rows=7, cols=3)
    tbl_ev.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ev.autofit = False
    tbl_ev.columns[0].width = Inches(2.3)
    tbl_ev.columns[1].width = Inches(1.1)
    tbl_ev.columns[2].width = Inches(3.1)
    set_table_borders(tbl_ev, color="CBD5E0", sz="4")

    hdr_ev = tbl_ev.rows[0]
    for c in hdr_ev.cells:
        set_cell_background(c, "003366")
    for idx, text in enumerate(["Rubro Evaluativo", "Ponderación", "Criterio de Suficiencia Técnica (≥ 75 pts)"]):
        p = hdr_ev.cells[idx].paragraphs[0]
        style_p(p, space_before=0, space_after=0)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    eval_items = [
        ("Proyecto Práctico Módulo 1 (Redes)", "15%", "Captura de paquetes en Wireshark y configuración correcta de reglas en firewall"),
        ("Proyecto Práctico Módulo 2 (Hardening)", "15%", "Auditoría LinPEAS sin vulnerabilidades críticas y aplicación de directivas GPO"),
        ("Proyecto Práctico Módulo 3 (SIEM)", "15%", "Triaje de 5 alertas de seguridad en Wazuh con tickets debidamente documentados"),
        ("Proyecto Práctico Módulo 4 (Threats)", "15%", "Análisis de cabeceras de correo phishing y reporte de escaneo de vulnerabilidades"),
        ("Proyecto Capstone / CTF (Módulo 5)", "25%", "Simulación de sala de crisis: contención de intrusión, preservación y sustentación oral"),
        ("Bitácora Berichtsheft y Ética Kinal", "15%", "Entrega semanal de bitácora técnica, asistencia puntual y cumplimiento del NDA")
    ]

    for idx, (rub, pond, crit) in enumerate(eval_items, start=1):
        row = tbl_ev.rows[idx]
        if idx % 2 == 1:
            for c in row.cells:
                set_cell_background(c, "F7FAFC")
        for c_idx, val in enumerate([rub, pond, crit]):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=60, bottom=60, left=70, right=70)
            p = cell.paragraphs[0]
            style_p(p, space_before=0, space_after=0)
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(8.5)
            r.font.color.rgb = COLOR_DARK
            if c_idx in [0, 1]:
                r.font.bold = True

    # =========================================================================
    # 8. ESTRATEGIA DE EMPLEABILIDAD INMEDIATA
    # =========================================================================
    h1_8 = doc.add_paragraph()
    style_heading(h1_8, "8. Estrategia de Inserción Laboral y Articulación Institucional", level=1)

    p_emp = doc.add_paragraph()
    style_p(p_emp, space_before=0, space_after=6)
    r = p_emp.add_run("Para maximizar la inserción laboral de los estudiantes desempleados (18 a 45 años), el programa articula tres mecanismos directos:\n"
                      "1. Bolsa de Empleo de Fundación Kinal: Vinculación directa de los egresados con empresas aliadas en el sector financiero, telecomunicaciones y empresas de seguridad gestionada (MSSP) que buscan analistas de monitoreo 24/7.\n"
                      "2. Portafolio de Casos Prácticos (Write-ups): Cada alumno compila durante el curso una bitácora técnica demostrable en GitHub/LinkedIn con los reportes de incidentes resueltos, demostrando experiencia técnica real a los reclutadores.\n"
                      "3. Simulación de Entrevistas Técnicas: Talleres de preparación para pruebas de contratación (preguntas técnicas sobre protocolos, comandos de Linux, lectura de logs y escenarios de contingencia bajo presión).")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_DARK

    section = doc.sections[0]
    footer = section.footer
    footer_p = footer.paragraphs[0]
    style_p(footer_p, space_before=0, space_after=0)
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_f = footer_p.add_run("Fundación Kinal | Propuesta Formativa Oficial — Fundamentos de Ciberseguridad")
    run_f.font.name = "Calibri"
    run_f.font.size = Pt(9)
    run_f.font.color.rgb = COLOR_SECONDARY

    out_dir = os.path.join("Cursos", "Ciberseguridad")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "Propuesta_Curso_Ciberseguridad_Kinal.docx")
    doc.save(out_path)
    print(f"Propuesta guardada con éxito en: {out_path}")

if __name__ == "__main__":
    generate_propuesta()
