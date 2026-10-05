import os
import re
import html
import json
import docx

WORKSPACE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CURSOS_DIR = os.path.join(WORKSPACE_DIR, "Cursos")
WEB_DIR = os.path.join(CURSOS_DIR, "Web")
os.makedirs(WEB_DIR, exist_ok=True)

def parse_docx(path):
    doc = docx.Document(path)
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    tables = []
    for t in doc.tables:
        rows = []
        for r in t.rows:
            rows.append([c.text.strip() for c in r.cells])
        tables.append(rows)
    return paras, tables

def clean_html(text):
    if not text:
        return ""
    return html.escape(text).replace('\n', '<br>')

# ==============================================================================
# 1. PARSE MECATRÓNICA
# ==============================================================================
meca_prop_p, meca_prop_t = parse_docx(os.path.join(CURSOS_DIR, "Mecatrónica", "Propuesta_Curso_Mecatronica_Kinal.docx"))
meca_dos_p, meca_dos_t = parse_docx(os.path.join(CURSOS_DIR, "Mecatrónica", "Dosificacion_y_Secuencia_Didactica_Mecatronica_Kinal.docx"))

meca_ficha = meca_prop_t[0]
meca_dqr = meca_prop_t[5]
meca_mod_prop = meca_prop_t[6]
meca_eval = meca_prop_t[7]
meca_rubric = meca_dos_t[22]

meca_sessions = []
meca_months = [
    ("Febrero", "Módulo 1: CAD & Láser", range(1, 5)),
    ("Marzo", "Módulo 2: FDM Mecanismos", range(5, 9)),
    ("Abril", "Módulo 3: Fresado CNC", range(9, 13)),
    ("Mayo", "Módulo 4: Sensórica & Control", range(13, 17)),
    ("Junio", "Módulo 5: Celda Capstone", range(17, 21))
]

for idx in range(2, 22):
    t = meca_dos_t[idx]
    s_num = idx - 1
    t1_row = meca_dos_t[1][s_num] if s_num < len(meca_dos_t[1]) else ["", "", "", "", ""]
    
    obj = t[0][1]
    ap = t[1][1]
    des = t[2][1]
    cie = t[3][1]
    evi = t[4][1]
    eq = t[5][1]
    
    mod_name = t1_row[2] if len(t1_row) > 2 else f"Módulo {((s_num-1)//4)+1}"
    title_short = t1_row[3] if len(t1_row) > 3 else f"Sesión {s_num}"
    
    month_name = "Febrero"
    if s_num > 16: month_name = "Junio"
    elif s_num > 12: month_name = "Mayo"
    elif s_num > 8: month_name = "Abril"
    elif s_num > 4: month_name = "Marzo"

    meca_sessions.append({
        "num": s_num,
        "code": f"S-{s_num:02d}",
        "month": month_name,
        "sabado_num": ((s_num - 1) % 4) + 1,
        "horario": "Sábado 8:00 a 12:30 hrs",
        "module": mod_name,
        "title": title_short,
        "duration": "4.5 Horas (270 min)",
        "objective": obj,
        "apertura": ap,
        "desarrollo": des,
        "cierre": cie,
        "evidencias": evi,
        "equipamiento": eq
    })

meca_modules_data = [
    {
        "num": 1,
        "title": "Diseño CAD 3D, Metrología de Taller y Prototipado con Cortadora Láser",
        "period": "MÓDULO 1 (Febrero • Sesiones 1 a 4)",
        "hours": "18 Horas de Taller",
        "logro": "Dominar el croquizado paramétrico 3D, metrología dimensional de precisión (±0.1 mm) y fabricación de gabinetes industriales mediante corte láser CO2.",
        "topics": [
            {
                "num": 1,
                "title": "Fundamentos de Mecatrónica Industrial y CAD Paramétrico",
                "cat": "software",
                "cat_name": "🔵 Modelado & Software",
                "cat_badge": "bg-blue-100 text-blue-900 border border-blue-200",
                "desc": "Metrología dimensional aplicada con pie de rey digital y micrómetro de exteriores. Restricciones geométricas, croquizado paramétrico y operaciones de extrusión 3D en Autodesk Fusion 360 / SolidWorks."
            },
            {
                "num": 2,
                "title": "Ingeniería Inversa y Diseño para Manufactura por Corte Láser (DFM)",
                "cat": "software",
                "cat_name": "🔵 Modelado & Software",
                "cat_badge": "bg-blue-100 text-blue-900 border border-blue-200",
                "desc": "Replicación y digitalización dimensional de piezas desgastadas o discontinuadas. Compensación de sangría de corte (kerf), parametrización de potencias por material y diseño de encastres autoblocantes (finger joints)."
            },
            {
                "num": 3,
                "title": "Operación Segura y Parametrización de Cortadora Láser CO2",
                "cat": "maquinaria",
                "cat_name": "🟡 Maquinaria & Taller",
                "cat_badge": "bg-amber-100 text-amber-900 border border-amber-200",
                "desc": "Preparación de vectores por capas (marcado, grabado y corte vectorial en DXF/SVG). Calibración de distancia focal, asistencia de aire y extracción de humos sobre acrílico industrial, Delrin y MDF técnico."
            },
            {
                "num": 4,
                "title": "Fabricación de Gabinetes Mecatrónicos y Evaluación Práctica",
                "cat": "ensamble",
                "cat_name": "🟢 Manufactura & Ensamble",
                "cat_badge": "bg-emerald-100 text-emerald-900 border border-emerald-200",
                "desc": "Montaje mecánico y verificación de rigidez estructural de chasis para alojar tarjetas electrónicas y rieles DIN. Verificación de tolerancias de ajuste y entrega del producto terminado del Módulo 1."
            }
        ]
    },
    {
        "num": 2,
        "title": "Fabricación Aditiva Avanzada (Impresión 3D FDM) para Mecanismos Industriales",
        "period": "MÓDULO 2 (Marzo • Sesiones 5 a 8)",
        "hours": "18 Horas de Taller",
        "logro": "Diseñar y manufacturar mecanismos cinemáticos articulados optimizados para manufactura aditiva FDM con polímeros técnicos y tolerancias dinámicas.",
        "topics": [
            {
                "num": 5,
                "title": "Cinemática Aplicada y Ciencia de Polímeros para FDM",
                "cat": "software",
                "cat_name": "🔵 Modelado & Software",
                "cat_badge": "bg-blue-100 text-blue-900 border border-blue-200",
                "desc": "Cálculo y diseño de engranes rectos, poleas sincrónicas GT2 y levas mecánicas. Propiedades mecánicas y térmicas de filamentos técnicos: PLA+, PETG industrial, ABS, TPU (flexible) y Nylon reforzado."
            },
            {
                "num": 6,
                "title": "Diseño para Fabricación Aditiva (DFAM) y Parámetros en Laminador",
                "cat": "software",
                "cat_name": "🔵 Modelado & Software",
                "cat_badge": "bg-blue-100 text-blue-900 border border-blue-200",
                "desc": "Orientación de piezas respecto a líneas de esfuerzo para mitigar anisotropía mecánica. Optimización de voladizos (overhangs), puentes y configuración de rellenos estructurales (giroide y cúbico)."
            },
            {
                "num": 7,
                "title": "Tolerancias Print-in-Place e Integración de Insertos Roscados",
                "cat": "maquinaria",
                "cat_name": "🟡 Maquinaria & Taller",
                "cat_badge": "bg-amber-100 text-amber-900 border border-amber-200",
                "desc": "Modelado de tolerancias dinámicas para mecanismos ensamblados en cama (0.2 a 0.4 mm). Instalación de insertos roscados de latón en caliente (heat-set inserts), calibración de flujo y nivelación de camas térmicas."
            },
            {
                "num": 8,
                "title": "Montaje de Subensambles Mecánicos y Evaluación Práctica",
                "cat": "ensamble",
                "cat_name": "🟢 Manufactura & Ensamble",
                "cat_badge": "bg-emerald-100 text-emerald-900 border border-emerald-200",
                "desc": "Integración de baleros lineales, ejes rectificados de acero y tornillería métrica. Pruebas de movimiento continuo, lubricación de engranes y evaluación del reto: Garra Robótica Articulada Funcional."
            }
        ]
    },
    {
        "num": 3,
        "title": "Mecanizado Sustractivo y Fresado CNC: CAM, Código G y Fabricación de Piezas",
        "period": "MÓDULO 3 (Abril • Sesiones 9 a 12)",
        "hours": "18 Horas de Taller",
        "logro": "Generar trayectorias de maquinado en software CAM, operar fresadora CNC mediante Código G y verificar dimensionalmente placas base mecatrónicas.",
        "topics": [
            {
                "num": 9,
                "title": "Cinemática CNC y Estructura del Código G y M",
                "cat": "software",
                "cat_name": "🔵 Modelado & Software",
                "cat_badge": "bg-blue-100 text-blue-900 border border-blue-200",
                "desc": "Ejes coordenados cartesianos (X, Y, Z), coordenadas de máquina (G53) y de trabajo (G54-G59). Sintaxis de comandos modales de movimiento (G00, G01, G02, G03) y comandos misceláneos de control de husillo (M03, M05)."
            },
            {
                "num": 10,
                "title": "Programación en Software CAM (Estrategias 2.5D y 3D)",
                "cat": "software",
                "cat_name": "🔵 Modelado & Software",
                "cat_badge": "bg-blue-100 text-blue-900 border border-blue-200",
                "desc": "Definición de bloque de material (stock), plano de seguridad y origen WCS. Estrategias de planeado, cajeras (pockets), contorneado exterior y cálculo riguroso de avances y RPM (Vc, fz)."
            },
            {
                "num": 11,
                "title": "Alistamiento, Montaje de Herramientas y Puesta a Cero en CNC",
                "cat": "maquinaria",
                "cat_name": "🟡 Maquinaria & Taller",
                "cat_badge": "bg-amber-100 text-amber-900 border border-amber-200",
                "desc": "Protocolos rigurosos de seguridad en corte sustractivo. Montaje y verificación de excentricidad en boquillas ER con fresas de carburo. Calibración de ceros de pieza (X, Y) y eje Z con sonda electrónica."
            },
            {
                "num": 12,
                "title": "Mecanizado Práctico, Metrología de Verificación y Evaluación",
                "cat": "ensamble",
                "cat_name": "🟢 Manufactura & Ensamble",
                "cat_badge": "bg-emerald-100 text-emerald-900 border border-emerald-200",
                "desc": "Fresado efectivo en banco sobre placas de Delrin (POM) y aluminio mecanizable. Inspección dimensional con micrómetro y escuadra de precisión; desbarbado y acabado de la Bancada Mecatrónica Kinal."
            }
        ]
    },
    {
        "num": 4,
        "title": "Sensórica Industrial, Actuación Electromecánica/Neumática y Control de Ejes",
        "period": "MÓDULO 4 (Mayo • Sesiones 13 a 16)",
        "hours": "18 Horas de Taller",
        "logro": "Integrar sensores industriales PNP/NPN, gobernar motores paso a paso mediante drivers y cablear circuitos electroneumáticos bajo normas Kinal.",
        "topics": [
            {
                "num": 13,
                "title": "Transducción y Sensórica Industrial",
                "cat": "control",
                "cat_name": "🟣 Control & Automatización",
                "cat_badge": "bg-purple-100 text-purple-900 border border-purple-200",
                "desc": "Principio operativo de sensores inductivos, capacitivos y fotoeléctricos (réflex y barrera). Conexión e identificación de salidas a 3 y 4 hilos con lógica PNP vs. NPN y finales de carrera magnéticos Reed."
            },
            {
                "num": 14,
                "title": "Control de Movimiento con Motores a Pasos y Drivers Industriales",
                "cat": "maquinaria",
                "cat_name": "🟡 Maquinaria & Taller",
                "cat_badge": "bg-amber-100 text-amber-900 border border-amber-200",
                "desc": "Curvas de torque vs. velocidad en motores NEMA 17 y NEMA 23. Configuración de dip-switches para microstepping, ajuste de límites de corriente y conexionado de señales de control lógico (PUL, DIR, ENA)."
            },
            {
                "num": 15,
                "title": "Actuación Electroneumática y Normas de Cableado Kinal",
                "cat": "control",
                "cat_name": "🟣 Control & Automatización",
                "cat_badge": "bg-purple-100 text-purple-900 border border-purple-200",
                "desc": "Electroválvulas de 24 VDC monoestables y biestables, y actuadores neumáticos de simple/doble efecto. Enrutamiento estético en ductos ranurados, peinado de cables con espirales y terminales de ferrul."
            },
            {
                "num": 16,
                "title": "Programación de Perfiles de Movimiento y Evaluación Práctica",
                "cat": "ensamble",
                "cat_name": "🟢 Manufactura & Ensamble",
                "cat_badge": "bg-emerald-100 text-emerald-900 border border-emerald-200",
                "desc": "Lógica de rampas de aceleración y desaceleración para evitar pérdida de pasos. Integración de parada de emergencia con corte seguro y evaluación: Sistema de Posicionamiento Lineal Motorizado."
            }
        ]
    },
    {
        "num": 5,
        "title": "Integración de Celda Mecatrónica, Automatización y Proyecto Capstone Industrial",
        "period": "MÓDULO 5 (Junio • Sesiones 17 a 20)",
        "hours": "18 Horas de Taller",
        "logro": "Integrar todos los subsistemas en una celda mecatrónica funcional automatizada y demostrar su operación autónoma ante jurado técnico evaluador.",
        "topics": [
            {
                "num": 17,
                "title": "Arquitectura de Integración Mecatrónica e Interfaz de Operador",
                "cat": "control",
                "cat_name": "🟣 Control & Automatización",
                "cat_badge": "bg-purple-100 text-purple-900 border border-purple-200",
                "desc": "Diseño de tableros de control con botoneras industriales (arranque, paro, rearme) y luces piloto. Diagramas de flujo y autómatas de estado finito (FSM) para ciclos secuenciales industriales coordinados."
            },
            {
                "num": 18,
                "title": "Ensamble Mecánico y Cableado Definitivo de la Celda",
                "cat": "ensamble",
                "cat_name": "🟢 Manufactura & Ensamble",
                "cat_badge": "bg-emerald-100 text-emerald-900 border border-emerald-200",
                "desc": "Montaje integral sobre bancada CNC de piezas cinemáticas 3D y gabinetes cortados en láser. Verificación punto a punto de continuidad eléctrica, aislamiento de seguridad y calibración de bandas dentadas."
            },
            {
                "num": 19,
                "title": "Puesta en Marcha, Sincronización y Diagnóstico Sistemático (Troubleshooting)",
                "cat": "control",
                "cat_name": "🟣 Control & Automatización",
                "cat_badge": "bg-purple-100 text-purple-900 border border-purple-200",
                "desc": "Carga del programa secuencial en el PLC Siemens S7-1200 y pruebas de marcha en vacío (dry run). Protocolos de localización de averías y resolución de fallas inducidas en sensores y actuadores."
            },
            {
                "num": 20,
                "title": "Evaluación Terminal Integrada (Capstone) y Clausura",
                "cat": "ensamble",
                "cat_name": "🟢 Manufactura & Ensamble",
                "cat_badge": "bg-emerald-100 text-emerald-900 border border-emerald-200",
                "desc": "Demostración de funcionamiento continuo (mínimo 10 ciclos completos autónomos). Sustentación técnica individual y grupal, revisión de la Bitácora Berichtsheft y acreditación final (umbral ≥ 75 pts)."
            }
        ]
    }
]

# ==============================================================================
# 2. PARSE CIBERSEGURIDAD
# ==============================================================================
ciber_prop_p, ciber_prop_t = parse_docx(os.path.join(CURSOS_DIR, "Ciberseguridad", "Propuesta_Curso_Ciberseguridad_Kinal.docx"))
ciber_dos_p, ciber_dos_t = parse_docx(os.path.join(CURSOS_DIR, "Ciberseguridad", "Dosificacion_y_Secuencia_Didactica_Ciberseguridad.docx"))

ciber_ficha = ciber_prop_t[0]
ciber_dqr = ciber_prop_t[5]
ciber_mod_prop = ciber_prop_t[6]
ciber_eval = ciber_prop_t[7]
ciber_rubric = ciber_dos_t[22]

ciber_sessions = []
# Presencial en Kinal: Sesiones 1, 8, 14, 20 (20% de 20 sesiones = 4 presenciales)
presenciales_ciber = {1, 8, 14, 20}

for idx in range(2, 22):
    t = ciber_dos_t[idx]
    s_num = idx - 1
    t1_row = ciber_dos_t[1][s_num] if s_num < len(ciber_dos_t[1]) else ["", "", "", ""]
    
    obj = t[0][1]
    ap = t[1][1]
    des = t[2][1]
    cie = t[3][1]
    evi = t[4][1]
    eq = t[5][1]
    
    mod_name = t1_row[1] if len(t1_row) > 1 else f"Módulo {s_num}"
    title_short = t1_row[2] if len(t1_row) > 2 else f"Sesión {s_num}"
    
    is_pres = s_num in presenciales_ciber
    modalidad = "Presencial en Sede Kinal" if is_pres else "Virtual Síncrona"
    semana_num = ((s_num - 1) // 2) + 1

    ciber_sessions.append({
        "num": s_num,
        "code": f"S-{s_num:02d}",
        "semana": f"Semana {semana_num}",
        "modalidad": modalidad,
        "is_presencial": is_pres,
        "horario": "Encuentro de 4 Horas (240 min)",
        "module": mod_name,
        "title": title_short,
        "duration": "4 Horas (240 min)",
        "objective": obj,
        "apertura": ap,
        "desarrollo": des,
        "cierre": cie,
        "evidencias": evi,
        "equipamiento": eq
    })

ciber_modules_data = [
    {
        "num": 1,
        "title": "Sistemas y Gestión Informática para la Seguridad",
        "period": "MÓDULO 1 (Sesiones 1 a 3)",
        "hours": "12 Horas Pedagógicas",
        "logro": "Representar componentes, flujos de red y dependencias de infraestructura aplicando la tríada CIA y estándares ITIL en la empresa del caso transversal.",
        "topics": [
            {
                "num": 1,
                "title": "Fundamentos de Seguridad de la Información y Tríada CIA",
                "cat": "sistemas",
                "cat_name": "🔵 Sistemas & Redes TI",
                "cat_badge": "bg-blue-100 text-blue-900 border border-blue-200",
                "desc": "Confidencialidad, integridad y disponibilidad aplicadas al negocio. Identificación de activos de información, niveles de criticidad y construcción de la Matriz CIA en la empresa ficticia."
            },
            {
                "num": 2,
                "title": "Soporte Técnico, Principios ITIL y Gestión de Servicios",
                "cat": "sistemas",
                "cat_name": "🔵 Sistemas & Redes TI",
                "cat_badge": "bg-blue-100 text-blue-900 border border-blue-200",
                "desc": "Clasificación de incidentes, solicitudes de servicio y problemas. Registro trazable de eventos, gestión de cambios autorizados y documentación de tickets de soporte técnico con criterio defensivo."
            },
            {
                "num": 3,
                "title": "Arquitectura de Sistemas, Topología de Red y Flujos de Datos",
                "cat": "sistemas",
                "cat_name": "🔵 Sistemas & Redes TI",
                "cat_badge": "bg-blue-100 text-blue-900 border border-blue-200",
                "desc": "Diagramación de servidores, terminales de usuario, dispositivos de borde y servicios SaaS. Segmentación de red (LAN vs. DMZ) e identificación de puntos únicos de falla (SPOF)."
            }
        ]
    },
    {
        "num": 2,
        "title": "Identidad y Control de Acceso en Aplicaciones y Sistemas",
        "period": "MÓDULO 2 (Sesiones 4 a 6)",
        "hours": "12 Horas Pedagógicas",
        "logro": "Diseñar esquemas de autenticación robustos, políticas de contraseñas, doble factor (MFA) y matrices de acceso basado en roles (RBAC).",
        "topics": [
            {
                "num": 4,
                "title": "Autenticación, Validación y Ciclo de Vida de la Sesión",
                "cat": "identidad",
                "cat_name": "🟣 Identidad & Acceso (IAM)",
                "cat_badge": "bg-purple-100 text-purple-900 border border-purple-200",
                "desc": "Factores de autenticación (conocimiento, posesión, inherencia). Almacenamiento seguro de tokens y credenciales, fijación de sesiones, tiempos de expiración y cierre seguro de conexiones."
            },
            {
                "num": 5,
                "title": "Autenticación Multifactor (MFA) y Políticas de Contraseñas",
                "cat": "identidad",
                "cat_name": "🟣 Identidad & Acceso (IAM)",
                "cat_badge": "bg-purple-100 text-purple-900 border border-purple-200",
                "desc": "Configuración de MFA basado en TOTP y llaves de seguridad FIDO2. Mitigación de fatiga de MFA, directrices de longitud y aleatoriedad, y restricciones de bloqueo de cuenta."
            },
            {
                "num": 6,
                "title": "Single Sign-On (SSO), Protocolos Federados y Matriz RBAC",
                "cat": "identidad",
                "cat_name": "🟣 Identidad & Acceso (IAM)",
                "cat_badge": "bg-purple-100 text-purple-900 border border-purple-200",
                "desc": "Conceptos de LDAP, Active Directory, SAML y OAuth 2.0 / OpenID Connect. Elaboración de la Matriz RBAC para altas, traslados departamentales y bajas de personal (Offboarding)."
            }
        ]
    },
    {
        "num": 3,
        "title": "Diseño de Controles y Responsabilidades de Seguridad",
        "period": "MÓDULO 3 (Sesiones 7 a 8)",
        "hours": "8 Horas Pedagógicas",
        "logro": "Formular controles operacionales con alcance explícito, responsable asignado, procedimiento de ejecución y evidencia comprobable de cumplimiento.",
        "topics": [
            {
                "num": 7,
                "title": "Principio de Menor Privilegio y Segregación de Funciones",
                "cat": "gobierno",
                "cat_name": "🟢 Gobierno & Cumplimiento",
                "cat_badge": "bg-emerald-100 text-emerald-900 border border-emerald-200",
                "desc": "Justificación técnica de privilegios administrativos mínimos. Eliminación de accesos compartidos y separación de funciones entre desarrollo, operación y auditoría en el caso empresarial."
            },
            {
                "num": 8,
                "title": "Diseño y Asignación de Controles Operacionales Auditables",
                "cat": "gobierno",
                "cat_name": "🟢 Gobierno & Cumplimiento",
                "cat_badge": "bg-emerald-100 text-emerald-900 border border-emerald-200",
                "desc": "Estructura formal de controles: alcance, custodio del control, frecuencia de ejecución, métrica de eficacia y registro de evidencia verificable para auditorías internas."
            }
        ]
    },
    {
        "num": 4,
        "title": "Seguridad de Comunicaciones Web y Protección de APIs",
        "period": "MÓDULO 4 (Sesiones 9 a 11)",
        "hours": "12 Horas Pedagógicas",
        "logro": "Analizar protocolos criptográficos, transporte seguro TLS/HTTPS, inspeccionar tráfico con Wireshark y verificar la protección de endpoints web.",
        "topics": [
            {
                "num": 9,
                "title": "Criptografía Aplicada, Algoritmos, Claves y Funciones Hash",
                "cat": "sistemas",
                "cat_name": "🔵 Sistemas & Redes TI",
                "cat_badge": "bg-blue-100 text-blue-900 border border-blue-200",
                "desc": "Cifrado simétrico (AES-256) vs. asimétrico (RSA/ECC). Funciones resumen unidireccionales (SHA-256, SHA-3) para integridad de datos, firma digital y no repudio."
            },
            {
                "num": 10,
                "title": "Arquitectura PKI y Certificados Digitales TLS/HTTPS",
                "cat": "sistemas",
                "cat_name": "🔵 Sistemas & Redes TI",
                "cat_badge": "bg-blue-100 text-blue-900 border border-blue-200",
                "desc": "Jerarquía de autoridades certificadoras (CA), emisión de certificados x509, ciclo de vida y revocación (CRL/OCSP). Verificación de la negociación de cifrado en el navegador."
            },
            {
                "num": 11,
                "title": "Inspección de Tráfico y Seguridad en el Recorrido Web",
                "cat": "operaciones",
                "cat_name": "🟡 Operaciones & Diagnóstico (SOC)",
                "cat_badge": "bg-amber-100 text-amber-900 border border-amber-200",
                "desc": "Inspección pasiva de paquetes con Wireshark, análisis de encabezados HTTP seguros (HSTS, Content Security Policy) y verificación de autenticación en endpoints de APIs REST."
            }
        ]
    },
    {
        "num": 5,
        "title": "Ciclo de Vida y Gestión de Vulnerabilidades",
        "period": "MÓDULO 5 (Sesiones 12 a 14)",
        "hours": "12 Horas Pedagógicas",
        "logro": "Interpretar avisos CVE y puntuación CVSS, priorizar remediaciones y gestionar actualizaciones sin generar disrupción en las operaciones de negocio.",
        "topics": [
            {
                "num": 12,
                "title": "Gestión del Ciclo de Vida de Hardware y Sistemas Obsoletos",
                "cat": "sistemas",
                "cat_name": "🔵 Sistemas & Redes TI",
                "cat_badge": "bg-blue-100 text-blue-900 border border-blue-200",
                "desc": "Control de fin de soporte (EOL/EOS) en conmutadores, servidores y sistemas operativos. Evaluación de riesgos de continuidad operativa y justificación de reemplazo tecnológico."
            },
            {
                "num": 13,
                "title": "Gestión de Actualizaciones y Parches en Ambientes de Prueba",
                "cat": "operaciones",
                "cat_name": "🟡 Operaciones & Diagnóstico (SOC)",
                "cat_badge": "bg-amber-100 text-amber-900 border border-amber-200",
                "desc": "Metodología de pruebas en entornos de laboratorio o staging antes del despliegue en producción. Ventanas de mantenimiento, procedimientos de reversión (rollback) y respaldo previo."
            },
            {
                "num": 14,
                "title": "Interpretación de Avisos CVE y Puntuación CVSS v3/v4",
                "cat": "operaciones",
                "cat_name": "🟡 Operaciones & Diagnóstico (SOC)",
                "cat_badge": "bg-amber-100 text-amber-900 border border-amber-200",
                "desc": "Lectura de boletines oficiales de vulnerabilidades del NIST NVD. Desglose del vector de ataque CVSS (red, complejidad, privilegios, interacción) y formulación de planes de mitigación temporal."
            }
        ]
    },
    {
        "num": 6,
        "title": "Documentación, Políticas y Gestión del Riesgo de Seguridad",
        "period": "MÓDULO 6 (Sesiones 15 a 16)",
        "hours": "8 Horas Pedagógicas",
        "logro": "Redactar políticas organizacionales claras, procedimientos operativos estándar (SOP) ejecutables y solicitudes formales de excepción técnica.",
        "topics": [
            {
                "num": 15,
                "title": "Estructura Documental y Redacción de Políticas de Seguridad",
                "cat": "gobierno",
                "cat_name": "🟢 Gobierno & Cumplimiento",
                "cat_badge": "bg-emerald-100 text-emerald-900 border border-emerald-200",
                "desc": "Jerarquía normativa organizacional: políticas marco, directrices técnicas, estándares y procedimientos. Redacción sin ambigüedades de reglas de uso aceptable y control de accesos."
            },
            {
                "num": 16,
                "title": "Procedimientos Operativos Estándar (SOP) y Solicitudes de Excepción",
                "cat": "gobierno",
                "cat_name": "🟢 Gobierno & Cumplimiento",
                "cat_badge": "bg-emerald-100 text-emerald-900 border border-emerald-200",
                "desc": "Elaboración de guías de ejecución paso a paso verificables. Gestión de excepciones autorizadas de seguridad con justificación de negocio, controles compensatorios y fecha de caducidad."
            }
        ]
    },
    {
        "num": 7,
        "title": "Estándares Internacionales, Regulaciones e Informes de Aseguramiento",
        "period": "MÓDULO 7 (Sesiones 17 a 18)",
        "hours": "8 Horas Pedagógicas",
        "logro": "Mapear controles frente a marcos internacionales de cumplimiento (PCI-DSS, HIPAA, ISO 27001, SOC 2) y preparar auditorías formales de TI.",
        "topics": [
            {
                "num": 17,
                "title": "Estándar PCI-DSS y Protección de Datos de Medios de Pago",
                "cat": "gobierno",
                "cat_name": "🟢 Gobierno & Cumplimiento",
                "cat_badge": "bg-emerald-100 text-emerald-900 border border-emerald-200",
                "desc": "Delimitación del entorno de datos del tarjetahabiente (CDE). Requisitos esenciales de segmentación, cifrado de transmisiones, controles de acceso físico y monitoreo de redes."
            },
            {
                "num": 18,
                "title": "Regulación HIPAA (ePHI) e Informes SOC 1 y SOC 2",
                "cat": "gobierno",
                "cat_name": "🟢 Gobierno & Cumplimiento",
                "cat_badge": "bg-emerald-100 text-emerald-900 border border-emerald-200",
                "desc": "Reglas de privacidad y seguridad de información de salud. Principios de servicios de confianza AICPA (Seguridad, Disponibilidad, Integridad, Confidencialidad) y lectura de reportes SOC Tipo II."
            }
        ]
    },
    {
        "num": 8,
        "title": "Higiene Digital, Respuesta Inicial a Incidentes y Cierre del Caso Transversal",
        "period": "MÓDULO 8 (Sesiones 19 a 20)",
        "hours": "8 Horas Pedagógicas",
        "logro": "Reconocer vectores de ingeniería social, aplicar protocolos iniciales de contención de incidentes y sustentar el Expediente Terminal del caso transversal.",
        "topics": [
            {
                "num": 19,
                "title": "Concientización en Seguridad, Phishing y Respuesta a Incidentes",
                "cat": "operaciones",
                "cat_name": "🟡 Operaciones & Diagnóstico (SOC)",
                "cat_badge": "bg-amber-100 text-amber-900 border border-amber-200",
                "desc": "Identificación de señales de correo fraudulento e ingeniería social. Procedimiento de aislamiento de estaciones comprometidas, reporte inmediato y preservación de registros."
            },
            {
                "num": 20,
                "title": "Consolidación y Sustentación del Expediente de Seguridad",
                "cat": "gobierno",
                "cat_name": "🟢 Gobierno & Cumplimiento",
                "cat_badge": "bg-emerald-100 text-emerald-900 border border-emerald-200",
                "desc": "Presentación y defensa técnica del portafolio terminal: Matriz CIA, diseño RBAC, políticas SOP, informe de gestión de vulnerabilidades y lecciones aprendidas ante el jurado Kinal (≥ 75 pts)."
            }
        ]
    }
]

print("Construyendo visualizadores...")

# ==============================================================================
# GENERADORES DE TARJETAS MODULARES PARA TEMARIO
# ==============================================================================
def build_modular_temario_html(modules):
    cards_html = ""
    for mod in modules:
        topics_items = ""
        for t in mod["topics"]:
            topics_items += f"""
            <div class="content-item item-{t['cat']} p-3.5 rounded-xl text-xs space-y-1.5 transition border border-slate-200/90 bg-white hover:bg-slate-50 shadow-sm" data-category="{t['cat']}">
                <div class="flex flex-wrap justify-between items-center gap-1.5 font-bold">
                    <span class="text-slate-900 text-sm flex items-center gap-2">
                        <span class="w-5 h-5 rounded-full bg-slate-100 text-slate-700 text-xs font-black flex items-center justify-center shrink-0">{t['num']}</span>
                        {clean_html(t['title'])}
                    </span>
                    <span class="px-2.5 py-0.5 rounded-full text-[10px] font-extrabold uppercase {t['cat_badge']}">
                        {clean_html(t['cat_name'])}
                    </span>
                </div>
                <p class="text-slate-600 leading-relaxed pl-7">{clean_html(t['desc'])}</p>
            </div>
            """
        
        cards_html += f"""
        <div class="bg-white rounded-2xl border border-slate-300 shadow-sm overflow-hidden flex flex-col hover:shadow-md transition">
            <div class="bg-slate-900 text-white p-4 sm:p-5 border-b border-slate-800 space-y-1">
                <div class="flex flex-wrap justify-between items-center text-xs text-amber-400 font-semibold mb-1 gap-2">
                    <span>{clean_html(mod['period'])}</span>
                    <span class="px-2 py-0.5 rounded bg-blue-900/60 text-blue-200 border border-blue-400/30 text-[11px] font-bold">{clean_html(mod['hours'])}</span>
                </div>
                <h4 class="text-base sm:text-lg font-bold text-white leading-snug">{clean_html(mod['title'])}</h4>
                <p class="text-xs text-slate-300 pt-1 leading-relaxed">
                    <strong class="text-amber-300">Logro esperado:</strong> {clean_html(mod['logro'])}
                </p>
            </div>
            <div class="p-4 sm:p-5 divide-y divide-slate-100 flex-1 space-y-2.5 bg-slate-50/50">
                {topics_items}
            </div>
        </div>
        """
    return cards_html

meca_modular_cards = build_modular_temario_html(meca_modules_data)
ciber_modular_cards = build_modular_temario_html(ciber_modules_data)

# ==============================================================================
# GENERADORES DE FICHA TÉCNICA UNIFICADA (UNA SOLA TARJETA)
# ==============================================================================
def build_single_card_ficha(ficha_rows, code):
    rows_html = ""
    for r in ficha_rows:
        rows_html += f"""
        <div class="py-3 sm:grid sm:grid-cols-3 sm:gap-4 flex flex-col">
            <dt class="font-bold text-slate-900 text-xs sm:col-span-1">{clean_html(r[0])}</dt>
            <dd class="text-slate-700 text-xs sm:col-span-2 mt-0.5 sm:mt-0 leading-relaxed">{clean_html(r[1])}</dd>
        </div>
        """
    
    return f"""
    <div class="bg-white rounded-3xl border border-slate-200 p-6 sm:p-8 shadow-sm space-y-5">
        <div class="flex flex-wrap items-center justify-between border-b border-slate-200 pb-4 gap-2">
            <div class="flex items-center gap-2">
                <span class="w-8 h-8 rounded-lg bg-blue-50 text-blue-900 flex items-center justify-center font-black">
                    <i data-lucide="file-badge" class="w-4 h-4"></i>
                </span>
                <h3 class="text-xl font-black text-slate-900">
                    Ficha Técnica Oficial del Programa
                </h3>
            </div>
            <span class="px-3 py-1 rounded-full bg-blue-50 text-blue-900 text-xs font-bold border border-blue-200">
                Código: {clean_html(code)}
            </span>
        </div>
        <div class="divide-y divide-slate-100">
            {rows_html}
        </div>
    </div>
    """

meca_ficha_single_card = build_single_card_ficha(meca_ficha, "MEC-2026-KINAL")
ciber_ficha_single_card = build_single_card_ficha(ciber_ficha, "CIBER-2026-KINAL")

# ==============================================================================
# GENERADOR DEL CALENDARIO CRONOLÓGICO VISUAL
# ==============================================================================
def build_meca_calendar_html(sessions):
    months = ["Febrero", "Marzo", "Abril", "Mayo", "Junio"]
    months_html = ""
    for m in months:
        m_sessions = [s for s in sessions if s["month"] == m]
        cards = ""
        for s in m_sessions:
            cards += f"""
            <div class="bg-white rounded-xl border border-slate-200 p-4 shadow-sm hover:border-blue-600 hover:shadow-md transition flex flex-col justify-between space-y-3">
                <div class="space-y-1.5">
                    <div class="flex justify-between items-center text-[11px]">
                        <span class="font-black px-2 py-0.5 rounded bg-blue-900 text-white">{s['code']}</span>
                        <span class="font-bold text-slate-500">Sábado {s['sabado_num']}</span>
                    </div>
                    <h5 class="font-bold text-slate-900 text-xs leading-snug">{clean_html(s['title'])}</h5>
                    <p class="text-[11px] text-slate-500 line-clamp-2">{clean_html(s['objective'])}</p>
                </div>
                <div class="pt-2 border-t border-slate-100 flex items-center justify-between text-[11px]">
                    <span class="text-blue-800 font-semibold">{s['horario']}</span>
                    <button onclick="goToSession({s['num']})" class="text-blue-600 font-bold hover:underline flex items-center gap-0.5">
                        Ver Secuencia →
                    </button>
                </div>
            </div>
            """
        months_html += f"""
        <div class="bg-slate-50 rounded-2xl border border-slate-200 p-5 space-y-3">
            <div class="flex items-center justify-between border-b border-slate-200 pb-2">
                <h4 class="font-black text-slate-900 text-base flex items-center gap-2">
                    <i data-lucide="calendar" class="w-4 h-4 text-blue-700"></i> Mes de {m} (4 Sábados)
                </h4>
                <span class="text-xs font-bold text-slate-500">18 Horas de Taller</span>
            </div>
            <div class="grid sm:grid-cols-2 lg:grid-cols-4 gap-3">
                {cards}
            </div>
        </div>
        """
    return months_html

def build_ciber_calendar_html(sessions):
    weeks_html = ""
    # 10 Semanas (2 sesiones por semana)
    for w in range(1, 11):
        w_sessions = [s for s in sessions if s["semana"] == f"Semana {w}"]
        cards = ""
        for s in w_sessions:
            badge_mod = '<span class="px-2 py-0.5 rounded bg-emerald-100 text-emerald-800 font-bold text-[10px]">🏢 Presencial en Kinal</span>' if s["is_presencial"] else '<span class="px-2 py-0.5 rounded bg-indigo-50 text-indigo-800 font-bold text-[10px]">💻 Virtual Síncrona</span>'
            cards += f"""
            <div class="bg-white rounded-xl border border-slate-200 p-4 shadow-sm hover:border-indigo-600 hover:shadow-md transition flex flex-col justify-between space-y-3">
                <div class="space-y-1.5">
                    <div class="flex justify-between items-center text-[11px]">
                        <span class="font-black px-2 py-0.5 rounded bg-indigo-900 text-white">{s['code']}</span>
                        {badge_mod}
                    </div>
                    <h5 class="font-bold text-slate-900 text-xs leading-snug">{clean_html(s['title'])}</h5>
                    <p class="text-[11px] text-slate-500 line-clamp-2">{clean_html(s['objective'])}</p>
                </div>
                <div class="pt-2 border-t border-slate-100 flex items-center justify-between text-[11px]">
                    <span class="text-slate-600 font-semibold">{s['duration']}</span>
                    <button onclick="goToSession({s['num']})" class="text-indigo-600 font-bold hover:underline flex items-center gap-0.5">
                        Ver Secuencia →
                    </button>
                </div>
            </div>
            """
        weeks_html += f"""
        <div class="bg-slate-50 rounded-2xl border border-slate-200 p-4 space-y-3">
            <div class="flex items-center justify-between border-b border-slate-200 pb-2">
                <h4 class="font-black text-slate-900 text-sm flex items-center gap-2">
                    <i data-lucide="clock" class="w-4 h-4 text-indigo-700"></i> Semana {w}
                </h4>
                <span class="text-xs font-bold text-slate-500">2 Sesiones (8 Horas)</span>
            </div>
            <div class="grid sm:grid-cols-2 gap-3">
                {cards}
            </div>
        </div>
        """
    return f'<div class="grid lg:grid-cols-2 gap-4">{weeks_html}</div>'

meca_calendar_content = build_meca_calendar_html(meca_sessions)
ciber_calendar_content = build_ciber_calendar_html(ciber_sessions)

# Rúbricas
meca_rubric_rows_html = "".join([f"""
    <tr class="hover:bg-slate-50/80 transition">
        <td class="py-3.5 px-4 font-bold text-slate-900 bg-slate-50 border-r border-slate-200 align-top">{clean_html(row[0])}</td>
        <td class="py-3 px-4 text-emerald-950 bg-emerald-50/30 border-r border-slate-200 align-top">{clean_html(row[1])}</td>
        <td class="py-3 px-4 text-blue-950 bg-blue-50/30 border-r border-slate-200 align-top">{clean_html(row[2])}</td>
        <td class="py-3 px-4 text-amber-950 bg-amber-50/30 border-r border-slate-200 align-top font-medium">{clean_html(row[3])}</td>
        <td class="py-3 px-4 text-rose-950 bg-rose-50/30 align-top text-xs">{clean_html(row[4])}</td>
    </tr>
""" for r_idx, row in enumerate(meca_rubric) if r_idx > 0])

ciber_rubric_rows_html = "".join([f"""
    <tr class="hover:bg-slate-50/80 transition">
        <td class="py-3.5 px-4 font-bold text-slate-900 bg-slate-50 border-r border-slate-200 align-top">{clean_html(row[0])}</td>
        <td class="py-3 px-4 text-emerald-950 bg-emerald-50/30 border-r border-slate-200 align-top">{clean_html(row[1])}</td>
        <td class="py-3 px-4 text-blue-950 bg-blue-50/30 border-r border-slate-200 align-top">{clean_html(row[2])}</td>
        <td class="py-3 px-4 text-amber-950 bg-amber-50/30 border-r border-slate-200 align-top font-medium">{clean_html(row[3])}</td>
        <td class="py-3 px-4 text-rose-950 bg-rose-50/30 align-top text-xs">{clean_html(row[4])}</td>
    </tr>
""" for r_idx, row in enumerate(ciber_rubric) if r_idx > 0])

print("Generando HTML de Mecatrónica...")

# ==============================================================================
# HTML BUILDER: MECATRONICA.HTML
# ==============================================================================
meca_sessions_json = json.dumps(meca_sessions, ensure_ascii=False)

mecatronica_html = f"""<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Mecatrónica Industrial y Fabricación Digital | Kinal 2026</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <script src="https://unpkg.com/lucide@latest"></script>
    <style>
        @import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700;800&family=Space+Grotesk:wght@500;700;800;900&display=swap');
        body {{ font-family: 'Inter', sans-serif; background-color: #f8fafc; color: #0f172a; }}
        h1, h2, h3, h4, .font-heading {{ font-family: 'Space Grotesk', sans-serif; }}
        .hero-pattern {{
            background: linear-gradient(135deg, #09172e 0%, #0f2d59 50%, #1e3a8a 100%);
        }}
        .tab-btn.active {{
            background-color: #003366;
            color: #ffffff;
            font-weight: 700;
        }}
        .session-item.active {{
            background-color: #003366;
            color: #ffffff;
            border-color: #003366;
        }}
        .session-item.active span, .session-item.active p {{
            color: #e2e8f0;
        }}
        .session-item.active .badge-code {{
            background-color: #f59e0b;
            color: #0f172a;
        }}
    </style>
</head>
<body class="antialiased min-h-screen flex flex-col">

    <!-- Top Navigation Header -->
    <header class="bg-slate-950 text-white border-b border-slate-800 sticky top-0 z-50 shadow-md">
        <div class="max-w-7xl mx-auto px-4 sm:px-6 lg:px-8 py-3 flex flex-wrap justify-between items-center gap-4">
            <div class="flex items-center gap-3">
                <a href="index.html" class="bg-slate-800 hover:bg-slate-700 text-slate-200 text-xs font-bold px-3 py-1.5 rounded-lg transition flex items-center gap-1.5">
                    <i data-lucide="arrow-left" class="w-3.5 h-3.5"></i> Cursos
                </a>
                <span class="text-slate-600">|</span>
                <div class="text-slate-200 font-bold text-xs sm:text-sm flex items-center gap-2">
                    <i data-lucide="cpu" class="w-4 h-4 text-blue-400"></i>
                    Mecatrónica Industrial & Fabricación Digital
                </div>
            </div>
            
            <!-- Quick Download Buttons -->
            <div class="flex items-center gap-2 text-xs">
                <span class="text-slate-400 hidden sm:inline">Word:</span>
                <a href="../Mecatrónica/Propuesta_Curso_Mecatronica_Kinal.docx" download class="px-2.5 py-1 rounded bg-blue-900/60 hover:bg-blue-800 text-blue-200 transition font-medium" title="Descargar Propuesta Oficial">
                    Propuesta
                </a>
                <a href="../Mecatrónica/Temario_Curso_Mecatronica_Kinal.docx" download class="px-2.5 py-1 rounded bg-emerald-900/60 hover:bg-emerald-800 text-emerald-200 transition font-medium" title="Descargar Temario Oficial">
                    Temario
                </a>
                <a href="../Mecatrónica/Dosificacion_y_Secuencia_Didactica_Mecatronica_Kinal.docx" download class="px-2.5 py-1 rounded bg-amber-900/60 hover:bg-amber-800 text-amber-200 transition font-medium" title="Descargar Dosificación Oficial">
                    Dosificación
                </a>
            </div>
        </div>
    </header>

    <!-- Hero Course Banner -->
    <section class="hero-pattern text-white py-10 px-4 sm:px-6 lg:px-8">
        <div class="max-w-7xl mx-auto space-y-3">
            <div class="flex flex-wrap items-center gap-2">
                <span class="px-3 py-1 rounded-full bg-blue-500/20 border border-blue-400/30 text-blue-300 text-xs font-bold tracking-wide uppercase">
                    Especialidad Presencial en Taller
                </span>
                <span class="px-3 py-1 rounded-full bg-amber-500/20 border border-amber-400/30 text-amber-300 text-xs font-bold">
                    DQR Nivel 4 - 5
                </span>
                <span class="px-3 py-1 rounded-full bg-emerald-500/20 border border-emerald-400/30 text-emerald-300 text-xs font-bold">
                    Aprobación: ≥ 75 Pts
                </span>
            </div>
            <h1 class="text-2xl sm:text-4xl font-black tracking-tight leading-tight">
                Mecatrónica Industrial y Fabricación Digital Aplicada
            </h1>
            <p class="text-slate-300 text-xs sm:text-sm max-w-4xl leading-relaxed">
                Diseño CAD paramétrico 3D, ingeniería inversa, manufactura aditiva FDM, corte láser CO2, fresado CNC, sensórica industrial PNP/NPN, electroneumática y control con PLC Siemens S7-1200 para la automatización de líneas de producción.
            </p>
        </div>
    </section>

    <!-- Tab Bar Navigation (UNA SOLA PALABRA, SIN NUMERACIÓN) -->
    <div class="bg-white border-b border-slate-200 sticky top-14 z-40 shadow-sm">
        <div class="max-w-7xl mx-auto px-4 sm:px-6 lg:px-8">
            <nav class="flex space-x-2 sm:space-x-3 py-2.5 overflow-x-auto text-xs font-bold">
                <button onclick="switchTab('propuesta')" id="btn-propuesta" class="tab-btn active px-4 py-2 rounded-xl transition flex items-center gap-1.5 shrink-0">
                    <i data-lucide="file-text" class="w-3.5 h-3.5"></i> Propuesta
                </button>
                <button onclick="switchTab('temario')" id="btn-temario" class="tab-btn px-4 py-2 rounded-xl bg-slate-100 hover:bg-slate-200 text-slate-700 transition flex items-center gap-1.5 shrink-0">
                    <i data-lucide="list-tree" class="w-3.5 h-3.5"></i> Temario
                </button>
                <button onclick="switchTab('calendario')" id="btn-calendario" class="tab-btn px-4 py-2 rounded-xl bg-slate-100 hover:bg-slate-200 text-slate-700 transition flex items-center gap-1.5 shrink-0">
                    <i data-lucide="calendar" class="w-3.5 h-3.5"></i> Calendario
                </button>
                <button onclick="switchTab('secuencias')" id="btn-secuencias" class="tab-btn px-4 py-2 rounded-xl bg-slate-100 hover:bg-slate-200 text-slate-700 transition flex items-center gap-1.5 shrink-0">
                    <i data-lucide="layers" class="w-3.5 h-3.5"></i> Secuencias
                </button>
                <button onclick="switchTab('evaluacion')" id="btn-evaluacion" class="tab-btn px-4 py-2 rounded-xl bg-slate-100 hover:bg-slate-200 text-slate-700 transition flex items-center gap-1.5 shrink-0">
                    <i data-lucide="award" class="w-3.5 h-3.5"></i> Evaluación
                </button>
            </nav>
        </div>
    </div>

    <!-- Main Container -->
    <main class="max-w-7xl mx-auto px-4 sm:px-6 lg:px-8 py-8 flex-1 space-y-8">

        <!-- ================================================================== -->
        <!-- PESTAÑA 1: PROPUESTA -->
        <!-- ================================================================== -->
        <div id="tab-propuesta" class="tab-content space-y-8">
            
            <!-- Ficha Técnica (UNA SOLA TARJETA) -->
            {meca_ficha_single_card}

            <!-- Marco Filosófico Kinal -->
            <div class="grid md:grid-cols-2 gap-6">
                <div class="bg-blue-50 border-l-4 border-blue-900 p-6 rounded-r-3xl space-y-3 shadow-sm">
                    <div class="flex items-center gap-2 text-blue-900 font-extrabold text-sm uppercase">
                        <i data-lucide="flag" class="w-4 h-4"></i> Misión Institucional de Fundación Kinal
                    </div>
                    <p class="text-slate-700 text-xs italic leading-relaxed">
                        «Formar a jóvenes y adultos a través de una educación integral, con énfasis en las áreas técnicas y tecnológicas, influyendo positivamente en su trabajo, su familia y la sociedad».
                    </p>
                    <div class="pt-2 text-xs font-bold text-blue-950">
                        Valores Nucleares: <span class="font-normal italic">Visión cristiana de la vida. Respeto a la dignidad de la persona. Espíritu de servicio. Trabajo bien hecho. Libertad personal responsable.</span>
                    </div>
                </div>

                <div class="bg-amber-50 border-l-4 border-amber-600 p-6 rounded-r-3xl space-y-3 shadow-sm">
                    <div class="flex items-center gap-2 text-amber-900 font-extrabold text-sm uppercase">
                        <i data-lucide="sparkle" class="w-4 h-4"></i> El Principio Rector del «Trabajo Bien Hecho»
                    </div>
                    <p class="text-slate-700 text-xs leading-relaxed">
                        En el taller de Mecatrónica, el trabajo bien hecho se expresa en <strong>precisión dimensional milimétrica</strong> (±0.1 mm), orden estricto en el conexionado eléctrico bajo normativa IEC, enrutamiento estético en ductos ranurados, desbarbado impecable y respeto inquebrantable por las normas de seguridad industrial y protocolos LOTO.
                    </p>
                </div>
            </div>

            <!-- Matriz DQR -->
            <div class="bg-white rounded-3xl border border-slate-200 p-8 shadow-sm space-y-6">
                <div class="space-y-1">
                    <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                        <i data-lucide="award" class="w-5 h-5 text-blue-700"></i>
                        Alineación con el Marco Alemán de Cualificaciones (DQR Nivel 4-5)
                    </h3>
                    <p class="text-slate-600 text-xs">
                        Desglose de competencias integrales de acción (Handlungskompetenz) desarrolladas a lo largo del programa.
                    </p>
                </div>

                <div class="overflow-x-auto">
                    <table class="w-full text-xs text-left border-collapse">
                        <thead>
                            <tr class="bg-slate-100 text-slate-700 uppercase font-black border-b border-slate-200">
                                <th class="py-3 px-4">Dimensión DQR</th>
                                <th class="py-3 px-4">Subdimensión</th>
                                <th class="py-3 px-4">Evidencia de Desempeño en el Curso</th>
                            </tr>
                        </thead>
                        <tbody class="divide-y divide-slate-100 text-slate-700">
                            { "".join([f'<tr><td class="py-3 px-4 font-bold bg-slate-50">{clean_html(r[0])}</td><td class="py-3 px-4 font-semibold text-blue-900">{clean_html(r[1])}</td><td class="py-3 px-4 leading-relaxed">{clean_html(r[2])}</td></tr>' for r in meca_dqr[1:]]) }
                        </tbody>
                    </table>
                </div>
            </div>

            <!-- Perfiles de Ingreso y Egreso -->
            <div class="grid md:grid-cols-2 gap-6">
                <div class="bg-white rounded-3xl border border-slate-200 p-7 shadow-sm space-y-4">
                    <h4 class="text-lg font-black text-slate-900 flex items-center gap-2">
                        <i data-lucide="log-in" class="w-5 h-5 text-blue-700"></i> Perfil de Ingreso
                    </h4>
                    <p class="text-slate-600 text-xs leading-relaxed">
                        Dirigido a egresados o estudiantes de último año de carreras técnicas de nivel medio (Perito en Electricidad Industrial, Electrónica, Mecánica General o Automotriz), técnicos de mantenimiento industrial y operarios en activo con nociones básicas de taller.
                    </p>
                    <ul class="text-xs space-y-2 text-slate-700">
                        <li class="flex items-start gap-2"><i data-lucide="check" class="w-4 h-4 text-emerald-600 mt-0.5"></i> Manejo básico de computación e interés por el diseño 3D.</li>
                        <li class="flex items-start gap-2"><i data-lucide="check" class="w-4 h-4 text-emerald-600 mt-0.5"></i> Conocimiento elemental de circuitos eléctricos (AC/DC).</li>
                        <li class="flex items-start gap-2"><i data-lucide="check" class="w-4 h-4 text-emerald-600 mt-0.5"></i> Compromiso con las normas de seguridad de taller y uso de EPP.</li>
                    </ul>
                </div>

                <div class="bg-white rounded-3xl border border-slate-200 p-7 shadow-sm space-y-4">
                    <h4 class="text-lg font-black text-slate-900 flex items-center gap-2">
                        <i data-lucide="log-out" class="w-5 h-5 text-emerald-700"></i> Perfil de Egreso
                    </h4>
                    <p class="text-slate-600 text-xs leading-relaxed">
                        El graduado estará cualificado para diseñar piezas y ensambles en CAD 3D, operar cortadoras láser CO2 e impresoras 3D industriales, mecanizar bancadas en Router CNC, conectar sensórica PNP/NPN y programar rutinas de automatización en PLC Siemens S7-1200.
                    </p>
                    <ul class="text-xs space-y-2 text-slate-700">
                        <li class="flex items-start gap-2"><i data-lucide="check-circle" class="w-4 h-4 text-emerald-600 mt-0.5"></i> Autonomía operativa para fabricar fixtures y gabinetes mecánicos.</li>
                        <li class="flex items-start gap-2"><i data-lucide="check-circle" class="w-4 h-4 text-emerald-600 mt-0.5"></i> Diagnóstico sistemático de fallas en sensores, drivers y actuadores.</li>
                        <li class="flex items-start gap-2"><i data-lucide="check-circle" class="w-4 h-4 text-emerald-600 mt-0.5"></i> Cumplimiento estricto de las 5S y estándares internacionales de cableado.</li>
                    </ul>
                </div>
            </div>

            <!-- Arquitectura de Proyectos por Módulo -->
            <div class="bg-white rounded-3xl border border-slate-200 p-8 shadow-sm space-y-6">
                <div class="space-y-1">
                    <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                        <i data-lucide="blocks" class="w-5 h-5 text-blue-700"></i>
                        Arquitectura Curricular: 5 Módulos y Proyectos Prácticos
                    </h3>
                    <p class="text-slate-600 text-xs">
                        Cada mes el participante construye un componente tangible de la celda mecatrónica final.
                    </p>
                </div>
                <div class="overflow-x-auto">
                    <table class="w-full text-xs text-left border-collapse">
                        <thead>
                            <tr class="bg-slate-100 text-slate-700 uppercase font-black border-b border-slate-200">
                                <th class="py-3 px-4">Módulo / Mes</th>
                                <th class="py-3 px-4">Eje Temático Principal</th>
                                <th class="py-3 px-4">Tecnología Kinal</th>
                                <th class="py-3 px-4">Proyecto Práctico Tangible</th>
                            </tr>
                        </thead>
                        <tbody class="divide-y divide-slate-100 text-slate-700">
                            { "".join([f'<tr><td class="py-3 px-4 font-bold bg-slate-50 text-blue-900">{clean_html(r[0])}</td><td class="py-3 px-4 font-medium">{clean_html(r[1])}</td><td class="py-3 px-4">{clean_html(r[2])}</td><td class="py-3 px-4 font-bold text-emerald-800">{clean_html(r[3]) if len(r)>3 else ""}</td></tr>' for r in meca_mod_prop[1:]]) }
                        </tbody>
                    </table>
                </div>
            </div>

        </div>

        <!-- ================================================================== -->
        <!-- PESTAÑA 2: TEMARIO (DISTRIBUCIÓN MODULAR AUDITADA) -->
        <!-- ================================================================== -->
        <div id="tab-temario" class="tab-content hidden space-y-8">
            
            <!-- Barra Superior de Control y Filtros -->
            <div class="flex flex-col md:flex-row justify-between items-start md:items-center gap-4 bg-white rounded-3xl p-6 border border-slate-200 shadow-sm">
                <div class="space-y-1">
                    <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                        <i data-lucide="list-tree" class="w-5 h-5 text-blue-700"></i>
                        Distribución Modular con Código de Color
                    </h3>
                    <p class="text-xs text-slate-500">
                        Estructura organizada en 5 módulos y 20 temas de taller clasificados por naturaleza técnica y logros esperados.
                    </p>
                </div>

                <!-- Botones de Filtro Interactivo -->
                <div class="flex flex-wrap items-center gap-2 text-xs font-semibold">
                    <button onclick="filterMecaTemario('all')" id="btn-meca-all" class="meca-filter-btn px-3 py-1.5 rounded-lg bg-slate-900 text-white shadow-sm transition">Todos (20)</button>
                    <button onclick="filterMecaTemario('software')" id="btn-meca-software" class="meca-filter-btn px-3 py-1.5 rounded-lg bg-blue-100 text-blue-800 hover:bg-blue-200 transition">🔵 Modelado & Software</button>
                    <button onclick="filterMecaTemario('control')" id="btn-meca-control" class="meca-filter-btn px-3 py-1.5 rounded-lg bg-purple-100 text-purple-800 hover:bg-purple-200 transition">🟣 Control & Automatización</button>
                    <button onclick="filterMecaTemario('maquinaria')" id="btn-meca-maquinaria" class="meca-filter-btn px-3 py-1.5 rounded-lg bg-amber-100 text-amber-800 hover:bg-amber-200 transition">🟡 Maquinaria & Taller</button>
                    <button onclick="filterMecaTemario('ensamble')" id="btn-meca-ensamble" class="meca-filter-btn px-3 py-1.5 rounded-lg bg-emerald-100 text-emerald-800 hover:bg-emerald-200 transition">🟢 Manufactura & Ensamble</button>
                </div>
            </div>

            <!-- Código de Leyenda Explicativo -->
            <div class="grid grid-cols-1 sm:grid-cols-2 lg:grid-cols-4 gap-3 text-xs">
                <div class="p-3.5 rounded-xl border border-blue-200 bg-blue-50/70 text-blue-950 flex items-start gap-2.5">
                    <span class="inline-block w-3 h-3 rounded-full bg-blue-600 mt-0.5 shrink-0"></span>
                    <div>
                        <strong class="text-blue-900">Color Azul (Modelado & Software):</strong> CAD paramétrico 3D, tolerancias DFM, estrategias CAM, Código G y software laminador FDM.
                    </div>
                </div>
                <div class="p-3.5 rounded-xl border border-purple-200 bg-purple-50/70 text-purple-950 flex items-start gap-2.5">
                    <span class="inline-block w-3 h-3 rounded-full bg-purple-600 mt-0.5 shrink-0"></span>
                    <div>
                        <strong class="text-purple-900">Color Púrpura (Control & Automatización):</strong> Sensórica industrial PNP/NPN, electroválvulas Festo, diagramas FSM y PLC Siemens S7-1200.
                    </div>
                </div>
                <div class="p-3.5 rounded-xl border border-amber-200 bg-amber-50/70 text-amber-950 flex items-start gap-2.5">
                    <span class="inline-block w-3 h-3 rounded-full bg-amber-600 mt-0.5 shrink-0"></span>
                    <div>
                        <strong class="text-amber-900">Color Ámbar (Maquinaria & Taller):</strong> Operación de cortadora láser CO2, granja FDM, fresadora CNC y drivers de motores paso a paso.
                    </div>
                </div>
                <div class="p-3.5 rounded-xl border border-emerald-200 bg-emerald-50/70 text-emerald-950 flex items-start gap-2.5">
                    <span class="inline-block w-3 h-3 rounded-full bg-emerald-600 mt-0.5 shrink-0"></span>
                    <div>
                        <strong class="text-emerald-900">Color Verde (Manufactura & Ensamble):</strong> Prácticas terminales de taller, metrología de tolerancias (±0.1 mm) y celda Capstone integrada.
                    </div>
                </div>
            </div>

            <!-- Cuadrícula Modular de Temarios -->
            <div class="grid lg:grid-cols-2 gap-8 pt-2">
                {meca_modular_cards}
            </div>

            <!-- Maquinaria, Software y EPP -->
            <div class="grid md:grid-cols-3 gap-6 pt-4">
                <div class="bg-white rounded-2xl border border-slate-200 p-6 shadow-sm space-y-3">
                    <h4 class="font-black text-slate-900 text-base flex items-center gap-2">
                        <i data-lucide="wrench" class="w-4 h-4 text-blue-700"></i> Maquinaria & Equipos Kinal
                    </h4>
                    <ul class="text-xs space-y-2 text-slate-600">
                        <li>• Cortadora Láser CO2 de 80W a 100W con asistencia de aire y chiller.</li>
                        <li>• Granja de Impresoras 3D FDM industriales (cama ≥ 250×250 mm).</li>
                        <li>• Router / Fresadora CNC de 3 ejes con control G-Code.</li>
                        <li>• Banco de neumática y compresor silencioso con unidad FRL.</li>
                        <li>• Entrenadores modulares con PLC Siemens S7-1200 y fuentes de 24 VDC.</li>
                    </ul>
                </div>

                <div class="bg-white rounded-2xl border border-slate-200 p-6 shadow-sm space-y-3">
                    <h4 class="font-black text-slate-900 text-base flex items-center gap-2">
                        <i data-lucide="laptop" class="w-4 h-4 text-emerald-700"></i> Software Utilizado
                    </h4>
                    <ul class="text-xs space-y-2 text-slate-600">
                        <li>• Autodesk Fusion 360 / SolidWorks (CAD paramétrico 3D y CAM).</li>
                        <li>• RDWorks / LightBurn (Control y corte láser CO2).</li>
                        <li>• PrusaSlicer / Cura (Laminación aditiva avanzada para FDM).</li>
                        <li>• Siemens TIA Portal V18 / V19 (Programación Ladder de PLC).</li>
                        <li>• Festo FluidSIM (Simulación de circuitos electroneumáticos).</li>
                    </ul>
                </div>

                <div class="bg-white rounded-2xl border border-slate-200 p-6 shadow-sm space-y-3">
                    <h4 class="font-black text-slate-900 text-base flex items-center gap-2">
                        <i data-lucide="shield-alert" class="w-4 h-4 text-amber-700"></i> EPP y Seguridad Industrial
                    </h4>
                    <ul class="text-xs space-y-2 text-slate-600">
                        <li>• Gafas de seguridad certificadas ANSI Z87.1 (obligatorias).</li>
                        <li>• Calzado industrial con puntera de protección dieléctrica.</li>
                        <li>• Cabello recogido y cero uso de joyas o ropa holgada cerca de husillos.</li>
                        <li>• Tarjetas y candados de consignación LOTO en bancos eléctricos.</li>
                        <li>• Cumplimiento permanente de limpieza y orden 5S en cada puesto.</li>
                    </ul>
                </div>
            </div>

        </div>

        <!-- ================================================================== -->
        <!-- PESTAÑA 3: CALENDARIO CRONOLÓGICO -->
        <!-- ================================================================== -->
        <div id="tab-calendario" class="tab-content hidden space-y-6">
            
            <div class="flex flex-col sm:flex-row justify-between items-start sm:items-center gap-3 bg-white rounded-3xl p-6 border border-slate-200 shadow-sm">
                <div>
                    <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                        <i data-lucide="calendar" class="w-5 h-5 text-blue-700"></i>
                        Calendario Cronológico de Taller Sabatino
                    </h3>
                    <p class="text-slate-600 text-xs mt-1">
                        Programa sabatino continuo de 20 semanas (Febrero a Junio) • Sábados de 8:00 a 12:30 horas (4.5h por encuentro).
                    </p>
                </div>
                <div class="flex items-center gap-2 text-xs font-bold text-blue-900 bg-blue-50 px-3 py-1.5 rounded-xl border border-blue-200">
                    <i data-lucide="map-pin" class="w-4 h-4 text-blue-700"></i> Talleres Centrales Kinal
                </div>
            </div>

            <!-- Grid Mensual de Calendario -->
            <div class="space-y-6">
                {meca_calendar_content}
            </div>

        </div>

        <!-- ================================================================== -->
        <!-- PESTAÑA 4: SECUENCIAS (MASTER-DETAIL SIN ALARGAR LA PÁGINA) -->
        <!-- ================================================================== -->
        <div id="tab-secuencias" class="tab-content hidden space-y-6">
            
            <div class="flex flex-col sm:flex-row justify-between items-start sm:items-center gap-3 bg-white rounded-3xl p-5 border border-slate-200 shadow-sm">
                <div>
                    <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                        <i data-lucide="layers" class="w-5 h-5 text-blue-700"></i>
                        Secuencia Didáctica Interactiva Sesión a Sesión
                    </h3>
                    <p class="text-slate-600 text-xs mt-0.5">
                        Selecciona una sesión en el panel izquierdo para consultar su microdiseño en 3 momentos (Apertura, Desarrollo y Cierre).
                    </p>
                </div>
                <div class="flex items-center gap-2 w-full sm:w-auto">
                    <div class="relative w-full sm:w-56">
                        <i data-lucide="search" class="w-3.5 h-3.5 text-slate-400 absolute left-3 top-2.5"></i>
                        <input type="text" id="meca-sec-search" oninput="filterMecaList()" placeholder="Buscar sesión..." class="w-full pl-8 pr-3 py-1.5 text-xs rounded-xl border border-slate-200 bg-slate-50 focus:bg-white focus:outline-none focus:ring-2 focus:ring-blue-600 transition">
                    </div>
                </div>
            </div>

            <!-- Master-Detail Container -->
            <div class="grid lg:grid-cols-12 gap-6 items-start">
                
                <!-- LISTADO LATERAL (MASTER) -->
                <div class="lg:col-span-4 bg-white rounded-3xl border border-slate-200 p-4 shadow-sm space-y-2">
                    <div class="text-[11px] font-black uppercase tracking-wider text-slate-400 px-2 pb-1 border-b border-slate-100 flex justify-between">
                        <span>Listado de Sesiones</span>
                        <span id="meca-count-badge">20 Sesiones</span>
                    </div>
                    <div id="meca-sessions-list" class="space-y-1.5 max-h-[560px] overflow-y-auto pr-1">
                        <!-- Rendered by JS -->
                    </div>
                </div>

                <!-- DETALLE DE LA SESIÓN SELECCIONADA (DETAIL) -->
                <div class="lg:col-span-8 bg-white rounded-3xl border border-slate-200 p-6 sm:p-8 shadow-sm space-y-6" id="meca-detail-card">
                    <!-- Rendered by JS -->
                </div>

            </div>

        </div>

        <!-- ================================================================== -->
        <!-- PESTAÑA 5: EVALUACIÓN -->
        <!-- ================================================================== -->
        <div id="tab-evaluacion" class="tab-content hidden space-y-8">
            
            <!-- Esquema Resumen -->
            <div class="bg-white rounded-3xl border border-slate-200 p-8 shadow-sm space-y-6">
                <div class="flex items-center justify-between border-b border-slate-200 pb-3">
                    <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                        <i data-lucide="check-square" class="w-5 h-5 text-blue-700"></i>
                        Esquema de Evaluación Institucional & Acreditación
                    </h3>
                    <span class="px-3 py-1 rounded-full bg-amber-100 text-amber-900 font-black text-xs">
                        Umbral Mínimo: 75 / 100 Puntos
                    </span>
                </div>
                <div class="grid md:grid-cols-4 gap-4 text-xs">
                    <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 text-center space-y-1">
                        <div class="text-slate-500 font-bold uppercase text-[10px]">Evaluaciones Modulares</div>
                        <div class="text-3xl font-black text-blue-900">40%</div>
                        <div class="text-slate-600">4 Proyectos de Taller (10% c/u)</div>
                    </div>
                    <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 text-center space-y-1">
                        <div class="text-slate-500 font-bold uppercase text-[10px]">Proyecto Capstone Final</div>
                        <div class="text-3xl font-black text-emerald-700">35%</div>
                        <div class="text-slate-600">Puesta en Marcha de la Celda</div>
                    </div>
                    <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 text-center space-y-1">
                        <div class="text-slate-500 font-bold uppercase text-[10px]">Bitácora Berichtsheft</div>
                        <div class="text-3xl font-black text-purple-700">15%</div>
                        <div class="text-slate-600">Registro Semanal de Destrezas</div>
                    </div>
                    <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 text-center space-y-1">
                        <div class="text-slate-500 font-bold uppercase text-[10px]">Orden 5S & Seguridad</div>
                        <div class="text-3xl font-black text-amber-700">10%</div>
                        <div class="text-slate-600">EPP, LOTO y Trabajo Bien Hecho</div>
                    </div>
                </div>
            </div>

            <!-- Rúbrica Terminal -->
            <div class="bg-white rounded-3xl border border-slate-200 p-8 shadow-sm space-y-6">
                <div class="flex flex-col sm:flex-row justify-between items-start sm:items-center gap-2 border-b border-slate-200 pb-4">
                    <div>
                        <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                            <i data-lucide="award" class="w-5 h-5 text-blue-700"></i>
                            Rúbrica Analítica Terminal de Evaluación Práctica
                        </h3>
                        <p class="text-slate-600 text-xs mt-1">
                            Evaluación objetiva del desempeño técnico con ponderaciones explícitas y umbral institucional de 75 puntos.
                        </p>
                    </div>
                    <span class="px-3 py-1.5 rounded-xl bg-amber-100 text-amber-900 font-extrabold text-xs">
                        Aprobación Mínima: ≥ 75 Pts
                    </span>
                </div>

                <div class="overflow-x-auto">
                    <table class="w-full text-xs text-left border-collapse">
                        <thead>
                            <tr class="bg-slate-100 text-slate-700 uppercase font-black border-b border-slate-200">
                                <th class="py-3 px-4 w-1/5">Criterio Técnico</th>
                                <th class="py-3 px-4 w-1/5 text-emerald-800">Excelente (90 - 100)</th>
                                <th class="py-3 px-4 w-1/5 text-blue-800">Muy Bueno (80 - 89)</th>
                                <th class="py-3 px-4 w-1/5 text-amber-800">Aprobado Mínimo (75 - 79)</th>
                                <th class="py-3 px-4 w-1/5 text-rose-800">No Aprobado (&lt; 75)</th>
                            </tr>
                        </thead>
                        <tbody class="divide-y divide-slate-100 text-slate-700">
                            {meca_rubric_rows_html}
                        </tbody>
                    </table>
                </div>
            </div>

            <!-- Bitácora Berichtsheft -->
            <div class="bg-white rounded-3xl border border-slate-200 p-8 shadow-sm space-y-6">
                <div class="flex justify-between items-center border-b border-slate-200 pb-4">
                    <div>
                        <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                            <i data-lucide="book-open" class="w-5 h-5 text-blue-700"></i>
                            Formato de la Bitácora de Taller («Berichtsheft»)
                        </h3>
                        <p class="text-slate-600 text-xs mt-1">
                            Instrumento obligatorio del Sistema Dual para el registro semanal de horas y verificación de competencias.
                        </p>
                    </div>
                    <button onclick="window.print()" class="px-3.5 py-1.5 rounded-xl bg-slate-100 hover:bg-slate-200 text-slate-800 font-bold text-xs transition flex items-center gap-1.5">
                        <i data-lucide="printer" class="w-3.5 h-3.5"></i> Imprimir Bitácora
                    </button>
                </div>

                <div class="p-6 rounded-2xl bg-slate-50 border border-slate-300 space-y-4 text-xs font-mono">
                    <div class="border-b border-slate-300 pb-3 font-bold text-slate-900 text-center uppercase tracking-wider">
                        FUNDACIÓN KINAL • REGISTRO SEMANAL DE PRÁCTICA DUAL (BERICHTSHEFT)
                    </div>
                    <div class="grid grid-cols-2 gap-4">
                        <div><strong>Estudiante:</strong> ____________________________________</div>
                        <div><strong>Carné:</strong> ______________</div>
                        <div><strong>Sesión No.:</strong> _____ | <strong>Fecha:</strong> ____/____/2026</div>
                        <div><strong>Instructor de Taller:</strong> __________________________</div>
                    </div>
                    <div class="space-y-1 pt-2">
                        <strong>1. Máquinas y Equipos Operados en la Sesión:</strong>
                        <div class="flex gap-4 text-slate-700 pt-1">
                            <label><input type="checkbox" disabled> Cortadora Láser CO2</label>
                            <label><input type="checkbox" disabled> Impresora 3D FDM</label>
                            <label><input type="checkbox" disabled> Router CNC</label>
                            <label><input type="checkbox" disabled> Banco Festo</label>
                            <label><input type="checkbox" disabled> PLC S7-1200</label>
                        </div>
                    </div>
                    <div class="space-y-1 pt-2">
                        <strong>2. Descripción de la Tarea Técnica Realizada:</strong>
                        <div class="h-16 border border-slate-300 rounded bg-white p-2 text-slate-400">Espacio para detallar la actividad práctica y mediciones obtenidas...</div>
                    </div>
                    <div class="grid grid-cols-2 gap-6 pt-6 text-center">
                        <div class="border-t border-slate-400 pt-2">Firma del Estudiante</div>
                        <div class="border-t border-slate-400 pt-2">Visto Bueno Instructor Kinal (Sello/Firma)</div>
                    </div>
                </div>
            </div>

        </div>

    </main>

    <!-- Footer -->
    <footer class="bg-slate-950 text-slate-400 py-6 border-t border-slate-800 text-xs">
        <div class="max-w-7xl mx-auto px-4 sm:px-6 lg:px-8 flex flex-col sm:flex-row justify-between items-center gap-4 text-center sm:text-left">
            <div>
                <p class="font-bold text-slate-200">Fundación Kinal • Curso de Mecatrónica Industrial y Fabricación Digital</p>
                <p class="text-slate-500">Diseño Curricular, Didáctico e Instruccional 2026 | Sistema Dual & DQR 4-5</p>
            </div>
            <div class="flex items-center gap-4">
                <a href="index.html" class="hover:text-amber-400 transition">Catálogo de Cursos</a>
                <span class="text-slate-700">•</span>
                <a href="ciberseguridad.html" class="hover:text-amber-400 transition">Ver Ciberseguridad</a>
                <span class="text-slate-700">•</span>
                <a href="../../index.html" class="hover:text-amber-400 transition">Portal Maestro</a>
            </div>
        </div>
    </footer>

    <!-- Interactive Logic Script -->
    <script>
        const mecaSessions = {meca_sessions_json};
        let currentSessionNum = 1;

        function switchTab(tabId) {{
            document.querySelectorAll('.tab-content').forEach(el => el.classList.add('hidden'));
            document.querySelectorAll('.tab-btn').forEach(el => {{
                el.classList.remove('active');
                el.classList.add('bg-slate-100', 'text-slate-700');
            }});
            
            const target = document.getElementById('tab-' + tabId);
            if (target) target.classList.remove('hidden');
            const activeBtn = document.getElementById('btn-' + tabId);
            if (activeBtn) {{
                activeBtn.classList.add('active');
                activeBtn.classList.remove('bg-slate-100', 'text-slate-700');
            }}
            window.scrollTo({{ top: 220, behavior: 'smooth' }});
        }}

        function filterMecaTemario(category) {{
            document.querySelectorAll('.meca-filter-btn').forEach(b => {{
                b.classList.remove('bg-slate-900', 'text-white', 'shadow-sm');
            }});
            const activeBtn = document.getElementById('btn-meca-' + category);
            if (activeBtn) {{
                activeBtn.classList.add('bg-slate-900', 'text-white', 'shadow-sm');
            }}

            const items = document.querySelectorAll('#tab-temario .content-item');
            items.forEach(it => {{
                if (category === 'all' || it.getAttribute('data-category') === category) {{
                    it.style.display = 'block';
                }} else {{
                    it.style.display = 'none';
                }}
            }});
        }}

        function renderMecaSessionList(list) {{
            const listContainer = document.getElementById('meca-sessions-list');
            if (!listContainer) return;
            
            listContainer.innerHTML = list.map(s => `
                <div onclick="selectMecaSession(${{s.num}})" id="meca-list-item-${{s.num}}" class="session-item p-3 rounded-2xl border border-slate-200 bg-slate-50/70 hover:bg-slate-100 transition cursor-pointer flex items-center justify-between text-xs ${{s.num === currentSessionNum ? 'active' : ''}}">
                    <div class="flex items-center gap-2.5 overflow-hidden">
                        <span class="badge-code w-7 h-7 rounded-lg bg-blue-900 text-white font-black text-[11px] flex items-center justify-center shrink-0">
                            ${{s.code}}
                        </span>
                        <div class="truncate">
                            <span class="text-[10px] uppercase font-bold text-blue-700 block">${{s.month}}</span>
                            <div class="font-bold text-slate-900 truncate">${{s.title}}</div>
                        </div>
                    </div>
                    <i data-lucide="chevron-right" class="w-4 h-4 text-slate-400 shrink-0"></i>
                </div>
            `).join('');
            lucide.createIcons();
        }}

        function renderMecaSessionDetail(s) {{
            const detailContainer = document.getElementById('meca-detail-card');
            if (!detailContainer) return;

            detailContainer.innerHTML = `
                <div class="space-y-4">
                    <div class="flex flex-wrap justify-between items-start gap-2 border-b border-slate-100 pb-3">
                        <div class="space-y-1">
                            <div class="flex items-center gap-2">
                                <span class="px-2.5 py-0.5 rounded-md bg-blue-900 text-white font-black text-xs">${{s.code}}</span>
                                <span class="px-2.5 py-0.5 rounded-md bg-blue-50 text-blue-900 font-bold text-xs border border-blue-200">${{s.module}}</span>
                                <span class="text-xs font-bold text-slate-500">${{s.month}} • Sábado ${{s.sabado_num}}</span>
                            </div>
                            <h4 class="text-xl sm:text-2xl font-black text-slate-900 pt-1">${{s.title}}</h4>
                        </div>
                        <span class="px-3 py-1 rounded-full bg-slate-100 text-slate-700 font-bold text-xs flex items-center gap-1">
                            <i data-lucide="clock" class="w-3.5 h-3.5 text-slate-500"></i> ${{s.duration}}
                        </span>
                    </div>

                    <!-- Objetivo -->
                    <div class="p-4 rounded-2xl bg-blue-50/80 border border-blue-100 text-xs space-y-1">
                        <strong class="text-blue-950 font-bold flex items-center gap-1.5 text-sm">
                            <i data-lucide="target" class="w-4 h-4 text-blue-700"></i> Objetivo Pedagógico de la Sesión
                        </strong>
                        <p class="text-slate-800 leading-relaxed text-xs">${{s.objective}}</p>
                    </div>

                    <!-- 3 Momentos -->
                    <div class="grid md:grid-cols-3 gap-3.5 text-xs pt-1">
                        <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 space-y-1.5">
                            <div class="flex items-center gap-1.5 font-black text-slate-900 text-xs">
                                <span class="w-5 h-5 rounded-full bg-blue-100 text-blue-900 text-[10px] flex items-center justify-center font-black">1</span>
                                Apertura (30 min)
                            </div>
                            <p class="text-slate-600 leading-relaxed text-[11px]">${{s.apertura}}</p>
                        </div>

                        <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 space-y-1.5">
                            <div class="flex items-center gap-1.5 font-black text-slate-900 text-xs">
                                <span class="w-5 h-5 rounded-full bg-emerald-100 text-emerald-900 text-[10px] flex items-center justify-center font-black">2</span>
                                Desarrollo (210 min)
                            </div>
                            <p class="text-slate-600 leading-relaxed text-[11px]">${{s.desarrollo}}</p>
                        </div>

                        <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 space-y-1.5">
                            <div class="flex items-center gap-1.5 font-black text-slate-900 text-xs">
                                <span class="w-5 h-5 rounded-full bg-amber-100 text-amber-900 text-[10px] flex items-center justify-center font-black">3</span>
                                Cierre (30 min)
                            </div>
                            <p class="text-slate-600 leading-relaxed text-[11px]">${{s.cierre}}</p>
                        </div>
                    </div>

                    <!-- Footer: Evidencias y Equipamiento -->
                    <div class="grid sm:grid-cols-2 gap-3 pt-3 text-xs border-t border-slate-100">
                        <div class="p-3 rounded-xl bg-slate-50 border border-slate-200 space-y-1">
                            <span class="text-[10px] uppercase font-bold text-slate-500 block">Evidencia de Taller</span>
                            <div class="text-slate-800 font-medium flex items-start gap-1.5">
                                <i data-lucide="check-circle" class="w-4 h-4 text-emerald-600 mt-0.5 shrink-0"></i>
                                <span>${{s.evidencias}}</span>
                            </div>
                        </div>
                        <div class="p-3 rounded-xl bg-slate-50 border border-slate-200 space-y-1">
                            <span class="text-[10px] uppercase font-bold text-slate-500 block">Equipamiento Kinal</span>
                            <div class="text-slate-800 font-medium flex items-start gap-1.5">
                                <i data-lucide="wrench" class="w-4 h-4 text-blue-600 mt-0.5 shrink-0"></i>
                                <span>${{s.equipamiento}}</span>
                            </div>
                        </div>
                    </div>

                    <!-- Navegación entre sesiones -->
                    <div class="flex justify-between items-center pt-2">
                        <button onclick="navigateMecaSession(-1)" ${{s.num === 1 ? 'disabled class="opacity-40 cursor-not-allowed"' : 'class="hover:bg-slate-100"'}} px-3 py-1.5 rounded-lg border border-slate-200 text-xs font-bold text-slate-700 flex items-center gap-1 transition">
                            <i data-lucide="arrow-left" class="w-3.5 h-3.5"></i> Sesión Anterior
                        </button>
                        <span class="text-xs font-bold text-slate-400">Sesión ${{s.num}} de 20</span>
                        <button onclick="navigateMecaSession(1)" ${{s.num === 20 ? 'disabled class="opacity-40 cursor-not-allowed"' : 'class="hover:bg-slate-100"'}} px-3 py-1.5 rounded-lg border border-slate-200 text-xs font-bold text-slate-700 flex items-center gap-1 transition">
                            Siguiente Sesión <i data-lucide="arrow-right" class="w-3.5 h-3.5"></i>
                        </button>
                    </div>
                </div>
            `;
            lucide.createIcons();
        }}

        function selectMecaSession(num) {{
            currentSessionNum = num;
            document.querySelectorAll('.session-item').forEach(el => el.classList.remove('active'));
            const activeItem = document.getElementById('meca-list-item-' + num);
            if (activeItem) activeItem.classList.add('active');

            const s = mecaSessions.find(item => item.num === num);
            if (s) renderMecaSessionDetail(s);
        }}

        function navigateMecaSession(direction) {{
            const newNum = currentSessionNum + direction;
            if (newNum >= 1 && newNum <= 20) {{
                selectMecaSession(newNum);
            }}
        }}

        function goToSession(num) {{
            switchTab('secuencias');
            selectMecaSession(num);
        }}

        function filterMecaList() {{
            const q = document.getElementById('meca-sec-search').value.toLowerCase();
            const filtered = mecaSessions.filter(s => 
                s.title.toLowerCase().includes(q) || 
                s.code.toLowerCase().includes(q) || 
                s.objective.toLowerCase().includes(q)
            );
            renderMecaSessionList(filtered);
            document.getElementById('meca-count-badge').innerText = filtered.length + ' Sesiones';
        }}

        document.addEventListener('DOMContentLoaded', () => {{
            renderMecaSessionList(mecaSessions);
            selectMecaSession(1);
            lucide.createIcons();
        }});
    </script>
</body>
</html>
"""

with open(os.path.join(WEB_DIR, "mecatronica.html"), "w", encoding="utf-8") as f:
    f.write(mecatronica_html)
print("Generado con éxito: Cursos/Web/mecatronica.html")

print("Generando HTML de Ciberseguridad...")

# ==============================================================================
# HTML BUILDER: CIBERSEGURIDAD.HTML
# ==============================================================================
ciber_sessions_json = json.dumps(ciber_sessions, ensure_ascii=False)

ciberseguridad_html = f"""<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Ciberseguridad y Fundamentos de Seguridad | Kinal 2026</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <script src="https://unpkg.com/lucide@latest"></script>
    <style>
        @import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700;800&family=Space+Grotesk:wght@500;700;800;900&display=swap');
        body {{ font-family: 'Inter', sans-serif; background-color: #f8fafc; color: #0f172a; }}
        h1, h2, h3, h4, .font-heading {{ font-family: 'Space Grotesk', sans-serif; }}
        .hero-pattern {{
            background: linear-gradient(135deg, #09172e 0%, #1e1b4b 50%, #312e81 100%);
        }}
        .tab-btn.active {{
            background-color: #312e81;
            color: #ffffff;
            font-weight: 700;
        }}
        .session-item.active {{
            background-color: #312e81;
            color: #ffffff;
            border-color: #312e81;
        }}
        .session-item.active span, .session-item.active p {{
            color: #e2e8f0;
        }}
        .session-item.active .badge-code {{
            background-color: #f59e0b;
            color: #0f172a;
        }}
    </style>
</head>
<body class="antialiased min-h-screen flex flex-col">

    <!-- Top Navigation Header -->
    <header class="bg-slate-950 text-white border-b border-slate-800 sticky top-0 z-50 shadow-md">
        <div class="max-w-7xl mx-auto px-4 sm:px-6 lg:px-8 py-3 flex flex-wrap justify-between items-center gap-4">
            <div class="flex items-center gap-3">
                <a href="index.html" class="bg-slate-800 hover:bg-slate-700 text-slate-200 text-xs font-bold px-3 py-1.5 rounded-lg transition flex items-center gap-1.5">
                    <i data-lucide="arrow-left" class="w-3.5 h-3.5"></i> Cursos
                </a>
                <span class="text-slate-600">|</span>
                <div class="text-slate-200 font-bold text-xs sm:text-sm flex items-center gap-2">
                    <i data-lucide="shield-check" class="w-4 h-4 text-indigo-400"></i>
                    Ciberseguridad & Fundamentos de Seguridad
                </div>
            </div>
            
            <!-- Quick Download Buttons -->
            <div class="flex items-center gap-2 text-xs">
                <span class="text-slate-400 hidden sm:inline">Word:</span>
                <a href="../Ciberseguridad/Propuesta_Curso_Ciberseguridad_Kinal.docx" download class="px-2.5 py-1 rounded bg-indigo-900/60 hover:bg-indigo-800 text-indigo-200 transition font-medium" title="Descargar Propuesta Oficial">
                    Propuesta
                </a>
                <a href="../Ciberseguridad/Temario_Curso_Ciberseguridad_Kinal.docx" download class="px-2.5 py-1 rounded bg-emerald-900/60 hover:bg-emerald-800 text-emerald-200 transition font-medium" title="Descargar Temario Oficial">
                    Temario
                </a>
                <a href="../Ciberseguridad/Dosificacion_y_Secuencia_Didactica_Ciberseguridad.docx" download class="px-2.5 py-1 rounded bg-amber-900/60 hover:bg-amber-800 text-amber-200 transition font-medium" title="Descargar Dosificación Oficial">
                    Dosificación
                </a>
            </div>
        </div>
    </header>

    <!-- Hero Course Banner -->
    <section class="hero-pattern text-white py-10 px-4 sm:px-6 lg:px-8">
        <div class="max-w-7xl mx-auto space-y-3">
            <div class="flex flex-wrap items-center gap-2">
                <span class="px-3 py-1 rounded-full bg-indigo-500/20 border border-indigo-400/30 text-indigo-300 text-xs font-bold tracking-wide uppercase">
                    Modalidad Híbrida (80% Virtual / 20% Presencial)
                </span>
                <span class="px-3 py-1 rounded-full bg-amber-500/20 border border-amber-400/30 text-amber-300 text-xs font-bold">
                    DQR Nivel 4 - 5
                </span>
                <span class="px-3 py-1 rounded-full bg-emerald-500/20 border border-emerald-400/30 text-emerald-300 text-xs font-bold">
                    Aprobación: ≥ 75 Pts
                </span>
            </div>
            <h1 class="text-2xl sm:text-4xl font-black tracking-tight leading-tight">
                Ciberseguridad y Fundamentos de Seguridad de la Información
            </h1>
            <p class="text-slate-300 text-xs sm:text-sm max-w-4xl leading-relaxed">
                Protección integral de infraestructuras corporativas, gestión de identidad y accesos (IAM), cifrado y firma digital, respuesta inicial a incidentes en SOC/SIEM, gestión de vulnerabilidades y cumplimiento de estándares internacionales (ISO 27001, SOC 2, HIPAA).
            </p>
        </div>
    </section>

    <!-- Tab Bar Navigation (UNA SOLA PALABRA, SIN NUMERACIÓN) -->
    <div class="bg-white border-b border-slate-200 sticky top-14 z-40 shadow-sm">
        <div class="max-w-7xl mx-auto px-4 sm:px-6 lg:px-8">
            <nav class="flex space-x-2 sm:space-x-3 py-2.5 overflow-x-auto text-xs font-bold">
                <button onclick="switchTab('propuesta')" id="btn-propuesta" class="tab-btn active px-4 py-2 rounded-xl transition flex items-center gap-1.5 shrink-0">
                    <i data-lucide="file-text" class="w-3.5 h-3.5"></i> Propuesta
                </button>
                <button onclick="switchTab('temario')" id="btn-temario" class="tab-btn px-4 py-2 rounded-xl bg-slate-100 hover:bg-slate-200 text-slate-700 transition flex items-center gap-1.5 shrink-0">
                    <i data-lucide="list-tree" class="w-3.5 h-3.5"></i> Temario
                </button>
                <button onclick="switchTab('calendario')" id="btn-calendario" class="tab-btn px-4 py-2 rounded-xl bg-slate-100 hover:bg-slate-200 text-slate-700 transition flex items-center gap-1.5 shrink-0">
                    <i data-lucide="calendar" class="w-3.5 h-3.5"></i> Calendario
                </button>
                <button onclick="switchTab('secuencias')" id="btn-secuencias" class="tab-btn px-4 py-2 rounded-xl bg-slate-100 hover:bg-slate-200 text-slate-700 transition flex items-center gap-1.5 shrink-0">
                    <i data-lucide="layers" class="w-3.5 h-3.5"></i> Secuencias
                </button>
                <button onclick="switchTab('evaluacion')" id="btn-evaluacion" class="tab-btn px-4 py-2 rounded-xl bg-slate-100 hover:bg-slate-200 text-slate-700 transition flex items-center gap-1.5 shrink-0">
                    <i data-lucide="award" class="w-3.5 h-3.5"></i> Evaluación
                </button>
            </nav>
        </div>
    </div>

    <!-- Main Container -->
    <main class="max-w-7xl mx-auto px-4 sm:px-6 lg:px-8 py-8 flex-1 space-y-8">

        <!-- ================================================================== -->
        <!-- PESTAÑA 1: PROPUESTA -->
        <!-- ================================================================== -->
        <div id="tab-propuesta" class="tab-content space-y-8">
            
            <!-- Ficha Técnica (UNA SOLA TARJETA) -->
            {ciber_ficha_single_card}

            <!-- Marco Filosófico Kinal -->
            <div class="grid md:grid-cols-2 gap-6">
                <div class="bg-indigo-50 border-l-4 border-indigo-900 p-6 rounded-r-3xl space-y-3 shadow-sm">
                    <div class="flex items-center gap-2 text-indigo-900 font-extrabold text-sm uppercase">
                        <i data-lucide="flag" class="w-4 h-4"></i> Misión Institucional de Fundación Kinal
                    </div>
                    <p class="text-slate-700 text-xs italic leading-relaxed">
                        «Formar a jóvenes y adultos a través de una educación integral, con énfasis en las áreas técnicas y tecnológicas, influyendo positivamente en su trabajo, su familia y la sociedad».
                    </p>
                    <div class="pt-2 text-xs font-bold text-indigo-950">
                        Valores Nucleares: <span class="font-normal italic">Visión cristiana de la vida. Respeto a la dignidad de la persona. Espíritu de servicio. Trabajo bien hecho. Libertad personal responsable.</span>
                    </div>
                </div>

                <div class="bg-amber-50 border-l-4 border-amber-600 p-6 rounded-r-3xl space-y-3 shadow-sm">
                    <div class="flex items-center gap-2 text-amber-900 font-extrabold text-sm uppercase">
                        <i data-lucide="sparkle" class="w-4 h-4"></i> El «Trabajo Bien Hecho» en Seguridad de la Información
                    </div>
                    <p class="text-slate-700 text-xs leading-relaxed">
                        En ciberseguridad, el «trabajo bien hecho» demanda <strong>rigor en la asignación de controles</strong> (alcance, responsable y evidencia comprobable), <strong>trazabilidad y verdad en los registros de auditoría</strong> sin alterar evidencias, y confidencialidad cotidiana estricta en el manejo de credenciales y datos sensibles.
                    </p>
                </div>
            </div>

            <!-- Fundamentación y Enfoque Pedagógico -->
            <div class="bg-white rounded-3xl border border-slate-200 p-8 shadow-sm space-y-4">
                <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                    <i data-lucide="compass" class="w-5 h-5 text-indigo-700"></i>
                    Fundamentación Técnica y Enfoque Constructivo (No Explotativo)
                </h3>
                <div class="grid md:grid-cols-3 gap-4 text-xs">
                    <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 space-y-2">
                        <strong class="text-indigo-950 font-bold block text-sm">Caso Práctico Transversal</strong>
                        <p class="text-slate-600 leading-relaxed">
                            Una empresa ficticia evoluciona durante los 8 módulos con sus estaciones de trabajo, servidores locales, base de datos y servicios en nube, permitiendo aplicar cada control en un contexto corporativo verosímil.
                        </p>
                    </div>
                    <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 space-y-2">
                        <strong class="text-emerald-950 font-bold block text-sm">Enfoque Defensivo</strong>
                        <p class="text-slate-600 leading-relaxed">
                            Las actividades se enfocan en leer configuraciones, interpretar logs, diseñar matrices RBAC y auditar cumplimientos, sin exigir ataques intrusivos ni herramientas de explotación destructiva.
                        </p>
                    </div>
                    <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 space-y-2">
                        <strong class="text-purple-950 font-bold block text-sm">Empleabilidad Inmediata</strong>
                        <p class="text-slate-600 leading-relaxed">
                            Prepara a los participantes para roles de soporte técnico con criterio de seguridad, operadores de consola SOC Nivel 1 y auxiliares de cumplimiento y auditoría de TI.
                        </p>
                    </div>
                </div>
            </div>

            <!-- Matriz DQR -->
            <div class="bg-white rounded-3xl border border-slate-200 p-8 shadow-sm space-y-6">
                <div class="space-y-1">
                    <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                        <i data-lucide="award" class="w-5 h-5 text-indigo-700"></i>
                        Alineación con el Marco Alemán de Cualificaciones (DQR Nivel 4-5)
                    </h3>
                    <p class="text-slate-600 text-xs">
                        Desglose de competencias integrales de acción técnica y personal según el marco europeo.
                    </p>
                </div>

                <div class="overflow-x-auto">
                    <table class="w-full text-xs text-left border-collapse">
                        <thead>
                            <tr class="bg-slate-100 text-slate-700 uppercase font-black border-b border-slate-200">
                                <th class="py-3 px-4">Dimensión DQR</th>
                                <th class="py-3 px-4">Subdimensión</th>
                                <th class="py-3 px-4">Evidencia de Desempeño en el Programa</th>
                            </tr>
                        </thead>
                        <tbody class="divide-y divide-slate-100 text-slate-700">
                            { "".join([f'<tr><td class="py-3 px-4 font-bold bg-slate-50">{clean_html(r[0])}</td><td class="py-3 px-4 font-semibold text-indigo-900">{clean_html(r[1])}</td><td class="py-3 px-4 leading-relaxed">{clean_html(r[2])}</td></tr>' for r in ciber_dqr[1:]]) }
                        </tbody>
                    </table>
                </div>
            </div>

            <!-- Perfiles de Ingreso y Egreso -->
            <div class="grid md:grid-cols-2 gap-6">
                <div class="bg-white rounded-3xl border border-slate-200 p-7 shadow-sm space-y-4">
                    <h4 class="text-lg font-black text-slate-900 flex items-center gap-2">
                        <i data-lucide="log-in" class="w-5 h-5 text-indigo-700"></i> Perfil de Ingreso
                    </h4>
                    <p class="text-slate-600 text-xs leading-relaxed">
                        Jóvenes y adultos desempleados de 18 a 45 años con conocimientos previos de informática básica, peritos en computación, estudiantes de ingeniería o técnicos en redes que deseen reconvertirse al sector de seguridad digital.
                    </p>
                    <ul class="text-xs space-y-2 text-slate-700">
                        <li class="flex items-start gap-2"><i data-lucide="check" class="w-4 h-4 text-emerald-600 mt-0.5"></i> Manejo fluido del sistema operativo y navegación en red.</li>
                        <li class="flex items-start gap-2"><i data-lucide="check" class="w-4 h-4 text-emerald-600 mt-0.5"></i> Capacidad analítica y hábito de lectura técnica.</li>
                        <li class="flex items-start gap-2"><i data-lucide="check" class="w-4 h-4 text-emerald-600 mt-0.5"></i> Ética comprobada y respeto por la privacidad de la información.</li>
                    </ul>
                </div>

                <div class="bg-white rounded-3xl border border-slate-200 p-7 shadow-sm space-y-4">
                    <h4 class="text-lg font-black text-slate-900 flex items-center gap-2">
                        <i data-lucide="log-out" class="w-5 h-5 text-emerald-700"></i> Perfil de Egreso
                    </h4>
                    <p class="text-slate-600 text-xs leading-relaxed">
                        El graduado será capaz de diagnosticar debilidades en entornos de TI, redactar políticas y procedimientos SOP, estructurar matrices de control de acceso RBAC, analizar eventos en consolas de monitoreo y responder metódicamente a incidentes.
                    </p>
                    <ul class="text-xs space-y-2 text-slate-700">
                        <li class="flex items-start gap-2"><i data-lucide="check-circle" class="w-4 h-4 text-emerald-600 mt-0.5"></i> Interpretación solvente de boletines CVE y métricas CVSS.</li>
                        <li class="flex items-start gap-2"><i data-lucide="check-circle" class="w-4 h-4 text-emerald-600 mt-0.5"></i> Mapeo de controles frente a estándares ISO 27001, SOC 2 y HIPAA.</li>
                        <li class="flex items-start gap-2"><i data-lucide="check-circle" class="w-4 h-4 text-emerald-600 mt-0.5"></i> Comunicación asertiva entre el área técnica y la administración.</li>
                    </ul>
                </div>
            </div>

            <!-- Arquitectura Curricular: 8 Módulos -->
            <div class="bg-white rounded-3xl border border-slate-200 p-8 shadow-sm space-y-6">
                <div class="space-y-1">
                    <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                        <i data-lucide="blocks" class="w-5 h-5 text-indigo-700"></i>
                        Arquitectura Curricular: 8 Módulos Formativos (80 Horas)
                    </h3>
                    <p class="text-slate-600 text-xs">
                        Distribución horaria y áreas formativas del programa.
                    </p>
                </div>
                <div class="overflow-x-auto">
                    <table class="w-full text-xs text-left border-collapse">
                        <thead>
                            <tr class="bg-slate-100 text-slate-700 uppercase font-black border-b border-slate-200">
                                <th class="py-3 px-4">Módulo</th>
                                <th class="py-3 px-4">Área de Formación</th>
                                <th class="py-3 px-4">Carga Horaria</th>
                            </tr>
                        </thead>
                        <tbody class="divide-y divide-slate-100 text-slate-700">
                            { "".join([f'<tr><td class="py-3 px-4 font-bold bg-slate-50 text-indigo-900">{clean_html(r[0])}</td><td class="py-3 px-4 font-medium">{clean_html(r[1])}</td><td class="py-3 px-4 font-black text-slate-900">{clean_html(r[2])}</td></tr>' for r in ciber_mod_prop[1:]]) }
                        </tbody>
                    </table>
                </div>
            </div>

        </div>

        <!-- ================================================================== -->
        <!-- PESTAÑA 2: TEMARIO (DISTRIBUCIÓN MODULAR AUDITADA) -->
        <!-- ================================================================== -->
        <div id="tab-temario" class="tab-content hidden space-y-8">
            
            <!-- Barra Superior de Control y Filtros -->
            <div class="flex flex-col md:flex-row justify-between items-start md:items-center gap-4 bg-white rounded-3xl p-6 border border-slate-200 shadow-sm">
                <div class="space-y-1">
                    <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                        <i data-lucide="shield-check" class="w-5 h-5 text-indigo-700"></i>
                        Distribución Modular con Código de Color
                    </h3>
                    <p class="text-xs text-slate-500">
                        Estructura organizada en 8 módulos formativos clasificados por disciplina de seguridad, logros esperados y temas analizados.
                    </p>
                </div>

                <!-- Botones de Filtro Interactivo -->
                <div class="flex flex-wrap items-center gap-2 text-xs font-semibold">
                    <button onclick="filterCiberTemario('all')" id="btn-ciber-all" class="ciber-filter-btn px-3 py-1.5 rounded-lg bg-slate-900 text-white shadow-sm transition">Todos (20)</button>
                    <button onclick="filterCiberTemario('sistemas')" id="btn-ciber-sistemas" class="ciber-filter-btn px-3 py-1.5 rounded-lg bg-blue-100 text-blue-800 hover:bg-blue-200 transition">🔵 Sistemas & Redes TI</button>
                    <button onclick="filterCiberTemario('identidad')" id="btn-ciber-identidad" class="ciber-filter-btn px-3 py-1.5 rounded-lg bg-purple-100 text-purple-800 hover:bg-purple-200 transition">🟣 Identidad & Acceso (IAM)</button>
                    <button onclick="filterCiberTemario('operaciones')" id="btn-ciber-operaciones" class="ciber-filter-btn px-3 py-1.5 rounded-lg bg-amber-100 text-amber-800 hover:bg-amber-200 transition">🟡 Operaciones & SOC</button>
                    <button onclick="filterCiberTemario('gobierno')" id="btn-ciber-gobierno" class="ciber-filter-btn px-3 py-1.5 rounded-lg bg-emerald-100 text-emerald-800 hover:bg-emerald-200 transition">🟢 Gobierno & Cumplimiento</button>
                </div>
            </div>

            <!-- Código de Leyenda Explicativo -->
            <div class="grid grid-cols-1 sm:grid-cols-2 lg:grid-cols-4 gap-3 text-xs">
                <div class="p-3.5 rounded-xl border border-blue-200 bg-blue-50/70 text-blue-950 flex items-start gap-2.5">
                    <span class="inline-block w-3 h-3 rounded-full bg-blue-600 mt-0.5 shrink-0"></span>
                    <div>
                        <strong class="text-blue-900">Color Azul (Sistemas & Redes TI):</strong> Tríada CIA, infraestructura de red, principios ITIL, criptografía simétrica/asimétrica y PKI/TLS.
                    </div>
                </div>
                <div class="p-3.5 rounded-xl border border-purple-200 bg-purple-50/70 text-purple-950 flex items-start gap-2.5">
                    <span class="inline-block w-3 h-3 rounded-full bg-purple-600 mt-0.5 shrink-0"></span>
                    <div>
                        <strong class="text-purple-900">Color Púrpura (Identidad & Acceso - IAM):</strong> Autenticación multifactor (MFA), ciclo de vida de sesiones, SSO federado (SAML/OAuth) y matrices RBAC.
                    </div>
                </div>
                <div class="p-3.5 rounded-xl border border-amber-200 bg-amber-50/70 text-amber-950 flex items-start gap-2.5">
                    <span class="inline-block w-3 h-3 rounded-full bg-amber-600 mt-0.5 shrink-0"></span>
                    <div>
                        <strong class="text-amber-900">Color Ámbar (Operaciones & Diagnóstico SOC):</strong> Inspección de tráfico web, gestión de vulnerabilidades CVE/CVSS, ventanas de parches y respuesta a incidentes.
                    </div>
                </div>
                <div class="p-3.5 rounded-xl border border-emerald-200 bg-emerald-50/70 text-emerald-950 flex items-start gap-2.5">
                    <span class="inline-block w-3 h-3 rounded-full bg-emerald-600 mt-0.5 shrink-0"></span>
                    <div>
                        <strong class="text-emerald-900">Color Verde (Gobierno & Cumplimiento):</strong> Principio de menor privilegio, redacción de SOPs, estándares PCI-DSS, HIPAA, SOC 2 y sustentación del Expediente.
                    </div>
                </div>
            </div>

            <!-- Cuadrícula Modular de Temarios -->
            <div class="grid lg:grid-cols-2 gap-8 pt-2">
                {ciber_modular_cards}
            </div>

            <!-- Laboratorios y Herramientas -->
            <div class="grid md:grid-cols-3 gap-6 pt-4">
                <div class="bg-white rounded-2xl border border-slate-200 p-6 shadow-sm space-y-3">
                    <h4 class="font-black text-slate-900 text-base flex items-center gap-2">
                        <i data-lucide="terminal" class="w-4 h-4 text-indigo-700"></i> Entornos & Sistemas
                    </h4>
                    <ul class="text-xs space-y-2 text-slate-600">
                        <li>• Máquinas virtuales Linux (Ubuntu Server / Debian) para inspección.</li>
                        <li>• Estaciones de trabajo Windows con perfiles de usuario estándar.</li>
                        <li>• Directorio de usuarios y simulación de LDAP / Keycloak.</li>
                        <li>• Entornos SaaS simulados para configuración de MFA y SSO.</li>
                    </ul>
                </div>

                <div class="bg-white rounded-2xl border border-slate-200 p-6 shadow-sm space-y-3">
                    <h4 class="font-black text-slate-900 text-base flex items-center gap-2">
                        <i data-lucide="activity" class="w-4 h-4 text-emerald-700"></i> Herramientas de Análisis
                    </h4>
                    <ul class="text-xs space-y-2 text-slate-600">
                        <li>• Wireshark (Captura e inspección pedagógica de paquetes de red).</li>
                        <li>• OpenVAS / Greenbone (Escaneo controlado de vulnerabilidades).</li>
                        <li>• OpenSearch / Graylog / Wazuh (Correlación y visualización de logs).</li>
                        <li>• OpenSSL (Generación de certificados x509 y pares de claves).</li>
                    </ul>
                </div>

                <div class="bg-white rounded-2xl border border-slate-200 p-6 shadow-sm space-y-3">
                    <h4 class="font-black text-slate-900 text-base flex items-center gap-2">
                        <i data-lucide="file-check-2" class="w-4 h-4 text-amber-700"></i> Marcos Normativos
                    </h4>
                    <ul class="text-xs space-y-2 text-slate-600">
                        <li>• ISO/IEC 27001:2022 (Controles organizacionales, físicos y tecnológicos).</li>
                        <li>• AICPA SOC 1 y SOC 2 (Criterios de servicios de confianza).</li>
                        <li>• Estándar PCI-DSS v4.0 (Protección de datos de tarjetahabientes).</li>
                        <li>• Regulación HIPAA / HITECH (Seguridad y privacidad de ePHI).</li>
                    </ul>
                </div>
            </div>

        </div>

        <!-- ================================================================== -->
        <!-- PESTAÑA 3: CALENDARIO CRONOLÓGICO -->
        <!-- ================================================================== -->
        <div id="tab-calendario" class="tab-content hidden space-y-6">
            
            <div class="flex flex-col sm:flex-row justify-between items-start sm:items-center gap-3 bg-white rounded-3xl p-6 border border-slate-200 shadow-sm">
                <div>
                    <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                        <i data-lucide="calendar" class="w-5 h-5 text-indigo-700"></i>
                        Calendario de Encuentros y Cronograma Formativo
                    </h3>
                    <p class="text-slate-600 text-xs mt-1">
                        Programa híbrido de 10 semanas (20 sesiones de 4 horas) • 80% Virtual (16 sesiones) y 20% Presencial en Kinal (4 sesiones).
                    </p>
                </div>
                <div class="flex items-center gap-2 text-xs font-bold text-indigo-900 bg-indigo-50 px-3 py-1.5 rounded-xl border border-indigo-200">
                    <i data-lucide="shield" class="w-4 h-4 text-indigo-700"></i> Formato Híbrido 80/20
                </div>
            </div>

            <!-- Grid Semanal de Calendario -->
            <div class="space-y-6">
                {ciber_calendar_content}
            </div>

        </div>

        <!-- ================================================================== -->
        <!-- PESTAÑA 4: SECUENCIAS (MASTER-DETAIL SIN ALARGAR LA PÁGINA) -->
        <!-- ================================================================== -->
        <div id="tab-secuencias" class="tab-content hidden space-y-6">
            
            <div class="flex flex-col sm:flex-row justify-between items-start sm:items-center gap-3 bg-white rounded-3xl p-5 border border-slate-200 shadow-sm">
                <div>
                    <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                        <i data-lucide="layers" class="w-5 h-5 text-indigo-700"></i>
                        Secuencia Didáctica Interactiva Sesión a Sesión
                    </h3>
                    <p class="text-slate-600 text-xs mt-0.5">
                        Selecciona una sesión en el panel izquierdo para consultar su microdiseño en 3 momentos (Apertura, Desarrollo y Cierre).
                    </p>
                </div>
                <div class="flex items-center gap-2 w-full sm:w-auto">
                    <div class="relative w-full sm:w-56">
                        <i data-lucide="search" class="w-3.5 h-3.5 text-slate-400 absolute left-3 top-2.5"></i>
                        <input type="text" id="ciber-sec-search" oninput="filterCiberList()" placeholder="Buscar sesión..." class="w-full pl-8 pr-3 py-1.5 text-xs rounded-xl border border-slate-200 bg-slate-50 focus:bg-white focus:outline-none focus:ring-2 focus:ring-indigo-600 transition">
                    </div>
                </div>
            </div>

            <!-- Master-Detail Container -->
            <div class="grid lg:grid-cols-12 gap-6 items-start">
                
                <!-- LISTADO LATERAL (MASTER) -->
                <div class="lg:col-span-4 bg-white rounded-3xl border border-slate-200 p-4 shadow-sm space-y-2">
                    <div class="text-[11px] font-black uppercase tracking-wider text-slate-400 px-2 pb-1 border-b border-slate-100 flex justify-between">
                        <span>Listado de Sesiones</span>
                        <span id="ciber-count-badge">20 Sesiones</span>
                    </div>
                    <div id="ciber-sessions-list" class="space-y-1.5 max-h-[560px] overflow-y-auto pr-1">
                        <!-- Rendered by JS -->
                    </div>
                </div>

                <!-- DETALLE DE LA SESIÓN SELECCIONADA (DETAIL) -->
                <div class="lg:col-span-8 bg-white rounded-3xl border border-slate-200 p-6 sm:p-8 shadow-sm space-y-6" id="ciber-detail-card">
                    <!-- Rendered by JS -->
                </div>

            </div>

        </div>

        <!-- ================================================================== -->
        <!-- PESTAÑA 5: EVALUACIÓN -->
        <!-- ================================================================== -->
        <div id="tab-evaluacion" class="tab-content hidden space-y-8">
            
            <!-- Esquema Resumen -->
            <div class="bg-white rounded-3xl border border-slate-200 p-8 shadow-sm space-y-6">
                <div class="flex items-center justify-between border-b border-slate-200 pb-3">
                    <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                        <i data-lucide="check-square" class="w-5 h-5 text-indigo-700"></i>
                        Esquema de Evaluación Institucional & Acreditación
                    </h3>
                    <span class="px-3 py-1 rounded-full bg-amber-100 text-amber-900 font-black text-xs">
                        Umbral Mínimo: 75 / 100 Puntos
                    </span>
                </div>
                <div class="grid md:grid-cols-4 gap-4 text-xs">
                    <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 text-center space-y-1">
                        <div class="text-slate-500 font-bold uppercase text-[10px]">Expediente del Caso</div>
                        <div class="text-3xl font-black text-indigo-900">40%</div>
                        <div class="text-slate-600">Entregable Terminal Integrado</div>
                    </div>
                    <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 text-center space-y-1">
                        <div class="text-slate-500 font-bold uppercase text-[10px]">Talleres Prácticos</div>
                        <div class="text-3xl font-black text-emerald-700">30%</div>
                        <div class="text-slate-600">Laboratorios y Diagramación</div>
                    </div>
                    <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 text-center space-y-1">
                        <div class="text-slate-500 font-bold uppercase text-[10px]">Comprobaciones Teóricas</div>
                        <div class="text-3xl font-black text-purple-700">15%</div>
                        <div class="text-slate-600">Evaluaciones de Criterio Técnico</div>
                    </div>
                    <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 text-center space-y-1">
                        <div class="text-slate-500 font-bold uppercase text-[10px]">Bitácora & Asistencia</div>
                        <div class="text-3xl font-black text-amber-700">15%</div>
                        <div class="text-slate-600">Registro Dual y 100% Presencial</div>
                    </div>
                </div>
            </div>

            <!-- Rúbrica Terminal -->
            <div class="bg-white rounded-3xl border border-slate-200 p-8 shadow-sm space-y-6">
                <div class="flex flex-col sm:flex-row justify-between items-start sm:items-center gap-2 border-b border-slate-200 pb-4">
                    <div>
                        <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                            <i data-lucide="award" class="w-5 h-5 text-indigo-700"></i>
                            Rúbrica Analítica Terminal de Evaluación (Normativa Kinal)
                        </h3>
                        <p class="text-slate-600 text-xs mt-1">
                            Evaluación de criterios conceptuales, precisión en controles, diagnóstico de riesgos y trazabilidad.
                        </p>
                    </div>
                    <span class="px-3 py-1.5 rounded-xl bg-amber-100 text-amber-900 font-extrabold text-xs">
                        Aprobación Mínima: ≥ 75 Pts
                    </span>
                </div>

                <div class="overflow-x-auto">
                    <table class="w-full text-xs text-left border-collapse">
                        <thead>
                            <tr class="bg-slate-100 text-slate-700 uppercase font-black border-b border-slate-200">
                                <th class="py-3 px-4 w-1/5">Criterio Evaluado</th>
                                <th class="py-3 px-4 w-1/5 text-emerald-800">Excelente (90 - 100)</th>
                                <th class="py-3 px-4 w-1/5 text-blue-800">Muy Bueno (80 - 89)</th>
                                <th class="py-3 px-4 w-1/5 text-amber-800">Aprobado Mínimo (75 - 79)</th>
                                <th class="py-3 px-4 w-1/5 text-rose-800">No Aprobado (&lt; 75)</th>
                            </tr>
                        </thead>
                        <tbody class="divide-y divide-slate-100 text-slate-700">
                            {ciber_rubric_rows_html}
                        </tbody>
                    </table>
                </div>
            </div>

            <!-- Expediente de Seguridad del Caso Transversal -->
            <div class="bg-white rounded-3xl border border-slate-200 p-8 shadow-sm space-y-6">
                <div class="flex justify-between items-center border-b border-slate-200 pb-4">
                    <div>
                        <h3 class="text-xl font-black text-slate-900 flex items-center gap-2">
                            <i data-lucide="folder-lock" class="w-5 h-5 text-indigo-700"></i>
                            Estructura del Expediente de Seguridad del Caso Transversal
                        </h3>
                        <p class="text-slate-600 text-xs mt-1">
                            Portafolio de evidencias construido por el aprendiz durante las 20 sesiones del programa.
                        </p>
                    </div>
                    <button onclick="window.print()" class="px-3.5 py-1.5 rounded-xl bg-slate-100 hover:bg-slate-200 text-slate-800 font-bold text-xs transition flex items-center gap-1.5">
                        <i data-lucide="printer" class="w-3.5 h-3.5"></i> Imprimir Portafolio
                    </button>
                </div>

                <div class="p-6 rounded-2xl bg-slate-50 border border-slate-300 space-y-4 text-xs font-mono">
                    <div class="border-b border-slate-300 pb-3 font-bold text-slate-900 text-center uppercase tracking-wider">
                        FUNDACIÓN KINAL • EXPEDIENTE TERMINAL DE SEGURIDAD EMPRESARIAL
                    </div>
                    <div class="grid grid-cols-2 gap-4">
                        <div><strong>Estudiante:</strong> ____________________________________</div>
                        <div><strong>Carné:</strong> ______________</div>
                        <div><strong>Organización Simulada:</strong> Distribuidora Médica S.A.</div>
                        <div><strong>Instructor Evaluador:</strong> __________________________</div>
                    </div>
                    <div class="space-y-2 pt-2 text-slate-800">
                        <strong>Componentes Documentales Auditables:</strong>
                        <ul class="space-y-1.5 list-disc pl-5">
                            <li><strong>Artefacto 1:</strong> Inventario de Activos Críticos y Matriz de Tríada CIA (Confidencialidad, Integridad, Disponibilidad).</li>
                            <li><strong>Artefacto 2:</strong> Matriz de Control de Accesos Basado en Roles (RBAC) y Políticas de Contraseñas / MFA.</li>
                            <li><strong>Artefacto 3:</strong> Procedimiento Operativo Estándar (SOP) para el ciclo de vida de credenciales y altas/bajas de personal.</li>
                            <li><strong>Artefacto 4:</strong> Informe de Evaluación de Vulnerabilidades CVE/CVSS y Plan de Mitigación sin Explotación.</li>
                            <li><strong>Artefacto 5:</strong> Procedimiento de Notificación y Escalación de Incidentes de Seguridad.</li>
                        </ul>
                    </div>
                    <div class="grid grid-cols-2 gap-6 pt-6 text-center">
                        <div class="border-t border-slate-400 pt-2">Firma del Aprendiz</div>
                        <div class="border-t border-slate-400 pt-2">Evaluación Terminal Kinal (Nota: _____ / 100)</div>
                    </div>
                </div>
            </div>

        </div>

    </main>

    <!-- Footer -->
    <footer class="bg-slate-950 text-slate-400 py-6 border-t border-slate-800 text-xs">
        <div class="max-w-7xl mx-auto px-4 sm:px-6 lg:px-8 flex flex-col sm:flex-row justify-between items-center gap-4 text-center sm:text-left">
            <div>
                <p class="font-bold text-slate-200">Fundación Kinal • Curso de Ciberseguridad y Fundamentos de Seguridad</p>
                <p class="text-slate-500">Diseño Curricular, Didáctico e Instruccional 2026 | Sistema Dual & DQR 4-5</p>
            </div>
            <div class="flex items-center gap-4">
                <a href="index.html" class="hover:text-amber-400 transition">Catálogo de Cursos</a>
                <span class="text-slate-700">•</span>
                <a href="mecatronica.html" class="hover:text-amber-400 transition">Ver Mecatrónica</a>
                <span class="text-slate-700">•</span>
                <a href="../../index.html" class="hover:text-amber-400 transition">Portal Maestro</a>
            </div>
        </div>
    </footer>

    <!-- Interactive Logic Script -->
    <script>
        const ciberSessions = {ciber_sessions_json};
        let currentSessionNum = 1;

        function switchTab(tabId) {{
            document.querySelectorAll('.tab-content').forEach(el => el.classList.add('hidden'));
            document.querySelectorAll('.tab-btn').forEach(el => {{
                el.classList.remove('active');
                el.classList.add('bg-slate-100', 'text-slate-700');
            }});
            
            const target = document.getElementById('tab-' + tabId);
            if (target) target.classList.remove('hidden');
            const activeBtn = document.getElementById('btn-' + tabId);
            if (activeBtn) {{
                activeBtn.classList.add('active');
                activeBtn.classList.remove('bg-slate-100', 'text-slate-700');
            }}
            window.scrollTo({{ top: 220, behavior: 'smooth' }});
        }}

        function filterCiberTemario(category) {{
            document.querySelectorAll('.ciber-filter-btn').forEach(b => {{
                b.classList.remove('bg-slate-900', 'text-white', 'shadow-sm');
            }});
            const activeBtn = document.getElementById('btn-ciber-' + category);
            if (activeBtn) {{
                activeBtn.classList.add('bg-slate-900', 'text-white', 'shadow-sm');
            }}

            const items = document.querySelectorAll('#tab-temario .content-item');
            items.forEach(it => {{
                if (category === 'all' || it.getAttribute('data-category') === category) {{
                    it.style.display = 'block';
                }} else {{
                    it.style.display = 'none';
                }}
            }});
        }}

        function renderCiberSessionList(list) {{
            const listContainer = document.getElementById('ciber-sessions-list');
            if (!listContainer) return;
            
            listContainer.innerHTML = list.map(s => `
                <div onclick="selectCiberSession(${{s.num}})" id="ciber-list-item-${{s.num}}" class="session-item p-3 rounded-2xl border border-slate-200 bg-slate-50/70 hover:bg-slate-100 transition cursor-pointer flex items-center justify-between text-xs ${{s.num === currentSessionNum ? 'active' : ''}}">
                    <div class="flex items-center gap-2.5 overflow-hidden">
                        <span class="badge-code w-7 h-7 rounded-lg bg-indigo-900 text-white font-black text-[11px] flex items-center justify-center shrink-0">
                            ${{s.code}}
                        </span>
                        <div class="truncate">
                            <span class="text-[10px] uppercase font-bold text-indigo-700 block">${{s.semana}}</span>
                            <div class="font-bold text-slate-900 truncate">${{s.title}}</div>
                        </div>
                    </div>
                    <i data-lucide="chevron-right" class="w-4 h-4 text-slate-400 shrink-0"></i>
                </div>
            `).join('');
            lucide.createIcons();
        }}

        function renderCiberSessionDetail(s) {{
            const detailContainer = document.getElementById('ciber-detail-card');
            if (!detailContainer) return;

            const badgeMod = s.is_presencial ? 
                '<span class="px-2 py-0.5 rounded-md bg-emerald-100 text-emerald-800 font-bold text-xs">🏢 Presencial en Kinal</span>' : 
                '<span class="px-2 py-0.5 rounded-md bg-indigo-50 text-indigo-800 font-bold text-xs border border-indigo-200">💻 Virtual Síncrona</span>';

            detailContainer.innerHTML = `
                <div class="space-y-4">
                    <div class="flex flex-wrap justify-between items-start gap-2 border-b border-slate-100 pb-3">
                        <div class="space-y-1">
                            <div class="flex items-center gap-2">
                                <span class="px-2.5 py-0.5 rounded-md bg-indigo-900 text-white font-black text-xs">${{s.code}}</span>
                                <span class="px-2.5 py-0.5 rounded-md bg-indigo-50 text-indigo-900 font-bold text-xs border border-indigo-200">${{s.module}}</span>
                                ${{badgeMod}}
                            </div>
                            <h4 class="text-xl sm:text-2xl font-black text-slate-900 pt-1">${{s.title}}</h4>
                        </div>
                        <span class="px-3 py-1 rounded-full bg-slate-100 text-slate-700 font-bold text-xs flex items-center gap-1">
                            <i data-lucide="clock" class="w-3.5 h-3.5 text-slate-500"></i> ${{s.duration}}
                        </span>
                    </div>

                    <!-- Objetivo -->
                    <div class="p-4 rounded-2xl bg-indigo-50/80 border border-indigo-100 text-xs space-y-1">
                        <strong class="text-indigo-950 font-bold flex items-center gap-1.5 text-sm">
                            <i data-lucide="target" class="w-4 h-4 text-indigo-700"></i> Objetivo Pedagógico de la Sesión
                        </strong>
                        <p class="text-slate-800 leading-relaxed text-xs">${{s.objective}}</p>
                    </div>

                    <!-- 3 Momentos -->
                    <div class="grid md:grid-cols-3 gap-3.5 text-xs pt-1">
                        <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 space-y-1.5">
                            <div class="flex items-center gap-1.5 font-black text-slate-900 text-xs">
                                <span class="w-5 h-5 rounded-full bg-indigo-100 text-indigo-900 text-[10px] flex items-center justify-center font-black">1</span>
                                Apertura (30 min)
                            </div>
                            <p class="text-slate-600 leading-relaxed text-[11px]">${{s.apertura}}</p>
                        </div>

                        <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 space-y-1.5">
                            <div class="flex items-center gap-1.5 font-black text-slate-900 text-xs">
                                <span class="w-5 h-5 rounded-full bg-emerald-100 text-emerald-900 text-[10px] flex items-center justify-center font-black">2</span>
                                Desarrollo (180 min)
                            </div>
                            <p class="text-slate-600 leading-relaxed text-[11px]">${{s.desarrollo}}</p>
                        </div>

                        <div class="p-4 rounded-2xl bg-slate-50 border border-slate-200 space-y-1.5">
                            <div class="flex items-center gap-1.5 font-black text-slate-900 text-xs">
                                <span class="w-5 h-5 rounded-full bg-amber-100 text-amber-900 text-[10px] flex items-center justify-center font-black">3</span>
                                Cierre (30 min)
                            </div>
                            <p class="text-slate-600 leading-relaxed text-[11px]">${{s.cierre}}</p>
                        </div>
                    </div>

                    <!-- Footer: Evidencias y Equipamiento -->
                    <div class="grid sm:grid-cols-2 gap-3 pt-3 text-xs border-t border-slate-100">
                        <div class="p-3 rounded-xl bg-slate-50 border border-slate-200 space-y-1">
                            <span class="text-[10px] uppercase font-bold text-slate-500 block">Producto Observable</span>
                            <div class="text-slate-800 font-medium flex items-start gap-1.5">
                                <i data-lucide="check-circle" class="w-4 h-4 text-emerald-600 mt-0.5 shrink-0"></i>
                                <span>${{s.evidencias}}</span>
                            </div>
                        </div>
                        <div class="p-3 rounded-xl bg-slate-50 border border-slate-200 space-y-1">
                            <span class="text-[10px] uppercase font-bold text-slate-500 block">Entorno / Caso Transversal</span>
                            <div class="text-slate-800 font-medium flex items-start gap-1.5">
                                <i data-lucide="shield" class="w-4 h-4 text-indigo-600 mt-0.5 shrink-0"></i>
                                <span>${{s.equipamiento}}</span>
                            </div>
                        </div>
                    </div>

                    <!-- Navegación entre sesiones -->
                    <div class="flex justify-between items-center pt-2">
                        <button onclick="navigateCiberSession(-1)" ${{s.num === 1 ? 'disabled class="opacity-40 cursor-not-allowed"' : 'class="hover:bg-slate-100"'}} px-3 py-1.5 rounded-lg border border-slate-200 text-xs font-bold text-slate-700 flex items-center gap-1 transition">
                            <i data-lucide="arrow-left" class="w-3.5 h-3.5"></i> Sesión Anterior
                        </button>
                        <span class="text-xs font-bold text-slate-400">Sesión ${{s.num}} de 20</span>
                        <button onclick="navigateCiberSession(1)" ${{s.num === 20 ? 'disabled class="opacity-40 cursor-not-allowed"' : 'class="hover:bg-slate-100"'}} px-3 py-1.5 rounded-lg border border-slate-200 text-xs font-bold text-slate-700 flex items-center gap-1 transition">
                            Siguiente Sesión <i data-lucide="arrow-right" class="w-3.5 h-3.5"></i>
                        </button>
                    </div>
                </div>
            `;
            lucide.createIcons();
        }}

        function selectCiberSession(num) {{
            currentSessionNum = num;
            document.querySelectorAll('.session-item').forEach(el => el.classList.remove('active'));
            const activeItem = document.getElementById('ciber-list-item-' + num);
            if (activeItem) activeItem.classList.add('active');

            const s = ciberSessions.find(item => item.num === num);
            if (s) renderCiberSessionDetail(s);
        }}

        function navigateCiberSession(direction) {{
            const newNum = currentSessionNum + direction;
            if (newNum >= 1 && newNum <= 20) {{
                selectCiberSession(newNum);
            }}
        }}

        function goToSession(num) {{
            switchTab('secuencias');
            selectCiberSession(num);
        }}

        function filterCiberList() {{
            const q = document.getElementById('ciber-sec-search').value.toLowerCase();
            const filtered = ciberSessions.filter(s => 
                s.title.toLowerCase().includes(q) || 
                s.code.toLowerCase().includes(q) || 
                s.objective.toLowerCase().includes(q)
            );
            renderCiberSessionList(filtered);
            document.getElementById('ciber-count-badge').innerText = filtered.length + ' Sesiones';
        }}

        document.addEventListener('DOMContentLoaded', () => {{
            renderCiberSessionList(ciberSessions);
            selectCiberSession(1);
            lucide.createIcons();
        }});
    </script>
</body>
</html>
"""

with open(os.path.join(WEB_DIR, "ciberseguridad.html"), "w", encoding="utf-8") as f:
    f.write(ciberseguridad_html)
print("Generado con éxito: Cursos/Web/ciberseguridad.html")

print("\n¡Ambos cursos actualizados con las nuevas especificaciones de navegación, calendario y secuencias!")
