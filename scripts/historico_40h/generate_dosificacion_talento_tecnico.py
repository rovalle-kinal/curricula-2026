import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=50, bottom=50, left=70, right=70):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E0", sz="4"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_callout(doc, text, bold_prefix=""):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    
    cell = table.cell(0, 0)
    set_cell_background(cell, "F7FAFC")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="003366"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(9.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 51, 102)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(9.5)
    r.font.italic = True
    r.font.color.rgb = RGBColor(45, 55, 72)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def style_heading(p, text, level=1):
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.bold = True
    if level == 1:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        r.font.size = Pt(13.5)
        r.font.color.rgb = RGBColor(0, 51, 102)
    elif level == 2:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        r.font.size = Pt(11.5)
        r.font.color.rgb = RGBColor(0, 51, 102)
    elif level == 3:
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(74, 85, 104)

def style_p(p, space_before=0, space_after=4, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def generate_dosificacion_doc():
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    COLOR_PRIMARY = RGBColor(0, 51, 102)      # Azul Kinal
    COLOR_SECONDARY = RGBColor(74, 85, 104)    # Gris pizarra
    COLOR_DARK = RGBColor(26, 32, 44)         # Texto oscuro

    # =========================================================================
    # ENCABEZADO INSTITUCIONAL
    # =========================================================================
    p_inst = doc.add_paragraph()
    style_p(p_inst, space_before=0, space_after=2)
    r_inst = p_inst.add_run("FUNDACIÓN KINAL — ESCUELA TÉCNICA SUPERIOR")
    r_inst.font.name = "Calibri"
    r_inst.font.size = Pt(11)
    r_inst.font.bold = True
    r_inst.font.color.rgb = COLOR_PRIMARY

    p_tit = doc.add_paragraph()
    style_p(p_tit, space_before=0, space_after=4)
    r_tit = p_tit.add_run("DOSIFICACIÓN DIDÁCTICA Y SECUENCIACIÓN PEDAGÓGICA:\nDIPLOMADO DE CERTIFICACIÓN GERENCIAL EN GESTIÓN ESTRATÉGICA DEL TALENTO TÉCNICO")
    r_tit.font.name = "Calibri"
    r_tit.font.size = Pt(17)
    r_tit.font.bold = True
    r_tit.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    style_p(p_sub, space_before=0, space_after=10)
    r_sub = p_sub.add_run("Cronograma Pedagógico Sesión por Sesión | 4 Semanas Intensivas Consecutivas | Modalidad Híbrida Ejecutiva (12 Sesiones Directas: 8 Virtuales Teams + 4 Sábados Consecutivos en Kinal)")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(10.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = COLOR_SECONDARY

    add_callout(doc, 
        "«El aprender en el proceso del trabajo es base y a la vez objetivo de la acción formativa». Las 32 horas de contacto directo se organizan en 4 semanas continuas e intensivas (1 mes), articulando 8 sesiones virtuales sincrónicas de 2 horas en Microsoft Teams (Martes y Jueves de 19:00 a 21:00 hrs) y 4 talleres presenciales vivenciales de 4 horas en el campus de Fundación Kinal (4 Sábados consecutivos sin interrupción, de 8:00 a 12:00 hrs). El programa capacita a directores de RRHH, gerentes generales y reclutadores para capitalizar el valor del Técnico Superior Universitario (TSU) y desplegar la cultura del «Factor Humano» y el «trabajo bien hecho» en sus cuadrillas operativas. Nota mínima aprobatoria institucional: 75 puntos sobre 100.",
        bold_prefix="Organización Pedagógica Dual e Ideario Kinal:")

    # =========================================================================
    # 1. MATRIZ GENERAL CRONOLÓGICA DE LAS 12 SESIONES (4 SEMANAS CONSECUTIVAS)
    # =========================================================================
    h1_1 = doc.add_paragraph()
    style_heading(h1_1, "1. Matriz General de Dosificación Cronológica (4 Semanas Consecutivas — 12 Sesiones)", level=1)

    sesiones_info = [
        # SEMANA 1
        ("Sem. 1 / S1 (Virt)", "Mód. 1: Brecha del TSU", "Radiografía del empleo técnico en Guatemala y costos del 'Missing Middle'", "Mapa de pérdidas operativas por rotación"),
        ("Sem. 1 / S2 (Virt)", "Mód. 1: Brecha del TSU", "Anatomía del TSU: competencias DQR 5-6 frente a ingenieros y bachilleres", "Matriz comparativa de roles y modelo de ROI"),
        ("Sem. 1 / S3 (Pres)", "Mód. 1: Taller Kinal 1", "SÁBADO 1 KINAL: Inmersión en laboratorios TSU y disección de parada de planta", "Business Case de TSU validado in situ"),
        
        # SEMANA 2
        ("Sem. 2 / S4 (Virt)", "Mód. 2: Reclutamiento Dual", "Reingeniería de perfiles técnicos: de requisitos a competencias demostrables", "Perfil de puesto basado en competencias"),
        ("Sem. 2 / S5 (Virt)", "Mód. 2: Reclutamiento Dual", "Entrevistas técnicas situacionales STAR y Sistema Dual (Ausbilder en planta)", "Guía de entrevista situacional y esquema dual"),
        ("Sem. 2 / S6 (Pres)", "Mód. 2: Taller Kinal 2", "SÁBADO 2 KINAL: Auditoría de competencias en vivo y diseño del semillero", "Plan de Semillero Dual empresarial"),
        
        # SEMANA 3
        ("Sem. 3 / S7 (Virt)", "Mód. 3: Factor Humano", "Antropología del trabajo de Kinal y los 4 Pilares del Trabajo Bien Hecho", "Diagnóstico de clima y código de excelencia"),
        ("Sem. 3 / S8 (Virt)", "Mód. 3: Factor Humano", "Erradicación del maltrato en cuadrillas, liderazgo empático y canales éticos", "Protocolo de tolerancia cero al maltrato"),
        ("Sem. 3 / S9 (Pres)", "Mód. 3: Taller Kinal 3", "SÁBADO 3 KINAL: Laboratorio de liderazgo y 'Charlas de 10 min de Factor Humano'", "Catálogo de 5 charlas operativas de valores"),
        
        # SEMANA 4
        ("Sem. 4 / S10 (Virt)", "Mód. 4: Proyecto Capstone", "Metodología de despliegue corporativo e integración en inducción/desempeño", "Cronograma de implementación corporativa"),
        ("Sem. 4 / S11 (Virt)", "Mód. 4: Proyecto Capstone", "Cuadro de Mando Integral (KPIs: OEE, SSO, retención), Kinal y ensayo previo", "Dashboard de KPIs y borrador final Capstone"),
        ("Sem. 4 / S12 (Pres)", "Mód. 4: Taller Kinal 4", "SÁBADO 4 KINAL: DEFENSA CAPSTONE ANTE TRIBUNAL, GRADUACIÓN Y CLAUSURA", "Plan Capstone sustentado (>= 75 Pts)")
    ]

    tbl_matriz = doc.add_table(rows=len(sesiones_info) + 1, cols=4)
    tbl_matriz.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_matriz.autofit = False
    tbl_matriz.columns[0].width = Inches(1.2)
    tbl_matriz.columns[1].width = Inches(1.6)
    tbl_matriz.columns[2].width = Inches(2.3)
    tbl_matriz.columns[3].width = Inches(1.4)
    set_table_borders(tbl_matriz, color="CBD5E0", sz="4")

    hdr_m = tbl_matriz.rows[0]
    for c in hdr_m.cells:
        set_cell_background(c, "003366")
    headers_mat = ["Sesión / Formato", "Módulo / Eje", "Contenido Central del Encuentro", "Producto Observable"]
    for idx, text in enumerate(headers_mat):
        p = hdr_m.cells[idx].paragraphs[0]
        style_p(p, space_before=0, space_after=0)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for idx, (ses, mod, cont, prod) in enumerate(sesiones_info, start=1):
        row = tbl_matriz.rows[idx]
        if "Pres" in ses:
            for c in row.cells:
                set_cell_background(c, "EBF8FF")  # Resaltado azul claro para sábados presenciales consecutivos
        elif idx % 2 == 1:
            for c in row.cells:
                set_cell_background(c, "F7FAFC")
        for c_idx, val in enumerate([ses, mod, cont, prod]):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=45, bottom=45, left=55, right=55)
            p = cell.paragraphs[0]
            style_p(p, space_before=0, space_after=0)
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(8)
            r.font.color.rgb = COLOR_DARK
            if c_idx in [0, 3] or "Pres" in ses:
                r.font.bold = True

    doc.add_page_break()

    # =========================================================================
    # 2. MICRODISEÑO DIDÁCTICO DETALLADO SESIÓN POR SESIÓN (12 SESIONES)
    # =========================================================================
    h1_2 = doc.add_paragraph()
    style_heading(h1_2, "2. Microdiseño Didáctico Detallado por Sesión (12 Encuentros)", level=1)

    p_micro_intro = doc.add_paragraph()
    style_p(p_micro_intro, space_before=0, space_after=8)
    r = p_micro_intro.add_run("Cada una de las 12 sesiones cuenta con dosificación temporal estricta: 2 horas (120 min) para sesiones virtuales en Microsoft Teams y 4 horas (240 min) para talleres presenciales vivenciales en Fundación Kinal (4 Sábados consecutivos sin interrupciones):")
    r.font.name = "Calibri"
    r.font.size = Pt(9.5)
    r.font.color.rgb = COLOR_DARK

    sesiones_detalladas = [
        # SEMANA 1
        ("Sesión 1 (Virtual Teams — Martes)", "Módulo 1: Brecha del TSU y Retorno de Inversión", 
         "Radiografía del Talento Operativo en Guatemala y Costos del 'Missing Middle'",
         "Identifica los costos invisibles de rotación y mala asignación técnica en su empresa mediante un mapa de pérdidas operativas.",
         "Apertura (15 min): Encuadre ejecutivo y pregunta detonante: '¿Cuánto le cuesta a su empresa un supervisor técnico vacante durante 3 meses?'. Reflexión sobre las fallas comunes en la industria.",
         "Desarrollo (90 min con pausa activa de 10 min): Exposición sobre el 'Missing Middle' en Guatemala. Comparativa de rotación entre industrias. Trabajo en salas de Teams: costeo de rotación e incidencias operativas en las empresas de los participantes.",
         "Cierre (15 min): Plenaria de hallazgos, síntesis metodológica y asignación de la lectura directiva del Módulo 1 en Kinal.academy.",
         "Microsoft Teams, salas privadas de Teams, Kinal.academy, plantilla de pérdidas operativas.",
         "Borrador del mapa de pérdidas operativas y costos ocultos de rotación.", "2 Horas (120 min)"),

        ("Sesión 2 (Virtual Teams — Jueves)", "Módulo 1: Brecha del TSU y Retorno de Inversión",
         "El Técnico Superior Universitario (TSU), Marco DQR 5-6 y Modelado del ROI",
         "Diferencia las competencias DQR 5-6 del TSU frente a ingenieros y bachilleres, y formula el modelo financiero de retorno de inversión.",
         "Apertura (15 min): Sondeo interactivo: '¿Por qué los ingenieros renuncian a los puestos de supervisión de planta?'. Presentación de objetivos.",
         "Desarrollo (90 min con pausa activa de 10 min): Estructura del TSU de Kinal (3 años, 1,500+ horas prácticas). Correspondencia DQR nivel 6 (Meister). Modelado del ROI: ahorro por reducción de MTTR, menor merma y retención a largo plazo.",
         "Cierre (15 min): Conclusiones sobre la complementariedad operativa (Ingeniero vs TSU vs Operario) y consigna para la visita presencial del Sábado 1.",
         "Microsoft Teams, presentación ejecutiva DQR, calculadora de ROI en Excel, Kinal.academy.",
         "Matriz comparativa de 3 perfiles técnicos y cálculo inicial del ROI.", "2 Horas (120 min)"),

        ("Sesión 3 (Presencial Sábado 1 Kinal)", "Módulo 1: Brecha del TSU y Retorno de Inversión",
         "SÁBADO 1 KINAL: Inmersión en Laboratorios TSU y Disección de Paradas de Planta",
         "Comprueba en campo la solvencia técnica del TSU de Kinal y valida el Business Case de ahorro operativo de su organización.",
         "Apertura (30 min): Bienvenida por autoridades de Fundación Kinal en Auditorio Ejecutivo. Café de networking entre directores y gerentes.",
         "Desarrollo (180 min con pausa de 20 min): Recorrido guiado por Laboratorios TSU (Automatización Siemens, Redes y Mecatrónica). Los estudiantes de último año ejecutan demostraciones en vivo. Mesa redonda 'Disección de una parada de planta': análisis de casos reales de downtime.",
         "Cierre (30 min): Síntesis ejecutiva, retroalimentación in situ del facilitador de Kinal y entrega del Business Case M1.",
         "Campus Kinal (Auditorio y Complejo de Laboratorios TSU), guías de análisis de caso, refrigerio.",
         "Business Case del TSU validado con datos operacionales de planta.", "4 Horas (240 min)"),

        # SEMANA 2
        ("Sesión 4 (Virtual Teams — Martes)", "Módulo 2: Reclutamiento por Competencias y Formación Dual",
         "Reingeniería del Perfil Técnico: De Requisitos a Competencias Demostrables",
         "Transforma descriptores de puesto burocráticos en matrices de competencias operativas observables y niveles de autonomía técnica.",
         "Apertura (15 min): Revisión de un perfil de puesto tradicional 'real': análisis de incongruencias y requisitos irrelevantes.",
         "Desarrollo (90 min con pausa activa de 10 min): Metodología de redacción de competencias por resultados de aprendizaje (saber, saber hacer, saber ser). Definición de niveles de autonomía en planta. Taller en Teams: reescritura de un puesto técnico crítico.",
         "Cierre (15 min): Socialización de matrices mejoradas y asignación de la lectura sobre el Sistema Dual alemán en Kinal.academy.",
         "Microsoft Teams, pizarras colaborativas Miro/Teams, Kinal.academy, guía de redacción DQR.",
         "Descriptor de puesto técnico transformado a competencias demostrables.", "2 Horas (120 min)"),

        ("Sesión 5 (Virtual Teams — Jueves)", "Módulo 2: Reclutamiento por Competencias y Formación Dual",
         "Entrevistas Situacionales (STAR) y el Sistema Dual Escuela-Empresa (Ausbilder)",
         "Diseña guías de entrevista situacional con método STAR y estructura un modelo de aprendizaje dual con tutores en planta.",
         "Apertura (15 min): Discusión guiada: '¿Por qué los tests psicométricos de oficina no predicen la destreza de un técnico en una avería?'.",
         "Desarrollo (90 min con pausa activa de 10 min): Técnica STAR adaptada al entorno técnico. Elaboración de rúbricas objetivas de evaluación. Fundamentos del Sistema Dual: alternancia formativa, marco legal y el rol del mentor Ausbilder en planta.",
         "Cierre (15 min): Síntesis de buenas prácticas de reclutamiento y preparación para el panel de simulación del Sábado 2 en Kinal.",
         "Microsoft Teams, formato de entrevista STAR técnica, lecturas de formación dual, Kinal.academy.",
         "Guía de entrevista técnica situacional y borrador de rotación dual.", "2 Horas (120 min)"),

        ("Sesión 6 (Presencial Sábado 2 Kinal)", "Módulo 2: Reclutamiento por Competencias y Formación Dual",
         "SÁBADO 2 KINAL: Auditoría de Competencias en Vivo y Diseño del Semillero Dual",
         "Aplica entrevistas situacionales a estudiantes de TSU, audita pruebas técnicas en banco y co-diseña un semillero dual para su empresa.",
         "Apertura (20 min): Encuadre metodológico en sala ejecutiva y asignación de roles para el panel de selección.",
         "Desarrollo (190 min con pausa de 20 min): Dinámica en dos bloques: 1) Panel de entrevistas técnicas reales con estudiantes de último año de TSU. 2) Auditoría en laboratorio: los gerentes observan y califican con rúbricas una prueba práctica de comisionamiento de circuito en tiempo real. Taller de co-diseño de convenios duales.",
         "Cierre (30 min): Plenaria de reflexiones: contrastar la percepción teórica con la competencia demostrada por los jóvenes técnicos de Kinal.",
         "Auditorio y Laboratorios Kinal, rúbricas impresas, expedientes de aprendices TSU, refrigerio.",
         "Plan de Semillero Dual empresarial y rúbricas de auditoría técnica diligenciadas.", "4 Horas (240 min)"),

        # SEMANA 3
        ("Sesión 7 (Virtual Teams — Martes)", "Módulo 3: Cultura del Factor Humano en la Industria",
         "Antropología del Trabajo de Kinal y los Cuatro Pilares del Trabajo Bien Hecho",
         "Fundamenta la cultura del Factor Humano en valores laicos universales y formula el código de excelencia operativa para su empresa.",
         "Apertura (15 min): Presentación del video testimonial: 'El técnico de Kinal en la industria'. Discusión sobre el valor intrínseco del trabajo.",
         "Desarrollo (90 min con pausa activa de 10 min): La filosofía humanista de Kinal: el trabajo como servicio y perfeccionamiento personal. Los 4 pilares: Respeto a la Dignidad, Trabajo Bien Hecho (5S, precisión moral), Servicio y Libertad Responsable. Ejercicio: diagnóstico de brechas culturales en planta.",
         "Cierre (15 min): Reflexión grupal y pautas para el protocolo de erradicación del maltrato en Kinal.academy.",
         "Microsoft Teams, presentaciones multimedia, Kinal.academy, guía de los 4 pilares de Kinal.",
         "Diagnóstico de clima operativo y borrador del Código del Trabajo Bien Hecho.", "2 Horas (120 min)"),

        ("Sesión 8 (Virtual Teams — Jueves)", "Módulo 3: Cultura del Factor Humano en la Industria",
         "Erradicación del Maltrato en Cuadrillas, Liderazgo Empático y Canales Éticos",
         "Estructura protocolos de cero tolerancia al maltrato en planta, seguridad psicológica y canales seguros de reporte ético (Whistleblowing).",
         "Apertura (15 min): Análisis de caso: 'El capataz agresivo que nadie despide porque saca la producción'. Discusión sobre los costos ocultos del maltrato.",
         "Desarrollo (90 min con pausa activa de 10 min): Transformación del liderazgo técnico: de la intimidación al liderazgo de servicio. Seguridad psicológica como requisito de SSO y productividad. Diseño de políticas de tolerancia cero al maltrato y canales seguros de denuncia ética.",
         "Cierre (15 min): Conclusiones y preparación para el Laboratorio de Liderazgo del Sábado 3 en Kinal.",
         "Microsoft Teams, estudio de caso ético, modelo de canal ético, Kinal.academy.",
         "Protocolo corporativo de erradicación del maltrato y canales de denuncia.", "2 Horas (120 min)"),

        ("Sesión 9 (Presencial Sábado 3 Kinal)", "Módulo 3: Cultura del Factor Humano en la Industria",
         "SÁBADO 3 KINAL: Laboratorio de Liderazgo y 'Charlas de 10 Minutos de Factor Humano'",
         "Resuelve dilemas éticos reales de planta mediante dramatizaciones y diseña un catálogo de charlas operativas de Factor Humano.",
         "Apertura (20 min): Encuadre del taller vivencial de liderazgo en Auditorio Kinal.",
         "Desarrollo (190 min con pausa de 20 min): Role-playing de resolución de dilemas éticos (presión para omitir protocolos de seguridad, encubrimiento de piezas defectuosas, maltrato verbal a un operario nuevo). Taller práctico: diseño y simulación en vivo de las 'Charlas de 10 minutos de Factor Humano' para el arranque de turno.",
         "Cierre (30 min): Retroalimentación de los instructores de Factor Humano de Kinal y afinación de los entregables individuales.",
         "Auditorio Kinal, dinámicas teatrales y guías de dilemas éticos, Kinal.academy, refrigerio.",
         "Catálogo de 5 micro-charlas operativas de Factor Humano para su empresa.", "4 Horas (240 min)"),

        # SEMANA 4
        ("Sesión 10 (Virtual Teams — Martes)", "Módulo 4: Creación y Defensa del Proyecto Capstone",
         "Metodología de Despliegue Corporativo e Integración en Inducción y Desempeño",
         "Estructura el plan integral de despliegue de Factor Humano en su organización, vinculándolo a procesos de inducción y evaluación.",
         "Apertura (15 min): Presentación de la estructura final obligatoria del Proyecto Directivo Capstone.",
         "Desarrollo (90 min con pausa activa de 10 min): Ruta de despliegue: diagnóstico de planta, sensibilización de mandos medios, plan piloto y escalamiento. Integración de valores en la inducción de personal operativo y en la evaluación semestral. Pautas para el cuadro de mando integral.",
         "Cierre (15 min): Aclaración de dudas metodológicas y calendarización de la sustentación del Sábado 4.",
         "Microsoft Teams, plantilla oficial del Proyecto Directivo Capstone, Kinal.academy.",
         "Plan de despliegue corporativo estructurado por fases y cronograma.", "2 Horas (120 min)"),

        ("Sesión 11 (Virtual Teams — Jueves)", "Módulo 4: Creación y Defensa del Proyecto Capstone",
         "Cuadro de Mando Integral (KPIs), Alianzas con Kinal y Ensayo Previo en Teams",
         "Define indicadores de impacto (OEE, SSO, retención) y realiza el ensayo de su sustentación oral con retroalimentación formativa.",
         "Apertura (15 min): Revisión de las métricas clave de éxito directivo (OEE, tasa de accidentes, índice de rotación voluntaria).",
         "Desarrollo (90 min con pausa activa de 10 min): Construcción del Dashboard de Talento Técnico. Ecosistema de vinculación con Kinal (convenios marco, acceso a bolsa de trabajo, homologación). Salas privadas de Teams: ensayo de presentación ejecutiva (pitch de 10 min) entre pares y retroalimentación de facilitadores.",
         "Cierre (15 min): Pautas protocolares finales para la Defensa Capstone del Sábado 4 en Kinal.",
         "Microsoft Teams, salas privadas de ensayo, rúbrica de sustentación oral, Kinal.academy.",
         "Dashboard de KPIs de talento técnico y diapositivas finales del Capstone.", "2 Horas (120 min)"),

        ("Sesión 12 (Presencial Sábado 4 Kinal)", "Módulo 4: Creación y Defensa del Proyecto Capstone",
         "SÁBADO 4 KINAL: DEFENSA CAPSTONE ANTE TRIBUNAL, GRADUACIÓN Y CLAUSURA",
         "Sustenta con solvencia directiva su Plan Integral de Talento y Factor Humano ante el tribunal institucional, obteniendo la certificación Kinal.",
         "Apertura (20 min): Instalación del Tribunal Evaluador Directivo conformado por autoridades de Fundación Kinal y líderes empresariales.",
         "Desarrollo (180 min con pausa de 20 min): Rondas individuales de sustentación oral (15 min de exposición + 5 min de preguntas del tribunal). Calificación con la Rúbrica Analítica Institucional (Umbral aprobatorio: nota mínima de 75 puntos sobre 100).",
         "Cierre (40 min): Deliberación del tribunal, proclamación de resultados, acto protocolario de entrega de diplomas de certificación y brindis de honor de networking directivo.",
         "Auditorio Ejecutivo Kinal, equipamiento audiovisual, actas de evaluación, pergaminos de certificación, cóctel.",
         "Proyecto Directivo Capstone defendido con éxito y certificación gerencial obtenida (>= 75 Pts).", "4 Horas (240 min)")
    ]

    for ses_nom, eje_nom, tema_nom, obj_esp, act_ap, act_des, act_cie, rec_did, prod_obs, dur_txt in sesiones_detalladas:
        p_tit = doc.add_paragraph()
        style_heading(p_tit, f"{ses_nom} — {tema_nom}", level=2)

        tbl_s = doc.add_table(rows=7, cols=2)
        tbl_s.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl_s.autofit = False
        tbl_s.columns[0].width = Inches(1.8)
        tbl_s.columns[1].width = Inches(4.7)
        set_table_borders(tbl_s, color="CBD5E0", sz="4")

        s_rows = [
            ("Módulo y Duración:", f"{eje_nom}  |  {dur_txt}"),
            ("Objetivo Específico:", obj_esp),
            ("Secuencia Didáctica — Apertura:", act_ap),
            ("Secuencia Didáctica — Desarrollo:", act_des),
            ("Secuencia Didáctica — Cierre:", act_cie),
            ("Recursos y Plataformas:", rec_did),
            ("Entregable y Evidencia:", prod_obs)
        ]

        for r_idx, (etiq, cont) in enumerate(s_rows):
            row = tbl_s.rows[r_idx]
            if "Presencial" in ses_nom or "SÁBADO" in ses_nom:
                set_cell_background(row.cells[0], "EBF8FF")
            else:
                set_cell_background(row.cells[0], "F7FAFC")
            set_cell_margins(row.cells[0], top=50, bottom=50, left=70, right=70)
            set_cell_margins(row.cells[1], top=50, bottom=50, left=70, right=70)
            
            p0 = row.cells[0].paragraphs[0]
            style_p(p0, space_before=0, space_after=0)
            r0 = p0.add_run(etiq)
            r0.font.name = "Calibri"
            r0.font.size = Pt(8.5)
            r0.font.bold = True
            r0.font.color.rgb = COLOR_PRIMARY

            p1 = row.cells[1].paragraphs[0]
            style_p(p1, space_before=0, space_after=0)
            r1 = p1.add_run(cont)
            r1.font.name = "Calibri"
            r1.font.size = Pt(8.5)
            r1.font.color.rgb = COLOR_DARK
            if r_idx in [0, 6]:
                r1.font.bold = True

        p_sp = doc.add_paragraph()
        style_p(p_sp, space_before=2, space_after=4)

    doc.add_page_break()

    # =========================================================================
    # 3. INSTRUMENTOS INSTITUCIONALES DE EVALUACIÓN Y RÚBRICA CAPSTONE
    # =========================================================================
    h1_3 = doc.add_paragraph()
    style_heading(h1_3, "3. Instrumentos de Evaluación Ejecutiva y Rúbrica Capstone", level=1)

    add_callout(doc,
        "La certificación gerencial emitida por Fundación Kinal exige alcanzar una calificación mínima de 75 puntos sobre 100 en la evaluación continua de entregables y en la defensa del Proyecto Directivo Capstone ante el tribunal institucional.",
        bold_prefix="Criterio de Evaluación y Umbral de Aprobación:")

    p_rub = doc.add_paragraph()
    style_heading(p_rub, "3.1 Rúbrica Analítica de Certificación para el Proyecto Directivo Capstone (100 Puntos)", level=2)

    rubrica_data = [
        ("Diagnóstico Estratégico y Business Case del TSU", "25 Puntos", "Demuestra con datos cuantitativos el ROI de incorporar perfiles TSU frente a la rotación de ingenieros y fallas de operarios empíricos.", "Diagnóstico superficial sin datos de costos o desconectado de la realidad operativa de la empresa (< 18.7 pts)."),
        ("Diseño del Perfil por Competencias y Semillero Dual", "25 Puntos", "Descriptor riguroso por competencias demostrables y plan estructurado de semillero dual con perfiles de tutores en planta.", "Descripción genérica de requisitos burocráticos sin esquema de rotación dual medible (< 18.7 pts)."),
        ("Despliegue del Factor Humano y Erradicación del Maltrato", "25 Puntos", "Propone políticas claras de cero tolerancia al maltrato en planta, seguridad psicológica y diseña 5 micro-charlas operativas de valores.", "Planteamiento abstracto de valores sin protocolos contra el maltrato ni aplicación práctica en cuadrillas (< 18.7 pts)."),
        ("Viabilidad de Implementación, KPIs y Sustentación", "25 Puntos", "Presenta un cronograma realista, dashboard de indicadores (OEE, SSO, retención) y defiende su propuesta con solvencia directiva.", "Incoherencia entre metas y recursos, ausencia de métricas de negocio o sustentación deficiente (< 18.7 pts).")
    ]

    tbl_rub = doc.add_table(rows=5, cols=4)
    tbl_rub.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_rub.autofit = False
    tbl_rub.columns[0].width = Inches(1.8)
    tbl_rub.columns[1].width = Inches(1.0)
    tbl_rub.columns[2].width = Inches(2.2)
    tbl_rub.columns[3].width = Inches(1.5)
    set_table_borders(tbl_rub, color="CBD5E0", sz="4")

    hdr_r = tbl_rub.rows[0]
    for c in hdr_r.cells:
        set_cell_background(c, "003366")
    for idx, text in enumerate(["Criterio Evaluado", "Ponderación", "Desempeño Excelente (Estándar Kinal)", "Desempeño Insuficiente (< 75%)"]):
        p = hdr_r.cells[idx].paragraphs[0]
        style_p(p, space_before=0, space_after=0)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for idx, (crit, pond, exc, ins) in enumerate(rubrica_data, start=1):
        row = tbl_rub.rows[idx]
        if idx % 2 == 1:
            for c in row.cells:
                set_cell_background(c, "F7FAFC")
        for c_idx, val in enumerate([crit, pond, exc, ins]):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
            p = cell.paragraphs[0]
            style_p(p, space_before=0, space_after=0)
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(8)
            r.font.color.rgb = COLOR_DARK
            if c_idx in [0, 1]:
                r.font.bold = True

    p_bf = doc.add_paragraph()
    style_heading(p_bf, "3.2 Ficha Ejecutiva del Plan de Implementación de Factor Humano", level=2)

    tbl_b_fmt = doc.add_table(rows=6, cols=2)
    tbl_b_fmt.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_b_fmt.autofit = False
    tbl_b_fmt.columns[0].width = Inches(2.2)
    tbl_b_fmt.columns[1].width = Inches(4.3)
    set_table_borders(tbl_b_fmt, color="CBD5E0", sz="4")

    formato_data = [
        ("Datos del Gerente y Empresa:", "Participante: _______________________ | Cargo: ________________ | Empresa: ________________"),
        ("Línea Operativa Piloto:", "Área técnica seleccionada (ej. 'Línea de envasado 2' / 'Taller de mantenimiento electromecánico')."),
        ("Problema de Clima / Técnico Detectado:", "Diagnóstico inicial de rotación, paradas de línea o estilo de supervisión en cuadrilla."),
        ("Acción de Factor Humano Aplicada:", "Protocolo de liderazgo respetuoso implementado y micro-charlas de valores programadas."),
        ("Resultado e Impacto Medido:", "Variación en OEE, reducción de quejas por maltrato, ausentismo o índice de accidentes SSO."),
        ("Firmas de Acreditación:", "Firma Gerente: ___________________  |  Vo.Bo. Tribunal Kinal: ___________________")
    ]

    for idx, (campo, lineas) in enumerate(formato_data):
        row = tbl_b_fmt.rows[idx]
        set_cell_background(row.cells[0], "F7FAFC")
        set_cell_margins(row.cells[0], top=55, bottom=55, left=75, right=75)
        set_cell_margins(row.cells[1], top=55, bottom=55, left=75, right=75)

        p0 = row.cells[0].paragraphs[0]
        style_p(p0, space_before=0, space_after=0)
        r0 = p0.add_run(campo)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_PRIMARY

        p1 = row.cells[1].paragraphs[0]
        style_p(p1, space_before=0, space_after=0)
        r1 = p1.add_run(lineas)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9)
        r1.font.color.rgb = COLOR_DARK

    section = doc.sections[0]
    footer = section.footer
    footer_p = footer.paragraphs[0]
    style_p(footer_p, space_before=0, space_after=0)
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_f = footer_p.add_run("Fundación Kinal | Diplomado Gerencial en Gestión del Talento Técnico — Dosificación Oficial (4 Semanas Consecutivas)")
    run_f.font.name = "Calibri"
    run_f.font.size = Pt(8.5)
    run_f.font.color.rgb = COLOR_SECONDARY

    out_dir = os.path.join("Cursos", "Talento_Tecnico")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "Dosificacion_y_Secuencia_Didactica_Talento_Tecnico_Kinal.docx")
    doc.save(out_path)
    print(f"Dosificación guardada con éxito en: {out_path}")

if __name__ == "__main__":
    generate_dosificacion_doc()
