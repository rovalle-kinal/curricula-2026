import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E0", sz="4"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_callout(doc, text, bold_prefix=""):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    
    cell = table.cell(0, 0)
    set_cell_background(cell, "F7FAFC")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="003366"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(9.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 51, 102)
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(9.5)
    r.font.italic = True
    r.font.color.rgb = RGBColor(45, 55, 72)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def style_heading(p, text, level=1):
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.bold = True
    if level == 1:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor(0, 51, 102)
    elif level == 2:
        p.paragraph_format.space_before = Pt(11)
        p.paragraph_format.space_after = Pt(4)
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor(0, 51, 102)
    elif level == 3:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(74, 85, 104)

def style_p(p, space_before=0, space_after=4, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def generate_temario_doc():
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    COLOR_PRIMARY = RGBColor(0, 51, 102)      # Azul Kinal
    COLOR_SECONDARY = RGBColor(74, 85, 104)    # Gris pizarra
    COLOR_DARK = RGBColor(26, 32, 44)         # Texto oscuro

    # ==========================================
    # ENCABEZADO INSTITUCIONAL
    # ==========================================
    title_p = doc.add_paragraph()
    style_p(title_p, space_before=0, space_after=4)
    run_title = title_p.add_run("Programa Ejecutivo de Vinculación y Gestión Estratégica del Talento Técnico")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    style_p(sub_p, space_before=0, space_after=12)
    run_sub = sub_p.add_run("Temario Analítico Oficial y Estructura Curricular por Ejes  |  Fundación Kinal — Dirección de Vinculación Empresarial\n1 Mes de Duración | 1 Sesión Semanal Virtual (Microsoft Teams) + 2 Jornadas Presenciales Especiales (Campus Kinal) | Entrega de Diplomas y Reconocimientos")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(10)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_SECONDARY

    add_callout(doc, 
        "«the ability and readiness of the individual to use knowledge, skills and personal, social and methodological competences and to behave in a considered, individual and socially responsible manner. Competence is understood in this sense as the comprehensive ability to act». El programa promueve un espacio de diálogo bidireccional entre la industria y Fundación Kinal para compartir herramientas de liderazgo, manejo de personal y mandos medios, integrando el Factor Humano y conociendo de primera mano las necesidades técnicas y humanas que demanda actualmente el sector productivo.", 
        bold_prefix="Definición de Competencia DQR e Ideario Kinal:")

    # ==========================================
    # 1. OBJETIVOS FORMATIVOS Y COMPETENCIAS
    # ==========================================
    h1_1 = doc.add_paragraph()
    style_heading(h1_1, "1. Competencia General y Resultados de Aprendizaje por Ejes", level=1)

    p_cg = doc.add_paragraph()
    style_p(p_cg, space_before=0, space_after=6)
    r = p_cg.add_run("Competencia General:\n"
                     "El participante implementa metodologías efectivas para el manejo, desarrollo y liderazgo de personal operativo y mandos medios, aplicando la cultura del Factor Humano y los estándares de excelencia del «trabajo bien hecho», participando en mesas de diálogo industria-academia para alinear las competencias del talento técnico con las necesidades actuales del mercado productivo.")
    r.font.name = "Calibri"
    r.font.size = Pt(9.5)
    r.font.bold = True
    r.font.color.rgb = COLOR_PRIMARY

    mod_comps = [
        ("Jornada Presencial Especial 1: Apertura e Inmersión en Kinal", 
         "Conoce la visión formativa integral de Kinal, recorre los laboratorios de última tecnología del TSU y establece redes de contacto directivo en la primera mesa de diálogo sobre empleo técnico."),
        ("Eje 1: Manejo y Desarrollo de Personal Técnico (Sesión Virtual 1)", 
         "Aplica metodologías para la atracción, inducción técnica acelerada y fidelización de personal operativo, mitigando los costos ocultos por rotación y el fenómeno del 'Missing Middle'."),
        ("Eje 2: Liderazgo y Factor Humano en Cuadrillas (Sesión Virtual 2)", 
         "Despliega protocolos de liderazgo humanista en plantas y talleres, erradicando el maltrato verbal y autoritarismo, y fortaleciendo la dignidad, la seguridad psicológica y el orgullo por el trabajo bien hecho."),
        ("Eje 3: Fortalecimiento de Competencias Técnicas y Formación Dual (Sesión Virtual 3)", 
         "Identifica la cualificación DQR 5-6 del Técnico Superior Universitario (TSU) y formula esquemas de semilleros duales escuela-empresa con tutores de planta (Ausbilder)."),
        ("Eje 4: Desarrollo de Mandos Medios y Necesidades Actuales (Sesión Virtual 4)", 
         "Equipa a supervisores de primera línea con herramientas de gestión operativa y resolución de conflictos, correlacionando el factor humano con KPIs de negocio (OEE, SSO y merma)."),
        ("Jornada Presencial Especial 2: Diálogo Industria-Academia y Clausura", 
         "Participa en la plenaria de síntesis de necesidades del sector empresarial, recibe el diploma de participación institucional y el reconocimiento corporativo para su empresa.")
    ]

    for m_tit, m_desc in mod_comps:
        p_m = doc.add_paragraph()
        style_p(p_m, space_before=2, space_after=4)
        p_m.paragraph_format.left_indent = Inches(0.2)
        r_b = p_m.add_run(f"• {m_tit}: ")
        r_b.font.name = "Calibri"
        r_b.font.size = Pt(9.5)
        r_b.font.bold = True
        r_b.font.color.rgb = COLOR_PRIMARY
        r_d = p_m.add_run(m_desc)
        r_d.font.name = "Calibri"
        r_d.font.size = Pt(9.5)
        r_d.font.color.rgb = COLOR_DARK

    # ==========================================
    # 2. TEMARIO ANALÍTICO DETALLADO (1 MES — 6 ENCUENTROS)
    # ==========================================
    h1_2 = doc.add_paragraph()
    style_heading(h1_2, "2. Temario Analítico Detallado por Encuentro", level=1)

    # -------------------------------------------------------------
    # ENCUENTRO 1: JORNADA PRESENCIAL DE APERTURA
    # -------------------------------------------------------------
    h2_e1 = doc.add_paragraph()
    style_heading(h2_e1, "JORNADA PRESENCIAL ESPECIAL 1: APERTURA INSTITUCIONAL Y RECORRIDO TÉCNICO (4 Horas en Campus Kinal)", level=2)

    temas_e1 = [
        ("1.1 Sesión Plenaria de Apertura y Presentación de la Iniciativa", [
            "Bienvenida protocolaria por autoridades de Fundación Kinal y presentación del propósito del programa de vinculación empresarial.",
            "La misión formativa de Kinal: 65 años promoviendo la dignificación de la persona a través del trabajo técnico de excelencia.",
            "Dinámica de networking directivo: presentación de las empresas participantes y sectores productivos representados."
        ]),
        ("1.2 Inmersión Vivencial: Recorrido Guiado por Laboratorios de Alta Tecnología TSU", [
            "Visita a los Laboratorios de Automatización Industrial y PLCs Siemens, Robótica y Mecatrónica CNC.",
            "Demostraciones en vivo a cargo de estudiantes de último año de TSU: calibración, diagnóstico analítico y comisionamiento.",
            "Recorrido por los centros de redes, electrónica aplicada, manufactura y tecnologías de información.",
            "Evidencia empírica de la madurez técnica y disciplina operativa del perfil TSU."
        ]),
        ("1.3 Mesa de Diálogo 1: Panorama Actual del Empleo Técnico en Guatemala", [
            "Diagnóstico compartido sobre la disponibilidad de personal técnico calificado en la industria guatemalteca.",
            "El dilema del 'Missing Middle': la brecha insatisfecha de mandos medios técnicos entre operarios básicos e ingenieros.",
            "Levantamiento preliminar de inquietudes y expectativas de los representantes empresariales."
        ])
    ]

    for sub_tit, sub_items in temas_e1:
        p_st = doc.add_paragraph()
        style_heading(p_st, sub_tit, level=3)
        for it in sub_items:
            p_it = doc.add_paragraph()
            style_p(p_it, space_before=1, space_after=2)
            p_it.paragraph_format.left_indent = Inches(0.25)
            r_b = p_it.add_run("• ")
            r_b.font.bold = True
            r_b.font.color.rgb = COLOR_PRIMARY
            r_txt = p_it.add_run(it)
            r_txt.font.name = "Calibri"
            r_txt.font.size = Pt(9.5)
            r_txt.font.color.rgb = COLOR_DARK

    # -------------------------------------------------------------
    # ENCUENTRO 2: SESIÓN VIRTUAL 1
    # -------------------------------------------------------------
    h2_e2 = doc.add_paragraph()
    style_heading(h2_e2, "SESIÓN VIRTUAL 1: MANEJO Y DESARROLLO DE PERSONAL TÉCNICO (2 Horas en Microsoft Teams)", level=2)

    temas_e2 = [
        ("2.1 Atracción y Reclutamiento Efectivo de Talento Operativo", [
            "Por qué las descripciones tradicionales de puesto fracasan al contratar personal de taller y planta.",
            "Traducción de perfiles burocráticos a matrices de competencias observables: saber, saber hacer y saber ser.",
            "Redacción de ofertas atractivas y canales idóneos para conectar con jóvenes técnicos calificados."
        ]),
        ("2.2 Curva de Aprendizaje, Inducción Técnica y Reducción de Rotación", [
            "El costo oculto de la rotación temprana: pérdida de conocimiento de planta, retrabajos y paradas no planificadas.",
            "Estructuración de procesos de onboarding técnico acelerado: cómo lograr que un nuevo técnico sea productivo y seguro en 30 días.",
            "Diseño de planes de desarrollo y crecimiento operativo dentro de la empresa para fidelizar al colaborador."
        ])
    ]

    for sub_tit, sub_items in temas_e2:
        p_st = doc.add_paragraph()
        style_heading(p_st, sub_tit, level=3)
        for it in sub_items:
            p_it = doc.add_paragraph()
            style_p(p_it, space_before=1, space_after=2)
            p_it.paragraph_format.left_indent = Inches(0.25)
            r_b = p_it.add_run("• ")
            r_b.font.bold = True
            r_b.font.color.rgb = COLOR_PRIMARY
            r_txt = p_it.add_run(it)
            r_txt.font.name = "Calibri"
            r_txt.font.size = Pt(9.5)
            r_txt.font.color.rgb = COLOR_DARK

    # -------------------------------------------------------------
    # ENCUENTRO 3: SESIÓN VIRTUAL 2
    # -------------------------------------------------------------
    h2_e3 = doc.add_paragraph()
    style_heading(h2_e3, "SESIÓN VIRTUAL 2: LIDERAZGO Y FACTOR HUMANO EN CUADRILLAS (2 Horas en Microsoft Teams)", level=2)

    temas_e3 = [
        ("3.1 La Filosofía del Factor Humano de Fundación Kinal", [
            "La persona en el centro de los procesos productivos: el trabajo como medio de realización y servicio social.",
            "Los 4 Pilares Operativos: Respeto a la Dignidad, Trabajo Bien Hecho, Espíritu de Servicio y Libertad Responsable.",
            "El «Trabajo Bien Hecho»: orden 5S, custodia de herramientas, veracidad en bitácoras y rigor profesional."
        ]),
        ("3.2 Erradicación del Maltrato en Cuadrillas y Seguridad Psicológica", [
            "Superación de la cultura del capataz: erradicación del maltrato verbal, el sarcasmo y el autoritarismo en planta.",
            "Inteligencia emocional aplicada a la supervisión bajo presión de paradas de línea o emergencias operativas.",
            "Seguridad psicológica: cómo lograr que los colaboradores reporten averías y cuasi-accidentes inmediatamente sin temor a represalias.",
            "Diseño de canales éticos y códigos de convivencia respetuosa en áreas operativas."
        ])
    ]

    for sub_tit, sub_items in temas_e3:
        p_st = doc.add_paragraph()
        style_heading(p_st, sub_tit, level=3)
        for it in sub_items:
            p_it = doc.add_paragraph()
            style_p(p_it, space_before=1, space_after=2)
            p_it.paragraph_format.left_indent = Inches(0.25)
            r_b = p_it.add_run("• ")
            r_b.font.bold = True
            r_b.font.color.rgb = COLOR_PRIMARY
            r_txt = p_it.add_run(it)
            r_txt.font.name = "Calibri"
            r_txt.font.size = Pt(9.5)
            r_txt.font.color.rgb = COLOR_DARK

    # -------------------------------------------------------------
    # ENCUENTRO 4: SESIÓN VIRTUAL 3
    # -------------------------------------------------------------
    h2_e4 = doc.add_paragraph()
    style_heading(h2_e4, "SESIÓN VIRTUAL 3: FORTALECIMIENTO DE COMPETENCIAS TÉCNICAS Y FORMACIÓN DUAL (2 Horas en Teams)", level=2)

    temas_e4 = [
        ("4.1 Estándares Internacionales (DQR 5-6) y el Retorno de Inversión del TSU", [
            "El Marco Alemán de Cualificaciones (DQR): niveles 5 y 6 (Maestro Técnico / Meister, Bachelor Professional y TSU).",
            "La trampa de la subcualificación (operarios empíricos sin diagnóstico analítico) vs sobrecualificación (ingenieros frustrados en piso).",
            "Análisis financiero del ROI del TSU: reducción del tiempo medio de reparación (MTTR), disponibilidad de línea y estabilidad laboral."
        ]),
        ("4.2 Modelos de Co-Formación Dual Escuela-Empresa", [
            "El concepto oficial de la Formación Dual: aprender en el proceso de trabajo como base y objetivo formativo.",
            "Estructuración de convenios de semilleros técnicos y pasantías planificadas.",
            "El rol del Tutor Empresarial (*Ausbilder*): pedagogía del trabajo y método de 4 pasos para enseñar tareas operativas."
        ])
    ]

    for sub_tit, sub_items in temas_e4:
        p_st = doc.add_paragraph()
        style_heading(p_st, sub_tit, level=3)
        for it in sub_items:
            p_it = doc.add_paragraph()
            style_p(p_it, space_before=1, space_after=2)
            p_it.paragraph_format.left_indent = Inches(0.25)
            r_b = p_it.add_run("• ")
            r_b.font.bold = True
            r_b.font.color.rgb = COLOR_PRIMARY
            r_txt = p_it.add_run(it)
            r_txt.font.name = "Calibri"
            r_txt.font.size = Pt(9.5)
            r_txt.font.color.rgb = COLOR_DARK

    # -------------------------------------------------------------
    # ENCUENTRO 5: SESIÓN VIRTUAL 4
    # -------------------------------------------------------------
    h2_e5 = doc.add_paragraph()
    style_heading(h2_e5, "SESIÓN VIRTUAL 4: DESARROLLO DE MANDOS MEDIOS Y NECESIDADES ACTUALES (2 Horas en Teams)", level=2)

    temas_e5 = [
        ("5.1 El Desarrollo Integral del Mando Medio en la Industria", [
            "El supervisor como eslabón crítico: articulación entre los objetivos estratégicos de la gerencia y la realidad del piso.",
            "Habilidades de comunicación asertiva, retroalimentación constructiva y gestión de caracteres difíciles.",
            "Diseño de 'Charlas de 10 minutos de Factor Humano': reflexiones breves de valores y seguridad al iniciar el turno."
        ]),
        ("5.2 Necesidades Tecnológicas Emergentes y Medición de Impacto", [
            "Demandas actuales de la Industria 4.0: automatización, ciberseguridad industrial, sensores IoT y mantenimiento predictivo.",
            "Vincular el Factor Humano con indicadores duros de negocio: OEE (Overall Equipment Effectiveness), tasas de accidentabilidad (SSO) y merma.",
            "Preparación de la ficha de diagnóstico empresarial para la jornada plenaria de cierre."
        ])
    ]

    for sub_tit, sub_items in temas_e5:
        p_st = doc.add_paragraph()
        style_heading(p_st, sub_tit, level=3)
        for it in sub_items:
            p_it = doc.add_paragraph()
            style_p(p_it, space_before=1, space_after=2)
            p_it.paragraph_format.left_indent = Inches(0.25)
            r_b = p_it.add_run("• ")
            r_b.font.bold = True
            r_b.font.color.rgb = COLOR_PRIMARY
            r_txt = p_it.add_run(it)
            r_txt.font.name = "Calibri"
            r_txt.font.size = Pt(9.5)
            r_txt.font.color.rgb = COLOR_DARK

    # -------------------------------------------------------------
    # ENCUENTRO 6: JORNADA PRESENCIAL DE CIERRE
    # -------------------------------------------------------------
    h2_e6 = doc.add_paragraph()
    style_heading(h2_e6, "JORNADA PRESENCIAL ESPECIAL 2: DIÁLOGO INDUSTRIA-ACADEMIA, DIPLOMAS Y RECONOCIMIENTOS (4 Horas en Kinal)", level=2)

    temas_e6 = [
        ("6.1 Mesa Redonda Plenaria: Demandas Actuales del Sector Empresarial", [
            "Sistematización de hallazgos: qué competencias técnicas y humanas demandan prioritariamente las empresas participantes.",
            "Identificación de brechas de formación y retroalimentación directa a los directores de carrera de Fundación Kinal.",
            "Estrategias conjuntas para fortalecer la empleabilidad y pertinencia de los estudiantes y egresados de TSU."
        ]),
        ("6.2 Acuerdos de Vinculación y Rutas de Cooperación Empresa-Kinal", [
            "Mecanismos de acceso directo y prioritario a la Bolsa de Egresados de Kinal.",
            "Oportunidades de convenios de semilleros duales y pasantías especializadas.",
            "Certificación de competencias laborales de colaboradores de planta en instalaciones de Kinal."
        ]),
        ("6.3 Acto Solemne de Clausura y Entrega de Distinciones", [
            "Entrega de Diplomas Institucionales de Participación a los representantes de las empresas.",
            "Entrega de Reconocimientos Empresariales Especiales a las empresas aliadas participantes.",
            "Brindis de honor y networking directivo de vinculación."
        ])
    ]

    for sub_tit, sub_items in temas_e6:
        p_st = doc.add_paragraph()
        style_heading(p_st, sub_tit, level=3)
        for it in sub_items:
            p_it = doc.add_paragraph()
            style_p(p_it, space_before=1, space_after=2)
            p_it.paragraph_format.left_indent = Inches(0.25)
            r_b = p_it.add_run("• ")
            r_b.font.bold = True
            r_b.font.color.rgb = COLOR_PRIMARY
            r_txt = p_it.add_run(it)
            r_txt.font.name = "Calibri"
            r_txt.font.size = Pt(9.5)
            r_txt.font.color.rgb = COLOR_DARK

    # ==========================================
    # 3. REFERENCIAS BIBLIOGRÁFICAS
    # ==========================================
    h1_3 = doc.add_paragraph()
    style_heading(h1_3, "3. Fuentes Bibliográficas y Documentos de Apoyo", level=1)

    referencias = [
        ("BIBB / DQR:", "German Qualifications Framework for Lifelong Learning (DQR) — Niveles 5 y 6 (Meister y Formación Dual)."),
        ("OIT / Cinterfor:", "Formación Profesional Dual y Vinculación Escuela-Empresa en América Latina: Hacia un Diálogo Social Técnico."),
        ("Fundación Kinal:", "Ideario Institucional: El Trabajo como Vocación y Servicio a la Persona Humana."),
        ("Fundación Kinal:", "Perfil del Egresado del Técnico Superior Universitario (TSU) y Modelos de Formación en Alternancia."),
        ("Fundación Kinal:", "Manual de Factor Humano: Liderazgo y Trabajo Bien Hecho en Cuadrillas Técnicas.")
    ]

    for org, det in referencias:
        p_r = doc.add_paragraph()
        style_p(p_r, space_before=1, space_after=2)
        p_r.paragraph_format.left_indent = Inches(0.2)
        r_b = p_r.add_run(f"• {org} ")
        r_b.font.name = "Calibri"
        r_b.font.size = Pt(9)
        r_b.font.bold = True
        r_b.font.color.rgb = COLOR_PRIMARY
        r_d = p_r.add_run(det)
        r_d.font.name = "Calibri"
        r_d.font.size = Pt(9)
        r_d.font.color.rgb = COLOR_DARK

    section = doc.sections[0]
    footer = section.footer
    footer_p = footer.paragraphs[0]
    style_p(footer_p, space_before=0, space_after=0)
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_f = footer_p.add_run("Fundación Kinal | Programa Ejecutivo de Vinculación en Talento Técnico — Temario Oficial")
    run_f.font.name = "Calibri"
    run_f.font.size = Pt(8.5)
    run_f.font.color.rgb = COLOR_SECONDARY

    for out_dir in [os.path.join("Cursos", "Talento_Tecnico"), "Talento_Tecnico"]:
        os.makedirs(out_dir, exist_ok=True)
        out_path = os.path.join(out_dir, "Temario_Diplomado_Talento_Tecnico_Kinal.docx")
        doc.save(out_path)
        print(f"Temario guardado con éxito en: {out_path}")

if __name__ == "__main__":
    generate_temario_doc()
