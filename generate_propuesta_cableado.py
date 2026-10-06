import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=70, bottom=70, left=90, right=90):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E0", sz="4"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_callout(doc, text, bold_prefix="", fill_hex="F7FAFC", border_color="003366"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    
    cell = tbl.cell(0, 0)
    set_cell_background(cell, fill_hex)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=120)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 51, 102)
        
    r_txt = p.add_run(text)
    r_txt.font.name = "Calibri"
    r_txt.font.size = Pt(10)
    r_txt.font.italic = True
    r_txt.font.color.rgb = RGBColor(45, 55, 72)
    
    p_post = doc.add_paragraph()
    style_p(p_post, space_before=0, space_after=6)

def style_heading(p, text, level=1, color=RGBColor(0, 51, 102)):
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.bold = True
    r.font.color.rgb = color
    if level == 1:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        r.font.size = Pt(15)
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        r.font.size = Pt(13)
    elif level == 3:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        r.font.size = Pt(11)

def style_p(p, space_before=0, space_after=6, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def generate_propuesta_doc():
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    COLOR_PRIMARY = RGBColor(0, 51, 102)      # Azul Kinal
    COLOR_SECONDARY = RGBColor(74, 85, 104)    # Gris pizarra
    COLOR_DARK = RGBColor(26, 32, 44)         # Texto oscuro

    # =========================================================================
    # PORTADA INSTITUCIONAL
    # =========================================================================
    p_inst = doc.add_paragraph()
    style_p(p_inst, space_before=20, space_after=2)
    r_inst = p_inst.add_run("FUNDACIÓN KINAL")
    r_inst.font.name = "Calibri"
    r_inst.font.size = Pt(14)
    r_inst.font.bold = True
    r_inst.font.color.rgb = COLOR_PRIMARY

    p_esc = doc.add_paragraph()
    style_p(p_esc, space_before=0, space_after=20)
    r_esc = p_esc.add_run("Escuela Técnica Superior — Coordinación de Formación Continua y Empleabilidad")
    r_esc.font.name = "Calibri"
    r_esc.font.size = Pt(11)
    r_esc.font.color.rgb = COLOR_SECONDARY

    p_tit = doc.add_paragraph()
    style_p(p_tit, space_before=14, space_after=10)
    r_tit = p_tit.add_run("PLANIFICACIÓN Y PROPUESTA FORMATIVA INSTITUCIONAL:\nCABLEADO ESTRUCTURADO Y REDES DE COBRE Y FIBRA ÓPTICA")
    r_tit.font.name = "Calibri"
    r_tit.font.size = Pt(18)
    r_tit.font.bold = True
    r_tit.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    style_p(p_sub, space_before=0, space_after=20)
    r_sub = p_sub.add_run("Programa de Certificación Técnica Profesional para la Canalización, Conectorización, Fusión y Certificación Instrumental de Telecomunicaciones (80 Horas — 20 Sesiones)")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(11.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = COLOR_SECONDARY

    # Ficha Técnica
    tbl_ficha = doc.add_table(rows=8, cols=2)
    tbl_ficha.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ficha.autofit = False
    tbl_ficha.columns[0].width = Inches(2.2)
    tbl_ficha.columns[1].width = Inches(4.3)
    set_table_borders(tbl_ficha, color="CBD5E0", sz="6")

    ficha_data = [
        ("Nivel Formativo de Referencia:", "Equivalencia DQR Nivel 4 - 5 (Competencia técnica especializada y autonomía operativa)"),
        ("Población Destinataria:", "Técnicos instaladores de redes, soporte IT, electricistas de baja tensión y contratistas de telecomunicaciones"),
        ("Duración Total:", "80 horas pedagógicas (20 sesiones de 4 horas / 240 min; 73h 20min de actividad neta de taller)"),
        ("Modalidad y Régimen:", "Presencial en Laboratorios de Redes de Kinal con apoyo digital en Kinal.academy (Moodle) (80% Taller / 20% Plataforma)"),
        ("Estructura de Horarios:", "Dos días entre semana (Martes y Jueves de 17:30 a 21:30 hrs), sincronizado con el instructor de TICs"),
        ("Orientación de Empleabilidad:", "Técnico especialista en cableado estructurado, instalador de fibra óptica, supervisor de infraestructura física"),
        ("Nota Mínima Aprobatoria:", "75 puntos sobre 100 en todas las comprobaciones prácticas de ponchado, fusión y certificación"),
        ("Sede Presencial:", "Fundación Kinal — Laboratorios de Redes y Telecomunicaciones, Sede Central, Zona 7, Ciudad de Guatemala")
    ]

    for idx, (label, val) in enumerate(ficha_data):
        row = tbl_ficha.rows[idx]
        set_cell_background(row.cells[0], "F7FAFC")
        set_cell_margins(row.cells[0], top=65, bottom=65, left=90, right=90)
        set_cell_margins(row.cells[1], top=65, bottom=65, left=90, right=90)
        
        p0 = row.cells[0].paragraphs[0]
        style_p(p0, space_before=0, space_after=0)
        r0 = p0.add_run(label)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_PRIMARY

        p1 = row.cells[1].paragraphs[0]
        style_p(p1, space_before=0, space_after=0)
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_DARK

    doc.add_page_break()

    # =========================================================================
    # 1. MARCO FILOSÓFICO E IDEARIO DE KINAL: ÉTICA Y EL TRABAJO BIEN HECHO
    # =========================================================================
    h1_1 = doc.add_paragraph()
    style_heading(h1_1, "1. Marco Filosófico e Ideario Institucional de Fundación Kinal", level=1)

    p_id_desc = doc.add_paragraph()
    style_p(p_id_desc, space_before=0, space_after=6)
    r = p_id_desc.add_run("El diseño formativo de esta certificación técnica se fundamenta de forma irrenunciable en la visión humanística y ética de Fundación Kinal, asegurando que el desarrollo de habilidades técnicas esté sustentado en la honradez profesional y el espíritu de servicio:")
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = COLOR_DARK

    add_callout(doc, "«Formar a jóvenes y adultos a través de una educación integral, con énfasis en las áreas técnicas y tecnológicas, influyendo positivamente en su trabajo, su familia y la sociedad».", bold_prefix="Misión Institucional de Kinal:")

    add_callout(doc, "«Visión cristiana de la vida. Respeto a la dignidad de la persona. Espíritu de servicio. Trabajo bien hecho. Libertad personal responsable».", bold_prefix="Valores Nucleares de Kinal:")

    p_et = doc.add_paragraph()
    style_heading(p_et, "1.1 El Principio del «Trabajo Bien Hecho» en el Cableado Estructurado y Redes de Alta Velocidad", level=2)

    p_et_txt = doc.add_paragraph()
    style_p(p_et_txt, space_before=0, space_after=6)
    r = p_et_txt.add_run("En la infraestructura de telecomunicaciones de alta frecuencia (donde las señales eléctricas viajan a 250 MHz en Cat 6 y 500 MHz en Cat 6A), la física no perdona la improvisación. Un milímetro de destrenzado excesivo o un cincho plástico demasiado apretado arruinan la comunicación corporativa. En este programa, el «trabajo bien hecho» se concreta en principios innegociables:\n"
                          "• Estética, Prolijidad y Respeto a la Física del Cable: Peinado impecable en mazos simétricos de 12 o 24 cables con cinchos textiles de velcro reutilizables (prohibición terminante de cinchos plásticos que estrangulan pares), respeto estricto a los radios mínimos de curvatura (4 veces el diámetro exterior) y cero tensión mecánica en los remates.\n"
                          "• Integridad, Verdad y Ética en la Certificación Instrumental: Cero tolerancia al maquillaje de reportes de escaneo. El estudiante aprende que alterar los límites de prueba en el certificador Fluke para forzar un 'Pass' falso es una falta ética grave. Todo enlace defectuoso debe ser diagnosticado mediante análisis de causa raíz (NEXT, Return Loss, Insertion Loss) y reconstruido con excelencia.\n"
                          "• Orden y Trazabilidad según ANSI/TIA-606-D: Rotulado indeleble y normalizado desde el primer momento en ambos extremos del cable, en puertos de patch panel y en tomas de pared (faceplates), evitando el desorden que incrementa exponencialmente los costos operativos de mantenimiento en las empresas.")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_DARK

    # =========================================================================
    # 2. FUNDAMENTACIÓN TÉCNICA Y CASO TRANSVERSAL DE EDIFICIO CORPORATIVO
    # =========================================================================
    h1_2 = doc.add_paragraph()
    style_heading(h1_2, "2. Fundamentación Técnica, Enfoque Pedagógico y Caso Transversal", level=1)

    p_met = doc.add_paragraph()
    style_p(p_met, space_before=0, space_after=6)
    r = p_met.add_run("El programa responde directamente a las carencias del mercado laboral en Guatemala, donde la mayoría de instaladores opera de manera empírica sin conocer los estándares internacionales ni disponer de instrumental de certificación:\n"
                      "• Caso Transversal de Edificio Corporativo (Torre Empresarial 3 Niveles): A lo largo de las 20 sesiones, los participantes abordan el diseño e implementación integral de la infraestructura pasiva para un edificio de 3 plantas que alberga 120 puestos de trabajo, telefonía VoIP, puntos de acceso Wi-Fi 6, cámaras de seguridad CCTV IP y control de acceso con tecnología PoE++ (IEEE 802.3bt). El caso abarca la Entrada de Servicios (EF), el Cuarto de Equipos principal (ER), dos Cuartos de Telecomunicaciones (TR), la canalización con charolas y tubería EMT, el tendido horizontal en Cat 6A, el backbone de fibra óptica multimodo/monomodo y el sistema de puesta a tierra TIA-607.\n"
                      "• Práctica Directa en Racks y Módulos de Taller: Cada estudiante ejecuta el montaje físico en bastidores de 19 pulgadas, conectorización en módulos Keystone Jack, ponchado en patch panels de 24 puertos, preparación y corte de fibra óptica con cleaver de precisión, empalme por fusión mediante arco voltaico y certificación de enlaces permanentes con escáner Fluke Networks.\n"
                      "• Metodología de la Acción Completa: Se aplica el ciclo pedagógico alemán de 6 fases: Informar, Planificar, Decidir, Ejecutar, Controlar y Valorar, fomentando autonomía operativa y responsabilidad profesional.")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_DARK

    # =========================================================================
    # 3. MARCO DQR Y SISTEMA DUAL DE APRENDIZAJE
    # =========================================================================
    h1_3 = doc.add_paragraph()
    style_heading(h1_3, "3. Marco Alemán de Cualificaciones (DQR Nivel 4-5) y Enfoque Dual", level=1)

    add_callout(doc, "«concepto de aprendizaje para jóvenes que tiene como objetivo preparar a los aprendices para su vida profesional, por lo que la formación se realiza en régimen de alternancia entre la escuela o el centro de formación y la empresa. El aprender en el proceso del trabajo es base y a la vez objetivo de la acción formativa».", bold_prefix="Definición de Formación Profesional Dual:")

    add_callout(doc, "«the ability and readiness of the individual to use knowledge, skills and personal, social and methodological competences and to behave in a considered, individual and socially responsible manner. Competence is understood in this sense as the comprehensive ability to act».", bold_prefix="Concepto de Competencia en el DQR:")

    p_dqr_txt = doc.add_paragraph()
    style_p(p_dqr_txt, space_before=0, space_after=6)
    r = p_dqr_txt.add_run("El programa articula las dimensiones de competencia técnica y humana del marco DQR:\n"
                          "1. Dimensión de Competencia Profesional (Professional Competence):\n"
                          "   • Conocimientos (Knowledge): Dominio riguroso de la suite de normas ANSI/TIA-568, 569, 606, 607, estándar ISO/IEC 11801, física de propagación de radiofrecuencia en par trenzado y óptica ondulatoria en núcleos de silicio.\n"
                          "   • Destrezas (Skills): Destreza manual de alta precisión para pelado de cubiertas sin marcar conductores, corte perpendicular de fibra con cleaver a 90° (+/- 0.5 grados), fusión de núcleo con pérdida inferior a 0.05 dB, remate 110 con mínimo destrenzado (< 13 mm) y operación diestra de equipos de certificación Tier 1.\n"
                          "2. Dimensión de Competencia Personal (Personal Competence):\n"
                          "   • Competencia Social (Social Competence): Coordinación en cuadrilla para tendido de cables sin torsión, comunicación asertiva con supervisores de obra civil y entrega profesional al cliente.\n"
                          "   • Autonomía (Autonomy): Capacidad para diagnosticar de forma independiente enlaces fallidos, interpretar gráficas de NEXT y Return Loss en el dominio del tiempo (TDNXT, TDR) y tomar decisiones de rediseño según normativas.")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_DARK

    # =========================================================================
    # 4. COMPETENCIA GENERAL Y PERFILES DEL PARTICIPANTE
    # =========================================================================
    h1_4 = doc.add_paragraph()
    style_heading(h1_4, "4. Competencia General y Perfiles del Participante", level=1)

    p_comp = doc.add_paragraph()
    style_p(p_comp, space_before=0, space_after=6)
    r = p_comp.add_run("Competencia General del Programa:\n"
                       "El egresado planifica, canaliza, tiende, remata, fusiona y certifica sistemas de cableado estructurado en cobre (Cat 6 / Cat 6A) y fibra óptica (Monomodo OS2 / Multimodo OM4), ejecutando el montaje pulcro de racks de 19 pulgadas y cuartos de telecomunicaciones, aterrizaje equipotencial bajo TIA-607, administración documental bajo TIA-606 y diagnóstico instrumental de enlaces con escáneres Fluke Networks, garantizando cero fallas de paradiafonía y cumplimiento del estándar institucional del trabajo bien hecho.")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = COLOR_PRIMARY

    p_perfiles = doc.add_paragraph()
    style_p(p_perfiles, space_before=4, space_after=6)
    r = p_perfiles.add_run("• Perfil de Ingreso: Técnicos de soporte informático, electricistas de baja tensión, instaladores de CCTV y seguridad electrónica, bachilleres o peritos técnicos y profesionales independientes con interés en especializarse en infraestructura de telecomunicaciones. Requiere habilidad manual motriz fina, discriminación visual de colores y conocimientos básicos de computación.\n"
                           "• Perfil de Egreso y Empleabilidad: Técnico especialista calificado para desempeñarse en empresas integradoras de telecomunicaciones, constructoras de obra corporativa, departamentos de IT corporativos (bancos, multinacionales, data centers) y cuadrillas de despliegue de fibra óptica, con solvencia para entregar proyectos certificados listos para garantía de fabricante de 25 años.")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_DARK

    # =========================================================================
    # 5. ESTRUCTURA MODULAR Y DOSIFICACIÓN DE LA CARGA HORARIA (80 HORAS)
    # =========================================================================
    h1_5 = doc.add_paragraph()
    style_heading(h1_5, "5. Estructura Modular y Distribución Formativa (80 Horas — 20 Sesiones)", level=1)

    modulos_data = [
        ("Módulo 1", "Normas ANSI/TIA, Espacios y Canalizaciones Físicas (TR, ER, Racks)", "20 Horas", "5 Sesiones", "Diseño de Espacios, Cálculo de Llenado TIA-569 y Ensamble de Bastidor de 19'' Nivelado"),
        ("Módulo 2", "Cableado de Cobre de Alto Rendimiento (Cat 6/6A) y Aterrizaje TIA-607", "20 Horas", "5 Sesiones", "Enlace Permanente Cat 6A Rematado en Patch Panel de 24p y Aterrizaje TGB Comprobado"),
        ("Módulo 3", "Infraestructura de Fibra Óptica, Conectorización y Empalme por Fusión", "20 Horas", "5 Sesiones", "Empalme por Fusión de Fibra Óptica con Pérdida < 0.05 dB y Conectorización LC/SC"),
        ("Módulo 4", "Certificación Instrumental con Fluke, Rotulado TIA-606 y Proyecto As-Built", "20 Horas", "5 Sesiones", "Certificación Completa Fluke DSX con Reportes PDF, Rotulado TIA-606 y Dossier As-Built")
    ]

    tbl_mod = doc.add_table(rows=5, cols=5)
    tbl_mod.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_mod.autofit = False
    tbl_mod.columns[0].width = Inches(1.0)
    tbl_mod.columns[1].width = Inches(2.2)
    tbl_mod.columns[2].width = Inches(0.8)
    tbl_mod.columns[3].width = Inches(0.9)
    tbl_mod.columns[4].width = Inches(1.6)
    set_table_borders(tbl_mod, color="CBD5E0", sz="4")

    hdr_mod = tbl_mod.rows[0]
    for c in hdr_mod.cells:
        set_cell_background(c, "003366")
    headers_mod_txt = ["Módulo", "Nombre de la Unidad Formativa", "Horas", "Sesiones", "Producto Terminal del Módulo"]
    for idx, text in enumerate(headers_mod_txt):
        p = hdr_mod.cells[idx].paragraphs[0]
        style_p(p, space_before=0, space_after=0)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for idx, (m_id, m_nom, m_hrs, m_ses, m_prod) in enumerate(modulos_data, start=1):
        row = tbl_mod.rows[idx]
        if idx % 2 == 1:
            for c in row.cells:
                set_cell_background(c, "F7FAFC")
        for c_idx, val in enumerate([m_id, m_nom, m_hrs, m_ses, m_prod]):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=55, bottom=55, left=65, right=65)
            p = cell.paragraphs[0]
            style_p(p, space_before=0, space_after=0)
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(8.5)
            r.font.color.rgb = COLOR_DARK
            if c_idx in [0, 2]:
                r.font.bold = True

    p_balance = doc.add_paragraph()
    style_p(p_balance, space_before=8, space_after=6)
    r = p_balance.add_run("Distribución del Tiempo Pedagógico y Compatibilidad Docente:\n"
                          "• Duración Exacta:** 80 horas pedagógicas distribuidas en 20 sesiones de 4 horas (240 minutos por encuentro; 220 minutos netos de actividad y 20 minutos de receso).\n"
                          "• Calendario Entre Semana (2 Días): Impartido en días alternos entre semana (ej. Martes y Jueves de 17:30 a 21:30 hrs), permitiendo que el mismo instructor titular de TICs de Kinal atienda este curso y el de Ciberseguridad sin traslapes ni fatiga pedagógica.\n"
                          "• Balance Formativo: 80% Práctica en Racks y Módulos de Taller (64 horas) y 20% Normativa, cálculo y diseño en software/plataforma digital (16 horas).")
    r.font.name = "Calibri"
    r.font.size = Pt(9.5)
    r.font.color.rgb = COLOR_DARK

    # =========================================================================
    # 6. ESTRATEGIA DIDÁCTICA EN TALLER Y BITÁCORA BERICHTSHEFT
    # =========================================================================
    h1_6 = doc.add_paragraph()
    style_heading(h1_6, "6. Estrategia Didáctica en Taller y Bitácora Semanal «Berichtsheft»", level=1)

    p_didact = doc.add_paragraph()
    style_p(p_didact, space_before=0, space_after=6)
    r = p_didact.add_run("Para garantizar que la destreza manual se traduzca en calidad de grado industrial, se implementan:\n"
                         "1. Estaciones de Trabajo por Parejas (Taller Vivo): Cada pareja dispone de un bastidor o sección de rack de 19'', patch panel modular, barra TGB de tierra, tramos de canalización, carretes de cable Cat 6A y bandejas de empalme óptico.\n"
                         "2. Inspección Escalonada de Calidad: Ningún cable se conecta al switch o se certifica con el Fluke sin antes pasar por la revisión de peinado, radio de curvatura y desforre por parte del instructor.\n"
                         "3. Bitácora Semanal Berichtsheft: Registro documental individual donde el estudiante anota: fecha, tramo instalado, estándares aplicados, problemas de paradiafonía o atenuación detectados, causa raíz identificada y corrección ejecutada bajo el lema del «trabajo bien hecho».")
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_DARK

    # =========================================================================
    # 7. SISTEMA DE EVALUACIÓN Y UMBRAL APROBATORIO INSTITUCIONAL
    # =========================================================================
    h1_7 = doc.add_paragraph()
    style_heading(h1_7, "7. Sistema de Evaluación Institucional y Criterio de Acreditación", level=1)

    add_callout(doc, "La acreditación del programa y la entrega del certificado institucional de Fundación Kinal exigen una calificación mínima de 75 puntos sobre 100 en todas las rúbricas de taller, comprobaciones de ponchado y fusión, y en la evaluación del proyecto integrador terminal.", bold_prefix="Umbral Aprobatorio Institucional Obligatorio:")

    eval_data = [
        ("Destreza de Canalización, Tendido y Aterrizaje TIA-607 (20%)", "Evaluación del montaje de racks, nivelación, curvatura de tubería EMT, instalación de charolas malla y conexión de barra TGB con conductor 6 AWG."),
        ("Conectorización y Peinado de Cobre Cat 6/6A (30%)", "Evaluación de remates en jacks y patch panels de 24 puertos: peinado con velcro, desforre sin mallas cortadas, destrenzado < 13 mm y cero pares divididos."),
        ("Preparación, Corte y Fusión de Fibra Óptica (20%)", "Limpieza de fibra con alcohol isopropílico, corte perpendicular con cleaver (< 1 grado), empalme por fusión con atenuación medida < 0.05 dB y horneado de manguito."),
        ("Bitácora Berichtsheft y Rotulado TIA-606 (10%)", "Revisión quincenal de bitácoras de taller y etiquetado industrial normalizado en cables, jacks y paneles con rotuladora profesional."),
        ("Certificación Instrumental Fluke y Proyecto As-Built (20%)", "Prueba terminal individual: certificación Tier 1 con Fluke DSX, diagnóstico de 2 fallas inducidas en menos de 45 min y entrega del dossier As-Built.")
    ]

    tbl_ev = doc.add_table(rows=6, cols=2)
    tbl_ev.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ev.autofit = False
    tbl_ev.columns[0].width = Inches(2.5)
    tbl_ev.columns[1].width = Inches(4.0)
    set_table_borders(tbl_ev, color="CBD5E0", sz="4")

    hdr_ev = tbl_ev.rows[0]
    for c in hdr_ev.cells:
        set_cell_background(c, "003366")
    for idx, text in enumerate(["Componente de Evaluación", "Criterio de Desempeño y Evidencia"]):
        p = hdr_ev.cells[idx].paragraphs[0]
        style_p(p, space_before=0, space_after=0)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    for idx, (comp, crit) in enumerate(eval_data, start=1):
        row = tbl_ev.rows[idx]
        if idx % 2 == 1:
            for c in row.cells:
                set_cell_background(c, "F7FAFC")
        for c_idx, val in enumerate([comp, crit]):
            cell = row.cells[c_idx]
            set_cell_margins(cell, top=55, bottom=55, left=65, right=65)
            p = cell.paragraphs[0]
            style_p(p, space_before=0, space_after=0)
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(8.5)
            r.font.color.rgb = COLOR_DARK
            if c_idx == 0:
                r.font.bold = True

    # =========================================================================
    # 8. INFRAESTRUCTURA Y SEGURIDAD INDUSTRIAL EN KINAL
    # =========================================================================
    h1_8 = doc.add_paragraph()
    style_heading(h1_8, "8. Equipamiento de Laboratorio y Requerimientos de Seguridad", level=1)

    p_infra = doc.add_paragraph()
    style_p(p_infra, space_before=0, space_after=6)
    r = p_infra.add_run("El curso se desarrolla en el Laboratorio de Redes y Telecomunicaciones de Fundación Kinal, equipado con:\n"
                        "• Bastidores de Telecomunicaciones: Racks abiertos de piso de 19'' (42U) y gabinetes de pared abatibles de 12U con organizadores horizontales de 1U y 2U.\n"
                        "• Aparamenta Pasiva: Patch panels modulares de 24 puertos Cat 6 y Cat 6A apantallados, módulos Keystone Jack RJ45, faceplates de 2 y 4 puertos, bandejas de distribución de fibra óptica (ODF) de 1U con acopladores dúplex LC y SC.\n"
                        "• Instrumental de Certificación: Certificador Fluke Networks DSX-5000 / DSX-8000 con adaptadores de Enlace Permanente y Canal Cat 6A, medidor de potencia óptica (OPM), fuente de luz calibrada y localizador visual de fallas (VFL láser rojo 650 nm).\n"
                        "• Equipos de Fusión Óptica: Fusionadoras por alineación de núcleo con electrodos calibrados, cortadoras de precisión con disco de diamante (cleavers) y peladoras de fibra de 3 posiciones.\n"
                        "• Normas de Seguridad en Taller: Uso obligatorio de gafas de policarbonato con protección lateral al cortar fibra óptica, contenedor sellado para recortes de fibra (prohibido arrojar residuos al piso por riesgo de punción cutánea), guantes de precisión mecánica y tapete negro de trabajo para contraste visual.")
    r.font.name = "Calibri"
    r.font.size = Pt(9.5)
    r.font.color.rgb = COLOR_DARK

    section = doc.sections[0]
    footer = section.footer
    footer_p = footer.paragraphs[0]
    style_p(footer_p, space_before=0, space_after=0)
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_f = footer_p.add_run("Fundación Kinal | Cableado Estructurado y Redes de Cobre/Fibra — Propuesta Curricular Oficial")
    run_f.font.name = "Calibri"
    run_f.font.size = Pt(8.5)
    run_f.font.color.rgb = COLOR_SECONDARY

    out_dir = os.path.join("Cursos", "Cableado_Estructurado")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "Propuesta_Curso_Cableado_Estructurado_Kinal.docx")
    doc.save(out_path)
    print(f"Propuesta guardada con éxito en: {out_path}")

if __name__ == "__main__":
    generate_propuesta_doc()
