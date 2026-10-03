import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = docx.Document()

    # Configuración de márgenes (1 pulgada / 2.54 cm estándar)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Estilos de colores institucionales Kinal
    HEX_PRIMARY = "003366"      # Azul institucional Kinal
    HEX_SECONDARY = "4A5568"    # Gris pizarra elegante
    HEX_DARK = "1A202C"         # Texto principal oscuro
    HEX_LIGHT_BG = "F7FAFC"     # Fondo sutil para recuadros
    HEX_BORDER = "CBD5E0"       # Borde sutil

    COLOR_PRIMARY = RGBColor(0, 51, 102)
    COLOR_SECONDARY = RGBColor(74, 85, 104)
    COLOR_DARK = RGBColor(26, 32, 44)

    # Función auxiliar para formatear párrafos
    def style_p(p, space_before=0, space_after=6, line_spacing=1.15):
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing

    # ==========================================
    # ENCABEZADO PRINCIPAL (# Nombre del curso)
    # ==========================================
    title_p = doc.add_paragraph()
    style_p(title_p, space_before=0, space_after=12)
    run_title = title_p.add_run("Mecatrónica Industrial y Fabricación Digital Aplicada")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    # Subtítulo institucional
    sub_p = doc.add_paragraph()
    style_p(sub_p, space_before=0, space_after=18)
    run_sub = sub_p.add_run("Programa Formativo Técnico Profesional — Fundación Kinal | Nivel DQR 4-5")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_SECONDARY

    # Separador sutil
    border_p = doc.add_paragraph()
    style_p(border_p, space_before=0, space_after=12)
    pPr = border_p._element.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="4" w:color="{HEX_PRIMARY}"/></w:pBdr>')
    pPr.append(pBdr)

    # ==========================================
    # ## Misión de Kinal
    # ==========================================
    h2_mision = doc.add_paragraph()
    style_p(h2_mision, space_before=10, space_after=4)
    run_h2_mision = h2_mision.add_run("Misión de Kinal")
    run_h2_mision.font.name = "Calibri"
    run_h2_mision.font.size = Pt(14)
    run_h2_mision.font.bold = True
    run_h2_mision.font.color.rgb = COLOR_PRIMARY

    p_mision = doc.add_paragraph()
    style_p(p_mision, space_before=0, space_after=12)
    run_mision = p_mision.add_run("«Formar a jóvenes y adultos a través de una educación integral, con énfasis en las áreas técnicas y tecnológicas, influyendo positivamente en su trabajo, su familia y la sociedad».")
    run_mision.font.name = "Calibri"
    run_mision.font.size = Pt(11)
    run_mision.font.italic = True
    run_mision.font.color.rgb = COLOR_DARK

    # ==========================================
    # ## Visión de Kinal
    # ==========================================
    h2_vision = doc.add_paragraph()
    style_p(h2_vision, space_before=10, space_after=4)
    run_h2_vision = h2_vision.add_run("Visión de Kinal")
    run_h2_vision.font.name = "Calibri"
    run_h2_vision.font.size = Pt(14)
    run_h2_vision.font.bold = True
    run_h2_vision.font.color.rgb = COLOR_PRIMARY

    p_vision = doc.add_paragraph()
    style_p(p_vision, space_before=0, space_after=12)
    run_vision = p_vision.add_run("Ser líderes en la formación técnica, tecnológica y humana de la región, propiciando la superación personal, laboral y social de nuestros estudiantes con un alto sentido ético y excelencia profesional.")
    run_vision.font.name = "Calibri"
    run_vision.font.size = Pt(11)
    run_vision.font.color.rgb = COLOR_DARK

    # ==========================================
    # ## Competencia del curso
    # ==========================================
    h2_comp = doc.add_paragraph()
    style_p(h2_comp, space_before=12, space_after=6)
    run_h2_comp = h2_comp.add_run("Competencia del curso")
    run_h2_comp.font.name = "Calibri"
    run_h2_comp.font.size = Pt(14)
    run_h2_comp.font.bold = True
    run_h2_comp.font.color.rgb = COLOR_PRIMARY

    # Datos generales con viñetas negritas requeridas
    items_generales = [
        ("Competencia general:", " Diseña, fabrica e integra sistemas mecatrónicos aplicados a la industria combinando manufactura digital (corte láser, impresión 3D y mecanizado CNC) con sensórica, actuadores y controladores, actuando con apego al ideario del trabajo bien hecho, seguridad industrial y normativas técnicas."),
        ("Duración total en horas:", " 90 horas pedagógicas presenciales (distribuidas en 20 sesiones sabatinas de 4.5 horas)."),
        ("Horario:", " Sábados de 8:00 a 12:30 horas (febrero a junio, 5 meses continuos)."),
        ("Modalidad:", " Presencial en instalaciones de Fundación Kinal (Laboratorios de Automatización y Talleres de Fabricación Digital, Sede Central, Zona 7, Ciudad de Guatemala)."),
        ("Perfil de ingreso:", " Dirigido a técnicos, electromecánicos, bachilleres industriales y estudiantes de ingeniería con nociones básicas de electricidad o mecánica, interesados en manufactura y automatización, con compromiso de puntualidad, respeto a las normas de seguridad del taller y actitud hacia el aprendizaje técnico disciplinado."),
        ("Perfil de egreso:", " El egresado manufactura piezas mecánicas de precisión mediante corte láser, impresión 3D y fresado CNC, conectando y calibrando sensores y actuadores electromecánicos para poner en marcha y mantener celdas mecatrónicas industriales bajo estándares internacionales, criterios de calidad rigurosos y el valor del trabajo bien hecho.")
    ]

    for label, text in items_generales:
        p_item = doc.add_paragraph()
        style_p(p_item, space_before=0, space_after=6)
        run_lbl = p_item.add_run(label + " ")
        run_lbl.font.name = "Calibri"
        run_lbl.font.size = Pt(11)
        run_lbl.font.bold = True
        run_lbl.font.color.rgb = COLOR_PRIMARY

        run_txt = p_item.add_run(text)
        run_txt.font.name = "Calibri"
        run_txt.font.size = Pt(11)
        run_txt.font.color.rgb = COLOR_DARK

    # ==========================================
    # ## Temario:
    # ==========================================
    doc.add_page_break()

    h2_temario = doc.add_paragraph()
    style_p(h2_temario, space_before=10, space_after=8)
    run_h2_temario = h2_temario.add_run("Temario:")
    run_h2_temario.font.name = "Calibri"
    run_h2_temario.font.size = Pt(16)
    run_h2_temario.font.bold = True
    run_h2_temario.font.color.rgb = COLOR_PRIMARY

    # Módulos estructurados
    modulos = [
        {
            "num": "1",
            "nombre": "Diseño CAD 3D, Metrología de Taller y Prototipado con Cortadora Láser (18 Horas — Febrero)",
            "indicador": "Modela componentes y gabinetes en software CAD 3D aplicando metrología de taller y parametrizando la cortadora láser para manufacturar ensambles rígidos que cumplan tolerancias mecánicas inferiores a ±0.2 mm.",
            "temas": [
                {
                    "titulo": "Fundamentos de Mecatrónica Industrial y CAD Paramétrico",
                    "subtemas": [
                        "Sinergia mecatrónica y normas de seguridad en el taller Kinal (EPP y protocolo 5S)",
                        "Metrología dimensional aplicada con pie de rey digital y micrómetro de exteriores",
                        "Restricciones geométricas, croquizado paramétrico y operaciones de extrusión 3D"
                    ]
                },
                {
                    "titulo": "Ingeniería Inversa y Diseño para Manufactura por Corte Láser (DFM)",
                    "subtemas": [
                        "Replicación y digitalización dimensional de piezas desgastadas o discontinuadas",
                        "Conceptos de sangría de corte (kerf), potencias y velocidades por material",
                        "Diseño de juntas mecánicas autoblocantes (finger joints y encastres a presión)"
                    ]
                },
                {
                    "titulo": "Operación Segura y Parametrización de Cortadora Láser CO2",
                    "subtemas": [
                        "Preparación de vectores en capas de marcado, grabado y corte (DXF y SVG)",
                        "Calibración de distancia focal, encendido de extracción de humos y asistencia de aire",
                        "Corte y grabado práctico sobre acrílico industrial, Delrin (POM) y MDF técnico"
                    ]
                },
                {
                    "titulo": "Fabricación de Gabinetes Mecatrónicos y Evaluación Práctica",
                    "subtemas": [
                        "Montaje mecánico y verificación de rigidez estructural de chasis y gabinetes",
                        "Ajustes dimensionales finales para alojar tarjetas electrónicas y rieles DIN",
                        "Evaluación práctica del Módulo 1 (Construcción de Gabinete y Fixture de Sujeción)"
                    ]
                }
            ]
        },
        {
            "num": "2",
            "nombre": "Fabricación Aditiva Avanzada (Impresión 3D FDM) para Mecanismos Industriales (18 Horas — Marzo)",
            "indicador": "Configura parámetros de laminación y opera impresoras 3D FDM seleccionando filamentos técnicos para fabricar piezas cinemáticas, engranes y actuadores articulados resistentes al esfuerzo dinámico.",
            "temas": [
                {
                    "titulo": "Cinemática Aplicada y Ciencia de Polímeros para FDM",
                    "subtemas": [
                        "Cálculo y diseño de engranes rectos, poleas sincrónicas GT2 y levas mecánicas",
                        "Arquitectura cinemática de impresoras 3D industriales (CoreXY y cartesianas)",
                        "Propiedades mecánicas y térmicas de filamentos: PLA+, PETG, ABS, TPU y Nylon"
                    ]
                },
                {
                    "titulo": "Diseño para Fabricación Aditiva (DFAM) y Parámetros en Laminador",
                    "subtemas": [
                        "Orientación de pieza respecto a líneas de esfuerzo y anisotropía mecánica",
                        "Optimización de voladizos (overhangs), puentes (bridging) y soportes solubles",
                        "Configuración de perímetros, densidad y patrones de relleno estructural (giroide y cúbico)"
                    ]
                },
                {
                    "titulo": "Tolerancias Print-in-Place e Integración de Insertos Roscados",
                    "subtemas": [
                        "Modelado de tolerancias dinámicas para mecanismos ensamblados en cama (0.2 a 0.4 mm)",
                        "Instalación de insertos roscados de latón en caliente (heat-set inserts)",
                        "Calibración de pasos por milímetro, nivelación de cama y compensación de flujo"
                    ]
                },
                {
                    "titulo": "Montaje de Subensambles Mecánicos y Evaluación Práctica",
                    "subtemas": [
                        "Integración de baleros lineales, ejes rectificados de acero y tornillería métrica",
                        "Ensayos de movimiento continuo, lubricación de engranes y verificación de fatiga",
                        "Evaluación práctica del Módulo 2 (Fabricación de Garra Robótica Articulada)"
                    ]
                }
            ]
        },
        {
            "num": "3",
            "nombre": "Mecanizado Sustractivo y Fresado CNC: CAM, Código G y Fabricación de Piezas (18 Horas — Abril)",
            "indicador": "Genera trayectorias de maquinado CAM seguras y opera el router/fresadora CNC de Kinal para producir bancadas y componentes mecánicos respetando tolerancias geométricas y dimensionales.",
            "temas": [
                {
                    "titulo": "Cinemática CNC y Estructura del Código G y M",
                    "subtemas": [
                        "Ejes coordenados cartesianos (X, Y, Z), coordenadas de máquina (G53) y de trabajo (G54-G59)",
                        "Sintaxis de comandos G (G00, G01, G02, G03) y comandos misceláneos M (M03, M05, M06)",
                        "Métodos de sujeción segura: prensas de precisión, bridas de fijación y prevención de colisiones"
                    ]
                },
                {
                    "titulo": "Programación en Software CAM (Estrategias 2.5D y 3D)",
                    "subtemas": [
                        "Definición de bloque de material (stock), plano de seguridad y origen de trabajo (WCS)",
                        "Estrategias de fresado: planeado, cajeras (pockets), contorneado exterior y taladrado ciclado",
                        "Cálculo de parámetros de corte: velocidad de corte (Vc), avance por diente (fz) y RPM"
                    ]
                },
                {
                    "titulo": "Alistamiento, Montaje de Herramientas y Puesta a Cero en CNC Kinal",
                    "subtemas": [
                        "Protocolos de seguridad industrial en maquinado sustractivo y parada de emergencia",
                        "Montaje y verificación de excentricidad en boquillas ER y fresas de carburo sólido",
                        "Establecimiento de ceros de pieza (X=0, Y=0) y calibración de eje Z con sonda electrónica"
                    ]
                },
                {
                    "titulo": "Mecanizado Práctico, Metrología de Verificación y Evaluación",
                    "subtemas": [
                        "Maquinado efectivo de piezas en materiales de ingeniería (Nylamid, Delrin y aluminio)",
                        "Inspección dimensional con micrómetro y escuadra de precisión, desbarbado y acabado superficial",
                        "Evaluación práctica del Módulo 3 (Mecanizado de Bancada / Placa Base Mecatrónica)"
                    ]
                }
            ]
        },
        {
            "num": "4",
            "nombre": "Sensórica Industrial, Actuación Electromecánica/Neumática y Control de Ejes (18 Horas — Mayo)",
            "indicador": "Interconecta sensores industriales y configura controladores de movimiento para gobernar motores a pasos y actuadores electroneumáticos bajo normas internacionales de cableado y seguridad eléctrica.",
            "temas": [
                {
                    "titulo": "Transducción y Sensórica Industrial",
                    "subtemas": [
                        "Principio de operación de sensores inductivos, capacitivos y fotoeléctricos (réflex y barrera)",
                        "Conexión e identificación de salidas a 3 y 4 hilos: lógica de conmutación PNP vs NPN",
                        "Finales de carrera mecánicos y sensores magnéticos Reed Switch para detección de posición"
                    ]
                },
                {
                    "titulo": "Control de Movimiento con Motores a Pasos y Drivers Industriales",
                    "subtemas": [
                        "Curvas de torque vs velocidad en motores NEMA y dimensionamiento de carga",
                        "Configuración de microstepping, límites de corriente y señales lógicas (PUL, DIR, ENA)",
                        "Control de servomotores industriales y gobernación en lazo abierto y cerrado"
                    ]
                },
                {
                    "titulo": "Actuación Electroneumática y Normas de Cableado Kinal",
                    "subtemas": [
                        "Electroválvulas de 24 VDC (monoestables y biestables) y actuadores de simple y doble efecto",
                        "Aplicación del principio del 'trabajo bien hecho': canalización estética en ductos ranurados",
                        "Uso obligatorio de casquillos (ferrules), separación de potenciales y rotulado alfanumérico (IEC 60617)"
                    ]
                },
                {
                    "titulo": "Programación de Perfiles de Movimiento y Evaluación Práctica",
                    "subtemas": [
                        "Lógica de rampas de aceleración y desaceleración para evitar pérdida de pasos",
                        "Integración de circuitos de parada de emergencia por corte de energía seguro",
                        "Evaluación práctica del Módulo 4 (Sistema de Posicionamiento Lineal Motorizado)"
                    ]
                }
            ]
        },
        {
            "num": "5",
            "nombre": "Integración de Celda Mecatrónica, Automatización y Proyecto Capstone Industrial (18 Horas — Junio)",
            "indicador": "Integra subsistemas mecánicos fabricados en láser, 3D y CNC con sensórica, actuadores y rutinas de control para operar una celda mecatrónica automática funcional para la industria guatemalteca.",
            "temas": [
                {
                    "titulo": "Arquitectura de Integración Mecatrónica e Interfaz de Operador",
                    "subtemas": [
                        "Diseño de tableros de control con botoneras industriales (arranque, paro, rearme) y luces piloto",
                        "Diagramas de flujo y autómatas de estado finito (FSM) para ciclos secuenciales industriales",
                        "Protocolos de comunicación y enlace entre controlador y módulos periféricos"
                    ]
                },
                {
                    "titulo": "Ensamble Mecánico y Cableado Definitivo de la Celda",
                    "subtemas": [
                        "Montaje integral de bancada CNC, partes cinemáticas 3D y gabinetes cortados en láser",
                        "Verificación punto a punto de continuidad eléctrica y aislamiento de seguridad",
                        "Calibración mecánica de tensión de bandas, escuadras y alineación de guías"
                    ]
                },
                {
                    "titulo": "Puesta en Marcha, Sincronización y Diagnóstico Sistemático (Troubleshooting)",
                    "subtemas": [
                        "Carga del programa secuencial y pruebas de marcha en vacío (dry run)",
                        "Protocolos de localización de averías y resolución de fallas inducidas en sensores y actuadores",
                        "Optimización de tiempos de ciclo y validación de seguridad operativa"
                    ]
                },
                {
                    "titulo": "Evaluación Terminal Integrada (Capstone) y Clausura",
                    "subtemas": [
                        "Pruebas de funcionamiento continuo en tiempo real (mínimo 10 ciclos completos autónomos)",
                        "Sustentación técnica individual y en equipo de la celda mecatrónica ante jurado evaluador",
                        "Entrega y revisión de la Bitácora de Taller (Berichtsheft) y evaluación final (umbral >= 75 pts)"
                    ]
                }
            ]
        }
    ]

    for mod in modulos:
        # Módulo X: {Nombre del módulo}
        h3_mod = doc.add_paragraph()
        style_p(h3_mod, space_before=14, space_after=4)
        run_h3 = h3_mod.add_run(f"Módulo {mod['num']}: {mod['nombre']}")
        run_h3.font.name = "Calibri"
        run_h3.font.size = Pt(13)
        run_h3.font.bold = True
        run_h3.font.color.rgb = COLOR_PRIMARY

        # Indicador de logro:
        p_ind = doc.add_paragraph()
        style_p(p_ind, space_before=0, space_after=6)
        run_ind_lbl = p_ind.add_run("Indicador de logro: ")
        run_ind_lbl.font.name = "Calibri"
        run_ind_lbl.font.size = Pt(11)
        run_ind_lbl.font.bold = True
        run_ind_lbl.font.color.rgb = COLOR_SECONDARY

        run_ind_txt = p_ind.add_run(mod["indicador"])
        run_ind_txt.font.name = "Calibri"
        run_ind_txt.font.size = Pt(11)
        run_ind_txt.font.color.rgb = COLOR_DARK

        # Temas y subtemas
        for t in mod["temas"]:
            p_tema = doc.add_paragraph()
            style_p(p_tema, space_before=3, space_after=2)
            p_tema.paragraph_format.left_indent = Inches(0.25)
            run_t_bullet = p_tema.add_run("• Tema: ")
            run_t_bullet.font.name = "Calibri"
            run_t_bullet.font.size = Pt(11)
            run_t_bullet.font.bold = True
            run_t_bullet.font.color.rgb = COLOR_PRIMARY

            run_t_title = p_tema.add_run(t["titulo"])
            run_t_title.font.name = "Calibri"
            run_t_title.font.size = Pt(11)
            run_t_title.font.bold = True
            run_t_title.font.color.rgb = COLOR_DARK

            for sub in t["subtemas"]:
                p_sub = doc.add_paragraph()
                style_p(p_sub, space_before=1, space_after=2)
                p_sub.paragraph_format.left_indent = Inches(0.55)
                run_sub_dash = p_sub.add_run("– ")
                run_sub_dash.font.name = "Calibri"
                run_sub_dash.font.size = Pt(10.5)
                run_sub_dash.font.color.rgb = COLOR_SECONDARY

                run_sub_txt = p_sub.add_run(sub)
                run_sub_txt.font.name = "Calibri"
                run_sub_txt.font.size = Pt(10.5)
                run_sub_txt.font.color.rgb = COLOR_DARK

    # ==========================================
    # ## Materiales necesarios
    # ==========================================
    doc.add_page_break()

    h2_mat = doc.add_paragraph()
    style_p(h2_mat, space_before=12, space_after=6)
    run_h2_mat = h2_mat.add_run("Materiales necesarios")
    run_h2_mat.font.name = "Calibri"
    run_h2_mat.font.size = Pt(16)
    run_h2_mat.font.bold = True
    run_h2_mat.font.color.rgb = COLOR_PRIMARY

    p_mat_intro = doc.add_paragraph()
    style_p(p_mat_intro, space_before=0, space_after=8)
    run_intro = p_mat_intro.add_run("Para el óptimo desarrollo teórico-práctico del curso se requiere la disposición y uso de la infraestructura y consumibles clasificados en los siguientes rubros:")
    run_intro.font.name = "Calibri"
    run_intro.font.size = Pt(11)
    run_intro.font.color.rgb = COLOR_DARK

    categorias_materiales = [
        ("Maquinaria y Equipos de Fabricación Digital (Instalaciones Kinal):", [
            "Cortadora y grabadora láser CO2 (área mínima de trabajo de 600x400 mm, tubo de 60W a 100W, con compresor air assist y extractor de humos operativo).",
            "Granja de impresoras 3D FDM (volumen de impresión de al menos 220x220x250 mm, cama térmica hasta 100°C, boquillas de 0.4 mm de acero endurecido o latón).",
            "Router / Fresadora CNC de 3 ejes (área de trabajo mínima de 300x400 mm, husillo con refrigeración y control de velocidad variable, pinzas ER11 o ER20 y sonda electrónica Z)."
        ]),
        ("Materiales Consumibles de Taller y Fabricación:", [
            "Láminas de acrílico cristal y de colores (espesores de 3 mm y 5 mm) y tableros de MDF técnico (3 mm y 5.5 mm).",
            "Placas de Delrin (POM), Nylamid o aluminio mecanizable para el fresado en CNC.",
            "Bobinas de filamento técnico para impresión 3D: PLA+, PETG industrial y TPU (flexible).",
            "Tornillería métrica de precisión (M3, M4, M5 en diversas longitudes), tuercas de seguridad, arandelas e insertos roscados de latón para inserción térmica.",
            "Kit de rodamientos lineales (LM8UU o similares), ejes de acero rectificado de 8 mm y baleros 608ZZ."
        ]),
        ("Componentes de Automatización, Sensórica y Control:", [
            "Fuentes de alimentación industriales conmutadas de 24 VDC / 5A a 10A con montaje para riel DIN.",
            "Motores a pasos híbridos NEMA 17 y NEMA 23 de alto torque.",
            "Drivers industriales para motores a pasos (tipo TB6600, DM542 o equivalentes) con microstepping.",
            "Microcontroladores / PLC compactos con salidas a relevador o transistor y puertos de comunicación.",
            "Kit de sensórica industrial: sensores inductivos de 24 VDC (PNP/NPN), sensores capacitivos, sensores ópticos réflex y finales de carrera mecánicos.",
            "Válvulas electroneumáticas 5/2 de 24 VDC, cilindros neumáticos compactos y manguera de poliuretano de 4 y 6 mm con conexiones rápidas."
        ]),
        ("Herramientas de Montaje, Cableado y EPP:", [
            "Herramientas de mano: pinzas pelacables de precisión, ponchadora de terminales tubulares (ferrules), juego de destornilladores dieléctricos y llaves Allen milimétricas.",
            "Instrumentos de medición: calibradores vernier digitales, micrómetros de exteriores y multímetros digitales autorrango.",
            "Consumibles de cableado: cable flexible de control calibre 18 y 20 AWG (código de colores), casquillos ferrules aislados, canaleta ranurada plástica y etiquetas alfanuméricas.",
            "Equipo de Protección Personal (EPP obligatorio en Kinal): lentes de seguridad transparentes de alto impacto, gabacha/bata de taller de algodón o lona, zapatos cerrados de suela de hule antideslizante (o de seguridad) y tapones auditivos para área CNC."
        ])
    ]

    for cat_title, items in categorias_materiales:
        p_cat = doc.add_paragraph()
        style_p(p_cat, space_before=8, space_after=3)
        run_cat = p_cat.add_run(f"• {cat_title}")
        run_cat.font.name = "Calibri"
        run_cat.font.size = Pt(11.5)
        run_cat.font.bold = True
        run_cat.font.color.rgb = COLOR_PRIMARY

        for it in items:
            p_it = doc.add_paragraph()
            style_p(p_it, space_before=1, space_after=2)
            p_it.paragraph_format.left_indent = Inches(0.35)
            run_dash = p_it.add_run("– ")
            run_dash.font.name = "Calibri"
            run_dash.font.size = Pt(10.5)
            run_dash.font.color.rgb = COLOR_SECONDARY

            run_text = p_it.add_run(it)
            run_text.font.name = "Calibri"
            run_text.font.size = Pt(10.5)
            run_text.font.color.rgb = COLOR_DARK

    # Pie de página institucional
    section = doc.sections[0]
    footer = section.footer
    footer_p = footer.paragraphs[0]
    style_p(footer_p, space_before=0, space_after=0)
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_f = footer_p.add_run("Fundación Kinal | Mecatrónica Industrial y Fabricación Digital Aplicada — Temario Oficial")
    run_f.font.name = "Calibri"
    run_f.font.size = Pt(9)
    run_f.font.color.rgb = COLOR_SECONDARY

    # Crear carpeta output si no existe
    os.makedirs("output", exist_ok=True)
    file_path = os.path.join("output", "Temario_Curso_Mecatronica_Kinal.docx")
    doc.save(file_path)
    print(f"Documento guardado con éxito en: {file_path}")

if __name__ == "__main__":
    create_document()
